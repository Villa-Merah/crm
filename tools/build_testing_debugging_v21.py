from __future__ import annotations

from pathlib import Path
import sys

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "storage" / "app" / "doc-tools"))

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = ROOT / "docs" / "DOKUMEN-TESTING-DAN-DEBUGGING-DAN-SECURITY-CRM-VILLA-MERAH-v2.2.docx"

NAVY = "0B2545"
BLUE = "1476B8"
BLUE_DARK = "075985"
MUTED = "64748B"
LIGHT = "F6F8FC"
PALE_BLUE = "E8EEF5"
PALE_GREEN = "ECFDF5"
PALE_AMBER = "FFFBEB"
PALE_RED = "FFF1F2"
GREEN = "047857"
AMBER = "9A6700"
RED = "B42318"
WHITE = "FFFFFF"
BORDER = "D7E1EC"
WIDTH_DXA = 9360
TABLE_INDENT = 120


def font(run, size=11, bold=False, color=NAVY, name="Calibri", italic=False):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)
    return run


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    node = tc_pr.find(qn("w:shd"))
    if node is None:
        node = OxmlElement("w:shd")
        tc_pr.append(node)
    node.set(qn("w:fill"), fill)


def cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    mar = tc_pr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tc_pr.append(mar)
    for key, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        el = mar.find(qn(f"w:{key}"))
        if el is None:
            el = OxmlElement(f"w:{key}")
            mar.append(el)
        el.set(qn("w:w"), str(value))
        el.set(qn("w:type"), "dxa")


def border(cell, color=BORDER, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:color"), color)


def set_table_geometry(table, widths, indent=TABLE_INDENT):
    assert sum(widths) == WIDTH_DXA
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tc_w)
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")


def cant_split(row):
    row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))


def repeat_header(row):
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    row._tr.get_or_add_trPr().append(el)


def keep_next(paragraph):
    paragraph._p.get_or_add_pPr().append(OxmlElement("w:keepNext"))


def para(doc, text="", bold=False, color=NAVY, size=11, after=6, before=0, italic=False, align=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.25
    if align is not None:
        p.alignment = align
    font(p.add_run(text), size=size, bold=bold, color=color, italic=italic)
    return p


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.375)
    p.paragraph_format.first_line_indent = Inches(-0.188)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    font(p.add_run(text), size=11)


def numbered(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.left_indent = Inches(0.375)
    p.paragraph_format.first_line_indent = Inches(-0.188)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    font(p.add_run(text), size=11)


def heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    keep_next(p)
    font(p.add_run(text), size={1: 16, 2: 13, 3: 12}[level], bold=True,
         color={1: BLUE, 2: BLUE, 3: BLUE_DARK}[level])
    return p


def table(doc, headers, rows, widths, compact=False):
    tbl = doc.add_table(rows=1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
    repeat_header(tbl.rows[0])
    cant_split(tbl.rows[0])
    for idx, text in enumerate(headers):
        cell = tbl.rows[0].cells[idx]
        shade(cell, PALE_BLUE)
        border(cell, BLUE_DARK)
        cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cell.paragraphs[0].paragraph_format.space_after = Pt(0)
        font(cell.paragraphs[0].add_run(text), size=9 if compact else 9.5, bold=True, color=BLUE_DARK)
    for ridx, values in enumerate(rows):
        row = tbl.add_row()
        cant_split(row)
        for idx, value in enumerate(values):
            cell = row.cells[idx]
            shade(cell, WHITE if ridx % 2 == 0 else LIGHT)
            border(cell)
            cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            cell.paragraphs[0].paragraph_format.space_after = Pt(0)
            cell.paragraphs[0].paragraph_format.line_spacing = 1.1
            font(cell.paragraphs[0].add_run(str(value)), size=8.6 if compact else 9.2)
    if len(rows) <= 6:
        for row in tbl.rows[:-1]:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    keep_next(paragraph)
    set_table_geometry(tbl, widths)
    para(doc, "", after=1)
    return tbl


def callout(doc, title, body, fill=PALE_BLUE, accent=BLUE_DARK):
    tbl = doc.add_table(rows=1, cols=1)
    cell = tbl.cell(0, 0)
    shade(cell, fill)
    border(cell, accent, "10")
    cell_margins(cell, 120, 150, 120, 150)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    font(p.add_run(title + "\n"), size=10, bold=True, color=accent)
    font(p.add_run(body), size=9.4)
    cant_split(tbl.rows[0])
    set_table_geometry(tbl, [WIDTH_DXA], indent=150)
    para(doc, "", after=1)


def code(doc, lines):
    tbl = doc.add_table(rows=1, cols=1)
    cell = tbl.cell(0, 0)
    shade(cell, "0F172A")
    border(cell, "0F172A")
    cell_margins(cell, 120, 150, 120, 150)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    for idx, line in enumerate(lines):
        font(p.add_run(line + ("\n" if idx < len(lines) - 1 else "")), size=8, color="E2E8F0", name="Cascadia Mono")
    cant_split(tbl.rows[0])
    set_table_geometry(tbl, [WIDTH_DXA], indent=150)
    para(doc, "", after=1)


def page_break(doc):
    # Pagination alami lebih stabil ketika tabel di halaman sebelumnya berubah tinggi.
    return None


def setup():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(NAVY)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    for level, size, before, after, color in ((1, 16, 18, 10, BLUE), (2, 13, 14, 7, BLUE), (3, 12, 10, 5, BLUE_DARK)):
        style = doc.styles[f"Heading {level}"]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    for style_name in ("List Bullet", "List Number"):
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    font(header.add_run("CRM_VILLA MERAH  |  TESTING, DEBUGGING & SECURITY  |  V2.2"), size=8, bold=True, color=MUTED)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font(footer.add_run("Dokumen internal - 2 September 2026  |  Halaman "), size=8, color=MUTED)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    return doc


def cover(doc):
    para(doc, "QUALITY ASSURANCE & OPERATIONS", bold=True, color=BLUE, size=11, after=28)
    para(doc, "DOKUMEN TESTING\nDAN DEBUGGING", bold=True, color=NAVY, size=28, after=10)
    para(doc, "CRM_VILLA MERAH / SIVMI CRM", bold=True, color=BLUE, size=16, after=18)
    para(doc, "Panduan pengujian, diagnosis masalah, regression gate, troubleshooting lokal-Docker-hosting, keamanan sistem, dan catatan hasil aktual versi 2.1.", color=MUTED, size=12, after=28)
    table(doc, ["Atribut", "Keterangan"], [
        ["Versi dokumen", "2.2"],
        ["Tanggal baseline", "2 September 2026"],
        ["Commit aplikasi", "6e3bdb8 dan hotfix terkait peta/dashboard"],
        ["Target pembaca", "Developer, QA/UAT, administrator, DevOps, dan pemilik proses"],
        ["Klasifikasi", "Dokumen internal; tidak memuat credential atau data siswa"],
    ], [2700, 6660])
    callout(doc, "STATUS QUALITY GATE: CONDITIONAL", "Composer dan build frontend lulus. Suite terbaru mencatat 110 dari 112 test lulus dengan 680 assertion. Dua temuan regresi masih terbuka dan harus diselesaikan sebelum status rilis diubah menjadi LULUS PENUH.", PALE_AMBER, AMBER)


def build():
    doc = setup()
    cover(doc)
    page_break(doc)

    heading(doc, "Daftar Isi", 1)
    for item in [
        "1. Tujuan dan prinsip", "2. Arsitektur pengujian", "3. Standar lingkungan dan keamanan",
        "4. Strategi serta jenis testing", "5. Prosedur testing lokal dan Docker",
        "6. Pengujian fitur utama", "7. Hasil aktual baseline 2.1", "8. Proses debugging standar",
        "9. Playbook troubleshooting", "10. Temuan debugging versi 2.1", "11. Regression dan release gate",
        "12. Pelaporan bug dan UAT", "13. Security dan keamanan sistem",
        "Lampiran A. Perintah cepat", "Lampiran B. Formulir hasil uji",
    ]:
        bullet(doc, item)
    callout(doc, "Cara menggunakan dokumen", "Mulai dari Bab 5 untuk menjalankan test. Jika test atau aplikasi gagal, ikuti Bab 8 dan playbook Bab 9. Gunakan Bab 11 sebagai gate sebelum deployment, Bab 13 untuk verifikasi keamanan, dan Lampiran B untuk menyimpan bukti.")

    heading(doc, "1. Tujuan dan Prinsip", 1)
    para(doc, "Dokumen ini menetapkan cara membuktikan bahwa perubahan CRM berfungsi, aman terhadap data, konsisten lintas role/cabang, dan dapat dijalankan pada lokal, Docker, serta shared hosting.")
    heading(doc, "1.1 Tujuan", 2)
    for text in [
        "Mencegah regresi pada alur lead, sekolah, follow-up, pendaftaran, pembayaran, laporan, dashboard, dan WhatsApp.",
        "Menentukan akar masalah dari bukti: test, log, query, status service, HTTP response, serta perilaku browser.",
        "Memisahkan masalah kode, data, konfigurasi, dependency, infrastruktur, dan layanan eksternal.",
        "Menyediakan keputusan rilis yang dapat diaudit tanpa mengekspos secret atau data pribadi.",
    ]:
        bullet(doc, text)
    heading(doc, "1.2 Prinsip", 2)
    table(doc, ["Prinsip", "Penerapan"], [
        ["Reproduksi minimum", "Gunakan data dan langkah paling sedikit yang tetap memunculkan masalah."],
        ["Isolasi data", "Test memakai SQLite :memory:; jangan arahkan PHPUnit ke database operasional."],
        ["Bukti sebelum patch", "Simpan error, timestamp, role, filter, query/log tersanitasi, dan expected vs actual."],
        ["Patch minimum", "Ubah komponen penyebab, lalu jalankan targeted test dan full regression."],
        ["Kesetaraan runtime", "Uji SQLite untuk kecepatan dan MySQL Docker untuk perbedaan SQL engine."],
        ["Tidak melemahkan keamanan", "Jangan menonaktifkan authorization, validasi, SQL mode, CSRF, atau signature webhook demi membuat test lulus."],
    ], [2200, 7160])

    page_break(doc)
    heading(doc, "2. Arsitektur Pengujian", 1)
    table(doc, ["Lapisan", "Teknologi", "Fokus"], [
        ["Unit", "PHPUnit", "Normalizer nomor, resolver koordinat, helper, enum, dan aturan murni."],
        ["Feature", "Laravel Test + Livewire", "HTTP, autentikasi, RBAC, form Filament, query, import/export, dan integrasi internal."],
        ["Integration", "MySQL 8.4 Docker", "Migrasi, foreign key, GROUP BY, transaksi, query agregat, dan perbedaan dari SQLite."],
        ["Build", "Vite / npm", "Kompilasi Tailwind, manifest, font, JS, dan asset produksi."],
        ["Service", "Docker Compose", "PHP-FPM, Nginx, MySQL, Redis, queue, scheduler, node, WAHA."],
        ["Smoke/UAT", "Browser + HTTP", "Login, halaman kritis, peta, API Sekolah, QR WAHA, dan alur pengguna."],
    ], [1700, 2300, 5360], compact=True)
    heading(doc, "2.1 Topologi runtime", 2)
    for text in [
        "Lokal full Docker: web :8000, Vite :5173, MySQL host :3307, WAHA :3000.",
        "Di jaringan Compose, Laravel memakai hostname database, redis, dan waha; bukan localhost.",
        "Shared hosting menjalankan Laravel/MySQL, sedangkan WAHA dapat berada pada Docker lokal melalui Cloudflare Tunnel.",
        "Leaflet 1.9.4 disajikan dari public/vendor/leaflet; tile peta tetap berasal dari OpenStreetMap.",
    ]:
        bullet(doc, text)

    heading(doc, "3. Standar Lingkungan dan Keamanan", 1)
    table(doc, ["Area", "Standar"], [
        ["PHP", "8.3+; proyek saat ini menggunakan Laravel 12 dan Filament 5."],
        ["Node", "Node.js 22 dan npm; build harus exit code 0."],
        ["Database", "SQLite :memory: untuk suite otomatis; MySQL 8.x untuk integration/runtime."],
        ["Queue", "sync pada test; worker aktif pada runtime WhatsApp dan pekerjaan async."],
        ["Secret", "APP_KEY, DB password, WAHA API key, webhook secret, cookie, dan token tidak boleh dicetak ke laporan."],
        ["Data", "Gunakan factory/demo. Anonimkan nama, nomor WhatsApp, percakapan, dan bukti pembayaran."],
    ], [2200, 7160])
    callout(doc, "PERINGATAN", "Jangan menjalankan migrate:fresh, db:wipe, down -v, truncate, atau seeder demo di server produksi. Backup dan verifikasi target database sebelum tindakan yang mengubah data.", PALE_RED, RED)

    page_break(doc)
    heading(doc, "4. Strategi dan Jenis Testing", 1)
    table(doc, ["Jenis", "Kapan", "Kriteria lulus"], [
        ["Smoke", "Setelah deploy/start", "Health, login, asset, dan halaman utama merespons benar."],
        ["Unit", "Setiap perubahan helper", "Semua kasus normal, alias, null, dan edge case lulus."],
        ["Feature", "Setiap perubahan alur", "Role, validasi, persistence, redirect, dan output sesuai."],
        ["Regression", "Sebelum merge/rilis", "Seluruh suite lulus tanpa failure/error."],
        ["Integration MySQL", "Query/migrasi berubah", "Tidak ada error SQL mode, FK, tipe kolom, atau transaksi."],
        ["UAT", "Release candidate", "Pengguna bisnis menerima alur dan bukti ditandatangani."],
        ["Security", "Auth/API/webhook berubah", "Scope cabang, permission, rate/auth/signature, dan sanitasi lulus."],
        ["Recovery", "Backup/deploy besar", "Restore diuji dan rollback dapat dilakukan."],
    ], [1800, 2600, 4960], compact=True)
    heading(doc, "4.1 Prioritas risiko", 2)
    table(doc, ["Prioritas", "Contoh"], [
        ["P0 - Kritis", "Kehilangan/kebocoran data, auth bypass, transaksi rusak, aplikasi tidak dapat digunakan."],
        ["P1 - Tinggi", "Dashboard/query gagal, lead tidak tersimpan, webhook/queue berhenti, scope cabang bocor."],
        ["P2 - Sedang", "Filter salah, import sebagian gagal, peta tidak lengkap, pesan validasi membingungkan."],
        ["P3 - Rendah", "Teks, spacing, warning dependency opsional, atau masalah kosmetik."],
    ], [2000, 7360])

    heading(doc, "5. Prosedur Testing Lokal dan Docker", 1)
    heading(doc, "5.1 Quality gate lokal", 2)
    code(doc, [
        "cd C:\\laragon\\www\\sivmi.crm",
        "composer validate --no-check-publish",
        "php artisan test",
        "npm run build",
        "git diff --check",
    ])
    heading(doc, "5.2 Targeted test", 2)
    code(doc, [
        "php artisan test --filter=DashboardTest",
        "php artisan test --filter=IndonesiaCityCoordinatesTest",
        "php artisan test --filter=SchoolApi",
        "php artisan test --filter=Waha",
    ])
    heading(doc, "5.3 Full Docker", 2)
    code(doc, [
        "Set-ExecutionPolicy -Scope Process Bypass",
        ".\\docker\\start-full.ps1",
        ".\\docker\\test-full.ps1",
        "docker compose --env-file docker\\.env -f docker\\compose.yml ps",
        "docker compose --env-file docker\\.env -f docker\\compose.yml exec -T app php artisan migrate:status",
    ])
    callout(doc, "Urutan aman", "Jalankan targeted test setelah patch, lalu full suite, build, MySQL integration, smoke browser, dan terakhir UAT. Targeted test saja tidak cukup menjadi gate rilis.")

    page_break(doc)
    heading(doc, "6. Pengujian Fitur Utama", 1)
    table(doc, ["Modul", "Skenario minimum", "Expected"], [
        ["Login/RBAC", "Guest, akun aktif/nonaktif, role dan cabang", "Redirect/403/scope sesuai policy."],
        ["Leads", "Tambah nama+WA, normalisasi, duplikat, edit, filter, bulk action", "Hanya field wajib divalidasi; nomor unik; scope aman."],
        ["Sekolah", "Cari API, pilih, persist lokal, manual fallback, bulk action", "NPSN mencegah duplikat; API gagal tidak memblokir input manual."],
        ["Follow-up", "Jadwal, status, histori, PIC", "Riwayat tersimpan dan data hanya terlihat sesuai scope."],
        ["Pendaftaran", "Lead ke siswa dan program", "Relasi serta transisi status konsisten."],
        ["Pembayaran", "Nominal, status, cabang, laporan", "Omzet hanya menghitung status yang sah."],
        ["Dashboard", "Periode, cabang, PIC, KPI, source, kota, peta", "Agregat akurat pada SQLite dan MySQL; peta graceful fallback."],
        ["API publik", "Auth, validasi, status HTTP, idempotensi", "401/403/422/2xx tepat dan data tersanitasi."],
        ["WAHA", "Health, session, QR, send, webhook HMAC, queue", "Session sehat, event tidak duplikat, pesan diproses worker."],
        ["Import/export", "CSV valid/invalid, formula injection, baris duplikat", "Pesan Indonesia jelas; data aman saat diekspor."],
    ], [1800, 3900, 3660], compact=True)
    heading(doc, "6.1 Uji khusus peta", 2)
    for text in [
        "Prioritaskan leads.kota_asal; gunakan sekolah.kota hanya jika kota asal kosong.",
        "Gabungkan variasi kota yang sama dan verifikasi jumlah marker/circle serta coverage.",
        "Pastikan aset /vendor/leaflet/leaflet.js dan leaflet.css berstatus 200.",
        "Pastikan tile.openstreetmap.org tidak diblokir CSP/firewall dan data-cfasync=false mencegah Rocket Loader merusak urutan skrip.",
        "Jalankan query agregat pada MySQL dengan ONLY_FULL_GROUP_BY aktif.",
    ]:
        bullet(doc, text)

    heading(doc, "7. Hasil Aktual Baseline 2.1", 1)
    callout(doc, "RINGKASAN EKSEKUSI 2 SEPTEMBER 2026", "Composer valid dan Vite build berhasil. Full suite menjalankan 112 test: 110 lulus, satu failure, satu error, total 680 assertion. Status rilis tetap CONDITIONAL sampai dua regresi ditutup.", PALE_AMBER, AMBER)
    table(doc, ["Gate", "Hasil", "Status"], [
        ["Composer", "composer.json valid", "LULUS"],
        ["PHPUnit penuh", "110/112 test lulus; 680 assertion", "BELUM LULUS"],
        ["Dashboard", "10 test; 81 assertion", "LULUS"],
        ["Resolver lokasi", "4 test; 23 assertion", "LULUS"],
        ["Vite build", "Build produksi selesai; warning fontaine opsional", "LULUS DENGAN CATATAN"],
        ["MySQL Dashboard", "Query role Admin JakPus dieksekusi pada MySQL Docker", "LULUS"],
        ["Produksi", "Leaflet tampil dan agregasi kota hotfix terverifikasi visual", "LULUS TERBATAS"],
    ], [2200, 5000, 2160], compact=True)
    heading(doc, "7.1 Regresi terbuka", 2)
    table(doc, ["ID", "Gejala", "Analisis awal", "Tindakan"], [
        ["TST-01", "SettingsTest mengharapkan nama aplikasi dinamis tampil di Dashboard.", "Brand topbar kini tetap CRM_VILLA MERAH sementara nilai setting tersimpan. Kontrak test dan keputusan branding tidak lagi selaras.", "Tentukan apakah nama harus dinamis. Perbaiki provider/view atau ubah expectation berdasarkan keputusan produk."],
        ["TST-02", "ReportTest gagal karena nomor WhatsApp duplikat.", "Fixture createTransaction memakai 08123456789 untuk dua cabang, sedangkan aturan unik global kini aktif.", "Buat nomor fixture unik per transaksi; pertahankan validasi duplikat produksi."],
    ], [1000, 2600, 3460, 2300], compact=True)
    callout(doc, "Keputusan rilis", "Kegagalan TST-02 merupakan masalah data test, tetapi tetap harus diperbaiki agar regression gate bersih. TST-01 memerlukan keputusan bisnis karena dapat berarti test usang atau fitur pengaturan identitas tidak lagi tercermin di UI.", PALE_RED, RED)

    page_break(doc)
    heading(doc, "8. Proses Debugging Standar", 1)
    for text in [
        "Catat gejala: waktu, environment, URL/perintah, role/cabang, filter, expected, actual, dan dampak.",
        "Reproduksi dengan data minimum pada test environment.",
        "Kumpulkan bukti tersanitasi: exception pertama, SQLSTATE, HTTP status, console/network, container status, dan commit.",
        "Klasifikasikan sumber: kode, test fixture, database/schema, cache, dependency, service, jaringan, atau layanan eksternal.",
        "Buat hipotesis tunggal dan lakukan pemeriksaan read-only untuk membuktikannya.",
        "Terapkan patch minimum tanpa melemahkan kontrol keamanan.",
        "Jalankan targeted test, full regression, build, MySQL integration, smoke, dan UAT terkait.",
        "Dokumentasikan akar masalah, patch, risiko, rollback, bukti, dan status penutupan.",
    ]:
        numbered(doc, text)
    heading(doc, "8.1 Perintah diagnosis umum", 2)
    code(doc, [
        "git rev-parse --short HEAD",
        "php artisan about",
        "php artisan migrate:status",
        "php artisan route:list",
        "php artisan optimize:clear",
        "tail -n 300 storage/logs/laravel.log",
        "docker compose --env-file docker\\.env -f docker\\compose.yml ps",
        "docker compose --env-file docker\\.env -f docker\\compose.yml logs --tail=200 app web queue waha",
    ])
    heading(doc, "8.2 Membaca error", 2)
    table(doc, ["Petunjuk", "Makna umum"], [
        ["SQLSTATE 42S22", "Kolom tidak ada; migrasi/cache/schema tidak sinkron."],
        ["SQLSTATE 42000 / 1055", "Query tidak kompatibel ONLY_FULL_GROUP_BY."],
        ["SQLSTATE 23000 / 1451", "Parent masih direferensikan foreign key."],
        ["cURL error 6", "DNS/hostname tidak dapat di-resolve dari runtime tersebut."],
        ["401 / 403", "Credential/signature/permission/scope salah."],
        ["419", "Session/CSRF/domain/cookie bermasalah."],
        ["422", "Validasi input gagal; periksa field dan pesan localization."],
        ["ERR_CONNECTION_REFUSED", "Service tidak berjalan atau port/binding salah."],
    ], [2900, 6460])

    heading(doc, "9. Playbook Troubleshooting", 1)
    heading(doc, "9.1 Dashboard merah atau semua angka nol", 2)
    code(doc, [
        "git rev-parse --short HEAD",
        "php artisan migrate:status",
        "tail -n 300 storage/logs/laravel.log | grep -E \"ERROR|SQLSTATE|Unknown column|Base table\"",
        "php artisan optimize:clear",
        "php artisan optimize",
    ])
    para(doc, "Jika log menyebut ONLY_FULL_GROUP_BY, perbaiki query menggunakan subquery/kolom deterministik. Jangan menonaktifkan SQL mode produksi.")
    heading(doc, "9.2 Peta kosong", 2)
    for text in [
        "Network: pastikan leaflet.js/css lokal dan tile OSM berstatus 200.",
        "Console: periksa L is not defined, CSP, mixed content, atau Rocket Loader.",
        "Data: periksa kota asal/kota sekolah dan daftar wilayah belum dikenali.",
        "Layout: pastikan container memiliki tinggi dan panggil invalidateSize setelah render Livewire.",
    ]:
        bullet(doc, text)
    heading(doc, "9.3 API Sekolah tidak menampilkan hasil", 2)
    code(doc, [
        "php artisan config:show services.school_api",
        "php artisan optimize:clear",
        "# Periksa SCHOOL_API_ENABLED, URL, timeout, DNS, TLS, dan response API",
        "# Jangan tampilkan token/credential di log atau screenshot.",
    ])
    para(doc, "Pastikan pencarian API hanya mengambil kandidat. Data sekolah baru disimpan ke database lokal setelah pengguna memilihnya; fallback sekolah lokal/manual harus tetap tersedia.")
    heading(doc, "9.4 WAHA/GOWS gagal", 2)
    for text in [
        "Periksa container/session/health dan pastikan Laravel memakai URL yang dapat dijangkau dari runtime.",
        "Shared hosting tidak dapat memakai localhost:3000 milik komputer lokal; gunakan tunnel bernama dengan domain stabil.",
        "Validasi API key dan webhook HMAC tanpa mencetak secret.",
        "Periksa queue worker dan failed_jobs jika webhook masuk tetapi Inbox tidak berubah.",
    ]:
        bullet(doc, text)
    heading(doc, "9.5 Import, hapus, dan duplikasi", 2)
    table(doc, ["Gejala", "Diagnosis", "Respons"], [
        ["validation.regex / required", "Mapping header, normalisasi, atau localization tidak sesuai.", "Tampilkan pesan Indonesia per baris dan sediakan contoh CSV."],
        ["Nomor WA sudah digunakan", "Nomor berbeda format menjadi nomor normalisasi yang sama.", "Cari lead lama; jangan membuat duplikat."],
        ["Delete gagal FK 1451", "Lead memiliki pendaftaran/follow-up/pesan terkait.", "Blokir dengan pesan bisnis atau gunakan arsip; jangan paksa delete."],
    ], [2600, 3200, 3560])

    page_break(doc)
    heading(doc, "10. Temuan Debugging Versi 2.1", 1)
    table(doc, ["ID", "Insiden", "Akar masalah", "Perbaikan / verifikasi"], [
        ["DBG-21-01", "Leaflet membentuk URL unpkg + tile OSM yang menyatu.", "Dependency CDN/Rocket Loader membuat urutan/URL asset tidak stabil.", "Leaflet disajikan lokal, query version, data-cfasync=false; peta tampil di produksi."],
        ["DBG-21-02", "Dashboard menampilkan banner gagal dan data nol.", "MySQL hosting menolak ekspresi kota pada GROUP BY (1055).", "Agregasi kota dipindah ke derived table; 10 Dashboard test dan MySQL role test lulus."],
        ["DBG-21-03", "Sebagian wilayah tidak terpetakan.", "Resolver belum mengenali provinsi/kabupaten dan variasi ejaan.", "Tambah Banten, Lembata/Lewoleba, Sumedang, Sumatera Utara serta alias."],
        ["DBG-21-04", "API sekolah tidak muncul di form Lead.", "Konfigurasi/jalur pencarian dan persistensi belum lengkap.", "SchoolApiClient + sinkronisasi lokal; data disimpan setelah dipilih."],
        ["DBG-21-05", "Sumber lead berbentuk teks bebas.", "Belum ada master/FK konsisten.", "Normalisasi ke sumber_lead_id dan fallback dashboard untuk transisi schema."],
    ], [1100, 2500, 2860, 2900], compact=True)
    heading(doc, "10.1 Pelajaran teknis", 2)
    for text in [
        "Test SQLite tidak cukup untuk membuktikan kompatibilitas query MySQL strict mode.",
        "Aset kritis yang kecil dapat disajikan lokal untuk mengurangi ketergantungan CDN.",
        "Fallback transisi schema membantu deployment bertahap, tetapi migrasi tetap wajib dijalankan.",
        "Coverage peta adalah indikator kualitas data lokasi, bukan sekadar masalah visual.",
        "Test fixture harus berkembang bersama aturan integritas seperti nomor WhatsApp unik.",
    ]:
        bullet(doc, text)

    heading(doc, "11. Regression dan Release Gate", 1)
    table(doc, ["Gate", "Wajib", "Bukti"], [
        ["Source", "Ya", "git diff --check; review file berubah; tidak ada .env/dump/secret."],
        ["Automated", "Ya", "Targeted dan full PHPUnit 100% lulus."],
        ["Build", "Ya", "composer validate dan npm run build exit 0."],
        ["Database", "Ya", "migrate:status sesuai; migrasi backup/rollback dipahami; query MySQL lulus."],
        ["Services", "Ya", "Container/service sehat; queue/scheduler aktif."],
        ["Smoke", "Ya", "Health, login, dashboard, CRUD penting, asset, API, dan WAHA diperiksa."],
        ["UAT", "Ya untuk rilis", "Perwakilan user menyetujui skenario bisnis."],
        ["Rollback", "Ya", "Commit/tag, backup, dan langkah rollback tersedia."],
    ], [2000, 1200, 6160], compact=True)
    callout(doc, "GATE SAAT INI", "BELUM LULUS PENUH karena dua regresi suite masih terbuka. Deployment darurat hanya boleh melalui persetujuan risiko tertulis, targeted verification, backup, dan rencana rollback.", PALE_RED, RED)
    heading(doc, "11.1 Setelah deployment hosting", 2)
    code(doc, [
        "cd /home/USER/domains/crm.sivmi.id/public_html",
        "git pull origin main",
        "composer install --no-dev --optimize-autoloader",
        "php artisan migrate --force",
        "php artisan optimize:clear",
        "php artisan optimize",
        "git rev-parse --short HEAD",
    ])
    para(doc, "Lakukan hard refresh dan purge cache Cloudflare bila asset lama masih terlihat. Hindari cache agresif untuk /admin/* dan endpoint Livewire.")

    page_break(doc)
    heading(doc, "12. Pelaporan Bug dan UAT", 1)
    table(doc, ["Field", "Isi"], [
        ["ID / judul", "Kode unik dan ringkasan gejala."],
        ["Environment", "Lokal, Docker, test, staging, atau production; commit dan waktu."],
        ["Aktor", "Role dan cabang; anonimkan identitas pribadi."],
        ["Precondition", "Data awal, status session, konfigurasi relevan tanpa secret."],
        ["Langkah", "Urutan bernomor yang dapat diulang."],
        ["Expected / actual", "Hasil yang seharusnya dan hasil aktual."],
        ["Bukti", "Screenshot, HTTP status, log tersanitasi, test, query, correlation ID."],
        ["Akar masalah", "Kode, test, data, config, dependency, infra, jaringan, atau integrasi."],
        ["Patch / rollback", "File/commit, risiko, compatibility, dan cara kembali."],
        ["Verifikasi", "Targeted, full suite, build, MySQL, smoke, UAT."],
        ["Status", "Open, fixed, monitoring, closed; pemilik dan approver."],
    ], [2500, 6860], compact=True)
    heading(doc, "12.1 Skenario UAT inti", 2)
    table(doc, ["ID", "Skenario", "Expected", "Hasil"], [
        ["UAT-01", "Login sesuai role/cabang", "Menu dan data sesuai kewenangan", ""],
        ["UAT-02", "Tambah lead hanya nama+WA", "Tersimpan; field lain opsional", ""],
        ["UAT-03", "Cari sekolah via API lalu pilih", "Sekolah tersimpan lokal dan terhubung ke lead", ""],
        ["UAT-04", "Filter dan bulk action leads/sekolah", "Hanya record terpilih diproses", ""],
        ["UAT-05", "Follow-up sampai pendaftaran", "Riwayat dan status konsisten", ""],
        ["UAT-06", "Dashboard dan peta", "KPI/filter/peta tampil tanpa banner error", ""],
        ["UAT-07", "WAHA QR, kirim, terima", "Session terhubung dan pesan masuk CRM", ""],
        ["UAT-08", "Laporan dan export", "Angka sesuai filter; file aman dibuka", ""],
    ], [1000, 3600, 3560, 1200], compact=True)

    doc.add_page_break()
    heading(doc, "13. Security dan Keamanan Sistem", 1)
    callout(
        doc,
        "TUJUAN BAB KEAMANAN",
        "Menetapkan baseline keamanan operasional untuk melindungi akun, data calon siswa, transaksi, integrasi API, dan layanan WhatsApp. Bab ini bukan sertifikasi keamanan atau pengganti penetration test independen.",
        PALE_BLUE,
        BLUE_DARK,
    )
    para(doc, "Pendekatan keamanan menggunakan prinsip defense in depth dan least privilege. Kontrol aplikasi, database, jaringan, deployment, monitoring, serta prosedur manusia harus saling mendukung; satu kontrol saja tidak cukup.")

    heading(doc, "13.1 Aset, ancaman, dan sasaran perlindungan", 2)
    table(doc, ["Aset", "Ancaman utama", "Sasaran kontrol"], [
        ["Akun pengguna", "Credential dicuri, akun nonaktif tetap masuk, eskalasi role", "Autentikasi kuat, status aktif, pembatasan role, pencabutan akses cepat."],
        ["Data lead/siswa", "Akses lintas cabang, ekspor tidak sah, data pribadi bocor", "Scope cabang, otorisasi per aksi, minimisasi data, audit dan retensi."],
        ["Pembayaran", "Perubahan nominal/status, laporan tidak sah, bukti terekspos", "Validasi bisnis, audit trail, pembatasan hak dan rekonsiliasi."],
        ["API publik", "API key bocor, brute force, replay, mass enumeration", "Key hash, ability, rate limit, idempotensi, request ID dan rotasi."],
        ["WAHA/webhook", "Webhook palsu, session diambil alih, pesan atau secret bocor", "HMAC SHA-512, API key, TLS/tunnel stabil, queue dan audit."],
        ["Server/backup", "Malware, salah konfigurasi, backup hilang atau dapat dibaca", "Patch, least privilege, enkripsi, backup teruji dan monitoring."],
    ], [2000, 3360, 4000], compact=True)

    heading(doc, "13.2 Kontrol keamanan yang sudah terlihat pada proyek", 2)
    table(doc, ["Kontrol", "Implementasi proyek", "Bukti verifikasi minimum"], [
        ["Akses panel", "User hanya dapat masuk bila aktif dan memiliki UserRole yang valid.", "Uji akun aktif/nonaktif dan setiap role."],
        ["Password", "Atribut password menggunakan cast hashed; password dan remember_token disembunyikan dari serialisasi.", "Pastikan hash tersimpan dan password tidak muncul pada respons/log."],
        ["Scope API", "API client memiliki ability leads:read/leads:write; write juga memeriksa hak operasional user.", "Uji 401, 403, role audit, dan akses lintas cabang."],
        ["API key", "Token minimal 32 karakter; database menyimpan SHA-256 key hash, bukan token mentah.", "Uji key salah/nonaktif dan rotasi tanpa downtime."],
        ["Rate limit", "Endpoint API v1 memakai throttle berbasis konfigurasi.", "Uji batas per menit dan respons 429."],
        ["Idempotensi", "POST lead wajib Idempotency-Key; payload berbeda dengan key sama ditolak 409.", "Uji replay identik, konflik payload, dan TTL."],
        ["Webhook WAHA", "Signature HMAC SHA-512 diverifikasi memakai hash_equals; secret kosong menghasilkan 503.", "Uji signature valid, hilang, salah, dan payload berubah."],
        ["Audit", "Aktivitas penting disimpan pada activity_logs dan tersedia pada menu Audit Aktivitas.", "Uji actor, event, waktu, objek, dan scope cabang."],
    ], [1900, 4460, 3000], compact=True)

    heading(doc, "13.3 Autentikasi, otorisasi, dan pengelolaan akun", 2)
    for text in [
        "Gunakan akun individual; jangan berbagi akun admin, API key, QR session, atau credential hosting.",
        "Terapkan least privilege: Superadmin hanya untuk pengelolaan global; Admin/Staff dibatasi pada cabang dan tugas operasionalnya; role audit bersifat baca saja.",
        "Nonaktifkan akun segera saat pegawai pindah fungsi atau keluar, lalu cabut session, token API, dan akses WAHA terkait.",
        "Gunakan password panjang dan unik, password manager, serta MFA pada hosting, Git provider, Cloudflare, email administrator, dan layanan backup. Jika panel aplikasi belum mendukung MFA, catat sebagai peningkatan prioritas tinggi.",
        "Uji direct URL dan request yang dimanipulasi. Menyembunyikan menu bukan pengganti policy, query scope, atau authorization pada server.",
        "Tinjau user, role, cabang, API client, dan akun WhatsApp minimal setiap bulan serta setelah perubahan organisasi.",
    ]:
        bullet(doc, text)

    heading(doc, "13.4 Perlindungan data pribadi dan database", 2)
    for text in [
        "Kumpulkan hanya data yang diperlukan untuk proses CRM. Nama dan nomor WhatsApp adalah data wajib minimum; data tambahan harus memiliki tujuan bisnis yang jelas.",
        "Gunakan HTTPS pada seluruh akses produksi. Cookie produksi harus Secure, HttpOnly, dan SameSite sesuai alur aplikasi; jangan mencampur HTTP dan HTTPS.",
        "Jangan menaruh dump database, .env, dokumen siswa, screenshot percakapan, atau bukti pembayaran di Git maupun direktori web publik.",
        "Backup database dan file penting secara terjadwal, terenkripsi, dipisahkan dari server utama, dibatasi aksesnya, serta diuji proses restore-nya.",
        "Ekspor CSV harus mengikuti scope pengguna dan mencegah formula injection. Catat siapa mengekspor, kapan, filter yang digunakan, dan tujuan penggunaan.",
        "Tetapkan retensi dan prosedur penghapusan/anonimisasi berdasarkan kebutuhan bisnis dan kewajiban hukum organisasi; relasi penting tidak boleh dihapus paksa dengan mematikan foreign key.",
    ]:
        bullet(doc, text)
    callout(doc, "DATA UJI", "Pengujian dan dokumentasi harus memakai factory/demo atau data yang telah dianonimkan. Jangan menyalin database produksi ke laptop atau layanan pihak ketiga tanpa persetujuan dan perlindungan yang setara.", PALE_RED, RED)

    heading(doc, "13.5 Keamanan API, API Sekolah, dan WAHA", 2)
    table(doc, ["Integrasi", "Kewajiban keamanan", "Kondisi gagal yang aman"], [
        ["API publik CRM", "TLS, key unik per client, ability minimum, rate limit, idempotensi, validasi, request ID, rotasi dan pencabutan.", "401/403/409/422 tanpa membocorkan detail internal atau data lain."],
        ["API Sekolah", "Batasi timeout, validasi struktur respons, encode query, jangan log credential, dan perlakukan respons sebagai data tidak tepercaya.", "Pencarian eksternal gagal tetapi pilihan lokal/manual tetap tersedia; kandidat belum disimpan sebelum dipilih."],
        ["WAHA API", "URL hanya dapat dijangkau pihak yang perlu, API key kuat, dashboard tidak terbuka publik, session dipisah per pengguna.", "Health gagal ditampilkan sebagai status layanan; jangan mencetak key atau QR session ke log."],
        ["Webhook WAHA", "Verifikasi HMAC atas raw body sebelum proses; TLS; deduplikasi event; pembatasan ukuran payload dan monitoring anomali.", "Signature hilang/salah ditolak sebelum menulis event atau menjalankan queue."],
        ["Cloudflare Tunnel", "Gunakan named tunnel, token dilindungi, DNS stabil, access policy bila dashboard diekspos, dan origin tidak dibuka langsung.", "Tunnel mati tidak menonaktifkan autentikasi atau memindahkan layanan ke endpoint tanpa proteksi."],
    ], [1600, 4660, 3100], compact=True)

    heading(doc, "13.6 Hardening Docker dan shared hosting", 2)
    for text in [
        "Jalankan image resmi/minimal, pin versi mayor, scan dependency dan image, serta rebuild berkala untuk memperoleh patch keamanan.",
        "Jangan menjalankan container sebagai privileged, jangan mount Docker socket, dan hindari publish port database/Redis/WAHA ke internet. Batasi port host hanya sesuai kebutuhan.",
        "Pisahkan secret dari image dan repository. Gunakan file environment dengan permission ketat atau secret manager; rotasi bila pernah tampil di terminal, log, screenshot, atau Git.",
        "Pada shared hosting, pastikan document root mengarah ke public, directory listing nonaktif, file .env tidak dapat diunduh, HTTPS dipaksa, dan versi PHP masih mendapat patch.",
        "Aktifkan backup sebelum migrasi, jalankan migrate --force hanya pada database yang telah diverifikasi, lalu periksa log, queue, scheduler, health, dan halaman kritis.",
        "Batasi SSH dengan key, nonaktifkan login password bila memungkinkan, aktifkan MFA panel hosting, dan tinjau lonjakan percobaan login gagal sebagai potensi serangan.",
    ]:
        bullet(doc, text)
    code(doc, [
        "docker compose --env-file docker\\.env -f docker\\compose.yml ps",
        "docker compose --env-file docker\\.env -f docker\\compose.yml config --services",
        "docker compose --env-file docker\\.env -f docker\\compose.yml logs --tail=200 app web queue waha",
        "composer audit",
        "npm audit --omit=dev",
        "php artisan about",
    ])

    heading(doc, "13.7 Pengujian keamanan minimum", 2)
    table(doc, ["Area", "Skenario", "Kriteria lulus"], [
        ["Authentication", "Login salah berulang, user nonaktif, session lama, reset password", "Tidak ada bypass; throttling/session bekerja; pesan tidak mengungkap akun."],
        ["Authorization", "Ubah ID, cabang, PIC, route langsung, bulk action, export", "Data/aksi di luar scope ditolak server dengan 403 atau hasil kosong yang tepat."],
        ["Input", "XSS, SQL meta-character, file/CSV berbahaya, nilai ekstrem", "Validasi dan encoding benar; tidak ada script/query tak terduga."],
        ["API", "Key salah/nonaktif, ability kurang, rate limit, replay idempotency", "401/403/429/409 sesuai kontrak dan respons tidak membocorkan secret."],
        ["Webhook", "Signature hilang/salah, body berubah, replay event, payload besar", "Ditolak sebelum persist; event valid tidak diproses ganda."],
        ["Configuration", "APP_DEBUG, HTTPS, cookie, CORS, error page, file sensitif", "Produksi tidak membocorkan stack trace, secret, atau file internal."],
        ["Dependency", "composer audit, npm audit, image scan", "Tidak ada kerentanan kritis/tinggi tanpa mitigasi dan persetujuan risiko."],
        ["Recovery", "Restore backup dan rotasi key", "RTO/RPO organisasi tercapai dan aplikasi dapat diverifikasi setelah pemulihan."],
    ], [1700, 4200, 3460], compact=True)

    heading(doc, "13.8 Respons insiden keamanan", 2)
    table(doc, ["Tahap", "Tindakan wajib"], [
        ["1. Deteksi", "Catat waktu, sistem, akun, indikator, sumber laporan, serta dampak tanpa mengubah bukti asli."],
        ["2. Containment", "Nonaktifkan akun/client terdampak, cabut session/token, batasi endpoint, isolasi host bila perlu, dan hentikan kebocoran tanpa menghapus log."],
        ["3. Preservasi bukti", "Salin log, request ID, audit trail, hash file, konfigurasi relevan, dan snapshot dalam media berakses terbatas."],
        ["4. Eradikasi", "Perbaiki akar masalah, patch dependency/configuration, rotasi seluruh secret yang mungkin terekspos, dan periksa persistence/backdoor."],
        ["5. Pemulihan", "Restore dari sumber tepercaya, jalankan regression/security test, pantau anomali, dan buka layanan bertahap dengan persetujuan pemilik sistem."],
        ["6. Tindak lanjut", "Eskalasi ke penanggung jawab, nilai kewajiban pemberitahuan, dokumentasikan keputusan, dan masukkan perbaikan ke backlog terukur."],
    ], [2200, 7160], compact=True)
    callout(doc, "JANGAN DILAKUKAN", "Jangan menghapus log, menjalankan pembersihan agresif, mengunggah bukti ke layanan publik, atau mengumumkan detail insiden sebelum otorisasi. Jangan menjanjikan bahwa insiden selesai sebelum rotasi secret, verifikasi, dan monitoring pascapemulihan dilakukan.", PALE_RED, RED)

    heading(doc, "13.9 Checklist security gate sebelum rilis", 2)
    for text in [
        "Review perubahan source, dependency, migrasi, route, policy, middleware, dan konfigurasi environment selesai.",
        "Full regression, targeted security test, build, composer audit, npm audit, serta integration test MySQL telah diperiksa.",
        "APP_ENV=production, APP_DEBUG=false, HTTPS aktif, cookie aman, CORS minimum, dan error page tidak membocorkan detail.",
        "Tidak ada .env, token, API key, dump, data siswa, file backup, atau credential di Git dan web root.",
        "Role/cabang, API abilities, rate limit, idempotensi, HMAC webhook, bulk action, export, dan audit log telah diuji.",
        "Backup terbaru tersedia, restore pernah diuji, rollback jelas, serta owner dan approver rilis tercatat.",
        "Temuan kritis/tinggi ditutup; pengecualian risiko memiliki alasan, mitigasi, pemilik, tenggat, dan persetujuan tertulis.",
    ]:
        bullet(doc, text)

    heading(doc, "13.10 Acuan dan batasan", 2)
    para(doc, "Baseline ini menggunakan kerangka berikut sebagai referensi praktik. Penyebutan acuan tidak berarti aplikasi telah tersertifikasi atau memenuhi seluruh persyaratan standar.")
    for text in [
        "OWASP Application Security Verification Standard (ASVS) 5.0.0 - https://owasp.org/www-project-application-security-verification-standard/",
        "NIST Cybersecurity Framework (CSF) 2.0 - https://www.nist.gov/cyberframework",
        "Laravel 12 Documentation: Authentication, Authorization, CSRF, Validation, Encryption, dan Deployment - https://laravel.com/docs/12.x",
        "CIS Docker Benchmark overview pada Docker Documentation - https://docs.docker.com/dhi/explore/security-concepts/cis/",
    ]:
        bullet(doc, text)
    callout(doc, "REKOMENDASI LANJUTAN", "Lakukan threat modeling per alur, code review keamanan, vulnerability scan terjadwal, dan penetration test independen sebelum penggunaan berskala besar atau ketika sistem mulai memproses data/akses dengan tingkat risiko lebih tinggi.", PALE_AMBER, AMBER)

    heading(doc, "Lampiran A - Perintah Cepat", 1)
    heading(doc, "A.1 Lokal", 2)
    code(doc, [
        "composer validate --no-check-publish",
        "php artisan test",
        "php artisan test --filter=NamaTest",
        "npm run build",
        "php -l app\\Path\\File.php",
    ])
    heading(doc, "A.2 Docker", 2)
    code(doc, [
        "docker compose --env-file docker\\.env -f docker\\compose.yml ps",
        "docker compose --env-file docker\\.env -f docker\\compose.yml exec -T app php artisan migrate:status",
        "docker compose --env-file docker\\.env -f docker\\compose.yml exec -T app php artisan test",
        "docker compose --env-file docker\\.env -f docker\\compose.yml logs --tail=200 app queue waha",
    ])
    heading(doc, "A.3 Hosting", 2)
    code(doc, [
        "git rev-parse --short HEAD",
        "php artisan migrate:status",
        "php artisan optimize:clear",
        "tail -n 300 storage/logs/laravel.log",
        "tail -n 300 storage/logs/laravel.log | grep -E \"ERROR|SQLSTATE|Unknown column|Base table\"",
    ])

    # `code()` menambahkan paragraf spacer. Pada posisi paling bawah halaman,
    # spacer itu dapat terdorong ke halaman baru lalu page-break judul lampiran
    # menghasilkan satu halaman kosong tambahan.
    spacer = doc.paragraphs[-1]
    spacer._element.getparent().remove(spacer._element)
    appendix_heading = heading(doc, "Lampiran B - Formulir Hasil Pengujian", 1)
    appendix_heading.paragraph_format.page_break_before = True
    table(doc, ["Informasi", "Isian"], [
        ["Tanggal / tester", ""], ["Commit / environment", ""], ["Modul / skenario", ""],
        ["Data uji", ""], ["Expected", ""], ["Actual", ""], ["Bukti", ""],
        ["Status", "PASS / FAIL / BLOCKED"], ["Defect ID", ""], ["Catatan / approver", ""],
    ], [2800, 6560], compact=True)
    callout(doc, "PRIVASI", "Dokumen dan lampiran tidak boleh memuat password, APP_KEY, API key, webhook secret, token, cookie, isi percakapan, atau data siswa yang dapat diidentifikasi.", PALE_RED, RED)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
