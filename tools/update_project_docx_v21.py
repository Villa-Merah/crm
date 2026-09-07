from __future__ import annotations

import argparse
import json
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.table import Table
from docx.text.paragraph import Paragraph


BLUE = "0B2C6B"
WHITE = "FFFFFF"
GRAY = "666666"
APPENDIX_TITLE = "LAMPIRAN E: PEMBARUAN VERSI 2.1"


def find_paragraph(doc: Document, text: str) -> Paragraph:
    for paragraph in doc.paragraphs:
        if paragraph.text.strip() == text:
            return paragraph
    raise ValueError(f"Paragraph not found: {text}")


def new_paragraph_before(anchor: Paragraph, text: str = "", style: str | None = None) -> Paragraph:
    paragraph = anchor.insert_paragraph_before(text)
    if style:
        paragraph.style = style
    return paragraph


def set_run_font(run, *, size: float | None = None, bold: bool | None = None,
                 color: str | None = None, name: str = "Arial") -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 100, start: int = 120, bottom: int = 100, end: int = 120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_widths(table: Table, widths: list[float]) -> None:
    table.autofit = False
    width_dxa = [round(width * 1440) for width in widths]
    for row in table.rows:
        for index, value in enumerate(width_dxa):
            row.cells[index].width = Inches(widths[index])
            tc_pr = row.cells[index]._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(value))
            tc_w.set(qn("w:type"), "dxa")
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(width_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for value in width_dxa:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(value))
        grid.append(column)


def insert_table_before(doc: Document, anchor: Paragraph, headers: list[str], rows: list[list[str]],
                        widths: list[float]) -> Table:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table._tbl.getparent().remove(table._tbl)
    anchor._p.addprevious(table._tbl)
    set_table_widths(table, widths)
    set_repeat_table_header(table.rows[0])
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = header
        set_cell_shading(cell, BLUE)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(0)
        for run in paragraph.runs:
            set_run_font(run, size=8.5, bold=True, color=WHITE)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for column_index, value in enumerate(values):
            cell = cells[column_index]
            cell.text = value
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index % 2:
                set_cell_shading(cell, "F4F7FB")
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if column_index in (0, len(headers) - 1) else WD_ALIGN_PARAGRAPH.LEFT
            for run in paragraph.runs:
                set_run_font(run, size=8.2)
    return table


def copy_heading(source: Paragraph, target: Paragraph, text: str) -> None:
    if source._p.pPr is not None:
        if target._p.pPr is not None:
            target._p.remove(target._p.pPr)
        target._p.insert(0, deepcopy(source._p.pPr))
    target.clear()
    target.paragraph_format.page_break_before = True
    run = target.add_run(text)
    source_run = source.runs[0] if source.runs else None
    set_run_font(
        run,
        size=source_run.font.size.pt if source_run and source_run.font.size else 14,
        bold=True,
        color=str(source_run.font.color.rgb) if source_run and source_run.font.color and source_run.font.color.rgb else BLUE,
    )


def add_bullet(doc: Document, anchor: Paragraph, text: str) -> Paragraph:
    paragraph = new_paragraph_before(anchor, text, "List Bullet")
    paragraph.paragraph_format.keep_together = True
    return paragraph


def add_code(doc: Document, anchor: Paragraph, text: str) -> Paragraph:
    paragraph = new_paragraph_before(anchor, text, "No Spacing")
    paragraph.paragraph_format.left_indent = Inches(0.25)
    paragraph.paragraph_format.space_after = Pt(2)
    for run in paragraph.runs:
        set_run_font(run, size=8.5, name="Consolas")
    return paragraph


def insert_toc_entry(doc: Document) -> None:
    if any(p.text.strip().startswith(APPENDIX_TITLE) for p in doc.paragraphs):
        return
    glossary_entry = next(
        (p for p in doc.paragraphs if p.text.strip().startswith("GLOSARIUM") and "\t" in p.text),
        None,
    )
    if glossary_entry is None:
        raise ValueError("TOC glossary entry not found")
    paragraph = glossary_entry.insert_paragraph_before()
    if glossary_entry._p.pPr is not None:
        paragraph._p.insert(0, deepcopy(glossary_entry._p.pPr))
    run = paragraph.add_run(f"{APPENDIX_TITLE}\t{{{{PAGE:{APPENDIX_TITLE}}}}}")
    set_run_font(run, size=9.5, bold=True)


def add_appendix(doc: Document) -> None:
    if any(p.text.strip() == APPENDIX_TITLE for p in doc.paragraphs):
        raise RuntimeError("Dokumen input sudah memuat pembaruan versi 2.1")
    glossary = find_paragraph(doc, "Glosarium")
    appendix_source = find_paragraph(doc, "LAMPIRAN D: USER ACCEPTANCE TESTING (UAT)")
    title = new_paragraph_before(glossary)
    copy_heading(appendix_source, title, APPENDIX_TITLE)

    intro = new_paragraph_before(
        glossary,
        "Lampiran ini memperbarui baseline as-built per 1 September 2026. Cakupan utamanya adalah integrasi API Sekolah Indonesia, peta persebaran leads, serta pilih banyak dan aksi massal Leads. Dokumen versi sebelumnya tetap dipertahankan sebagai arsip historis.",
        "Normal",
    )
    intro.paragraph_format.space_after = Pt(8)

    new_paragraph_before(glossary, "E.1 Ringkasan Fitur Terbaru", "Heading 1")
    for text in [
        "Sekolah dapat dicari berdasarkan nama atau NPSN dari form Lead dan master Sekolah, lalu disimpan ke database lokal berdasarkan NPSN.",
        "Dashboard menampilkan peta interaktif tingkat kota menggunakan Leaflet 1.9.4 dan tile OpenStreetMap; ukuran bubble merepresentasikan jumlah leads dalam filter aktif.",
        "Lokasi peta mengutamakan kota asal lead dan memakai kota sekolah sebagai fallback. Wilayah yang belum dikenali ditampilkan sebagai data yang perlu diperbaiki.",
        "Daftar Leads menyediakan checkbox, pilih seluruh hasil filter, bulk edit, dan bulk delete terkontrol. Direksi read-only dan penghapusan hanya tersedia bagi Superadmin.",
        "Lead yang telah memiliki pendaftaran tidak dapat dihapus. Pembaruan massal tetap melalui model agar observer, histori, dan audit aktivitas berjalan.",
    ]:
        add_bullet(doc, glossary, text)

    new_paragraph_before(glossary, "E.2 Dampak Arsitektur dan Data", "Heading 1")
    for text in [
        "Arsitektur tetap modular monolith Laravel; API Sekolah dan layanan peta merupakan dependency eksternal, bukan microservice internal.",
        "Database CRM tetap menjadi sumber operasional. API Sekolah hanya menjadi sumber referensi dan hasil pilihan pengguna dipersistenkan ke master sekolah lokal.",
        "Peta tidak menambah tabel database. Agregasi memakai leads.kota_asal dengan fallback sekolah.kota, sedangkan koordinat kota disediakan resolver statis IndonesiaCityCoordinates.",
        "Peta menampilkan agregasi kota, bukan alamat atau koordinat individu, sehingga paparan data pribadi lebih rendah.",
    ]:
        add_bullet(doc, glossary, text)

    new_paragraph_before(glossary, "E.3 Panduan Operasional", "Heading 1")
    insert_table_before(
        doc,
        glossary,
        ["Area", "Langkah", "Kontrol penting"],
        [
            ["API Sekolah", "Cari minimal tiga karakter nama/NPSN, verifikasi wilayah, pilih hasil, lalu gunakan data sekolah.", "Jika API gagal, gunakan sekolah lokal atau input manual; jangan membuat NPSN fiktif."],
            ["Peta Dashboard", "Gunakan filter periode, cabang, dan PIC; klik bubble untuk melihat kota dan jumlah leads.", "Pantau cakupan terpetakan dan perbaiki kota yang belum dikenali."],
            ["Bulk edit Leads", "Terapkan filter, pilih baris atau seluruh hasil filter, pilih field dan nilai, lalu konfirmasi.", "Periksa jumlah terpilih dan scope cabang sebelum menjalankan aksi."],
            ["Bulk delete", "Superadmin memilih lead, meninjau peringatan, lalu mengonfirmasi penghapusan.", "Backup tersedia; lead dengan pendaftaran harus tetap ditolak."],
        ],
        [1.25, 3.0, 3.05],
    )

    new_paragraph_before(glossary, "E.4 Deployment dan Troubleshooting", "Heading 1")
    new_paragraph_before(glossary, "Shared hosting/server:", "Normal").runs[0].bold = True
    for command in [
        "cd ~/domains/crm.sivmi.id/public_html",
        "git pull origin main",
        "composer install --no-dev --optimize-autoloader",
        "php artisan migrate --force",
        "php artisan optimize:clear",
        "php artisan optimize",
    ]:
        add_code(doc, glossary, command)
    new_paragraph_before(glossary, "Pastikan SCHOOL_API_* tersedia pada .env dan koneksi keluar HTTPS menuju API Sekolah, CDN Leaflet, serta tile OpenStreetMap tidak diblokir. Bila peta kosong, lakukan hard refresh dan periksa Console/Network browser. Perbaikan tinggi kontainer peta terbaru tidak membutuhkan migration atau build aset.", "Normal")

    new_paragraph_before(glossary, "E.5 UAT Tambahan Versi 2.1", "Heading 1")
    insert_table_before(
        doc,
        glossary,
        ["ID", "Skenario", "Hasil yang diharapkan", "Status"],
        [
            ["MAP-01", "Buka Dashboard dengan koneksi internet aktif.", "Tile, zoom, dan bubble kota tampil.", "___"],
            ["MAP-02", "Ubah periode/cabang/PIC.", "Peta, top kota, dan cakupan mengikuti filter serta scope role.", "___"],
            ["MAP-03", "Lead tanpa kota asal memiliki sekolah berkota.", "Lokasi memakai kota sekolah sebagai fallback.", "___"],
            ["MAP-04", "Kota belum tersedia pada resolver.", "Kota masuk daftar belum dikenali tanpa merusak Dashboard.", "___"],
            ["BULK-01", "Pilih beberapa Lead dan ubah status.", "Semua data terpilih berubah dan audit tercatat.", "___"],
            ["BULK-02", "Pilih seluruh hasil filter.", "Hanya hasil dalam filter dan scope pengguna diproses.", "___"],
            ["BULK-03", "Direksi mencoba bulk edit.", "Aksi tidak tersedia atau ditolak.", "___"],
            ["BULK-04", "Superadmin menghapus Lead tanpa pendaftaran.", "Data terhapus setelah konfirmasi.", "___"],
            ["BULK-05", "Hapus Lead yang telah memiliki pendaftaran.", "Penghapusan ditolak dengan alasan yang jelas.", "___"],
            ["SCH-01", "Cari dan pilih sekolah dari API.", "Sekolah tersimpan lokal dan terhubung ke Lead.", "___"],
            ["SCH-02", "Pilih kembali NPSN yang sama.", "Record diperbarui tanpa duplikasi.", "___"],
            ["SCH-03", "API timeout/tidak aktif.", "Fallback lokal/manual tetap berfungsi.", "___"],
        ],
        [0.85, 2.2, 3.55, 0.7],
    )

    new_paragraph_before(glossary, "E.6 Referensi Baseline", "Heading 1")
    for text in [
        "Commit API Sekolah Indonesia: 4d60ead.",
        "Commit peta dashboard dan aksi massal Leads: 31c387e.",
        "Pengujian fitur terkait: 20 test dan 155 assertion lulus.",
        "ERD terbaru: ERD_CRM_VILLA_MERAH_TERBARU_2026-09-01.drawio. Fitur peta tidak mengubah ERD karena tidak menambah tabel.",
    ]:
        add_bullet(doc, glossary, text)


def patch_page_map(input_path: Path, output_path: Path, page_map: dict[str, int]) -> None:
    doc = Document(input_path)
    token = f"{{{{PAGE:{APPENDIX_TITLE}}}}}"
    for paragraph in doc.paragraphs:
        if token not in paragraph.text:
            continue
        for run in paragraph.runs:
            if token in run.text:
                run.text = run.text.replace(token, str(page_map[APPENDIX_TITLE]))
    doc.save(output_path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("input", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--page-map", type=Path)
    args = parser.parse_args()
    if args.page_map:
        page_map = json.loads(args.page_map.read_text(encoding="utf-8"))
        patch_page_map(args.input, args.output, page_map)
        return
    doc = Document(args.input)
    insert_toc_entry(doc)
    add_appendix(doc)
    doc.core_properties.title = "Dokumentasi Project CRM_VILLA MERAH Versi 2.1"
    doc.core_properties.subject = "Dokumentasi terbaru API Sekolah, peta Dashboard, aksi massal Leads, SOP, deployment, dan UAT"
    doc.core_properties.comments = "Revisi 1 September 2026: pembaruan versi 2.1 dan Lampiran E."
    doc.settings.update_fields_on_open = True
    args.output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(args.output)


if __name__ == "__main__":
    main()
