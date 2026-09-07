from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


SOURCE = Path("docs/Dokumen Kerja CRM_VILLAMERAH 01092026.docx")
OUTPUT = Path("docs/Dokumen Kerja CRM_VILLAMERAH 01092026-REVISI-DAFTAR-ISI.docx")

BLUE = "123765"
GRAY_FILL = "EDEDED"
TEXT = "111827"
MUTED = "5B6573"


PAGES = [
    (
        [
            ("BAB I: PENDAHULUAN", [
                "Gambaran Umum Sistem",
                "Latar Belakang",
                "Identifikasi Masalah",
                "Rumusan Masalah",
                "Manfaat Sistem",
                "Tujuan Pembuatan",
                "Ruang Lingkup Sistem",
            ]),
            ("BAB II: LANDASAN TEORI", [
                "Sistem Informasi",
                "Aplikasi Berbasis Web",
                "PHP (Hypertext Preprocessor)",
                "MySQL",
                "XAMPP",
                "HTML, CSS, JavaScript, dan Bootstrap",
                "Enterprise Architecture (EA)",
                "Use Case Diagram",
                "Data Flow Diagram (DFD)",
                "Entity Relationship Diagram (ERD)",
                "Flowchart",
                "Activity Diagram",
                "Laravel",
                "User Acceptance Testing (UAT)",
                "Webhook dan API",
                "Artificial Intelligence (AI) dalam Pengembangan Perangkat Lunak",
            ]),
        ],
        [
            ("BAB III: ANALISA KEBUTUHAN", [
                "Kebutuhan Fungsional",
                "Aktor, Role, dan Scope",
                "Dashboard, Filter, KPI, dan Laporan",
                "Lead, Sekolah, dan Import/Export",
                "Follow-up, Pendaftaran, Pembayaran, dan Program Bimbel",
                "Target Kinerja dan Master Data",
                "Integrasi WhatsApp, Chat, dan Notifikasi",
                "Pengaturan dan Audit Aktivitas",
                "Aturan Bisnis Fungsional",
                "Kebutuhan Nonfungsional",
            ]),
            ("BAB IV: ARSITEKTUR BISNIS", [
                "Profil dan Visi Misi Perusahaan",
                "Struktur Organisasi dan Stakeholder",
                "Proses dan Alur Kerja Bisnis",
                "Pemodelan Aktor dan Key Performance Indicator (KPI)",
                "Perancangan Enterprise Architecture (Overview)",
                "Matriks Sasaran, Target Terukur (Objective), dan Indikator Kinerja (KPI)",
                "Aturan Bisnis",
            ]),
        ],
    ),
    (
        [
            ("BAB V: ARSITEKTUR DATA", [
                "RBAC (Role Based Access Control)",
                "Matriks RBAC",
                "Alur Proses Sistem",
                "Flowchart",
                "Use Case",
                "Activity Diagram",
                "Data Flow Diagram (DFD)",
                "Entity Relationship Diagram (ERD)",
                "Relasi dan Kardinalitas",
                "Siklus Data Lead",
                "Database dan Kamus Data",
                "Keamanan Hak Akses Data",
            ]),
            ("BAB VI: ARSITEKTUR APLIKASI", [
                "Pengertian",
                "Tujuan",
                "Ruang Lingkup",
                "Konteks Aplikasi",
                "Prinsip Arsitektur Aplikasi",
                "Gaya Arsitektur",
                "Layer dan Tanggung Jawab",
                "Struktur Komponen",
                "Peta Modul Aplikasi",
                "Alur Request Web",
                "Alur Follow-up dan Transaksi",
                "Arsitektur Data pada Aplikasi",
                "Arsitektur API",
                "Scheduler, Queue, dan Reminder",
                "Penyimpanan File",
                "Presentasi dan UI",
                "Error Handling dan Logging",
                "Roadmap Teknis",
            ]),
        ],
        [
            ("BAB VII: ARSITEKTUR TEKNOLOGI", [
                "Pengertian",
                "Baseline Teknologi",
                "Parameter PHP dan Web",
                "Lingkungan (Environment)",
                "Konfigurasi dan Secret",
                "Integrasi API Sekolah Indonesia",
                "Session, Cookie, dan Cache",
                "Queue dan Scheduler",
                "Penyimpanan File",
                "Backup, Restore, RPO, dan RTO",
                "Logging, Monitoring, dan Health Check",
                "Status As-Built dan Target Teknologi",
                "Dasar Verifikasi Project",
            ]),
            ("BAB VIII: MANAJEMEN DEBUGGING", [
                "Pengertian",
                "Tujuan",
                "Prinsip Debugging",
                "Siklus Debugging dan QA",
                "Deteksi dan Identifikasi",
                "Reproduksi Masalah",
                "Analisis Akar Masalah",
                "Perbaikan Kode dan Refactoring",
                "Pengujian Ulang",
                "Perintah Diagnostik",
                "Logging Aman",
                "Triage Insiden dan Pemulihan",
                "Temuan Debugging yang Telah Ditangani",
                "Baseline Verifikasi Tercatat",
                "Format Catatan Debugging",
                "Kriteria Penutupan Bug",
                "Rekomendasi Debugging dan QA",
            ]),
        ],
    ),
    (
        [
            ("BAB IX: DEPLOYMENT, MIGRASI, DAN PEMELIHARAAN SISTEM", [
                "Prinsip Pengendalian",
                "Peran dan Tanggung Jawab",
                "Matriks RACI Ringkas",
                "Lingkungan dan Baseline Teknologi",
                "Maintenance Mode",
                "Menjalankan Migrasi",
                "Cache, Storage, Queue, dan Scheduler",
                "Smoke Test Pasca-Deployment",
                "Observasi Setelah Release",
                "Migrasi Basis Data dan Data",
                "Rollback dan Pemulihan Release",
                "Backup dan Recovery Strategy",
                "Pemeliharaan Sistem",
                "Monitoring, Performa, dan Kapasitas",
                "Pemeliharaan Keamanan",
                "Penanganan Insiden dan Pemeliharaan Darurat",
                "Rekomendasi Pengembangan Operasional",
                "Kriteria Dokumen Dinyatakan Diterapkan",
            ]),
            ("BAB X: PANDUAN PENGGUNAAN", [
                "Login",
                "Hak Akses Role",
                "Navigasi dan Dashboard",
                "Data Leads",
                "Memilih Sekolah dari API Sekolah Indonesia",
                "Import Leads CSV",
                "Export Leads",
                "Follow-up",
                "Inbox WhatsApp",
                "Pendaftaran (Data Siswa)",
                "Pembayaran",
                "Program Bimbel",
                "Target Kinerja",
                "Laporan",
                "Master Data Cabang, Sekolah, dan Pengguna",
                "Koneksi WhatsApp",
                "Pengaturan",
                "Notifikasi, Audit Aktivitas, Profil, dan Logout",
            ]),
        ],
        [
            ("BAB XI: PANDUAN ADMINISTRATOR", [
                "Peran Administrator dan Ruang Lingkup",
                "Prasyarat dan Baseline Sistem",
                "Instalasi Development Full Docker",
                "Konfigurasi Environment dan Secret",
                "Database, Migration, Seeder, dan Storage",
                "Queue dan Scheduler",
                "WAHA/GOWS dan Webhook",
                "Client API dan Endpoint Integrasi",
                "Konfigurasi API Sekolah Indonesia",
                "Administrasi User, Role, dan Cabang",
                "Konfigurasi Aplikasi dan Program",
                "Backup",
                "Restore dan Rekonsiliasi",
                "Permission Server, File, dan Volume",
                "Maintenance Mode",
                "Deployment dan Update",
                "Monitoring, Logging, dan Health Check",
                "Troubleshooting, Pemulihan, dan Serah Terima",
            ]),
            ("BAB XII: STANDARD OPERATING PROCEDURE OPERASIONAL CRM", [
                "Standar Data dan Keamanan Operasional",
                "SOP Pembukaan Operasional Harian",
                "SOP Pengelolaan Akun dan Akses",
                "SOP Master Data",
                "SOP Input, Impor, dan Kualitas Lead",
                "SOP Penugasan, Follow-up, dan Status Lead",
                "SOP WhatsApp/WAHA dan Inbox",
                "SOP Pendaftaran dan Pembayaran",
                "SOP Target, Dashboard, dan Laporan",
                "SOP Koreksi, Penghapusan, Impor/Ekspor, dan Audit",
                "SOP Backup, Restore, dan Perubahan Produksi",
                "SOP Penanganan Insiden dan Eskalasi",
                "SOP Penutupan Harian, Mingguan, dan Bulanan",
                "KPI, SLA, dan Bukti Operasional",
                "Formulir dan Checklist Siap Pakai",
                "Referensi Internal dan Penutup",
            ]),
            ("LAMPIRAN", [
                "Lampiran A: API dan REST API",
                "Lampiran B: Laporan Testing dan Debugging",
                "Lampiran C: Matriks ERD, Relasi, dan Kardinalitas",
                "Lampiran D: User Acceptance Testing (UAT)",
                "Glosarium",
            ]),
        ],
    ),
]


def set_run_font(run, size: float, bold: bool = False, color: str = TEXT) -> None:
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = __import__("docx").shared.RGBColor.from_string(color)
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    r_fonts.set(qn("w:ascii"), "Arial")
    r_fonts.set(qn("w:hAnsi"), "Arial")


def shade_paragraph(paragraph, fill: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shading = p_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        p_pr.append(shading)
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=0, start=70, bottom=0, end=70) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def remove_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "nil")


def format_base_paragraph(paragraph, after=0, before=0, line=1.0) -> None:
    paragraph.style = "Normal"
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line
    paragraph.paragraph_format.keep_together = True


def add_toc_column(cell, chapters) -> None:
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    set_cell_margins(cell, start=90, end=130)
    first = cell.paragraphs[0]
    first._element.getparent().remove(first._element)
    for chapter_index, (title, items) in enumerate(chapters):
        heading = cell.add_paragraph()
        format_base_paragraph(heading, after=1, before=3 if chapter_index else 0, line=1.0)
        heading.paragraph_format.keep_with_next = True
        shade_paragraph(heading, GRAY_FILL)
        run = heading.add_run(title)
        set_run_font(run, 8.7, bold=True, color="111111")
        for number, item in enumerate(items, 1):
            paragraph = cell.add_paragraph()
            format_base_paragraph(paragraph, after=0, before=0, line=1.0)
            paragraph.paragraph_format.left_indent = Pt(24)
            paragraph.paragraph_format.first_line_indent = Pt(-16)
            number_run = paragraph.add_run(f"{number}.")
            set_run_font(number_run, 8.2, color=MUTED)
            text_run = paragraph.add_run(f"    {item}")
            set_run_font(text_run, 8.2, color="111111")


def set_table_geometry(table, total_width: int) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(total_width))
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    widths = (total_width // 2, total_width - total_width // 2)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for cell, width in zip(table.rows[0].cells, widths):
        tc_pr = cell._tc.get_or_add_tcPr()
        tc_w = tc_pr.find(qn("w:tcW"))
        if tc_w is None:
            tc_w = OxmlElement("w:tcW")
            tc_pr.append(tc_w)
        tc_w.set(qn("w:type"), "dxa")
        tc_w.set(qn("w:w"), str(width))


def insert_before(anchor, element) -> None:
    anchor._p.addprevious(element)


def add_page_break_before(document: Document, anchor) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.add_run().add_break(WD_BREAK.PAGE)
    insert_before(anchor, paragraph._p)


def build() -> None:
    document = Document(SOURCE)
    anchor = next(
        p for p in document.paragraphs
        if p.text.strip() == "BAB I: PENDAHULUAN"
    )

    body = document._element.body
    for element in list(body):
        if element is anchor._p:
            break
        body.remove(element)

    usable_width = int(
        (document.sections[0].page_width
         - document.sections[0].left_margin
         - document.sections[0].right_margin) / 635
    )

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(2)
    title.paragraph_format.keep_with_next = True
    run = title.add_run("DAFTAR ISI")
    set_run_font(run, 15, bold=True, color=BLUE)
    insert_before(anchor, title._p)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(8)
    subtitle.paragraph_format.keep_with_next = True
    run = subtitle.add_run("DOKUMEN KERJA CRM_VILLA MERAH")
    set_run_font(run, 8.5, bold=True, color=MUTED)
    insert_before(anchor, subtitle._p)

    for page_number, (left, right) in enumerate(PAGES):
        if page_number:
            continuation = document.add_paragraph()
            continuation.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            continuation.paragraph_format.space_after = Pt(5)
            continuation.paragraph_format.keep_with_next = True
            run = continuation.add_run("DAFTAR ISI - LANJUTAN")
            set_run_font(run, 8.2, bold=True, color=MUTED)
            insert_before(anchor, continuation._p)

        table = document.add_table(rows=1, cols=2)
        remove_table_borders(table)
        set_table_geometry(table, usable_width)
        add_toc_column(table.rows[0].cells[0], left)
        add_toc_column(table.rows[0].cells[1], right)
        insert_before(anchor, table._tbl)

        if page_number < len(PAGES) - 1:
            add_page_break_before(document, anchor)

    add_page_break_before(document, anchor)

    settings = document.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")

    core = document.core_properties
    core.title = "Dokumen Kerja CRM_VILLA MERAH - Revisi Daftar Isi"
    core.subject = "Daftar isi dua kolom sesuai referensi"
    document.save(OUTPUT)
    print(OUTPUT.resolve())


if __name__ == "__main__":
    build()
