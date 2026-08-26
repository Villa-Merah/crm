from __future__ import annotations

import math
import shutil
import sys
import textwrap
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


SOURCE = Path(sys.argv[1]).resolve()
OUTPUT = Path(sys.argv[2]).resolve()
WORK = Path(sys.argv[3]).resolve()
WORK.mkdir(parents=True, exist_ok=True)

BLUE = "#14357A"
PRIMARY = "#1476B8"
TEXT = "#334155"
MUTED = "#64748B"
GREEN = "#10B981"
AMBER = "#F59E0B"
RED = "#F43F5E"
PURPLE = "#A855F7"


def font(size: int, bold: bool = False):
    candidates = [
        Path("C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf"),
        Path("C:/Windows/Fonts/calibrib.ttf" if bold else "C:/Windows/Fonts/calibri.ttf"),
    ]
    for candidate in candidates:
        if candidate.exists():
            return ImageFont.truetype(str(candidate), size)
    return ImageFont.load_default()


def wrapped(draw: ImageDraw.ImageDraw, xy, text: str, size=24, bold=False, fill=TEXT, width=27, anchor="mm"):
    lines = "\n".join(textwrap.wrap(text, width=width, break_long_words=False))
    draw.multiline_text(xy, lines, font=font(size, bold), fill=fill, anchor=anchor, align="center", spacing=8)


def box(draw: ImageDraw.ImageDraw, bounds, title: str, detail: str, fill="#EAF4FB", stroke=PRIMARY):
    x1, y1, x2, y2 = bounds
    draw.rounded_rectangle(bounds, radius=22, fill=fill, outline=stroke, width=4)
    draw.text(((x1 + x2) / 2, y1 + 42), title, font=font(25, True), fill=BLUE, anchor="mm")
    wrapped(draw, ((x1 + x2) / 2, (y1 + y2) / 2 + 24), detail, 20, False, TEXT, max(18, int((x2 - x1) / 15)))


def arrow(draw: ImageDraw.ImageDraw, start, end, label="", color=MUTED):
    draw.line((start, end), fill=color, width=4)
    angle = math.atan2(end[1] - start[1], end[0] - start[0])
    length = 18
    wing = 8
    points = [
        end,
        (end[0] - length * math.cos(angle) + wing * math.sin(angle), end[1] - length * math.sin(angle) - wing * math.cos(angle)),
        (end[0] - length * math.cos(angle) - wing * math.sin(angle), end[1] - length * math.sin(angle) + wing * math.cos(angle)),
    ]
    draw.polygon(points, fill=color)
    if label:
        mid = ((start[0] + end[0]) / 2, (start[1] + end[1]) / 2)
        bounds = draw.textbbox(mid, label, font=font(15, True), anchor="mm")
        draw.rounded_rectangle((bounds[0] - 8, bounds[1] - 5, bounds[2] + 8, bounds[3] + 5), radius=6, fill="white")
        draw.text(mid, label, font=font(15, True), fill=TEXT, anchor="mm")


def canvas(title: str, subtitle: str):
    image = Image.new("RGB", (1600, 920), "white")
    draw = ImageDraw.Draw(image)
    draw.text((55, 42), title, font=font(38, True), fill=BLUE)
    draw.text((58, 94), subtitle, font=font(19), fill=MUTED)
    return image, draw


def architecture(path: Path):
    image, draw = canvas("Arsitektur Integrasi WhatsApp WAHA", "Mode hybrid Laragon + Docker dengan queue database")
    box(draw, (55, 210, 330, 400), "WhatsApp", "Nomor admin/staff dan perangkat tertaut", "#EAFBF5", GREEN)
    box(draw, (430, 180, 790, 430), "WAHA GOWS", "Session per user\nAPI X-Api-Key\nWebhook HMAC", "#EAF4FB", PRIMARY)
    box(draw, (900, 180, 1300, 430), "Laravel", "Controller webhook\nEvent idempotent\nLead matcher", "#FFF7E6", AMBER)
    box(draw, (930, 590, 1330, 820), "MySQL", "Akun, event, conversation, message, dan lead", "#FFF7E6", AMBER)
    box(draw, (420, 610, 800, 810), "Queue Worker", "ProcessWahaWebhook\nRetry dan backoff\nWajib aktif", "#F5EEFF", PURPLE)
    box(draw, (55, 610, 330, 810), "Filament", "Koneksi, QR, inbox, RBAC, dan resolusi lead", "#F6F8FC", BLUE)
    arrow(draw, (330, 300), (430, 300), "chat")
    arrow(draw, (790, 300), (900, 300), "webhook HMAC")
    arrow(draw, (1100, 430), (1120, 590), "event")
    arrow(draw, (930, 700), (800, 700), "job")
    arrow(draw, (420, 720), (330, 720), "inbox")
    arrow(draw, (900, 235), (790, 235), "API")
    image.save(path)


def message_flow(path: Path):
    image, draw = canvas("Alur Pesan Masuk hingga Lead", "Pipeline aman, asinkron, idempotent, dan branch-aware")
    items = [
        ("1. Terima", "GOWS menerima chat personal", "#EAF4FB", PRIMARY),
        ("2. Verifikasi", "HMAC SHA-512 dan validasi payload", "#EAFBF5", GREEN),
        ("3. Event", "Hash payload mencegah duplikat", "#FFF7E6", AMBER),
        ("4. Queue", "Worker memproses retry/backoff", "#F5EEFF", PURPLE),
        ("5. Normalisasi", "LID memakai SenderAlt; suffix :device dihapus", "#FFEFF2", RED),
        ("6. Match", "0=lead Cold; 1=tautkan; konflik=review", "#EAF4FB", PRIMARY),
        ("7. Inbox", "Pesan, unread, status, user, cabang", "#EAFBF5", GREEN),
    ]
    positions = [(45 + i * 385, 180, 350 + i * 385, 400) for i in range(4)] + [(45 + i * 385, 590, 350 + i * 385, 810) for i in range(3)]
    centers = []
    for (title, detail, fill, stroke), bounds in zip(items, positions):
        box(draw, bounds, title, detail, fill, stroke)
        centers.append(((bounds[0] + bounds[2]) / 2, (bounds[1] + bounds[3]) / 2))
    for index in range(3):
        arrow(draw, (positions[index][2], centers[index][1]), (positions[index + 1][0], centers[index + 1][1]))
    arrow(draw, (centers[3][0], positions[3][3]), (centers[4][0], positions[4][1]), "asinkron")
    for index in range(4, 6):
        arrow(draw, (positions[index][2], centers[index][1]), (positions[index + 1][0], centers[index + 1][1]))
    image.save(path)


def state_machine(path: Path):
    image, draw = canvas("State Machine Session WAHA", "Pemulihan otomatis dan pengambilan QR yang aman")
    nodes = {
        "Disconnected": (60, 210, 330, 350),
        "Starting": (450, 210, 720, 350),
        "Scan QR": (840, 210, 1110, 350),
        "Working": (450, 600, 720, 740),
        "Failed / Stopped": (1110, 590, 1470, 760),
    }
    box(draw, nodes["Disconnected"], "DISCONNECTED", "Session belum dimulai", "#F6F8FC", MUTED)
    box(draw, nodes["Starting"], "STARTING", "Engine menyiapkan session", "#FFF7E6", AMBER)
    box(draw, nodes["Scan QR"], "SCAN_QR_CODE", "QR siap dipindai", "#EAF4FB", PRIMARY)
    box(draw, nodes["Working"], "WORKING", "gRPC dan stream READY", "#EAFBF5", GREEN)
    box(draw, nodes["Failed / Stopped"], "FAILED / STOPPED", "Timeout, jaringan, atau session berhenti", "#FFEFF2", RED)
    arrow(draw, (330, 280), (450, 280), "Mulai")
    arrow(draw, (720, 280), (840, 280), "siap QR")
    arrow(draw, (975, 350), (650, 600), "scan berhasil")
    arrow(draw, (720, 670), (1110, 670), "disconnect")
    arrow(draw, (1280, 590), (700, 350), "restart otomatis")
    wrapped(draw, (800, 850), "CRM memeriksa status, me-restart FAILED/STOPPED, menunggu SCAN_QR_CODE, lalu mengirim PNG no-cache. Status belum siap menghasilkan halaman pemulihan 503, bukan error 500.", 18, False, TEXT, 115)
    image.save(path)


def erd(path: Path):
    image, draw = canvas("ERD Integrasi WhatsApp", "Entitas, foreign key, relasi logis, dan kardinalitas")
    nodes = {
        "users": (40, 150, 310, 330),
        "accounts": (410, 130, 790, 380),
        "events": (410, 600, 790, 830),
        "conversations": (900, 130, 1280, 390),
        "messages": (900, 600, 1280, 840),
        "leads": (1340, 335, 1580, 650),
    }
    box(draw, nodes["users"], "users", "PK id\nFK cabang_id\nrole", "#EAF4FB", PRIMARY)
    box(draw, nodes["accounts"], "user_whatsapp_accounts", "PK id\nFK user_id unique\nFK cabang_id\nphone + session + status", "#EAF4FB", PRIMARY)
    box(draw, nodes["events"], "waha_webhook_events", "PK id\npayload_hash unique\nevent + status + attempts", "#FFF7E6", AMBER)
    box(draw, nodes["conversations"], "whatsapp_conversations", "PK id\nFK account\nFK lead nullable\nchat_id + match_status", "#EAFBF5", GREEN)
    box(draw, nodes["messages"], "whatsapp_messages", "PK id\nFK conversation/account/lead\nexternal ID + direction + body", "#F5EEFF", PURPLE)
    box(draw, nodes["leads"], "leads", "PK lead_id\nno_wa_normalized\nkanal WhatsApp\nstatus Cold\nfirst/last seen", "#FFEFF2", RED)
    arrow(draw, (310, 240), (410, 240), "1 : 0..1")
    arrow(draw, (790, 250), (900, 250), "1 : 0..N")
    arrow(draw, (1090, 390), (1090, 600), "1 : 0..N")
    arrow(draw, (1280, 260), (1340, 410), "N : 0..1")
    arrow(draw, (1280, 700), (1340, 570), "N : 0..1")
    arrow(draw, (600, 380), (600, 600), "session logis")
    image.save(path)


def dfd(path: Path):
    image, draw = canvas("DFD Level 1 - Integrasi WhatsApp", "Session, webhook, queue, pencocokan lead, dan inbox")
    box(draw, (40, 170, 300, 300), "Pengirim", "Pesan WhatsApp", "#F6F8FC", MUTED)
    box(draw, (40, 650, 300, 790), "Admin / Staff", "Session, QR, dan inbox", "#F6F8FC", MUTED)
    processes = [
        (390, 130, 760, 250, "6.1 Session", "provision/start/restart/QR"),
        (390, 320, 760, 440, "6.2 Webhook", "HMAC dan idempotensi"),
        (390, 510, 760, 630, "6.3 Pesan", "queue dan normalisasi LID"),
        (390, 700, 760, 820, "6.4 Lead", "new/matched/conflict"),
    ]
    for x1, y1, x2, y2, title, detail in processes:
        box(draw, (x1, y1, x2, y2), title, detail, "#EAF4FB", PRIMARY)
    stores = [
        (930, 160, 1320, 260, "D6.1 Akun/Session", "pemilik, cabang, status"),
        (930, 350, 1320, 450, "D6.2 Webhook Events", "payload dan status proses"),
        (930, 540, 1320, 640, "D6.3 Inbox", "conversation dan message"),
        (930, 730, 1320, 830, "D1 Leads", "nomor dan siklus lead"),
    ]
    for x1, y1, x2, y2, title, detail in stores:
        box(draw, (x1, y1, x2, y2), title, detail, "#FFF7E6", AMBER)
    arrow(draw, (300, 720), (390, 190), "perintah")
    arrow(draw, (300, 235), (390, 380), "pesan")
    for index, (_, y1, _, y2, _, _) in enumerate(processes):
        sy = (y1 + y2) / 2
        ty = (stores[index][1] + stores[index][3]) / 2
        arrow(draw, (760, sy), (930, ty), "baca/simpan")
    arrow(draw, (580, 440), (580, 510), "job")
    arrow(draw, (580, 630), (580, 700), "nomor")
    arrow(draw, (930, 590), (300, 740), "inbox")
    image.save(path)


def replace_text(document: Document, replacements: dict[str, str]):
    containers = list(document.paragraphs)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                containers.extend(cell.paragraphs)
    for paragraph in containers:
        for run in paragraph.runs:
            for old, new in replacements.items():
                if old in run.text:
                    run.text = run.text.replace(old, new)


def set_cell_text(cell, text: str, bold=False, color=TEXT):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.font.name = "Aptos"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Aptos")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Aptos")
    run.font.size = Pt(9)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color.lstrip("#"))
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(document: Document, headers, rows, widths=None):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for index, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[index], header, True, "FFFFFF")
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), PRIMARY.lstrip("#"))
        table.rows[0].cells[index]._tc.get_or_add_tcPr().append(shd)
    for row_data in rows:
        row = table.add_row()
        for index, value in enumerate(row_data):
            set_cell_text(row.cells[index], str(value))
    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Inches(width)
    document.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_para(document: Document, text: str, bold_lead: str | None = None):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.space_after = Pt(6)
    if bold_lead and text.startswith(bold_lead):
        lead = paragraph.add_run(bold_lead)
        lead.bold = True
        paragraph.add_run(text[len(bold_lead):])
    else:
        paragraph.add_run(text)
    return paragraph


def add_bullets(document: Document, items):
    for item in items:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.add_run(item)


def add_steps(document: Document, items):
    for item in items:
        paragraph = document.add_paragraph(style="List Number")
        paragraph.add_run(item)


def add_figure(document: Document, path: Path, caption: str):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    paragraph.add_run().add_picture(str(path), width=Inches(6.65))
    cap = document.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(8)
    run = cap.add_run(caption)
    run.bold = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(BLUE.lstrip("#"))


cover = WORK / "cover-v1.5.png"
architecture_path = WORK / "waha-architecture.png"
flow_path = WORK / "waha-message-flow.png"
state_path = WORK / "waha-state-machine.png"
erd_path = WORK / "waha-erd.png"
dfd_path = WORK / "waha-dfd.png"

architecture(architecture_path)
message_flow(flow_path)
state_machine(state_path)
erd(erd_path)
dfd(dfd_path)

cover_source = Path(__file__).resolve().parents[1] / "docs" / "_diagram_v1_5" / "cover-v1_5.png"
if cover_source.exists():
    shutil.copyfile(cover_source, cover)

document = Document(SOURCE)
replace_text(document, {
    "Sementara 1.4": "Sementara 1.5",
    "v1.4": "v1.5",
    "40 test dengan 267 assertion": "60 test dengan 400 assertion",
    "40 test dan 267 assertion": "60 test dan 400 assertion",
})

if cover.exists() and document.inline_shapes:
    rel_id, _ = document.part.get_or_add_image(str(cover))
    blip = document.inline_shapes[0]._inline.graphic.graphicData.pic.blipFill.blip
    blip.set(qn("r:embed"), rel_id)

for paragraph in list(document.paragraphs):
    if paragraph.text.strip() == "--- Akhir dokumen ---":
        paragraph._element.getparent().remove(paragraph._element)

heading = document.add_heading("LAMPIRAN B: INTEGRASI WHATSAPP WAHA DAN INBOX", level=1)
heading.paragraph_format.page_break_before = True
add_para(document, "Lampiran ini mencatat implementasi as-built integrasi WhatsApp terbaru pada SIVMI CRM. Setiap admin/staff mempunyai nomor dan session WAHA sendiri. WAHA menggunakan engine GOWS dalam container Docker, sedangkan aplikasi Laravel dapat tetap berjalan melalui Laragon (hybrid) atau berada dalam Compose yang sama pada production.")

document.add_heading("B.1 Tujuan dan Ruang Lingkup", level=2)
add_bullets(document, [
    "Merekam pesan masuk WhatsApp secara otomatis dan menampilkannya pada Inbox WhatsApp sesuai role dan cabang.",
    "Mencegah duplikasi lead melalui normalisasi nomor, pencocokan dalam cabang, status konflik, dan idempotensi webhook.",
    "Mengelola lifecycle session, QR, restart, status koneksi, serta kepemilikan nomor per pengguna.",
    "Menjaga keamanan API dan webhook melalui API key, HMAC SHA-512, autentikasi panel, RBAC, dan branch scope.",
])

document.add_heading("B.2 Arsitektur Komponen", level=2)
add_figure(document, architecture_path, "Gambar B-1. Arsitektur integrasi WhatsApp WAHA pada konfigurasi hybrid")
add_table(document, ["Komponen", "Tanggung jawab", "Kontrol utama"], [
    ["WAHA GOWS", "Session WhatsApp, QR, koneksi perangkat, API, dan webhook", "API key, volume session, health check"],
    ["Laravel", "Validasi webhook, penyimpanan event, RBAC, inbox, dan pencocokan lead", "HMAC, validation, branch scope"],
    ["Queue worker", "Memproses event secara asinkron dan dapat diulang", "Retry, backoff, failed_jobs, monitoring"],
    ["MySQL", "Akun WA, event, percakapan, pesan, dan lead", "FK, unique index, transaction"],
    ["Filament", "Koneksi WA, QR terlindungi, inbox, dan resolusi konflik", "Authentication dan authorization"],
], [1.25, 3.35, 2.1])

document.add_heading("B.3 Lifecycle Session dan QR", level=2)
add_figure(document, state_path, "Gambar B-2. State machine session WAHA dan pemulihan QR")
add_para(document, "Saat halaman QR dibuka, CRM mengambil status session terlebih dahulu. Status FAILED atau STOPPED memicu restart, sedangkan STARTING dipantau sampai berubah. QR hanya diminta ketika status SCAN_QR_CODE. Apabila belum siap, pengguna memperoleh halaman pemulihan berstatus 503 tanpa stack trace atau error 500.")

document.add_heading("B.4 Alur Proses Pesan Masuk", level=2)
add_figure(document, flow_path, "Gambar B-3. Alur pesan masuk, normalisasi, dan pencocokan lead")
add_table(document, ["Tahap", "Proses"], [
    ["1", "WAHA GOWS menerima pesan personal pada session milik admin/staff."],
    ["2", "WAHA mengirim event message ke endpoint Laravel dengan signature HMAC SHA-512."],
    ["3", "Controller memvalidasi payload dan menyimpan event berdasarkan hash payload unik."],
    ["4", "Job ProcessWahaWebhook dimasukkan ke queue database."],
    ["5", "Worker membaca payload, mengabaikan group/newsletter/status, dan menormalisasi identitas pengirim."],
    ["6", "Nomor dibandingkan dengan lead pada cabang yang berwenang; hasilnya NewLead, Matched, Conflict, atau Unresolved."],
    ["7", "Conversation dan message disimpan; unread count serta waktu WhatsApp lead diperbarui."],
], [0.7, 6.0])

document.add_heading("B.5 DFD Level 1 Integrasi WhatsApp", level=2)
add_figure(document, dfd_path, "Gambar B-4. DFD Level 1 integrasi session, webhook, inbox, dan lead")

document.add_heading("B.6 ERD, Relasi, dan Kardinalitas", level=2)
add_figure(document, erd_path, "Gambar B-5. ERD integrasi WhatsApp dan pencocokan lead")
add_table(document, ["Relasi", "Kardinalitas", "Aturan"], [
    ["users - user_whatsapp_accounts", "1 : 0..1", "Satu pengguna maksimal satu koneksi aktif"],
    ["cabang - user_whatsapp_accounts", "1 : 0..N", "Akun mengikuti cabang pemilik"],
    ["account - conversations", "1 : 0..N", "Chat unik per akun dan chat_id"],
    ["conversation - messages", "1 : 0..N", "Riwayat pesan berurutan berdasarkan sent_at"],
    ["lead - conversations", "1 : 0..N", "Relasi opsional sampai proses matching selesai"],
    ["lead - messages", "1 : 0..N", "Pesan mengikuti lead yang sudah ditautkan"],
    ["account - webhook events", "1 : 0..N logis", "Dihubungkan melalui session_name"],
], [2.5, 1.2, 3.0])

document.add_heading("B.7 Normalisasi GOWS dan LID", level=2)
add_para(document, "GOWS dapat mengirim chat personal dengan ID berakhiran @lid. LID dipertahankan sebagai chat_id, tetapi bukan nomor telepon. Untuk pesan masuk, sistem mengambil nomor dari _data.Info.SenderAlt dan nama dari _data.Info.PushName. Suffix perangkat seperti :5 pada JID dihapus sebelum digit dinormalisasi ke format E.164 tanpa tanda plus.")
add_table(document, ["Field payload", "Contoh", "Pemakaian"], [
    ["from", "109719301685330@lid", "Identitas chat GOWS"],
    ["_data.Info.SenderAlt", "6285320008589:5@s.whatsapp.net", "Sumber nomor telepon"],
    ["_data.Info.PushName", "Nama Pengirim", "Nama kontak/lead awal"],
    ["id", "false_<lid>_<message-id>", "Idempotensi pesan per akun"],
], [2.1, 2.8, 1.8])

document.add_heading("B.8 Aturan Pencocokan dan Deduplikasi Lead", level=2)
add_table(document, ["Kondisi", "Hasil", "Dampak data"], [
    ["Tidak ada lead dengan nomor sama", "NewLead", "Buat lead Cold, kanal WhatsApp, milik session/cabang"],
    ["Tepat satu lead pada cabang sama", "Matched", "Tautkan tanpa membuat duplikat"],
    ["Lebih dari satu lead cocok", "Conflict", "Tidak mengubah lead; perlu resolusi manual"],
    ["Nomor cocok pada cabang lain", "Conflict", "Mencegah akses silang cabang"],
    ["Nomor tidak dapat diekstrak", "Unresolved", "Simpan chat, tunggu resolusi manual"],
], [2.2, 1.15, 3.35])

document.add_heading("B.9 Keamanan, RBAC, dan Audit", level=2)
add_bullets(document, [
    "API WAHA menggunakan X-Api-Key; webhook menggunakan secret HMAC terpisah.",
    "Request tanpa signature, algoritma selain SHA-512, atau signature tidak valid ditolak.",
    "Endpoint QR membutuhkan autentikasi dan hak mengelola data operasional.",
    "Staff hanya melihat conversation milik akun sendiri; role lebih tinggi tetap dibatasi cabang sesuai scope.",
    "Payload mentah yang disimpan dibatasi pada metadata yang diperlukan; secret tidak dicatat pada log aplikasi.",
])

document.add_heading("B.10 Queue, Scheduler, dan Operasional", level=2)
add_para(document, "Webhook disimpan lebih dahulu sebelum diproses. Karena QUEUE_CONNECTION=database, worker wajib berjalan terus. Event dapat terlihat di waha_webhook_events walaupun inbox kosong apabila worker berhenti.")
add_table(document, ["Lingkungan", "Cara menjalankan", "Ketahanan"], [
    ["Hybrid Laragon", "php artisan queue:work --tries=5 --timeout=120", "Terminal/service harus tetap hidup dan dijalankan ulang setelah restart"],
    ["Docker development", "docker compose --profile workers up -d", "Service queue dan scheduler mengikuti Compose"],
    ["Production", "Service queue pada compose.production.yml", "Restart policy, monitoring, dan alert failed job"],
], [1.45, 3.1, 2.15])
add_para(document, "Diagnosis cepat.", bold_lead="Diagnosis cepat.")
add_bullets(document, [
    "Event bertambah tetapi conversation/message nol: worker queue tidak berjalan atau job masih Pending.",
    "Event tidak bertambah: periksa webhook URL, host-gateway, firewall, HMAC, dan status session WORKING.",
    "Event Failed atau failed_jobs terisi: periksa last_error dan storage/logs/laravel.log sebelum retry.",
])

document.add_heading("B.11 Deployment Hybrid dan Full Docker", level=2)
add_table(document, ["Mode", "WAHA_BASE_URL", "WAHA_WEBHOOK_URL"], [
    ["Laragon + WAHA Docker", "http://127.0.0.1:3000", "http://sivmi.crm.test/api/webhooks/waha"],
    ["Full Docker", "http://waha:3000", "http://web/api/webhooks/waha"],
], [2.0, 2.25, 2.45])
add_bullets(document, [
    "Gunakan engine GOWS dan volume waha_sessions persisten.",
    "Jangan gunakan docker compose down -v kecuali seluruh session memang akan dihapus.",
    "Production menggunakan HTTPS, credential unik, dashboard/Swagger terbatas, dan port WAHA tidak dipublikasikan ke internet.",
])

document.add_heading("B.12 UAT dan Kriteria Penerimaan", level=2)
add_table(document, ["Skenario", "Hasil yang diharapkan", "Status as-built"], [
    ["Provision dan start session", "Session dibuat untuk user/cabang tepat", "Terverifikasi"],
    ["QR dan restart", "FAILED/STOPPED pulih; QR PNG muncul saat siap", "Automated + runtime"],
    ["Pesan GOWS LID", "SenderAlt dan PushName dikenali", "Terverifikasi"],
    ["Lead baru", "Cold + kanal WhatsApp + owner session", "Terverifikasi"],
    ["Nomor sama", "Tautkan atau Conflict tanpa duplikat", "Automated test"],
    ["RBAC inbox", "Staff hanya melihat conversation miliknya", "Automated test"],
    ["Queue berhenti", "Event tertahan dan dapat diproses ulang", "Terverifikasi"],
], [2.15, 3.15, 1.4])

document.add_heading("B.13 Snapshot Pengujian dan Changelog", level=2)
add_para(document, "Snapshot 24 Agustus 2026: WAHA 2026.8.1 CORE dengan GOWS berstatus WORKING; gRPC dan stream READY; webhook HMAC dapat menjangkau aplikasi hybrid; pesan berbasis LID tersimpan; nomor SenderAlt dinormalisasi; lead baru dibuat dan ditautkan. Suite aplikasi lulus 60 test dengan 400 assertion.")
add_table(document, ["Versi", "Perubahan utama"], [
    ["v1.5", "Integrasi WAHA/GOWS, inbox, HMAC, queue, LID, ERD/DFD/state machine, UAT, dan troubleshooting"],
    ["v1.4", "Struktur file/folder as-built"],
    ["v1.3", "Arsitektur, deployment, debugging, panduan, dan SOP diperluas"],
], [1.0, 5.7])

end = document.add_paragraph("--- Akhir dokumen ---")
end.alignment = WD_ALIGN_PARAGRAPH.CENTER
end.runs[0].font.size = Pt(8)
end.runs[0].font.color.rgb = RGBColor.from_string(MUTED.lstrip("#"))

document.core_properties.comments = "Versi sementara 1.5 - snapshot 24 Agustus 2026 - integrasi WAHA GOWS as-built"

image_alt_texts = [
    "Sampul dokumentasi dan pengembangan sistem CRM SIVMI versi 1.5.",
    "Diagram proses bisnis end-to-end CRM SIVMI.",
    "Diagram arsitektur bisnis CRM SIVMI.",
    "Flowchart proses lead hingga pembayaran.",
    "Flowmap proses lintas aktor CRM SIVMI.",
    "DFD level 0 atau diagram konteks CRM SIVMI.",
    "DFD level 1 proses utama CRM SIVMI.",
    "ERD operasional serta relasi data CRM SIVMI.",
    "Diagram arsitektur, sitemap, atau kontrol data CRM SIVMI.",
    "Arsitektur integrasi WhatsApp WAHA pada konfigurasi hybrid.",
    "State machine lifecycle session WAHA dan pemulihan QR.",
    "Alur pesan masuk, normalisasi GOWS, dan pencocokan lead.",
    "DFD level 1 integrasi WhatsApp.",
    "ERD integrasi WhatsApp, conversation, message, dan lead.",
]
for index, shape in enumerate(document.inline_shapes):
    description = image_alt_texts[index] if index < len(image_alt_texts) else f"Diagram dokumentasi SIVMI CRM {index + 1}."
    shape._inline.docPr.set("title", f"Gambar {index + 1} SIVMI CRM")
    shape._inline.docPr.set("descr", description)

for table in document.tables:
    if not table.rows:
        continue
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:tblHeader")) is None:
        header_marker = OxmlElement("w:tblHeader")
        header_marker.set(qn("w:val"), "true")
        tr_pr.append(header_marker)

settings = document.settings._element
update_fields = settings.find(qn("w:updateFields"))
if update_fields is None:
    update_fields = OxmlElement("w:updateFields")
    settings.append(update_fields)
update_fields.set(qn("w:val"), "true")

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
document.save(OUTPUT)
print(OUTPUT)
