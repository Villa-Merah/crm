from __future__ import annotations

import hashlib
import json
import sys
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


def emu_inches(value):
    return None if value is None else round(value.inches, 3)


def color_value(font):
    return None if font.color is None or font.color.rgb is None else str(font.color.rgb)


def style_record(style):
    pf = style.paragraph_format
    font = style.font
    return {
        "name": style.name,
        "type": str(style.type),
        "font": font.name,
        "size_pt": None if font.size is None else round(font.size.pt, 2),
        "bold": font.bold,
        "italic": font.italic,
        "color": color_value(font),
        "alignment": None if pf.alignment is None else str(pf.alignment),
        "space_before_pt": None if pf.space_before is None else round(pf.space_before.pt, 2),
        "space_after_pt": None if pf.space_after is None else round(pf.space_after.pt, 2),
        "line_spacing": None if pf.line_spacing is None else str(pf.line_spacing),
        "left_indent_in": emu_inches(pf.left_indent),
        "first_line_indent_in": emu_inches(pf.first_line_indent),
        "keep_with_next": pf.keep_with_next,
    }


def paragraph_record(paragraph, index):
    return {
        "index": index,
        "style": paragraph.style.name if paragraph.style else None,
        "text": paragraph.text.strip(),
        "alignment": None if paragraph.alignment is None else str(paragraph.alignment),
    }


def main():
    source = Path(sys.argv[1]).resolve()
    out_dir = Path(sys.argv[2]).resolve()
    out_dir.mkdir(parents=True, exist_ok=True)

    raw = source.read_bytes()
    document = Document(source)
    section = document.sections[0]
    headings = [
        paragraph_record(p, i)
        for i, p in enumerate(document.paragraphs)
        if p.text.strip() and p.style and p.style.name.startswith("Heading")
    ]
    opening = [paragraph_record(p, i) for i, p in enumerate(document.paragraphs[:80]) if p.text.strip()]
    style_names = ["Normal", "No Spacing", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3", "TOC 1", "TOC 2", "TOC 3"]
    styles = [style_record(document.styles[name]) for name in style_names if name in document.styles]

    package = []
    with zipfile.ZipFile(source) as archive:
        for info in archive.infolist():
            package.append({"path": info.filename, "size": info.file_size, "sha256": hashlib.sha256(archive.read(info.filename)).hexdigest()})

    evidence = {
        "reference": str(source),
        "sha256": hashlib.sha256(raw).hexdigest(),
        "bytes": len(raw),
        "paragraphs": len(document.paragraphs),
        "tables": len(document.tables),
        "inline_shapes": len(document.inline_shapes),
        "sections": len(document.sections),
        "page": {
            "width_in": emu_inches(section.page_width),
            "height_in": emu_inches(section.page_height),
            "left_margin_in": emu_inches(section.left_margin),
            "right_margin_in": emu_inches(section.right_margin),
            "top_margin_in": emu_inches(section.top_margin),
            "bottom_margin_in": emu_inches(section.bottom_margin),
            "header_distance_in": emu_inches(section.header_distance),
            "footer_distance_in": emu_inches(section.footer_distance),
            "different_first_page": section.different_first_page_header_footer,
        },
        "styles": styles,
        "opening": opening,
        "headings": headings,
        "header_text": [p.text for p in section.header.paragraphs if p.text.strip()],
        "first_page_header_text": [p.text for p in section.first_page_header.paragraphs if p.text.strip()],
        "footer_text": [p.text for p in section.footer.paragraphs if p.text.strip()],
        "package": package,
    }
    (out_dir / "reference-evidence.json").write_text(json.dumps(evidence, ensure_ascii=False, indent=2), encoding="utf-8")

    lines = [
        "# Artifact contract - Referensi Dokumentasi CRM",
        "",
        f"- Reference: `{source}`",
        f"- SHA-256: `{evidence['sha256']}`",
        f"- Paragraphs: {evidence['paragraphs']}; tables: {evidence['tables']}; sections: {evidence['sections']}",
        f"- Page: {evidence['page']['width_in']} x {evidence['page']['height_in']} in; margins L/R/T/B: {evidence['page']['left_margin_in']}/{evidence['page']['right_margin_in']}/{evidence['page']['top_margin_in']}/{evidence['page']['bottom_margin_in']} in",
        "",
        "## Typography roles",
    ]
    for style in styles:
        lines.append(f"- {style['name']}: {style['font']}, {style['size_pt']} pt, bold={style['bold']}, color={style['color']}, before={style['space_before_pt']} pt, after={style['space_after_pt']} pt")
    lines.extend(["", "## Content flow", "- Cover/title block", "- Revision metadata", "- Daftar isi berbasis Heading 1-3 dan PAGEREF", "- Bab dokumentasi berjenjang", "- Header/footer dan nomor halaman", "", "## Slot map", "- Seluruh isi proyek lama diganti dengan dokumentasi SIVMI; gaya, hierarki heading, page system, dan footer menjadi pola sumber.", "- Daftar isi harus diperbarui saat dokumen dibuka/diubah di Microsoft Word.", "", "## Fidelity gates", "- Ukuran halaman, margin, style heading, header/footer, dan sistem nomor halaman tetap berasal dari referensi.", "- Tidak ada teks proyek referensi yang tersisa di dokumen final.", "- Semua halaman final harus lolos inspeksi visual."])
    (out_dir / "artifact.md").write_text("\n".join(lines), encoding="utf-8")

    print(json.dumps({k: evidence[k] for k in ["sha256", "bytes", "paragraphs", "tables", "inline_shapes", "sections", "page", "styles", "opening", "headings", "header_text", "footer_text"]}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
