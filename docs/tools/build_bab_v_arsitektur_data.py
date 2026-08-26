from __future__ import annotations

from pathlib import Path
import sys

PROJECT_ROOT = Path(__file__).resolve().parents[2]
DOC_TOOLS = PROJECT_ROOT / "storage" / "app" / "doc-tools"
sys.path.insert(0, str(DOC_TOOLS))

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

import build_bab_iv_arsitektur_bisnis as b


OUTPUT = PROJECT_ROOT / "docs" / "BAB V - Arsitektur Data SIVMI CRM v1.1.docx"
ASSETS = PROJECT_ROOT / "docs" / "_bab_v_assets"
COVER = ASSETS / "cover-bab-v.png"
LIFECYCLE = ASSETS / "siklus-data-lead.png"
USE_CASE = ASSETS / "use-case-data.png"
ACTIVITY = ASSETS / "activity-data.png"
ERD_OPERATIONAL = ASSETS / "erd-operasional.png"
ERD_PERFORMANCE = ASSETS / "erd-kinerja-kontrol.png"
ERD_WHATSAPP = ASSETS / "erd-whatsapp.png"
DATABASE_LAYOUT = ASSETS / "susunan-tabel-database.png"


def font(name: str, size: int):
    return ImageFont.truetype(f"C:/Windows/Fonts/{name}", size)


def make_cover() -> None:
    ASSETS.mkdir(parents=True, exist_ok=True)
    width, height = 1600, 2200
    image = Image.new("RGB", (width, height), "#F6F8FC")
    draw = ImageDraw.Draw(image)
    for y in range(height):
        ratio = y / height
        start = (20, 118, 184)
        end = (7, 89, 133)
        color = tuple(round(start[i] * (1 - ratio) + end[i] * ratio) for i in range(3))
        draw.line((0, y, width, y), fill=color)
    draw.ellipse((1010, -180, 1770, 580), fill=(27, 139, 199))
    draw.ellipse((1140, 140, 1640, 640), fill=(8, 103, 151))
    draw.rounded_rectangle((120, 250, 1480, 1910), radius=44, fill=(255, 255, 255))
    draw.text((190, 350), "SIVMI CRM", font=font("arialbd.ttf", 34), fill="#1476B8")
    draw.text((190, 600), "BAB V", font=font("arialbd.ttf", 48), fill="#075985")
    draw.multiline_text((190, 690), "ARSITEKTUR\nDATA", font=font("arialbd.ttf", 102), fill="#0B2545", spacing=18)
    draw.text((190, 1040), "RBAC, alur data, DFD, ERD, kardinalitas,", font=font("arial.ttf", 34), fill="#475569")
    draw.text((190, 1092), "kamus data, siklus lead, dan keamanan akses", font=font("arial.ttf", 34), fill="#475569")
    draw.rounded_rectangle((190, 1245, 760, 1325), radius=40, fill="#EAF4FB")
    draw.text((230, 1267), "BASELINE AS-BUILT 24 AGUSTUS 2026", font=font("arial.ttf", 28), fill="#075985")
    draw.text((190, 1610), "Versi dokumen 1.1", font=font("arial.ttf", 28), fill="#64748B")
    draw.text((190, 1665), "Status: dokumen kerja untuk validasi manajemen", font=font("arial.ttf", 28), fill="#64748B")
    draw.text((190, 1775), "SIVMI CRM - Catat, Kelola, Closing", font=font("arialbd.ttf", 34), fill="#1476B8")
    image.save(COVER)


def rounded_box(draw, xy, fill, outline, title, lines, title_color="#0B2545"):
    draw.rounded_rectangle(xy, radius=24, fill=fill, outline=outline, width=4)
    x1, y1, x2, y2 = xy
    draw.text(((x1 + x2) / 2, y1 + 28), title, anchor="ma", font=font("arialbd.ttf", 30), fill=title_color)
    draw.multiline_text(((x1 + x2) / 2, y1 + 78), lines, anchor="ma", align="center", spacing=10, font=font("arial.ttf", 23), fill="#334155")


def arrow(draw, start, end, color="#64748B", width=5):
    draw.line((start, end), fill=color, width=width)
    x, y = end
    sx, sy = start
    if abs(x - sx) >= abs(y - sy):
        sign = 1 if x > sx else -1
        points = [(x, y), (x - sign * 18, y - 11), (x - sign * 18, y + 11)]
    else:
        sign = 1 if y > sy else -1
        points = [(x, y), (x - 11, y - sign * 18), (x + 11, y - sign * 18)]
    draw.polygon(points, fill=color)


def make_lifecycle() -> None:
    image = Image.new("RGB", (1800, 980), "white")
    draw = ImageDraw.Draw(image)
    draw.text((900, 45), "Siklus Data Lead SIVMI CRM", anchor="ma", font=font("arialbd.ttf", 46), fill="#14357A")
    draw.text((900, 105), "Dari akuisisi hingga retensi, dengan kontrol cabang, audit, dan pencocokan WhatsApp", anchor="ma", font=font("arial.ttf", 25), fill="#64748B")
    boxes = [
        ("1. Akuisisi", "Manual / CSV /\nWhatsApp", "#EAF4FB", "#1476B8"),
        ("2. Validasi", "Identitas, nomor,\nsekolah, program", "#ECFDF5", "#10B981"),
        ("3. Distribusi", "Cabang dan PIC\nditetapkan", "#FFFBEB", "#F59E0B"),
        ("4. Engagement", "Follow-up, chat,\nstatus pipeline", "#FAF5FF", "#A855F7"),
        ("5. Konversi", "Pendaftaran dan\nprogram", "#FFF1F2", "#F43F5E"),
        ("6. Transaksi", "Pembayaran dan\nverifikasi", "#EAF4FB", "#1476B8"),
        ("7. Monitoring", "KPI, target, audit,\nretensi", "#ECFDF5", "#10B981"),
    ]
    x0, y, w, h, gap = 45, 260, 225, 260, 25
    for idx, (title, lines, fill, border) in enumerate(boxes):
        x = x0 + idx * (w + gap)
        rounded_box(draw, (x, y, x + w, y + h), fill, border, title, lines)
        if idx < len(boxes) - 1:
            arrow(draw, (x + w, y + h / 2), (x + w + gap - 4, y + h / 2))
    draw.rounded_rectangle((260, 650, 1540, 865), radius=26, fill="#F8FAFC", outline="#CBD5E1", width=4)
    draw.text((900, 690), "Kontrol lintas siklus", anchor="ma", font=font("arialbd.ttf", 30), fill="#075985")
    draw.multiline_text((900, 745), "Normalisasi nomor dan deduplikasi | RBAC dan branch scope | validasi FK dan status\nindeks tanggal/status | audit event | notifikasi | backup dan kebijakan retensi", anchor="ma", align="center", spacing=12, font=font("arial.ttf", 25), fill="#334155")
    image.save(LIFECYCLE)


def make_use_case() -> None:
    image = Image.new("RGB", (1650, 1100), "white")
    draw = ImageDraw.Draw(image)
    draw.text((825, 45), "Use Case Pengelolaan Data", anchor="ma", font=font("arialbd.ttf", 45), fill="#14357A")
    actors = [("Superadmin", 115, 235), ("Admin Cabang", 115, 535), ("Staff", 115, 835), ("Direksi", 1535, 535), ("WAHA GOWS", 1535, 835)]
    for name, x, y in actors:
        draw.ellipse((x - 28, y - 75, x + 28, y - 19), outline="#1476B8", width=4)
        draw.line((x, y - 19, x, y + 65), fill="#1476B8", width=4)
        draw.line((x - 48, y + 8, x + 48, y + 8), fill="#1476B8", width=4)
        draw.line((x, y + 65, x - 45, y + 120), fill="#1476B8", width=4)
        draw.line((x, y + 65, x + 45, y + 120), fill="#1476B8", width=4)
        draw.text((x, y + 145), name, anchor="ma", font=font("arialbd.ttf", 24), fill="#334155")
    draw.rounded_rectangle((330, 125, 1320, 1030), radius=30, fill="#F8FAFC", outline="#CBD5E1", width=4)
    cases = [
        ("Kelola master, role, dan pengaturan", 825, 230),
        ("Input, impor, validasi, dan distribusi lead", 825, 365),
        ("Catat follow-up, pendaftaran, dan pembayaran", 825, 500),
        ("Kelola akun, conversation, dan message WhatsApp", 825, 635),
        ("Pantau KPI, target, laporan, notifikasi, dan audit", 825, 770),
        ("Terima webhook, cocokkan lead, dan simpan event", 825, 905),
    ]
    for label, x, y in cases:
        draw.ellipse((x - 365, y - 48, x + 365, y + 48), fill="#EAF4FB", outline="#1476B8", width=3)
        draw.text((x, y), label, anchor="mm", font=font("arialbd.ttf", 24), fill="#0B2545")
    for idx in range(5):
        arrow(draw, (175, 235 + (idx % 3) * 300), (450, 230 + idx * 135), width=3)
    arrow(draw, (1475, 535), (1190, 770), width=3)
    arrow(draw, (1475, 835), (1190, 905), width=3)
    image.save(USE_CASE)


def make_activity() -> None:
    image = Image.new("RGB", (1650, 1220), "white")
    draw = ImageDraw.Draw(image)
    draw.text((825, 45), "Activity Diagram Akuisisi dan Persistensi Lead", anchor="ma", font=font("arialbd.ttf", 43), fill="#14357A")
    lanes = [(70, 420, "Pengguna / WAHA"), (420, 920, "Aplikasi CRM"), (920, 1280, "Database"), (1280, 1580, "Pengguna penerima")]
    for x1, x2, label in lanes:
        draw.rectangle((x1, 125, x2, 1140), outline="#CBD5E1", width=3)
        draw.rectangle((x1, 125, x2, 195), fill="#EAF4FB", outline="#CBD5E1", width=3)
        draw.text(((x1 + x2) / 2, 160), label, anchor="mm", font=font("arialbd.ttf", 24), fill="#075985")
    steps = [
        (245, 285, "Kirim form, CSV,\natau pesan WhatsApp"),
        (670, 285, "Autentikasi kanal dan\nvalidasi payload"),
        (670, 450, "Normalisasi nomor,\nstatus, cabang, dan PIC"),
        (1100, 450, "Cari kandidat lead\ndan constraint unik"),
        (670, 625, "Tentukan NewLead, Matched,\nConflict, Unresolved, Ignored"),
        (1100, 625, "Simpan lead/event/\nconversation/message"),
        (670, 800, "Buat audit dan notifikasi;\nperbarui KPI"),
        (1430, 800, "Tinjau inbox, lead,\natau tindakan tertunda"),
        (670, 990, "Selesai / retry bila gagal"),
    ]
    for x, y, label in steps:
        fill = "#ECFDF5" if x == 1100 else "#EAF4FB"
        outline = "#10B981" if x == 1100 else "#1476B8"
        draw.rounded_rectangle((x - 165, y - 58, x + 165, y + 58), radius=20, fill=fill, outline=outline, width=3)
        draw.multiline_text((x, y), label, anchor="mm", align="center", font=font("arialbd.ttf", 21), fill="#334155", spacing=7)
    connections = [(245,285,505,285),(670,343,670,392),(835,450,935,450),(1100,508,835,625),(835,625,935,625),(1100,683,835,800),(835,800,1265,800),(670,858,670,932)]
    for x1,y1,x2,y2 in connections:
        arrow(draw,(x1,y1),(x2,y2),width=4)
    image.save(ACTIVITY)


def erd_card_height(fields: list[tuple[str, str]]) -> int:
    return 78 + len(fields) * 39 + 18


def erd_card(draw: ImageDraw.ImageDraw, xy: tuple[int, int, int, int], title: str,
             fields: list[tuple[str, str]], color: str) -> None:
    """Draw a high-contrast database table card with explicit key badges."""
    x1, y1, x2, y2 = xy
    draw.rounded_rectangle(xy, radius=18, fill="white", outline=color, width=4)
    draw.rounded_rectangle((x1, y1, x2, y1 + 62), radius=18, fill=color, outline=color, width=4)
    draw.rectangle((x1, y1 + 43, x2, y1 + 62), fill=color)
    draw.text((x1 + 20, y1 + 31), title, anchor="lm", font=font("arialbd.ttf", 27), fill="white")
    for row_index, (key, label) in enumerate(fields):
        top = y1 + 68 + row_index * 39
        if row_index % 2:
            draw.rectangle((x1 + 3, top, x2 - 3, top + 39), fill="#F8FAFC")
        if row_index:
            draw.line((x1 + 3, top, x2 - 3, top), fill="#E2E8F0", width=2)
        if key:
            badge_fill = {"PK": "#075985", "FK": "#A855F7", "UQ": "#F59E0B", "IDX": "#10B981"}.get(key, "#64748B")
            draw.rounded_rectangle((x1 + 12, top + 7, x1 + 66, top + 32), radius=8, fill=badge_fill)
            draw.text((x1 + 39, top + 20), key, anchor="mm", font=font("arialbd.ttf", 16), fill="white")
        draw.text((x1 + 78, top + 20), label, anchor="lm", font=font("arial.ttf", 22), fill="#1E293B")


def relation(draw: ImageDraw.ImageDraw, points: list[tuple[int, int]], label: str,
             color: str = "#64748B", dashed: bool = False,
             label_pos: tuple[int, int] | None = None) -> None:
    """Draw an orthogonal relation behind cards and place a readable cardinality badge."""
    for start, end in zip(points, points[1:]):
        if dashed:
            length = max(abs(end[0] - start[0]), abs(end[1] - start[1]))
            parts = max(1, length // 18)
            for index in range(0, parts, 2):
                a = index / parts
                z = min(1, (index + 1) / parts)
                p1 = (round(start[0] + (end[0] - start[0]) * a), round(start[1] + (end[1] - start[1]) * a))
                p2 = (round(start[0] + (end[0] - start[0]) * z), round(start[1] + (end[1] - start[1]) * z))
                draw.line((p1, p2), fill=color, width=4)
        else:
            draw.line((start, end), fill=color, width=4)
    start, end = points[-2], points[-1]
    arrow(draw, start, end, color=color, width=4)
    longest = max(zip(points, points[1:]), key=lambda pair: abs(pair[1][0] - pair[0][0]) + abs(pair[1][1] - pair[0][1]))
    lx = (longest[0][0] + longest[1][0]) // 2
    ly = (longest[0][1] + longest[1][1]) // 2
    if label_pos is not None:
        lx, ly = label_pos
    tw = max(92, int(draw.textlength(label, font=font("arialbd.ttf", 18))) + 24)
    draw.rounded_rectangle((lx - tw // 2, ly - 18, lx + tw // 2, ly + 18), radius=9, fill="white", outline="#CBD5E1", width=2)
    draw.text((lx, ly), label, anchor="mm", font=font("arialbd.ttf", 18), fill="#334155")


def diagram_title(draw: ImageDraw.ImageDraw, title: str, subtitle: str, width: int) -> None:
    draw.text((50, 42), title, anchor="la", font=font("arialbd.ttf", 39), fill="#075985")
    draw.text((50, 91), subtitle, anchor="la", font=font("arial.ttf", 23), fill="#64748B")
    draw.line((50, 122, width - 50, 122), fill="#D7E2EC", width=3)


def make_erd_operational() -> None:
    image = Image.new("RGB", (1600, 1420), "white")
    draw = ImageDraw.Draw(image)
    diagram_title(draw, "ERD Operasional: Lead sampai Pembayaran", "Kolom kunci ditampilkan; atribut lengkap tersedia pada kamus data.", 1600)
    cards = {
        "cabang": (40, 155, 330, 155 + erd_card_height([("PK", "cabang_id"), ("UQ", "kode_cabang"), ("", "nama"), ("IDX", "aktif")])),
        "users": (360, 155, 680, 155 + erd_card_height([("PK", "id"), ("FK", "cabang_id?"), ("UQ", "email"), ("", "role"), ("IDX", "aktif")])),
        "sekolah": (710, 155, 1015, 155 + erd_card_height([("PK", "sekolah_id"), ("UQ", "npsn?"), ("", "nama_sekolah"), ("", "status / kluster"), ("IDX", "aktif")])),
        "program": (1045, 155, 1555, 155 + erd_card_height([("PK", "id"), ("", "kategori / kelas / metode"), ("", "tipe / jenjang / tingkat"), ("UQ", "nama (kombinasi)"), ("", "biaya / diskon"), ("IDX", "aktif")])),
        "leads": (230, 560, 700, 560 + erd_card_height([("PK", "lead_id"), ("FK", "cabang_id"), ("FK", "sekolah_id?"), ("FK", "minat_program_id?"), ("FK", "inputter / PIC?"), ("IDX", "no_wa_normalized?"), ("IDX", "status / sumber / kanal"), ("IDX", "tanggal_masuk")])),
        "followup": (750, 560, 1120, 560 + erd_card_height([("PK", "id"), ("FK", "lead_id"), ("FK", "user_id"), ("IDX", "tanggal_followup"), ("", "metode / status_kontak"), ("IDX", "berikutnya?"), ("", "catatan?")])),
        "pendaftaran": (750, 1010, 1120, 1010 + erd_card_height([("PK", "id"), ("UQ", "lead_id"), ("FK", "program_id"), ("FK", "cabang_id"), ("IDX", "tanggal_daftar / status"), ("", "kelas_angkatan?")])),
        "pembayaran": (1190, 1010, 1555, 1010 + erd_card_height([("PK", "id"), ("FK", "pendaftaran_id"), ("IDX", "tanggal"), ("", "nominal / metode"), ("IDX", "status"), ("", "referensi?")])),
    }
    relation(draw, [(330, 270), (360, 270)], "1 : N", label_pos=(345, 270))
    relation(draw, [(185, 365), (185, 500), (260, 500), (260, 670), (230, 670)], "1 : N", label_pos=(205, 530))
    relation(draw, [(520, 388), (520, 560)], "1 : N", label_pos=(520, 480))
    relation(draw, [(860, 388), (860, 500), (620, 500), (620, 560)], "1 : N", label_pos=(790, 500))
    relation(draw, [(1280, 388), (1280, 505), (660, 505), (660, 560)], "1 : N", label_pos=(980, 505))
    relation(draw, [(700, 700), (750, 700)], "1 : N", label_pos=(725, 700))
    relation(draw, [(680, 305), (720, 305), (720, 630), (750, 630)], "1 : N", label_pos=(720, 445))
    relation(draw, [(700, 850), (725, 850), (725, 1125), (750, 1125)], "1 : 0..1", label_pos=(725, 990))
    relation(draw, [(1280, 388), (1320, 388), (1320, 970), (1080, 970), (1080, 1010)], "1 : N", label_pos=(1320, 680))
    relation(draw, [(1120, 1130), (1190, 1130)], "1 : N", label_pos=(1155, 1130))
    data = {
        "cabang": ("cabang", [("PK", "cabang_id"), ("UQ", "kode_cabang"), ("", "nama"), ("IDX", "aktif")], "#1476B8"),
        "users": ("users", [("PK", "id"), ("FK", "cabang_id?"), ("UQ", "email"), ("", "role"), ("IDX", "aktif")], "#1476B8"),
        "sekolah": ("sekolah", [("PK", "sekolah_id"), ("UQ", "npsn?"), ("", "nama_sekolah"), ("", "status / kluster"), ("IDX", "aktif")], "#1476B8"),
        "program": ("program_bimbel", [("PK", "id"), ("", "kategori / kelas / metode"), ("", "tipe / jenjang / tingkat"), ("UQ", "nama (kombinasi)"), ("", "biaya / diskon"), ("IDX", "aktif")], "#1476B8"),
        "leads": ("leads", [("PK", "lead_id"), ("FK", "cabang_id"), ("FK", "sekolah_id?"), ("FK", "minat_program_id?"), ("FK", "inputter / PIC?"), ("IDX", "no_wa_normalized?"), ("IDX", "status / sumber / kanal"), ("IDX", "tanggal_masuk")], "#10B981"),
        "followup": ("followup_riwayat", [("PK", "id"), ("FK", "lead_id"), ("FK", "user_id"), ("IDX", "tanggal_followup"), ("", "metode / status_kontak"), ("IDX", "berikutnya?"), ("", "catatan?")], "#10B981"),
        "pendaftaran": ("pendaftaran", [("PK", "id"), ("UQ", "lead_id"), ("FK", "program_id"), ("FK", "cabang_id"), ("IDX", "tanggal_daftar / status"), ("", "kelas_angkatan?")], "#F59E0B"),
        "pembayaran": ("pembayaran", [("PK", "id"), ("FK", "pendaftaran_id"), ("IDX", "tanggal"), ("", "nominal / metode"), ("IDX", "status"), ("", "referensi?")], "#F59E0B"),
    }
    for key, xy in cards.items():
        title, fields, color = data[key]
        erd_card(draw, xy, title, fields, color)
    image.save(ERD_OPERATIONAL)


def make_erd_performance() -> None:
    image = Image.new("RGB", (1600, 1180), "white")
    draw = ImageDraw.Draw(image)
    diagram_title(draw, "ERD Kinerja, Pengaturan, dan Audit", "Relasi kinerja dipisahkan secara visual dari tabel kontrol sistem.", 1600)
    definitions = {
        "cabang": ((45, 180, 330, 180 + erd_card_height([("PK", "cabang_id"), ("UQ", "kode_cabang"), ("", "nama"), ("IDX", "aktif")])), "cabang", [("PK", "cabang_id"), ("UQ", "kode_cabang"), ("", "nama"), ("IDX", "aktif")], "#1476B8"),
        "target_cabang": ((400, 180, 775, 180 + erd_card_height([("PK", "id"), ("FK", "cabang_id"), ("UQ", "tahun + bulan"), ("", "target_omzet"), ("", "target_siswa"), ("FK", "created_by?")])), "target_kinerja_cabang", [("PK", "id"), ("FK", "cabang_id"), ("UQ", "tahun + bulan"), ("", "target_omzet"), ("", "target_siswa"), ("FK", "created_by?")], "#10B981"),
        "target_user": ((845, 180, 1200, 180 + erd_card_height([("PK", "id"), ("FK", "target_kinerja_cabang_id"), ("FK", "user_id"), ("UQ", "target + user"), ("", "target_omzet"), ("", "target_siswa")])), "target_kinerja_user", [("PK", "id"), ("FK", "target_kinerja_cabang_id"), ("FK", "user_id"), ("UQ", "target + user"), ("", "target_omzet"), ("", "target_siswa")], "#10B981"),
        "users": ((1270, 180, 1555, 180 + erd_card_height([("PK", "id"), ("FK", "cabang_id?"), ("UQ", "email"), ("", "role"), ("IDX", "aktif")])), "users", [("PK", "id"), ("FK", "cabang_id?"), ("UQ", "email"), ("", "role"), ("IDX", "aktif")], "#1476B8"),
        "settings": ((210, 700, 570, 700 + erd_card_height([("PK", "key"), ("", "value? / type"), ("IDX", "group"), ("FK", "updated_by?")])), "crm_settings", [("PK", "key"), ("", "value? / type"), ("IDX", "group"), ("FK", "updated_by?")], "#A855F7"),
        "audit": ((640, 700, 1045, 700 + erd_card_height([("PK", "id"), ("FK", "user_id? / cabang_id?"), ("FK", "subject_type + subject_id?"), ("IDX", "event / created_at"), ("", "description / properties")])), "activity_logs", [("PK", "id"), ("FK", "user_id? / cabang_id?"), ("FK", "subject_type + subject_id?"), ("IDX", "event / created_at"), ("", "description / properties")], "#A855F7"),
        "notif": ((1115, 700, 1555, 700 + erd_card_height([("PK", "id (UUID)"), ("FK", "notifiable_type + id"), ("", "type / data"), ("IDX", "read_at? / created_at")])), "notifications", [("PK", "id (UUID)"), ("FK", "notifiable_type + id"), ("", "type / data"), ("IDX", "read_at? / created_at")], "#A855F7"),
    }
    relation(draw, [(330, 290), (400, 290)], "1 : N")
    relation(draw, [(775, 300), (845, 300)], "1 : N")
    relation(draw, [(1270, 330), (1235, 330), (1235, 410), (1200, 410)], "1 : N")
    relation(draw, [(1410, 413), (1410, 620), (390, 620), (390, 700)], "1 : N")
    relation(draw, [(1410, 413), (1410, 650), (840, 650), (840, 700)], "1 : N")
    relation(draw, [(190, 375), (190, 660), (720, 660), (720, 700)], "1 : N")
    relation(draw, [(1410, 413), (1410, 700)], "1 : N poly", dashed=True)
    for xy, title, fields, color in definitions.values():
        erd_card(draw, xy, title, fields, color)
    draw.rounded_rectangle((45, 1095, 1555, 1150), radius=15, fill="#F8FAFC", outline="#CBD5E1", width=2)
    draw.text((70, 1122), "Keterangan: garis putus-putus = relasi polymorphic; tanda ? = kolom nullable; UQ = unique constraint.", anchor="lm", font=font("arial.ttf", 21), fill="#475569")
    image.save(ERD_PERFORMANCE)


def make_erd_whatsapp() -> None:
    image = Image.new("RGB", (1600, 1360), "white")
    draw = ImageDraw.Draw(image)
    diagram_title(draw, "ERD Integrasi WhatsApp / WAHA", "Akun pengguna, percakapan, pesan, webhook, dan pencocokan lead.", 1600)
    definitions = {
        "users": ((45, 165, 350, 165 + erd_card_height([("PK", "id"), ("FK", "cabang_id?"), ("UQ", "email"), ("IDX", "aktif")])), "users", [("PK", "id"), ("FK", "cabang_id?"), ("UQ", "email"), ("IDX", "aktif")], "#1476B8"),
        "cabang": ((395, 165, 700, 165 + erd_card_height([("PK", "cabang_id"), ("UQ", "kode_cabang"), ("", "nama"), ("IDX", "aktif")])), "cabang", [("PK", "cabang_id"), ("UQ", "kode_cabang"), ("", "nama"), ("IDX", "aktif")], "#1476B8"),
        "leads": ((745, 165, 1090, 165 + erd_card_height([("PK", "lead_id"), ("FK", "cabang_id"), ("IDX", "no_wa_normalized?"), ("IDX", "status / kanal"), ("IDX", "WA first / last seen?")])), "leads", [("PK", "lead_id"), ("FK", "cabang_id"), ("IDX", "no_wa_normalized?"), ("IDX", "status / kanal"), ("IDX", "WA first / last seen?")], "#10B981"),
        "account": ((160, 560, 570, 560 + erd_card_height([("PK", "id"), ("UQ", "user_id / phone_e164"), ("FK", "cabang_id"), ("UQ", "session_name"), ("IDX", "status / aktif"), ("", "metadata? / last_error?")])), "user_whatsapp_accounts", [("PK", "id"), ("UQ", "user_id / phone_e164"), ("FK", "cabang_id"), ("UQ", "session_name"), ("IDX", "status / aktif"), ("", "metadata? / last_error?")], "#A855F7"),
        "conversation": ((650, 560, 1055, 560 + erd_card_height([("PK", "id"), ("FK", "account_id"), ("FK", "lead_id?"), ("UQ", "account + chat_id"), ("IDX", "contact_phone / match_status"), ("", "unread_count / last_message?")])), "whatsapp_conversations", [("PK", "id"), ("FK", "account_id"), ("FK", "lead_id?"), ("UQ", "account + chat_id"), ("IDX", "contact_phone / match_status"), ("", "unread_count / last_message?")], "#A855F7"),
        "message": ((1135, 560, 1555, 560 + erd_card_height([("PK", "id"), ("FK", "conversation_id / account_id"), ("FK", "lead_id?"), ("UQ", "account + external_message_id"), ("IDX", "direction / sent_at"), ("", "type / body / media?"), ("IDX", "processing_status")])), "whatsapp_messages", [("PK", "id"), ("FK", "conversation_id / account_id"), ("FK", "lead_id?"), ("UQ", "account + external_message_id"), ("IDX", "direction / sent_at"), ("", "type / body / media?"), ("IDX", "processing_status")], "#A855F7"),
        "webhook": ((565, 1000, 1035, 1000 + erd_card_height([("PK", "id"), ("UQ", "payload_hash"), ("IDX", "event_name / session_name"), ("IDX", "status / processed_at?"), ("", "payload / attempts / error?")])), "waha_webhook_events", [("PK", "id"), ("UQ", "payload_hash"), ("IDX", "event_name / session_name"), ("IDX", "status / processed_at?"), ("", "payload / attempts / error?")], "#F59E0B"),
    }
    relation(draw, [(200, 360), (200, 500), (300, 500), (300, 560)], "1 : 0..1", label_pos=(200, 430))
    relation(draw, [(550, 360), (550, 500), (470, 500), (470, 560)], "1 : N", label_pos=(550, 430))
    relation(draw, [(570, 710), (650, 710)], "1 : N", label_pos=(610, 710))
    relation(draw, [(915, 398), (915, 560)], "1 : N", label_pos=(915, 480))
    relation(draw, [(1055, 710), (1135, 710)], "1 : N", label_pos=(1095, 710))
    relation(draw, [(915, 398), (1095, 398), (1095, 785), (1135, 785)], "1 : N", label_pos=(1095, 590))
    relation(draw, [(365, 820), (365, 940), (1270, 940), (1270, 840)], "1 : N", label_pos=(800, 940))
    relation(draw, [(800, 1000), (800, 950), (850, 950), (850, 850)], "event", dashed=True, label_pos=(850, 900))
    relation(draw, [(900, 1000), (900, 930), (1390, 930), (1390, 870)], "event", dashed=True, label_pos=(1150, 930))
    for xy, title, fields, color in definitions.values():
        erd_card(draw, xy, title, fields, color)
    image.save(ERD_WHATSAPP)


def make_database_layout() -> None:
    image = Image.new("RGB", (1600, 1120), "white")
    draw = ImageDraw.Draw(image)
    diagram_title(draw, "Susunan 27 Tabel Database per Domain", "Urutan membaca: identitas dan master -> proses CRM -> transaksi/kinerja -> integrasi dan kontrol.", 1600)
    groups = [
        ("01  IDENTITAS & MASTER", ["cabang", "users", "sekolah", "program_bimbel", "crm_settings"], "#1476B8", (50, 165, 500, 480)),
        ("02  CRM & TRANSAKSI", ["leads", "followup_riwayat", "pendaftaran", "pembayaran"], "#10B981", (575, 165, 1025, 480)),
        ("03  KINERJA", ["target_kinerja_cabang", "target_kinerja_user"], "#F59E0B", (1100, 165, 1550, 480)),
        ("04  WHATSAPP / WAHA", ["user_whatsapp_accounts", "whatsapp_conversations", "whatsapp_messages", "waha_webhook_events"], "#A855F7", (50, 555, 500, 1020)),
        ("05  KONTROL & I/O", ["notifications", "activity_logs", "imports", "failed_import_rows", "exports"], "#F43F5E", (575, 555, 1025, 1020)),
        ("06  FRAMEWORK & QUEUE", ["sessions", "password_reset_tokens", "cache", "cache_locks", "jobs", "job_batches", "failed_jobs"], "#64748B", (1100, 555, 1550, 1020)),
    ]
    for title, tables, color, (x1, y1, x2, y2) in groups:
        draw.rounded_rectangle((x1, y1, x2, y2), radius=22, fill="#F8FAFC", outline=color, width=4)
        draw.rounded_rectangle((x1, y1, x2, y1 + 64), radius=22, fill=color, outline=color, width=4)
        draw.rectangle((x1, y1 + 43, x2, y1 + 64), fill=color)
        draw.text((x1 + 22, y1 + 32), title, anchor="lm", font=font("arialbd.ttf", 24), fill="white")
        top = y1 + 84
        for index, table_name in enumerate(tables):
            row_top = top + index * 49
            draw.rounded_rectangle((x1 + 20, row_top, x2 - 20, row_top + 38), radius=10, fill="white", outline="#D7E2EC", width=2)
            draw.text((x1 + 38, row_top + 19), f"{index + 1:02}.  {table_name}", anchor="lm", font=font("arial.ttf", 22), fill="#1E293B")
    image.save(DATABASE_LAYOUT)


def add_numbered(document: Document, items: list[str]) -> None:
    """Render urutan sebagai tabel langkah agar restart konsisten di Word."""
    table = document.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    b.set_repeat_table_header(table.rows[0])
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
    for index, header in enumerate(("Langkah", "Uraian")):
        cell = table.rows[0].cells[index]
        b.set_cell_shading(cell, b.PRIMARY)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.style = document.styles["Table Text"]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(header)
        b.set_run_font(run, 8.5, True, b.WHITE)
    for index, item in enumerate(items, start=1):
        row = table.add_row()
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        if index % 2 == 0:
            for cell in row.cells:
                b.set_cell_shading(cell, b.LIGHT)
        for column_index, value in enumerate((str(index), item)):
            cell = row.cells[column_index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            paragraph = cell.paragraphs[0]
            paragraph.style = document.styles["Table Text"]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if column_index == 0 else WD_ALIGN_PARAGRAPH.LEFT
            run = paragraph.add_run(value)
            b.set_run_font(run, 8.5, column_index == 0, b.INK)
    b.set_table_geometry(table, [0.72, 5.78])
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def add_figure(document: Document, filename: str, caption: str, width: float = 6.2) -> None:
    b.add_figure(document, ASSETS / filename, caption, width)


def table(document, caption, headers, rows, widths):
    b.add_table(document, headers, rows, widths, caption)


def build_document() -> None:
    make_cover()
    make_lifecycle()
    make_use_case()
    make_activity()
    make_erd_operational()
    make_erd_performance()
    make_erd_whatsapp()
    make_database_layout()
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    b.configure_styles(document)
    section.header.paragraphs[0].text = "SIVMI CRM  |  BAB V: ARSITEKTUR DATA"
    for run in section.header.paragraphs[0].runs:
        b.set_run_font(run, 8.5, True, b.PRIMARY)
    b.add_page_number(section.footer.paragraphs[0])

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    picture = p.add_run().add_picture(str(COVER), width=Inches(6.1))
    picture._inline.docPr.set("descr", "Sampul BAB V Arsitektur Data SIVMI CRM")
    document.add_page_break()

    b.add_heading(document, "KENDALI DOKUMEN", 1)
    table(document, "Tabel 0-1. Kendali dan baseline dokumen", ["Atribut", "Nilai"], [
        ["Judul", "BAB V: Arsitektur Data SIVMI CRM"],
        ["Versi", "1.1"],
        ["Tanggal baseline", "24 Agustus 2026"],
        ["Status", "Dokumen kerja terpisah - siap validasi manajemen"],
        ["Ruang lingkup", "RBAC, alur data, use case, activity, DFD, ERD, kardinalitas, siklus lead, katalog tabel, keamanan data"],
        ["Basis implementasi", "Laravel 12, Eloquent ORM, MySQL 8, Redis, Filament 5, Livewire 4, full Docker, WAHA 2026.8.1 CORE/GOWS"],
        ["Objek verifikasi", "27 tabel aplikasi/framework, 16 model bisnis/kontrol, 11 enum domain, migration sampai 2026-08-24"],
        ["Sumber", "Codebase as-built, migration, model, enum, seeder, test, diagram gabungan v1.4, dan dokumentasi induk v1.6"],
    ], [1.55, 4.95])
    b.add_callout(document, "Kedudukan dokumen", "BAB V ini berdiri sendiri dan menjelaskan arsitektur data as-built. Apabila terdapat perbedaan dengan diagram lama, migration dan model Eloquent terbaru menjadi sumber kebenaran teknis. Kebijakan organisasi yang belum tersedia ditandai sebagai keputusan yang masih diperlukan.")
    b.add_heading(document, "DAFTAR ISI", 1)
    b.add_toc(document)
    document.add_page_break()

    b.add_heading(document, "RINGKASAN EKSEKUTIF", 1)
    b.add_body(document, "Arsitektur data SIVMI CRM memakai model relasional terpusat dengan cabang sebagai batas kepemilikan data operasional. Data master, lead, aktivitas follow-up, pendaftaran, pembayaran, target, komunikasi WhatsApp, notifikasi, dan audit dihubungkan melalui primary key, foreign key, indeks, unique constraint, serta scope Eloquent berdasarkan role dan cabang.")
    b.add_body(document, "Data lead dapat masuk melalui kanal Manual, CSV, atau WhatsApp. Sumber pemasaran disimpan terpisah sebagai Iklan, Instagram, Kunjungan, Saudara, Facebook, atau Tiktok. Nomor WhatsApp asli dipertahankan untuk tampilan dan nomor ternormalisasi dipakai untuk pencocokan. Integrasi WAHA menyimpan akun per pengguna, conversation, message, serta event webhook idempoten sehingga pesan dapat dilacak dan lead ganda dapat ditandai tanpa penggabungan lintas cabang secara diam-diam.")
    b.add_callout(document, "Prinsip data utama", "Satu sumber kebenaran, cabang sebagai data boundary, histori tidak ditimpa, transaksi menggunakan FK, payload sensitif dibatasi, dan perubahan penting menghasilkan audit serta notifikasi.", b.LIGHT_GREEN)
    table(document, "Tabel 0-2. Ringkasan domain data", ["Domain", "Entitas utama", "Kontrol utama"], [
        ["Identitas & akses", "users, cabang", "Role enum, aktif, branch scope"],
        ["Master", "sekolah, program_bimbel, crm_settings", "NPSN opsional unik, enum Grade A-E, nama program otomatis"],
        ["Pipeline CRM", "leads, followup_riwayat", "Status, sumber, kanal, PIC, histori"],
        ["Konversi & transaksi", "pendaftaran, pembayaran", "Lead unik per pendaftaran, FK program/cabang, status verifikasi"],
        ["Kinerja", "target_kinerja_cabang, target_kinerja_user", "Unik cabang-periode dan alokasi user"],
        ["WhatsApp", "user_whatsapp_accounts, conversations, messages, webhook_events", "Session unik, event idempoten, pencocokan lead"],
        ["Kontrol", "notifications, activity_logs, imports/exports", "Unread, audit, hasil baris, kegagalan tervalidasi"],
    ], [1.35, 2.7, 2.45])

    b.add_heading(document, "BAB V: ARSITEKTUR DATA", 1)
    b.add_body(document, "Bab ini mendefinisikan struktur, aliran, kepemilikan, relasi, kualitas, keamanan, dan siklus hidup data SIVMI CRM. Fokusnya adalah bagaimana data dibuat, divalidasi, disimpan, dibaca, diubah, ditelusuri, dan diproteksi pada baseline aplikasi terbaru.")

    b.add_heading(document, "5.1 RBAC (Role Based Access Control)", 2)
    b.add_heading(document, "5.1.1 Model Role dan Cakupan Cabang", 3)
    b.add_body(document, "Role disimpan pada kolom users.role dan dicast menjadi UserRole. Superadmin dan Direksi memiliki akses cabang global; Admin Cabang dan Staff dibatasi pada cabang pengguna. Direksi bersifat observasional karena canManageOperationalData bernilai false.")
    table(document, "Tabel 5-1. Definisi role as-built", ["Role", "Scope data", "Kewenangan utama", "Batas"], [
        ["Superadmin", "Seluruh cabang", "Konfigurasi, master, pengguna, integrasi, audit", "Keputusan bisnis tetap perlu mandat organisasi"],
        ["Direksi", "Seluruh cabang", "Dashboard, KPI, laporan dan monitoring", "Tidak mengelola data operasional"],
        ["Admin Cabang", "Cabang sendiri", "Operasi cabang, distribusi PIC, target, transaksi", "Tidak melihat/mengubah cabang lain"],
        ["Staff", "Cabang dan data relevan miliknya", "Lead, follow-up, chat, pendaftaran, pencatatan pembayaran", "Tidak mengelola konfigurasi dan pengguna"],
    ], [1.05, 1.35, 2.45, 1.65])
    b.add_heading(document, "5.1.2 Kepemilikan dan Data Boundary", 3)
    table(document, "Tabel 5-2. Aturan kepemilikan data", ["Objek", "Kunci boundary", "Staff", "Admin Cabang", "Global"], [
        ["Lead", "leads.cabang_id + PIC/inputter", "PIC atau inputter", "Semua lead cabang", "Semua cabang"],
        ["Pendaftaran", "pendaftaran.cabang_id + lead", "Lead visibleTo", "Semua cabang sendiri", "Semua cabang"],
        ["Pembayaran", "melalui pendaftaran", "Pendaftaran visibleTo", "Semua cabang sendiri", "Semua cabang"],
        ["Target", "target_kinerja_cabang.cabang_id", "Alokasi user sendiri", "Target cabang", "Semua cabang"],
        ["Akun WA", "account.cabang_id + user_id", "Akun sendiri", "Akun cabang", "Semua akun"],
        ["Conversation", "account.cabang_id + user_id", "Session sendiri", "Session cabang", "Semua conversation"],
        ["Audit", "cabang_id + user_id", "Aktivitas sendiri", "Aktivitas cabang", "Semua log"],
    ], [1.2, 1.55, 1.2, 1.25, 1.3])
    b.add_heading(document, "5.1.3 Prinsip Least Privilege", 3)
    b.add_bullets(document, [
        "Akses panel hanya diberikan kepada pengguna aktif dengan role enum yang valid.",
        "Filter cabang tidak dipercaya sebagai satu-satunya kontrol; query resource menerapkan scope server-side.",
        "Admin Cabang hanya dapat membuat atau mengubah Staff pada cabangnya.",
        "Staff yang mencatat pembayaran dipaksa menggunakan status Menunggu sampai diverifikasi sesuai alur operasional.",
        "Akun WhatsApp hanya dapat dikaitkan kepada Admin Cabang atau Staff yang aktif dan memiliki cabang.",
    ])

    b.add_heading(document, "5.2 Matriks RBAC", 2)
    b.add_heading(document, "5.2.1 Matriks Modul dan Aksi", 3)
    table(document, "Tabel 5-3. Matriks akses modul", ["Modul", "Superadmin", "Direksi", "Admin Cabang", "Staff"], [
        ["Dashboard/Laporan", "R seluruh cabang", "R seluruh cabang", "R cabang", "R data relevan"],
        ["Leads/Follow-up", "CRUD", "R", "CRUD cabang", "CRUD penugasan"],
        ["Pendaftaran", "CRUD", "R", "CRUD cabang", "CRU data relevan"],
        ["Pembayaran", "CRUD/verifikasi", "R", "CRUD/verifikasi cabang", "CR status Menunggu"],
        ["Sekolah/Program", "CRUD", "R", "R/operasional sesuai resource", "R dan tambah sekolah via lead"],
        ["Target Kinerja", "CRUD", "R", "CRUD cabang", "R alokasi sendiri"],
        ["Pengguna", "CRUD", "-", "CRUD Staff cabang", "-"],
        ["Koneksi WhatsApp", "CRUD", "R bila diizinkan", "CRUD cabang", "R/operasi akun sendiri"],
        ["Inbox WhatsApp", "R/operasi", "R", "R/operasi cabang", "R/operasi session sendiri"],
        ["Pengaturan/Audit", "CRUD/R", "R", "R terbatas cabang", "R aktivitas sendiri"],
    ], [1.45, 1.25, 1.05, 1.45, 1.3])
    b.add_callout(document, "Notasi", "C = create, R = read, U = update, D = delete. Matriks ini merekam baseline resource dan scope saat ini; keputusan menu tidak boleh dianggap sebagai pengganti authorization server-side.")
    b.add_heading(document, "5.2.2 Matriks Operasi Data", 3)
    table(document, "Tabel 5-4. Matriks kontrol operasi data", ["Operasi", "Validasi wajib", "Audit", "Pembatasan"], [
        ["Create", "FK, required field, enum/status, branch/PIC", "Actor, event, subject, properties", "Role dan cabang"],
        ["Read", "visibleTo/query scope", "Tidak semua read dicatat", "Role, cabang, PIC"],
        ["Update", "Record tetap dalam scope; business rule", "Before/after yang aman", "Field sensitif/role"],
        ["Delete", "Relasi dan kebijakan retensi", "Event penghapusan", "Belum ada soft delete"],
        ["Import", "Validasi per baris, mapping relasi", "Import + failed rows", "Scope user/cabang"],
        ["Export", "Query terfilter", "Job export dan pemilik", "Staff tergantung setting"],
        ["Webhook", "HMAC, payload hash, session", "Event/status/error", "Idempoten dan queue"],
    ], [1.0, 2.35, 1.45, 1.7])

    b.add_heading(document, "5.3 Alur Proses Sistem", 2)
    b.add_heading(document, "5.3.1 Pola Akuisisi hingga Persistensi", 3)
    add_numbered(document, [
        "Terima input dari form, CSV, atau webhook WAHA dan identifikasi pengguna/session sumber.",
        "Validasi field wajib, format, foreign key, role, cabang, PIC, status, serta enum domain.",
        "Normalisasi nomor WhatsApp menjadi digit E.164 tanpa tanda untuk pencocokan deterministik.",
        "Cari kandidat lead pada scope yang diperbolehkan dan tentukan matched, new lead, conflict, unresolved, atau ignored.",
        "Simpan perubahan dalam tabel utama dan histori tanpa mengganti riwayat dengan agregat tunggal.",
        "Buat event audit/notifikasi, perbarui unread/KPI, dan kirim pekerjaan berat melalui queue.",
        "Jika gagal, simpan validasi/error yang aman dan jalankan retry hanya untuk operasi idempoten.",
    ])
    table(document, "Tabel 5-5. Input, proses, output, dan kontrol", ["Tahap", "Input", "Output", "Kontrol"], [
        ["Akuisisi", "Form/CSV/pesan", "Payload kandidat", "Auth, HMAC, ukuran/format"],
        ["Validasi", "Payload kandidat", "Data valid/failed row", "Required, enum, FK, unique"],
        ["Resolusi", "Nomor/cabang/PIC", "Status match", "Normalisasi, branch scope"],
        ["Persistensi", "Data valid", "Record + relasi", "Transaction, constraint, index"],
        ["Publikasi", "Event perubahan", "Notifikasi/KPI/audit", "Queue, observer, idempotensi"],
    ], [1.0, 1.7, 1.75, 2.05])
    b.add_heading(document, "5.3.2 Diagram Konteks", 3)
    add_figure(document, "dfd-level-0.png", "Gambar 5-1. DFD Level 0 - konteks pertukaran data SIVMI CRM")

    b.add_heading(document, "5.4 Flowchart", 2)
    b.add_body(document, "Flowchart data memisahkan sinyal akuisisi, validasi, pemilihan cabang/PIC, engagement, konversi, transaksi, dan monitoring. Keputusan data utama adalah validitas identitas, hasil pencocokan nomor, kesiapan pendaftaran, serta status verifikasi pembayaran.")
    add_figure(document, "siklus-data-lead.png", "Gambar 5-2. Siklus data lead dari akuisisi hingga monitoring")
    b.add_heading(document, "5.4.1 Titik Keputusan dan Jalur Gagal", 3)
    table(document, "Tabel 5-6. Titik keputusan flowchart data", ["Keputusan", "Ya", "Tidak/gagal", "Bukti"], [
        ["Payload valid?", "Lanjut normalisasi", "Tolak/failed row", "Validation error"],
        ["Nomor dapat dinormalisasi?", "Cari kandidat", "Unresolved/manual review", "no_wa_normalized"],
        ["Satu kandidat valid?", "Matched", "NewLead/Conflict", "match_status"],
        ["Lead siap daftar?", "Buat pendaftaran", "Lanjut follow-up", "status + histori"],
        ["Pembayaran valid?", "Lunas/Terverifikasi", "Menunggu/Gagal", "referensi + audit"],
        ["Webhook pernah diproses?", "Abaikan aman", "Queue processing", "payload_hash"],
    ], [1.55, 1.45, 1.8, 1.7])

    b.add_heading(document, "5.5 Use Case", 2)
    add_figure(document, "use-case-data.png", "Gambar 5-3. Use case pengelolaan data berdasarkan aktor")
    table(document, "Tabel 5-7. Katalog use case data", ["ID", "Use case", "Aktor", "Precondition", "Postcondition"], [
        ["UC-D01", "Kelola master dan role", "Superadmin", "Login aktif", "Master/akses tervalidasi"],
        ["UC-D02", "Input/impor lead", "Admin/Staff", "Cabang valid", "Lead dan hasil validasi tersimpan"],
        ["UC-D03", "Distribusi dan follow-up", "Admin/Staff", "Lead visible", "PIC/histori/reminder tersimpan"],
        ["UC-D04", "Konversi dan transaksi", "Admin/Staff", "Lead/program valid", "Pendaftaran/pembayaran terkait"],
        ["UC-D05", "Kelola session/chat", "Superadmin/Admin/Staff", "Akun WA aktif", "Conversation/message tersimpan"],
        ["UC-D06", "Monitor KPI dan audit", "Direksi/Superadmin/Admin", "Scope valid", "Insight terfilter"],
        ["UC-D07", "Terima webhook", "WAHA/System", "HMAC valid", "Event idempoten diproses"],
    ], [0.7, 1.55, 1.35, 1.45, 1.45])

    b.add_heading(document, "5.6 Activity Diagram", 2)
    add_figure(document, "activity-data.png", "Gambar 5-4. Activity diagram akuisisi, pencocokan, dan persistensi lead", 6.0)
    b.add_heading(document, "5.6.1 Transaksi dan Konsistensi", 3)
    b.add_bullets(document, [
        "Relasi wajib ditulis dengan FK yang valid; relasi opsional menggunakan nullable dan kebijakan nullOnDelete bila sesuai.",
        "Pembuatan lead otomatis dari WhatsApp tidak boleh menggabungkan konflik lintas cabang.",
        "Pendaftaran aktif memperbarui status lead menjadi Daftar melalui event model.",
        "Follow-up pertama pada lead Cold memperbarui status menjadi Warm tanpa menghapus histori.",
        "Perubahan konfigurasi massal crm_settings dilakukan dalam DB transaction lalu cache dihapus.",
    ])

    b.add_heading(document, "5.7 Data Flow Diagram (DFD)", 2)
    b.add_heading(document, "5.7.1 DFD Level 0", 3)
    b.add_body(document, "DFD Level 0 menempatkan SIVMI CRM sebagai satu proses yang bertukar data dengan calon siswa/orang tua, pengguna operasional, Superadmin, dan Direksi. Batas sistem mencakup input, transaksi, konfigurasi, tugas/notifikasi, dan keluaran KPI/laporan.")
    b.add_heading(document, "5.7.2 DFD Level 1 CRM", 3)
    add_figure(document, "dfd-level-1.png", "Gambar 5-5. DFD Level 1 - dekomposisi proses CRM")
    table(document, "Tabel 5-8. Data store DFD CRM", ["Store", "Cakupan tabel", "Produsen", "Konsumen"], [
        ["D1 Leads & Master", "cabang, users, sekolah, program, leads", "Form/import/webhook", "Lead, dashboard, laporan"],
        ["D2 Follow-up", "followup_riwayat", "Admin/Staff", "Reminder, KPI, audit"],
        ["D3 Pendaftaran", "pendaftaran", "Admin/Staff", "Pembayaran, siswa aktif"],
        ["D4 Pembayaran", "pembayaran", "Admin/Staff", "Omzet, laporan, target"],
        ["D5 Kontrol", "target, notifications, activity_logs", "System/manajemen", "Dashboard, audit, tindak lanjut"],
    ], [1.2, 2.4, 1.4, 1.5])
    b.add_heading(document, "5.7.3 DFD Integrasi WhatsApp", 3)
    add_figure(document, "dfd-whatsapp.png", "Gambar 5-6. DFD Level 1 - session, webhook, pesan, dan pencocokan lead")

    b.add_heading(document, "5.8 Entity Relationship Diagram (ERD)", 2)
    b.add_heading(document, "5.8.1 Peta Domain Entitas", 3)
    add_figure(document, "peta-entitas.png", "Gambar 5-7. Peta domain entitas inti SIVMI CRM")
    table(document, "Tabel 5-9. Domain dan aggregate root", ["Domain", "Aggregate/root", "Entitas turunan", "Kunci integrasi"], [
        ["Akses", "users/cabang", "account, audit", "user_id, cabang_id"],
        ["CRM", "leads", "followup, pendaftaran, conversation, message", "lead_id"],
        ["Program", "program_bimbel", "pendaftaran", "program_bimbel_id"],
        ["Transaksi", "pendaftaran", "pembayaran", "pendaftaran_id"],
        ["Kinerja", "target_kinerja_cabang", "target_kinerja_user", "target_kinerja_cabang_id"],
        ["WhatsApp", "user_whatsapp_accounts", "conversations, messages", "whatsapp_account_id"],
        ["Kontrol", "activity_logs/notifications", "polymorphic subject/notifiable", "type + id"],
    ], [1.0, 1.5, 2.55, 1.45])
    b.add_heading(document, "5.8.2 ERD Operasional", 3)
    b.add_body(document, "Notasi diagram: PK = primary key, FK = foreign key, UQ = unique constraint, IDX = indeks, dan tanda ? = kolom nullable. Atribut lengkap tetap mengikuti migration serta kamus data.")
    document.paragraphs[-1].paragraph_format.keep_with_next = True
    add_figure(document, "erd-operasional.png", "Gambar 5-8. ERD operasional lead-to-payment")
    b.add_heading(document, "5.8.3 ERD Kinerja dan Kontrol", 3)
    add_figure(document, "erd-kinerja-kontrol.png", "Gambar 5-9. ERD target kinerja, pengaturan, notifikasi, dan audit")
    b.add_heading(document, "5.8.4 ERD Integrasi WhatsApp", 3)
    add_figure(document, "erd-whatsapp.png", "Gambar 5-10. ERD akun WhatsApp, conversation, message, event, dan lead")

    b.add_heading(document, "5.9 Relasi dan Kardinalitas", 2)
    b.add_body(document, "Kardinalitas berikut berasal dari foreign key, unique constraint, nullable, dan relasi Eloquent. Notasi 1:N berarti satu parent dapat memiliki banyak child; 0..1 berarti relasi opsional tunggal.")
    relationships = [
        ["cabang", "users", "1:N", "users.cabang_id nullable", "nullOnDelete"],
        ["cabang", "leads", "1:N", "leads.cabang_id wajib", "Restrict/default"],
        ["cabang", "pendaftaran", "1:N", "pendaftaran.cabang_id wajib", "Restrict/default"],
        ["cabang", "target_kinerja_cabang", "1:N", "Unik cabang+tahun+bulan", "Restrict/default"],
        ["sekolah", "leads", "1:N", "lead boleh tanpa sekolah", "nullOnDelete"],
        ["program_bimbel", "leads", "1:N", "minat_program_id opsional", "nullOnDelete"],
        ["users(inputter)", "leads", "1:N", "diinput_oleh wajib", "Restrict/default"],
        ["users(PIC)", "leads", "1:N", "diserahkan_ke opsional", "nullOnDelete"],
        ["leads", "followup_riwayat", "1:N", "Histori per lead", "cascadeOnDelete"],
        ["users", "followup_riwayat", "1:N", "Pelaksana wajib", "Restrict/default"],
        ["leads", "pendaftaran", "1:0..1", "lead_id unique", "Restrict/default"],
        ["program_bimbel", "pendaftaran", "1:N", "Program wajib", "Restrict/default"],
        ["pendaftaran", "pembayaran", "1:N", "Cicilan/multi transaksi dimungkinkan", "cascadeOnDelete"],
        ["target_kinerja_cabang", "target_kinerja_user", "1:N", "Unik target+user", "cascadeOnDelete"],
        ["users", "target_kinerja_user", "1:N", "Alokasi user", "cascadeOnDelete"],
        ["users", "user_whatsapp_accounts", "1:0..1", "user_id unique", "cascadeOnDelete"],
        ["cabang", "user_whatsapp_accounts", "1:N", "Akun wajib cabang", "cascadeOnDelete"],
        ["user_whatsapp_accounts", "whatsapp_conversations", "1:N", "Unik account+chat_id", "cascadeOnDelete"],
        ["leads", "whatsapp_conversations", "1:N", "lead_id opsional", "nullOnDelete"],
        ["conversation", "whatsapp_messages", "1:N", "Pesan wajib conversation", "cascadeOnDelete"],
        ["account", "whatsapp_messages", "1:N", "Unik account+external_message_id", "cascadeOnDelete"],
        ["leads", "whatsapp_messages", "1:N", "lead_id opsional", "nullOnDelete"],
        ["users", "crm_settings", "1:N", "updated_by opsional", "nullOnDelete"],
        ["users/cabang", "activity_logs", "1:N", "Keduanya opsional", "nullOnDelete"],
        ["subject", "activity_logs", "1:N polymorphic", "subject_type+subject_id nullable", "Logical"],
        ["notifiable", "notifications", "1:N polymorphic", "notifiable_type+id", "Framework"],
        ["users", "imports/exports", "1:N", "Pemilik job wajib", "cascadeOnDelete"],
        ["imports", "failed_import_rows", "1:N", "Detail kegagalan", "cascadeOnDelete"],
    ]
    table(document, "Tabel 5-10. Matriks relasi dan kardinalitas", ["Parent", "Child", "Kardinalitas", "Constraint", "On delete"], relationships, [1.2, 1.55, 0.8, 2.05, 0.9])
    b.add_heading(document, "5.9.1 Risiko Kardinalitas", 3)
    b.add_bullets(document, [
        "Lead ke pendaftaran saat ini 1:0..1 karena lead_id diberi unique; perubahan menjadi multi-pendaftaran memerlukan migration dan revisi KPI.",
        "Pendaftaran ke pembayaran 1:N memungkinkan cicilan atau pembayaran bertahap, tetapi aturan rekonsiliasi dan refund belum menjadi modul penuh.",
        "Lead dapat terhubung ke banyak conversation/account; konflik nomor lintas cabang tidak boleh digabung otomatis.",
        "Cascade delete pada akun WA, conversation, message, pendaftaran, dan pembayaran perlu disejajarkan dengan kebijakan retensi sebelum produksi.",
    ])

    b.add_heading(document, "5.10 Siklus Data Lead", 2)
    b.add_heading(document, "5.10.1 Status Pipeline", 3)
    table(document, "Tabel 5-11. Status lead dan pemicu", ["Status", "Makna", "Pemicu umum", "Transisi utama"], [
        ["Cold", "Lead baru/belum dinilai", "Form/import/WA unik", "Warm, Hot, Tidak Berminat"],
        ["Warm", "Sudah dihubungi/berinteraksi", "Follow-up pertama dari Cold", "Hot, Daftar, Tidak Berminat"],
        ["Hot", "Minat/urgensi tinggi", "Hasil follow-up", "Daftar, Warm, Tidak Berminat"],
        ["Daftar", "Telah menjadi pendaftaran aktif", "Pendaftaran status Aktif", "Ditentukan kebijakan koreksi"],
        ["Tidak Berminat", "Lead tidak dilanjutkan", "Hasil follow-up", "Retensi/reaktivasi perlu kebijakan"],
    ], [1.0, 1.8, 1.8, 1.9])
    b.add_heading(document, "5.10.2 Kanal Masuk dan Sumber", 3)
    table(document, "Tabel 5-12. Kanal dan sumber lead", ["Dimensi", "Nilai", "Tujuan analitik", "Kolom"], [
        ["Kanal masuk", "Manual, CSV, WhatsApp", "Cara data masuk ke sistem", "kanal_masuk"],
        ["Sumber pemasaran", "Iklan, Instagram, Kunjungan, Saudara, Facebook, Tiktok", "Asal minat/kampanye", "sumber_lead"],
        ["Nomor tampilan", "Format input pengguna", "Komunikasi", "no_wa"],
        ["Nomor pencocokan", "Digit E.164 tanpa tanda", "Deduplikasi", "no_wa_normalized"],
        ["Jejak WhatsApp", "first_seen, last_seen", "Aktivitas komunikasi", "whatsapp_*_seen_at"],
    ], [1.2, 2.2, 1.7, 1.4])
    b.add_heading(document, "5.10.3 Pencocokan WhatsApp dan Deduplikasi", 3)
    table(document, "Tabel 5-13. Status pencocokan WhatsApp", ["Status", "Kondisi", "Tindakan", "Larangan"], [
        ["Matched", "Satu lead valid ditemukan", "Tautkan conversation/message", "Jangan ubah identitas tanpa audit"],
        ["NewLead", "Nomor belum ditemukan", "Buat/siapkan lead Cold", "Jangan membuat ganda pada event retry"],
        ["Conflict", "Lebih dari satu kandidat/lintas cabang", "Review Admin/Superadmin", "Dilarang merge otomatis"],
        ["Unresolved", "Nomor/identitas belum cukup", "Review manual", "Jangan menebak cabang/PIC"],
        ["Ignored", "Group/status/event non-lead", "Simpan atau abaikan sesuai event", "Jangan masuk funnel"],
    ], [1.0, 2.0, 1.85, 1.65])
    add_numbered(document, [
        "Normalisasi nomor dan validasi format 8-15 digit internasional.",
        "Prioritaskan kandidat pada account/cabang/PIC yang sah.",
        "Gunakan unique payload_hash dan external_message_id untuk mencegah proses ganda.",
        "Simpan hasil match dan jejak event; jangan menghapus percakapan saat link lead berubah.",
        "Eskalasi conflict dan unresolved kepada pemilik cabang atau Superadmin.",
    ])

    b.add_heading(document, "5.11 Tabel Database dan Kamus Data", 2)
    b.add_heading(document, "5.11.1 Susunan dan Katalog Tabel", 3)
    add_figure(document, "susunan-tabel-database.png", "Gambar 5-11. Susunan 27 tabel database berdasarkan domain dan urutan pemakaian")
    catalog = [
        ["01", "Identitas", "cabang", "-", "Master batas kepemilikan data"], ["02", "Identitas", "users", "cabang?", "Pengguna, role, dan status aktif"],
        ["03", "Master", "sekolah", "-", "Sekolah, NPSN opsional, Grade A-E"], ["04", "Master", "program_bimbel", "-", "Klasifikasi, nama otomatis, dan biaya"], ["05", "Master", "crm_settings", "users?", "Konfigurasi dinamis terkontrol"],
        ["06", "CRM", "leads", "cabang", "Aggregate calon siswa dan pipeline"], ["07", "CRM", "followup_riwayat", "leads", "Histori kontak dan reminder"], ["08", "Transaksi", "pendaftaran", "leads", "Konversi satu lead"], ["09", "Transaksi", "pembayaran", "pendaftaran", "Pembayaran/cicilan pendaftaran"],
        ["10", "Kinerja", "target_kinerja_cabang", "cabang", "Target per cabang dan periode"], ["11", "Kinerja", "target_kinerja_user", "target cabang", "Alokasi target per pengguna"],
        ["12", "WhatsApp", "user_whatsapp_accounts", "users, cabang", "Akun/session WAHA per user"], ["13", "WhatsApp", "whatsapp_conversations", "WA account", "Percakapan unik per account/chat"], ["14", "WhatsApp", "whatsapp_messages", "conversation", "Pesan inbound dan outbound"], ["15", "WhatsApp", "waha_webhook_events", "-", "Inbox event idempoten dan retry"],
        ["16", "Kontrol", "notifications", "notifiable", "Notifikasi database polymorphic"], ["17", "Kontrol", "activity_logs", "users/cabang?", "Audit aktivitas dan subject"],
        ["18", "I/O", "imports", "users", "Metadata job impor"], ["19", "I/O", "failed_import_rows", "imports", "Baris gagal dan pesan validasi"], ["20", "I/O", "exports", "users", "Metadata job ekspor"],
        ["21", "Framework", "sessions", "users?", "Session autentikasi"], ["22", "Framework", "password_reset_tokens", "email", "Reset kata sandi"], ["23", "Framework", "cache", "-", "Cache key/value"], ["24", "Framework", "cache_locks", "cache", "Lock cache"],
        ["25", "Queue", "jobs", "-", "Pekerjaan antrean"], ["26", "Queue", "job_batches", "jobs", "Batch pekerjaan"], ["27", "Queue", "failed_jobs", "jobs", "Pekerjaan gagal"],
    ]
    table(document, "Tabel 5-14. Katalog 27 tabel dalam urutan logis", ["No.", "Domain", "Tabel", "Parent utama", "Fungsi"], catalog, [0.45, 0.9, 1.85, 1.2, 2.1])

    dictionaries = [
        ("5.11.2 Identitas dan Master", "Tabel 5-15. Kamus data identitas dan master", [
            ["users", "id; email UQ", "cabang_id?", "name; password; role; aktif; timestamps", "Role cast"],
            ["cabang", "cabang_id; kode UQ", "-", "nama; alamat?; kota?; telepon?; email?; aktif", "BDG/JKS/JKP"],
            ["sekolah", "sekolah_id; npsn? UQ", "-", "nama; status; wilayah; kluster; aktif", "Grade A-E"],
            ["program_bimbel", "id; nama UQ", "-", "kategori; kelas; metode; tipe; jenjang; biaya; aktif", "Nama 3 unsur"],
            ["crm_settings", "key", "updated_by?", "value?; type; group IDX; timestamps", "Cached"],
        ]),
        ("5.11.3 Pipeline CRM", "Tabel 5-16. Kamus data lead dan follow-up", [
            ["leads", "lead_id", "cabang; sekolah?; program?; inputter; PIC?", "nama; nomor WA; status; sumber; kanal; tanggal; WA seen", "Aggregate CRM"],
            ["followup_riwayat", "id", "lead_id; user_id", "tanggal; metode; status kontak; berikutnya?; catatan?", "Histori append"],
        ]),
        ("5.11.4 Konversi dan Transaksi", "Tabel 5-17. Kamus data pendaftaran dan pembayaran", [
            ["pendaftaran", "id; lead_id UQ", "program_id; cabang_id", "tanggal; kelas_angkatan?; status; catatan?", "1 per lead"],
            ["pembayaran", "id", "pendaftaran_id", "tanggal; nominal; metode; status; referensi?; catatan?", "1:N/cicilan"],
        ]),
        ("5.11.5 Kinerja", "Tabel 5-18. Kamus data target", [
            ["target_kinerja_cabang", "id; UQ cabang+periode", "cabang_id; created_by?", "tahun; bulan; status; target omzet/siswa; catatan?", "Target induk"],
            ["target_kinerja_user", "id; UQ target+user", "target_cabang_id; user_id", "target omzet; target siswa; catatan?", "Alokasi user"],
        ]),
        ("5.11.6 Integrasi WhatsApp", "Tabel 5-19. Kamus data WhatsApp", [
            ["user_whatsapp_accounts", "id; user/phone/session UQ", "user_id; cabang_id", "WAHA id; status; aktif; connected/seen; error?; metadata?", "1 akun/user"],
            ["whatsapp_conversations", "id; UQ account+chat", "account_id; lead_id?", "contact; match_status; first/last message; unread", "Thread"],
            ["whatsapp_messages", "id; UQ account+external", "conversation; account; lead?", "direction; type; body/media; sent_at; processing; raw?", "Histori pesan"],
            ["waha_webhook_events", "id; payload_hash UQ", "-", "event; session; status; payload; attempts; error?; processed?", "Idempoten"],
        ]),
        ("5.11.7 Kontrol dan Pertukaran Data", "Tabel 5-20. Kamus data kontrol dan I/O", [
            ["notifications", "id UUID", "notifiable type+id", "type; data; read_at?; timestamps", "Polymorphic"],
            ["activity_logs", "id", "user?; cabang?; subject?", "event; description; properties; IP; user agent; created", "Audit"],
            ["imports", "id", "user_id", "file; importer; processed/total/success; completed?", "Job import"],
            ["failed_import_rows", "id", "import_id", "data JSON; validation_error?; timestamps", "Cascade"],
            ["exports", "id", "user_id", "disk; file; exporter; processed/total/success; completed?", "Job export"],
        ]),
    ]
    for heading, caption, rows in dictionaries:
        b.add_heading(document, heading, 3)
        table(document, caption, ["Tabel", "PK / Unique", "Foreign key", "Atribut utama", "Catatan"], rows, [1.25, 1.1, 1.25, 2.15, 0.75])

    b.add_heading(document, "5.11.8 Indeks, Unique, dan Integritas", 3)
    table(document, "Tabel 5-21. Constraint penting", ["Objek", "Constraint/index", "Tujuan"], [
        ["cabang", "kode_cabang unique", "Identitas cabang stabil"],
        ["sekolah", "npsn unique nullable", "Referensi nasional tanpa mewajibkan isi"],
        ["leads", "status, tanggal, sumber, kanal, WA last seen; cabang+tanggal", "Filter dashboard dan pencarian"],
        ["pendaftaran", "lead_id unique; cabang+tanggal", "Mencegah pendaftaran ganda pada baseline"],
        ["target cabang", "cabang+tahun+bulan unique", "Satu target per periode"],
        ["target user", "target_cabang+user unique", "Satu alokasi per user"],
        ["WA account", "user, phone, session unique", "Kepemilikan session deterministik"],
        ["Conversation", "account+chat unique; account+match index", "Satu thread per account"],
        ["Message", "account+external_id unique; lead+sent index", "Deduplikasi dan histori"],
        ["Webhook", "payload_hash unique", "Idempotensi event"],
    ], [1.4, 2.8, 2.3])

    b.add_heading(document, "5.12 Keamanan Hak Akses Data", 2)
    b.add_heading(document, "5.12.1 Klasifikasi Data", 3)
    table(document, "Tabel 5-22. Klasifikasi dan perlindungan data", ["Kelas", "Contoh", "Kontrol minimum", "Eksposur"], [
        ["Publik/internal ringan", "Nama program, cabang aktif", "Validasi dan integritas", "Panel sesuai role"],
        ["Data pribadi", "Nama, email, nomor WA, sekolah", "RBAC, scope, masking ekspor", "Hanya kebutuhan kerja"],
        ["Komunikasi", "Body chat, contact name, media", "Scope session/cabang, retensi, audit", "Inbox pemilik"],
        ["Transaksi", "Nominal, metode, referensi", "Verifikasi, audit, pembatasan update", "Operasional/management"],
        ["Secret", "APP_KEY, API key, HMAC secret", "Environment/secret manager", "Tidak boleh masuk DB/log"],
        ["Forensik", "raw payload, IP, user agent, error", "Redaksi, akses terbatas, retensi", "Superadmin terbatas"],
    ], [1.2, 1.85, 2.25, 1.2])
    b.add_heading(document, "5.12.2 Kontrol Teknis As-Built", 3)
    add_numbered(document, [
        "Autentikasi Filament dan status user aktif sebelum panel dapat diakses.",
        "UserRole enum untuk mencegah role bebas dan menentukan akses global/operasional.",
        "Scope visibleTo pada Lead, Pendaftaran, Pembayaran, Target, WhatsApp, dan ActivityLog.",
        "Foreign key, unique constraint, index, type cast, serta observer/model event menjaga konsistensi.",
        "Webhook WAHA diverifikasi HMAC, disimpan dengan payload_hash unik, dan diproses melalui queue.",
        "Secret WAHA/API tidak dicatat pada activity log atau metadata terbuka; log harus meredaksi payload sensitif.",
        "Database Docker persisten pada volume dan terpisah dari database Laragon; target VPS memerlukan backup terenkripsi.",
    ])
    b.add_heading(document, "5.12.3 Ancaman dan Mitigasi", 3)
    table(document, "Tabel 5-23. Risiko data dan mitigasi", ["Risiko", "Dampak", "Kontrol saat ini", "Target penguatan"], [
        ["Akses lintas cabang", "Kebocoran data", "visibleTo + resource query", "Policy test per resource"],
        ["Duplikasi lead", "Funnel/layanan ganda", "Normalisasi + match status", "Unique business key terukur"],
        ["Replay webhook", "Pesan/event ganda", "HMAC + payload_hash", "Timestamp window dan rate limit"],
        ["Raw payload sensitif", "Paparan data", "Akses aplikasi terbatas", "Redaksi/enkripsi/retensi"],
        ["Cascade delete", "Histori hilang", "FK eksplisit", "Soft delete/arsip sesuai kebijakan"],
        ["Export berlebih", "Eksfiltrasi", "Scope dan setting staff", "Watermark, audit download, DLP"],
        ["Backup tidak diuji", "Kehilangan data", "Volume persisten lokal", "RPO/RTO dan restore drill"],
    ], [1.35, 1.35, 1.75, 2.05])
    b.add_heading(document, "5.12.4 Retensi, Backup, dan Pemulihan", 3)
    table(document, "Tabel 5-24. Kebijakan retensi yang perlu disahkan", ["Data", "Baseline", "Rekomendasi keputusan", "Pemilik"], [
        ["Lead/follow-up", "Tidak ada soft delete/retensi formal", "Retensi aktif, arsip, dan reaktivasi", "Direksi/Operasional"],
        ["Chat/media/raw payload", "Tersimpan sesuai event", "Durasi, redaksi, akses dan penghapusan", "Direksi/Superadmin"],
        ["Audit log", "Created_at tanpa expiry", "Retensi minimum dan legal hold", "Superadmin"],
        ["Import/export file", "Disimpan pada disk job", "TTL file dan akses download", "Superadmin"],
        ["Backup database", "Volume lokal", "Backup terenkripsi, offsite, restore test", "Infrastruktur"],
        ["Secret/config", "Environment Docker", "Secret manager dan rotasi", "Infrastruktur"],
    ], [1.45, 1.65, 2.35, 1.05])

    b.add_heading(document, "5.13 Temuan, Gap, dan Keputusan Data", 2)
    table(document, "Tabel 5-25. Keputusan yang masih diperlukan", ["Keputusan", "Kondisi saat ini", "Risiko bila ditunda", "Pemilik rekomendasi"], [
        ["Definisi Grade D dan E", "Enum tersedia; deskripsi null", "Klasifikasi tidak seragam", "Direksi/Operasional"],
        ["Retensi lead/chat/media/audit", "Belum formal", "Kapasitas dan kepatuhan", "Direksi/Superadmin"],
        ["Soft delete/arsip", "Belum diterapkan", "Cascade dapat menghapus histori", "Superadmin/Operasional"],
        ["Multi-pendaftaran per lead", "lead_id unique", "Model tidak mendukung siklus ulang", "Direksi/Operasional"],
        ["Definisi siswa aktif", "Mengacu status pendaftaran", "KPI akademik dapat berbeda", "Direksi/Operasional"],
        ["Data wilayah sekolah baku", "Kota/provinsi string", "Duplikasi ejaan", "Operasional"],
        ["Rekonsiliasi/refund", "Pencatatan CRM dasar", "Omzet/koreksi tidak berlapis", "Direksi/Keuangan"],
        ["RPO/RTO dan restore drill", "Belum disahkan", "Pemulihan tidak terukur", "Direksi/Infrastruktur"],
        ["Enkripsi payload/media", "Belum dinyatakan", "Paparan data komunikasi", "Superadmin/Infrastruktur"],
        ["Kamus KPI resmi", "Rumus as-built", "Perbedaan laporan", "Direksi"],
    ], [1.65, 1.75, 1.65, 1.45])
    b.add_callout(document, "Kriteria BAB V dinyatakan disahkan", "Direksi dan pemilik proses menyetujui definisi data, role dan scope, kardinalitas, status, sumber/kanal, kamus KPI, klasifikasi keamanan, retensi, backup, RPO/RTO, serta keputusan atas semua gap material.", b.LIGHT_GREEN)

    b.add_heading(document, "PENUTUP BAB V", 1)
    b.add_body(document, "Arsitektur data SIVMI CRM telah memiliki fondasi relasional yang jelas, constraint penting, scope cabang, histori aktivitas, serta integrasi WhatsApp yang idempoten. Kesiapan menuju VPS terutama bergantung pada pengesahan retensi, backup/restore, perlindungan payload komunikasi, definisi KPI, dan prosedur koreksi data yang telah diidentifikasi pada dokumen ini.")
    b.add_heading(document, "SUMBER VERIFIKASI INTERNAL", 1)
    b.add_bullets(document, [
        "Migration database sampai 2026_08_24_000002_create_whatsapp_integration_tables.php.",
        "Model Eloquent pada app/Models dan enum domain pada app/Enums.",
        "Scope RBAC pada model dan Filament resources; seeder role/cabang Bandung, JakSel, dan JakPus.",
        "Diagram Gabungan SIVMI CRM v1.4.drawio sebagai referensi awal; ERD versi 1.1 digambar ulang dari migration dan model as-built.",
        "Dokumentasi induk SIVMI CRM v1.6 serta BAB IV Arsitektur Bisnis v1.0.",
        "Test fitur dan baseline quality gate full Docker yang tercatat pada proyek.",
    ])

    props = document.core_properties
    props.title = "BAB V: Arsitektur Data SIVMI CRM"
    props.subject = "RBAC, alur data, DFD, ERD, kardinalitas, kamus data, dan keamanan"
    props.author = "Tim Pengembangan SIVMI CRM"
    props.keywords = "SIVMI CRM, arsitektur data, RBAC, ERD, DFD, kamus data, WAHA"
    props.comments = "Versi 1.1 - penyempurnaan ERD dan susunan tabel database"
    document.settings.element.append(OxmlElement("w:updateFields"))
    update = document.settings.element[-1]
    update.set(qn("w:val"), "true")
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
