from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
DOC_TOOLS = ROOT / "storage" / "app" / "doc-tools"
sys.path.insert(0, str(DOC_TOOLS))

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


SOURCE = ROOT / "docs" / "Dokumentasi dan Pengembangan Sistem CRM SIVMI v1.5 - Sementara.docx"
OUTPUT = ROOT / "docs" / "Dokumentasi dan Pengembangan Sistem CRM SIVMI v1.6 - Sementara.docx"
ASSET_DIR = ROOT / "docs" / "_diagram_v1_6"
COVER = ASSET_DIR / "cover-v1.6.png"


def font(size: int, bold: bool = False):
    candidates = [
        Path(r"C:\Windows\Fonts\segoeuib.ttf" if bold else r"C:\Windows\Fonts\segoeui.ttf"),
        Path(r"C:\Windows\Fonts\arialbd.ttf" if bold else r"C:\Windows\Fonts\arial.ttf"),
    ]
    for candidate in candidates:
        if candidate.exists():
            return ImageFont.truetype(str(candidate), size)
    return ImageFont.load_default()


def make_cover() -> None:
    width, height = 1240, 1754
    image = Image.new("RGB", (width, height), "#F6F8FC")
    draw = ImageDraw.Draw(image)
    draw.rounded_rectangle((80, 80, width - 80, height - 80), 44, fill="#FFFFFF", outline="#D8E2EE", width=3)
    draw.rounded_rectangle((80, 80, width - 80, 610), 44, fill="#1476B8")
    draw.ellipse((830, -70, 1320, 420), fill="#1D87C5")
    draw.ellipse((950, 250, 1290, 590), fill="#075985")
    draw.rectangle((80, 560, width - 80, 610), fill="#075985")

    draw.text((145, 160), "SIVMI CRM", font=font(42, True), fill="#D9F0FF")
    draw.multiline_text(
        (145, 255),
        "DOKUMENTASI DAN\nPENGEMBANGAN SISTEM CRM",
        font=font(66, True),
        fill="#FFFFFF",
        spacing=14,
    )
    draw.text((145, 500), "Arsitektur • Data • Operasional • Deployment", font=font(28), fill="#D9F0FF")

    draw.text((145, 735), "VERSI SEMENTARA 1.6", font=font(43, True), fill="#14357A")
    draw.text((145, 815), "Snapshot as-built 24 Agustus 2026", font=font(29), fill="#64748B")

    cards = [
        ("FULL DOCKER", "Nginx, Laravel, MySQL, Redis, queue, scheduler"),
        ("WAHA GOWS", "Session per pengguna, webhook HMAC, inbox dan chat"),
        ("QUALITY GATE", "72 test lulus • 448 assertion"),
    ]
    y = 960
    colors = [("#EAF4FB", "#1476B8"), ("#EAFBF5", "#10B981"), ("#FAF5FF", "#A855F7")]
    for (title, description), (fill, stroke) in zip(cards, colors):
        draw.rounded_rectangle((145, y, width - 145, y + 155), 24, fill=fill, outline=stroke, width=3)
        draw.text((185, y + 28), title, font=font(30, True), fill="#0F172A")
        draw.text((185, y + 82), description, font=font(24), fill="#475569")
        y += 190

    draw.text((145, 1580), "SIVMI • CRM BIMBINGAN BELAJAR", font=font(24, True), fill="#64748B")
    image.save(COVER, quality=95)


def set_cell_shading(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), color)


def set_run_font(run, size: float = 10.5, bold: bool = False, color: str = "334155") -> None:
    run.font.name = "Aptos"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Aptos")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Aptos")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def replace_paragraph_text(paragraph, replacements: dict[str, str]) -> None:
    if paragraph.style and paragraph.style.name.lower().startswith("toc"):
        return
    current = paragraph.text
    updated = current
    for old, new in replacements.items():
        if old in updated:
            updated = updated.replace(old, new)
    if updated == current:
        return
    paragraph.clear()
    run = paragraph.add_run(updated)
    if not (paragraph.style and paragraph.style.name.startswith("Heading")):
        set_run_font(run, 10.5, False)


def add_body(document: Document, text: str, style: str | None = None):
    paragraph = document.add_paragraph(style=style)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if style in (None, "Normal") else WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(text)
    set_run_font(run)
    paragraph.paragraph_format.space_after = Pt(6)
    return paragraph


def add_heading(document: Document, text: str, level: int):
    paragraph = document.add_heading(text, level=level)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    return paragraph


def add_bullets(document: Document, items: list[str]) -> None:
    for item in items:
        add_body(document, item, "List Bullet")


def add_table(document: Document, headers: list[str], rows: list[list[str]], caption: str) -> None:
    caption_paragraph = document.add_paragraph(style="Caption")
    caption_paragraph.add_run(caption)
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = True
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = header
        set_cell_shading(cell, "1476B8")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            set_run_font(run, 9, True, "FFFFFF")
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for column_index, value in enumerate(values):
            cells[column_index].text = value
            cells[column_index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index % 2:
                set_cell_shading(cells[column_index], "F6F8FC")
            for paragraph in cells[column_index].paragraphs:
                for run in paragraph.runs:
                    set_run_font(run, 8.5)
    document.add_paragraph()


def add_figure(document: Document, path: Path, caption: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(path), width=Inches(6.35))
    caption_paragraph = document.add_paragraph(style="Caption")
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_paragraph.add_run(caption)


make_cover()
document = Document(SOURCE)

# Ganti sampul gambar tanpa mengubah section dan format badan dokumen.
cover_paragraph = document.paragraphs[0]
cover_paragraph.clear()
cover_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
cover_paragraph.add_run().add_picture(str(COVER), width=Inches(6.87))

replacements = {
    "Bagaimana membangun CRM bimbel yang mengelola siklus lead sampai transaksi, menyediakan informasi kinerja yang dapat dipertanggungjawabkan, mendukung empat role organisasi, dan dapat dijalankan konsisten pada Laragon maupun Docker tanpa mengorbankan keamanan data cabang?":
        "Bagaimana membangun CRM bimbel yang mengelola siklus lead sampai transaksi, menyediakan informasi kinerja yang dapat dipertanggungjawabkan, mendukung empat role organisasi, serta dapat dijalankan konsisten melalui full Docker tanpa mengorbankan keamanan data cabang?",
    "XAMPP merupakan paket Apache, MySQL/MariaDB, PHP, dan alat pendukung untuk pengembangan lokal. Proyek ini tidak bergantung pada XAMPP; baseline lokal aktual menggunakan Laragon, sedangkan Docker menjadi alternatif lingkungan yang lebih konsisten.":
        "XAMPP merupakan paket Apache, MySQL/MariaDB, PHP, dan alat pendukung untuk pengembangan lokal. Proyek ini tidak bergantung pada XAMPP maupun Laragon sebagai runtime; baseline lokal aktual menggunakan full Docker agar versi PHP, Nginx, MySQL, Redis, queue, scheduler, dan WAHA konsisten.",
    "Baseline aktual: PHP 8.4 pada image produksi, Laravel 12, Filament 5, Livewire 4, Node 22 untuk build, Nginx, MySQL, Redis, Docker Compose, serta Vite/Tailwind. Laragon tetap didukung untuk pengembangan lokal.":
        "Baseline aktual: PHP 8.4 pada image aplikasi, Laravel 12, Filament 5, Livewire 4, Node 22 untuk build, Nginx, MySQL, Redis, Docker Compose, Vite/Tailwind, serta WAHA 2026.8.1 CORE dengan engine GOWS. Full Docker menjadi jalur pengembangan utama; Laragon hanya opsi legacy.",
    "Development dapat memakai Laragon atau Docker. Testing memakai konfigurasi terisolasi. Staging meniru produksi untuk UAT. Production memakai image immutable, environment secret, APP_DEBUG=false, HTTPS, dan volume persisten.":
        "Development utama menggunakan full Docker. Testing dijalankan melalui docker/test-full.ps1 dengan APP_ENV=testing dan SQLite in-memory. Staging meniru topologi produksi untuk UAT. Production memakai image immutable, secret terkelola, APP_DEBUG=false, HTTPS, volume persisten, health check, backup, dan rollback teruji.",
    "Browser mengakses /admin melalui Nginx/Apache. Laravel menangani request, MySQL menyimpan data, Redis mendukung cache/queue, dan worker/scheduler menjalankan proses latar.":
        "Browser mengakses /admin melalui Nginx container. Laravel/PHP-FPM menangani request, MySQL menyimpan data, Redis mendukung cache/session/queue, sedangkan service queue dan scheduler menjalankan proses latar.",
    "Endpoint health /up tersedia. API publik dan webhook eksternal belum menjadi kontrak produksi. Target desain memerlukan versioning, autentikasi token, rate limit, idempotency key, signature webhook, retry, dan audit.":
        "Endpoint health /up dan webhook internal WAHA /api/webhooks/waha tersedia. API publik serta webhook untuk pihak ketiga belum menjadi kontrak produksi. Integrasi baru wajib memiliki versioning, autentikasi, rate limit, idempotency key, signature, retry, dan audit.",
    "Verifikasi bersumber dari composer.lock, package-lock.json, migration, route list, test, build Vite, Compose config, health endpoint, dan smoke test. Snapshot terakhir: 50 test dan 339 assertion lulus, build produksi berhasil, serta audit dependency tidak menemukan kerentanan yang diketahui.":
        "Verifikasi bersumber dari composer.lock, package-lock.json, migration, route list, test, build Vite, Compose config, health endpoint, dan smoke test. Snapshot terakhir: 72 test dan 448 assertion lulus pada environment testing terisolasi; build produksi serta stack full Docker berhasil dijalankan.",
    "Laragon: arahkan document root ke public, gunakan DB_HOST=127.0.0.1, siapkan database, composer install, npm install, key:generate, migrate --seed untuk lokal, npm run build/dev. Docker: salin docker/.env, lalu docker compose --env-file docker/.env -f docker/compose.yml up -d --build dan jalankan migration di service app.":
        "Instalasi development utama tidak memerlukan Laragon. Salin docker/.env.example menjadi docker/.env, jalankan .\\docker\\start-full.ps1, kemudian akses http://localhost:8000/admin/login. Skrip menyiapkan key lokal, membangun image, menjalankan migration dan seluruh service, serta menunggu health check. Gunakan .\\docker\\status-full.ps1, .\\docker\\logs-full.ps1, .\\docker\\test-full.ps1, dan .\\docker\\stop-full.ps1 untuk operasi harian.",
    "Lampiran ini mencatat implementasi as-built integrasi WhatsApp terbaru pada SIVMI CRM. Setiap admin/staff mempunyai nomor dan session WAHA sendiri. WAHA menggunakan engine GOWS dalam container Docker, sedangkan aplikasi Laravel dapat tetap berjalan melalui Laragon (hybrid) atau berada dalam Compose yang sama pada production.":
        "Lampiran ini mencatat implementasi as-built integrasi WhatsApp terbaru pada SIVMI CRM. Setiap admin/staff mempunyai nomor dan session WAHA sendiri. Baseline utama menjalankan WAHA engine GOWS bersama Laravel, Nginx, MySQL, Redis, queue, dan scheduler dalam satu jaringan Docker Compose; mode Laragon/hybrid hanya referensi legacy.",
    "B.11 Deployment Hybrid dan Full Docker": "B.11 Deployment Full Docker dan Referensi Hybrid",
    "Gambar B-1. Arsitektur integrasi WhatsApp WAHA pada konfigurasi hybrid": "Gambar B-1. Arsitektur integrasi WhatsApp WAHA pada konfigurasi full Docker",
    "Snapshot 24 Agustus 2026: WAHA 2026.8.1 CORE dengan GOWS berstatus WORKING; gRPC dan stream READY; webhook HMAC dapat menjangkau aplikasi hybrid; pesan berbasis LID tersimpan; nomor SenderAlt dinormalisasi; lead baru dibuat dan ditautkan. Suite aplikasi lulus 60 test dengan 400 assertion.":
        "Snapshot 24 Agustus 2026: baseline full Docker menjalankan Nginx, Laravel/PHP-FPM, MySQL, Redis, queue, scheduler, dan WAHA 2026.8.1 CORE dengan GOWS. Session dapat berstatus WORKING; webhook HMAC, normalisasi LID/SenderAlt, pencocokan lead, chat dua arah, ACK, unread badge, serta notifikasi sistem tersedia. Suite aplikasi lulus 72 test dengan 448 assertion pada environment testing terisolasi.",
}

tree_replacements = {
    "|   |   |-- SchoolGrade.php": "|   |   |-- SchoolGrade.php\n|   |   |-- SchoolNamePrefix.php\n|   |   |-- WhatsAppAccountStatus.php\n|   |   |-- WhatsAppMatchStatus.php\n|   |   |-- WhatsAppMessageDirection.php\n|   |   |-- WhatsAppMessageStatus.php",
    "|   |           `-- UserResource.php": "|   |           `-- UserResource.php\n|   |       |-- WhatsAppAccounts/\n|   |       |   |-- Pages/\n|   |       |   `-- WhatsAppAccountResource.php\n|   |       `-- WhatsAppConversations/\n|   |           |-- Pages/\n|   |           `-- WhatsAppConversationResource.php",
    "|   |       `-- Controller.php": "|   |       |-- Controller.php\n|   |       |-- WahaWebhookController.php\n|   |       `-- WhatsAppQrController.php",
    "|   |   `-- User.php": "|   |   |-- User.php\n|   |   |-- UserWhatsAppAccount.php\n|   |   |-- WahaWebhookEvent.php\n|   |   |-- WhatsAppConversation.php\n|   |   `-- WhatsAppMessage.php\n|   |-- Jobs/\n|   |   `-- ProcessWahaWebhook.php\n|   |-- Services/\n|   |   `-- WhatsApp/\n|   |       |-- LeadMatcher.php\n|   |       |-- WahaClient.php\n|   |       |-- WahaSessionService.php\n|   |       `-- WhatsAppChatService.php",
    "|   |   `-- 2026_08_24_000001_add_sumber_lead_to_leads_table.php": "|   |   |-- 2026_08_24_000001_add_sumber_lead_to_leads_table.php\n|   |   `-- 2026_08_24_000002_create_whatsapp_integration_tables.php",
    "|   `-- compose.yml": "|   |-- compose.yml\n|   |-- start-full.ps1\n|   |-- stop-full.ps1\n|   |-- status-full.ps1\n|   |-- logs-full.ps1\n|   `-- test-full.ps1",
    "|   |   `-- SettingsTest.php": "|   |   |-- SettingsTest.php\n|   |   |-- CreateSchoolFromLeadTest.php\n|   |   `-- WahaIntegrationTest.php",
}

replacements.update(tree_replacements)

for paragraph in document.paragraphs:
    replace_paragraph_text(paragraph, replacements)

# Ganti ilustrasi arsitektur WAHA hybrid lama dengan diagram full Docker terbaru.
for index, paragraph in enumerate(document.paragraphs):
    if paragraph.text.startswith("Gambar B-1."):
        for previous in range(index - 1, max(index - 5, -1), -1):
            image_paragraph = document.paragraphs[previous]
            if image_paragraph._p.xpath(".//a:blip"):
                image_paragraph.clear()
                image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                image_paragraph.add_run().add_picture(str(ASSET_DIR / "waha-architecture-full-docker.png"), width=Inches(6.35))
                break
        break

for section in document.sections:
    for paragraph in section.header.paragraphs:
        replace_paragraph_text(paragraph, {"1.5": "1.6"})
    for paragraph in section.footer.paragraphs:
        replace_paragraph_text(paragraph, {"1.5": "1.6"})

# Perbarui lembar kendali dokumen.
control = document.tables[0]
for row in control.rows:
    key = row.cells[0].text.strip()
    if key == "Versi":
        row.cells[1].text = "Sementara 1.6"
    elif key == "Acuan":
        row.cells[1].text = "Codebase as-built, migration, 72 test/448 assertion, diagram v1.4, konfigurasi full Docker, WAHA GOWS, dan dokumentasi v1.5"
    elif key == "Teknologi aktual":
        row.cells[1].text = "Laravel 12, Filament 5, Livewire 4, Tailwind CSS 4, Nginx, MySQL, Redis, Docker Compose, WAHA GOWS"
    for cell in row.cells:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                set_run_font(run, 8.5, key == "Atribut")

# Perbarui nilai tabel operasional lama tanpa mengubah struktur tabel.
for table in document.tables:
    for row in table.rows:
        for cell in row.cells:
            if cell.text.strip() == "docker compose --profile workers up -d":
                cell.text = r".\docker\start-full.ps1"
            elif cell.text.strip() == "Service queue dan scheduler mengikuti Compose":
                cell.text = "Seluruh service, health check, migration, queue, dan scheduler dikelola skrip full Docker"

document.add_page_break()
add_heading(document, "LAMPIRAN C: BASELINE FULL DOCKER, CHAT, NOTIFIKASI, DAN QUALITY GATE", 1)
add_body(document, "Lampiran ini memperbarui baseline operasional SIVMI CRM setelah perpindahan pengembangan lokal dari Laragon/hybrid ke full Docker. Seluruh komponen aplikasi, data, worker, scheduler, dan integrasi WAHA berjalan sebagai service Compose yang dapat direproduksi.")

add_heading(document, "C.1 Topologi Full Docker", 2)
add_bullets(document, [
    "web: Nginx pada port host 8000 yang meneruskan FastCGI ke app:9000.",
    "app: Laravel/PHP-FPM dengan Filament 5, Livewire 4, dan aset Tailwind/Vite hasil build.",
    "mysql dan redis: penyimpanan relasional, cache, session, serta queue pada volume persisten.",
    "queue dan scheduler: pemrosesan webhook, reminder, notifikasi, dan tugas terjadwal tanpa terminal interaktif.",
    "waha: WAHA 2026.8.1 CORE dengan engine GOWS, API internal http://waha:3000, dan webhook http://web/api/webhooks/waha.",
])
add_figure(document, ASSET_DIR / "deployment-full-docker.png", "Gambar C-1. Topologi development full Docker SIVMI CRM")

add_heading(document, "C.2 Operasi Harian Development", 2)
add_table(document, ["Kebutuhan", "Perintah"], [
    ["Menjalankan seluruh stack", r".\docker\start-full.ps1"],
    ["Memeriksa status service", r".\docker\status-full.ps1"],
    ["Membaca log terpusat", r".\docker\logs-full.ps1"],
    ["Menjalankan quality gate", r".\docker\test-full.ps1"],
    ["Menghentikan service", r".\docker\stop-full.ps1"],
], "Tabel C-1. Perintah operasional full Docker")
add_body(document, "Database Docker berada pada volume MySQL yang terpisah dari database Laragon. Penggantian runtime tidak menyinkronkan data secara otomatis. Jangan menjalankan docker compose down -v tanpa backup dan persetujuan karena tindakan tersebut menghapus volume persisten.")

add_heading(document, "C.3 Chat WhatsApp dan Notifikasi Sistem", 2)
add_body(document, "Pesan masuk diterima WAHA GOWS, dikirim ke webhook bertanda tangan HMAC, dicatat idempoten, lalu diproses oleh queue. Nomor JID/LID dan SenderAlt dinormalisasi sebelum sistem mencocokkan lead berdasarkan nomor, cabang, dan PIC. Percakapan tampil pada Inbox WhatsApp, unread badge diperbarui, balasan dikirim melalui session pemilik, dan notifikasi database diarahkan sesuai role serta cabang.")
add_figure(document, ASSET_DIR / "chat-notifikasi.png", "Gambar C-2. Alur chat WhatsApp dan notifikasi sistem")

add_heading(document, "C.4 Data dan Kardinalitas Integrasi", 2)
add_table(document, ["Relasi", "Kardinalitas", "Makna"], [
    ["users — user_whatsapp_accounts", "1 : 0..1", "Satu pengguna memiliki paling banyak satu nomor/session operasional."],
    ["user_whatsapp_accounts — whatsapp_conversations", "1 : 0..N", "Satu session melayani banyak percakapan."],
    ["whatsapp_conversations — whatsapp_messages", "1 : 0..N", "Satu percakapan berisi banyak pesan inbound/outbound."],
    ["leads — whatsapp_conversations", "1 : 0..N", "Satu lead dapat terhubung dengan beberapa percakapan sepanjang lifecycle."],
    ["user_whatsapp_accounts — waha_webhook_events", "1 : 0..N", "Event disimpan untuk audit, idempotensi, dan retry."],
], "Tabel C-2. Relasi dan kardinalitas integrasi WhatsApp")

add_heading(document, "C.5 Quality Gate Tercatat", 2)
add_figure(document, ASSET_DIR / "quality-gate.png", "Gambar C-3. Quality gate full Docker sebelum rilis")
add_table(document, ["Pemeriksaan", "Hasil snapshot 24 Agustus 2026"], [
    ["Automated test", "72 test lulus; 448 assertion; 0 failure"],
    ["Database testing", "SQLite in-memory dengan APP_ENV=testing; database development tidak berubah"],
    ["Container", "App, MySQL, Redis, WAHA, queue, scheduler, dan Nginx berjalan pada Compose"],
    ["Web", "Login tersedia pada http://localhost:8000/admin/login; health endpoint /up"],
    ["WAHA", "Engine GOWS; API internal; webhook HMAC; session dan volume persisten"],
], "Tabel C-3. Baseline verifikasi terbaru")

add_heading(document, "C.6 Kriteria Siap Menuju VPS", 2)
add_bullets(document, [
    "APP_ENV=production, APP_DEBUG=false, APP_URL HTTPS, dan secret unik tersimpan di secret manager atau environment server.",
    "Port MySQL, Redis, dan WAHA tidak dipublikasikan ke internet; akses admin dibatasi firewall/VPN bila diperlukan.",
    "TLS, backup MySQL/storage/session WAHA, restore rehearsal, RPO/RTO, monitoring, log rotation, dan alerting telah ditetapkan.",
    "Migration, build, automated test, smoke test, RBAC/cabang, chat WAHA, notifikasi, queue, scheduler, dan rollback diverifikasi di staging.",
    "Seeder data dummy tidak dijalankan pada production dan kredensial default wajib diganti sebelum rilis.",
])

add_heading(document, "C.7 Riwayat Perubahan Versi 1.6", 2)
add_table(document, ["Area", "Perubahan"], [
    ["Runtime", "Full Docker menjadi baseline utama; Laragon/hybrid menjadi referensi legacy."],
    ["Nginx", "Upstream PHP-FPM menggunakan resolver jaringan Docker untuk mencegah 502 setelah container direkreasi."],
    ["Testing", "Skrip test khusus memaksa APP_ENV=testing dan SQLite in-memory; 72 test/448 assertion lulus."],
    ["WhatsApp", "Topologi WAHA GOWS, webhook HMAC, session per pengguna, chat, unread badge, ACK, dan notifikasi didokumentasikan."],
    ["Diagram", "Diagram gabungan v1.4 berisi 18 halaman termasuk deployment full Docker, chat/notifikasi, dan quality gate."],
], "Tabel C-4. Changelog dokumentasi versi 1.6")

add_body(document, "--- Akhir pembaruan versi 1.6 ---")

# Word/LibreOffice diminta memperbarui field TOC saat dokumen dibuka.
settings = document.settings._element
update_fields = settings.find(qn("w:updateFields"))
if update_fields is None:
    update_fields = OxmlElement("w:updateFields")
    settings.append(update_fields)
update_fields.set(qn("w:val"), "true")

document.core_properties.title = "Dokumentasi dan Pengembangan Sistem CRM SIVMI v1.6"
document.core_properties.subject = "As-built full Docker, WAHA GOWS, chat, notifikasi, data, dan quality gate"
document.core_properties.keywords = "SIVMI, CRM, Laravel, Filament, Docker, WAHA, GOWS, RBAC, ERD, DFD"
document.core_properties.comments = "Snapshot as-built 24 Agustus 2026"
document.save(OUTPUT)
print(OUTPUT)
