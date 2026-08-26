from __future__ import annotations

from pathlib import Path
import sys

PROJECT_ROOT = Path(__file__).resolve().parents[2]
DOC_TOOLS = PROJECT_ROOT / "storage" / "app" / "doc-tools"
sys.path.insert(0, str(DOC_TOOLS))

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

import build_bab_iv_arsitektur_bisnis as b


OUTPUT = PROJECT_ROOT / "docs" / "BAB VI - Arsitektur Aplikasi SIVMI CRM v1.0.docx"
ASSETS = PROJECT_ROOT / "docs" / "_bab_vi_assets"
COVER = ASSETS / "cover-bab-vi.png"

PRIMARY = "#1476B8"
PRIMARY_DARK = "#075985"
INK = "#0B2545"
MUTED = "#64748B"
BG = "#F6F8FC"
GREEN = "#10B981"
AMBER = "#F59E0B"
ROSE = "#F43F5E"
PURPLE = "#A855F7"


def font(name: str, size: int):
    return ImageFont.truetype(f"C:/Windows/Fonts/{name}", size)


def canvas(title: str, subtitle: str, size=(1800, 1050)):
    image = Image.new("RGB", size, "white")
    draw = ImageDraw.Draw(image)
    draw.text((size[0] // 2, 50), title, anchor="ma", font=font("arialbd.ttf", 45), fill="#14357A")
    draw.text((size[0] // 2, 108), subtitle, anchor="ma", font=font("arial.ttf", 24), fill=MUTED)
    return image, draw


def box(draw, xy, title, lines, fill="#EAF4FB", outline=PRIMARY, title_size=27, body_size=21):
    x1, y1, x2, y2 = xy
    draw.rounded_rectangle(xy, radius=22, fill=fill, outline=outline, width=4)
    draw.text(((x1 + x2) // 2, y1 + 32), title, anchor="ma", font=font("arialbd.ttf", title_size), fill=INK)
    draw.multiline_text(((x1 + x2) // 2, y1 + 82), lines, anchor="ma", align="center", spacing=8,
                        font=font("arial.ttf", body_size), fill="#334155")


def arrow(draw, start, end, color="#64748B", width=5):
    draw.line((start, end), fill=color, width=width)
    x, y = end
    sx, sy = start
    if abs(x - sx) >= abs(y - sy):
        sign = 1 if x > sx else -1
        pts = [(x, y), (x - sign * 18, y - 11), (x - sign * 18, y + 11)]
    else:
        sign = 1 if y > sy else -1
        pts = [(x, y), (x - 11, y - sign * 18), (x + 11, y - sign * 18)]
    draw.polygon(pts, fill=color)


def make_cover():
    ASSETS.mkdir(parents=True, exist_ok=True)
    image = Image.new("RGB", (1600, 2200), BG)
    draw = ImageDraw.Draw(image)
    for y in range(2200):
        ratio = y / 2200
        a, z = (20, 118, 184), (7, 89, 133)
        draw.line((0, y, 1600, y), fill=tuple(round(a[i] * (1-ratio) + z[i] * ratio) for i in range(3)))
    draw.ellipse((1010, -180, 1770, 580), fill=(27, 139, 199))
    draw.ellipse((1140, 140, 1640, 640), fill=(8, 103, 151))
    draw.rounded_rectangle((120, 250, 1480, 1910), radius=44, fill="white")
    draw.text((190, 350), "SIVMI CRM", font=font("arialbd.ttf", 34), fill=PRIMARY)
    draw.text((190, 600), "BAB VI", font=font("arialbd.ttf", 48), fill=PRIMARY_DARK)
    draw.multiline_text((190, 690), "ARSITEKTUR\nAPLIKASI", font=font("arialbd.ttf", 96), fill=INK, spacing=18)
    draw.text((190, 1050), "Konteks, layer, komponen, modul, request, API,", font=font("arial.ttf", 34), fill="#475569")
    draw.text((190, 1102), "queue, penyimpanan, UI, logging, dan pengujian", font=font("arial.ttf", 34), fill="#475569")
    draw.rounded_rectangle((190, 1245, 760, 1325), radius=40, fill="#EAF4FB")
    draw.text((230, 1267), "BASELINE AS-BUILT 24 AGUSTUS 2026", font=font("arial.ttf", 28), fill=PRIMARY_DARK)
    draw.text((190, 1610), "Versi dokumen 1.0", font=font("arial.ttf", 28), fill=MUTED)
    draw.text((190, 1665), "Status: dokumen kerja untuk validasi teknis", font=font("arial.ttf", 28), fill=MUTED)
    draw.text((190, 1775), "SIVMI CRM - Catat, Kelola, Closing", font=font("arialbd.ttf", 34), fill=PRIMARY)
    image.save(COVER)


def make_context():
    image, draw = canvas("Konteks Aplikasi SIVMI CRM", "Batas aplikasi, kanal pengguna, dan layanan pendukung")
    box(draw, (610, 260, 1190, 720), "SIVMI CRM", "Admin panel Filament 5\nLivewire 4 + Tailwind CSS 4\nLaravel 12 + Eloquent\nRBAC dan branch scope", "#EAF4FB", PRIMARY, 38, 26)
    externals = [
        ((70, 230, 440, 450), "Pengguna", "Superadmin\nDireksi\nAdmin Cabang\nStaff", "#F8FAFC", PRIMARY),
        ((70, 620, 440, 840), "Calon siswa", "Form, CSV, telepon,\ndan WhatsApp", "#ECFDF5", GREEN),
        ((1360, 190, 1730, 390), "WAHA / GOWS", "Session per pengguna\nWebhook dan sendText", "#FAF5FF", PURPLE),
        ((1360, 450, 1730, 650), "Data & state", "MySQL 8.4\nRedis 7.4\nFilesystem", "#FFFBEB", AMBER),
        ((1360, 710, 1730, 910), "Operasional", "Queue worker\nScheduler\nLog dan health check", "#FFF1F2", ROSE),
    ]
    for xy, title, lines, fill, color in externals:
        box(draw, xy, title, lines, fill, color)
    for s, e in [((440,340),(610,390)),((440,730),(610,610)),((1190,360),(1360,290)),((1190,490),(1360,550)),((1190,630),(1360,805))]:
        arrow(draw, s, e)
    image.save(ASSETS / "konteks-aplikasi.png")


def make_layers():
    image, draw = canvas("Layer dan Tanggung Jawab", "Dependensi bergerak dari presentation menuju infrastructure melalui kontrak aplikasi")
    layers = [
        ("1. Presentation", "Filament panel, Resource, Page, Form, Table, Widget, Blade, Livewire", "#EAF4FB", PRIMARY),
        ("2. Application", "Page action, controller, command, observer, notifier, import/export", "#ECFDF5", GREEN),
        ("3. Domain service", "LeadMatcher, CreateSchoolFromLead, WhatsAppChatService, WahaSessionService", "#FFFBEB", AMBER),
        ("4. Persistence", "Eloquent model, relationship, scope visibleTo, transaction, query builder", "#FAF5FF", PURPLE),
        ("5. Infrastructure", "MySQL, Redis, filesystem, queue, scheduler, HTTP WAHA, Nginx/PHP-FPM", "#FFF1F2", ROSE),
    ]
    y = 190
    for idx, (title, lines, fill, color) in enumerate(layers):
        draw.rounded_rectangle((180, y, 1620, y+125), radius=22, fill=fill, outline=color, width=4)
        draw.text((230, y+34), title, anchor="lm", font=font("arialbd.ttf", 28), fill=INK)
        draw.text((620, y+63), lines, anchor="lm", font=font("arial.ttf", 23), fill="#334155")
        if idx < len(layers)-1:
            arrow(draw, (900, y+125), (900, y+155))
        y += 160
    image.save(ASSETS / "layer-aplikasi.png")


def make_components():
    image, draw = canvas("Struktur Komponen", "Komponen as-built dan arah interaksi utama", (1800, 1180))
    columns = [
        (55, 390, "Antarmuka", ["Dashboard & laporan", "12 resource/master", "Inbox & chat WA", "Notifikasi database", "Import/export CSV"], PRIMARY),
        (455, 790, "Orkestrasi", ["Filament actions", "Controller webhook/QR", "Console commands", "Observers", "CrmNotifier"], GREEN),
        (855, 1245, "Domain & data", ["16 model Eloquent", "11 enum domain", "visibleTo scopes", "DB transaction", "Lead matching"], AMBER),
        (1310, 1745, "Integrasi", ["WahaClient HTTP", "ProcessWahaWebhook", "Queue & scheduler", "MySQL / Redis", "Filesystem / logs"], PURPLE),
    ]
    for x1, x2, title, items, color in columns:
        draw.rounded_rectangle((x1, 180, x2, 1050), radius=24, fill="#F8FAFC", outline=color, width=4)
        draw.rectangle((x1, 180, x2, 250), fill=color)
        draw.text(((x1+x2)//2, 215), title, anchor="mm", font=font("arialbd.ttf", 29), fill="white")
        y = 300
        for item in items:
            draw.rounded_rectangle((x1+24, y, x2-24, y+105), radius=18, fill="white", outline="#D8E1EC", width=3)
            draw.multiline_text(((x1+x2)//2, y+52), item, anchor="mm", align="center", font=font("arialbd.ttf", 22), fill=INK)
            y += 140
    for x in (390, 790, 1245):
        arrow(draw, (x+5, 615), (x+60, 615))
    image.save(ASSETS / "struktur-komponen.png")


def make_module_map():
    image, draw = canvas("Peta Modul Aplikasi", "Modul UI dipetakan ke domain dan layanan bersama", (1800, 1100))
    center=(900,540)
    box(draw,(690,385,1110,695),"Layanan bersama","RBAC & branch scope\nAudit & notifikasi\nImport/export\nQueue & scheduler", "#EAF4FB", PRIMARY,32,23)
    modules=[
        (120,190,480,380,"Akuisisi","Leads\nSekolah\nImport CSV",GREEN),
        (120,700,480,890,"Engagement","Follow-up\nInbox WhatsApp\nChat",PURPLE),
        (1320,190,1680,380,"Konversi","Pendaftaran\nProgram Bimbel",AMBER),
        (1320,700,1680,890,"Transaksi","Pembayaran\nVerifikasi",ROSE),
        (580,160,1220,310,"Manajemen","Dashboard | Target | Laporan",PRIMARY_DARK),
        (580,820,1220,970,"Administrasi","Cabang | Pengguna | Pengaturan | Audit",PRIMARY),
    ]
    for x1,y1,x2,y2,t,l,c in modules:
        box(draw,(x1,y1,x2,y2),t,l,"#F8FAFC",c,27,21)
        arrow(draw,((x1+x2)//2,(y1+y2)//2),center,color="#94A3B8",width=4)
    image.save(ASSETS / "peta-modul.png")


def make_web_request():
    image, draw = canvas("Alur Request Web", "Siklus request server-side dan pembaruan reaktif Livewire", (1800, 1000))
    items=[
        ("1", "Browser", "HTTP/HTTPS\nasset Vite"),
        ("2", "Nginx", "public/index.php\nsecurity headers"),
        ("3", "Laravel", "routing, session,\nCSRF, auth"),
        ("4", "Filament / Livewire", "resource/page/action\nvalidation"),
        ("5", "Service / Model", "authorization, scope,\ntransaction"),
        ("6", "MySQL / Redis", "persist, cache,\nqueue state"),
    ]
    x=45; y=285; w=245; gap=52
    for i,(n,t,l) in enumerate(items):
        box(draw,(x+i*(w+gap),y,x+i*(w+gap)+w,y+310),f"{n}. {t}",l,"#F8FAFC",[PRIMARY,GREEN,AMBER,PURPLE,ROSE,PRIMARY_DARK][i],25,20)
        if i < len(items)-1: arrow(draw,(x+i*(w+gap)+w,y+155),(x+(i+1)*(w+gap)-8,y+155))
    draw.rounded_rectangle((300,720,1500,880),radius=24,fill="#EAF4FB",outline=PRIMARY,width=3)
    draw.multiline_text((900,800),"Respons HTML awal atau payload Livewire dikembalikan ke browser.\nError validasi tampil pada form; exception API dipaksa menjadi JSON.",anchor="mm",align="center",font=font("arialbd.ttf",23),fill=INK)
    image.save(ASSETS / "alur-request-web.png")


def make_business_flow():
    image, draw = canvas("Alur Follow-up dan Transaksi", "State bisnis utama dari lead hingga pengakuan omzet", (1800, 980))
    stages=[
        ("Lead", "Cold / Warm / Hot\nTidak Berminat", PRIMARY),
        ("Follow-up", "Metode, hasil,\njadwal berikutnya", PURPLE),
        ("Pendaftaran", "Lead + program\nstatus siswa", GREEN),
        ("Pembayaran", "Menunggu / Lunas\nTerverifikasi", AMBER),
        ("Kinerja", "Omzet, siswa aktif,\ntarget dan laporan", ROSE),
    ]
    x=70; y=260; w=285; gap=65
    for i,(t,l,c) in enumerate(stages):
        box(draw,(x+i*(w+gap),y,x+i*(w+gap)+w,y+300),t,l,"#F8FAFC",c,29,21)
        if i<len(stages)-1: arrow(draw,(x+i*(w+gap)+w,y+150),(x+(i+1)*(w+gap)-8,y+150))
    draw.rounded_rectangle((280,680,1520,860),radius=24,fill="#F8FAFC",outline="#CBD5E1",width=3)
    draw.multiline_text((900,770),"Kontrol lintas proses: role + cabang + PIC | validasi status | histori tidak ditimpa\nobserver audit/notifikasi | reminder terjadwal | query KPI terfilter",anchor="mm",align="center",font=font("arialbd.ttf",23),fill=INK)
    image.save(ASSETS / "alur-followup-transaksi.png")


def make_api_sequence():
    image, draw = canvas("Arsitektur API dan Webhook WAHA", "Inbound aman, idempoten, dan asynchronous; outbound melalui client terpusat", (1800, 1120))
    lanes=[(110,"WAHA/GOWS"),(420,"API Laravel"),(760,"Queue worker"),(1110,"Domain/Data"),(1490,"Pengguna")]
    for x,label in lanes:
        draw.text((x,175),label,anchor="ma",font=font("arialbd.ttf",25),fill=PRIMARY_DARK)
        draw.line((x,210,x,1040),fill="#CBD5E1",width=3)
    steps=[
        (110,420,270,"POST /api/webhooks/waha\n+ HMAC SHA-512"),
        (420,760,385,"Validasi + simpan event\npayload_hash unik"),
        (760,1110,500,"ProcessWahaWebhook\ntries 5 + backoff"),
        (1110,1110,615,"Normalisasi, match lead,\nsimpan conversation/message"),
        (1110,1490,730,"Notifikasi pesan baru\ndan unread count"),
        (1490,420,855,"Kirim chat dari UI\nWhatsAppChatService"),
        (420,110,970,"POST /api/sendText\nX-Api-Key"),
    ]
    for x1,x2,y,label in steps:
        direction_end=(x2,y)
        arrow(draw,(x1,y),direction_end,color=PRIMARY if x2>=x1 else PURPLE,width=4)
        draw.rounded_rectangle(((x1+x2)//2-155,y-56,(x1+x2)//2+155,y-8),radius=12,fill="white",outline="#D8E1EC",width=2)
        draw.multiline_text(((x1+x2)//2,y-32),label,anchor="mm",align="center",font=font("arial.ttf",18),fill=INK)
    image.save(ASSETS / "arsitektur-api-waha.png")


def make_async():
    image, draw = canvas("Scheduler, Queue, dan Reminder", "Pemisahan pekerjaan interaktif, terjadwal, dan asynchronous", (1800, 1000))
    box(draw,(80,230,430,450),"Trigger","Webhook WAHA\nFilament import/export\nSchedule", "#EAF4FB", PRIMARY)
    box(draw,(600,210,1010,470),"Queue","Database/Redis driver\nProcessWahaWebhook\nfailed_jobs", "#FAF5FF", PURPLE)
    box(draw,(1220,180,1720,500),"Worker","Retry dan backoff\nPersist event/message\nNotifikasi", "#ECFDF5", GREEN)
    arrow(draw,(430,340),(600,340)); arrow(draw,(1010,340),(1220,340))
    box(draw,(80,650,520,870),"Scheduler","Setiap 5 menit: sync WAHA\nSetiap jam: reminder CRM", "#FFFBEB", AMBER)
    box(draw,(700,620,1160,900),"Command","SyncWahaSessions\nSendCrmReminders\nchunkById(100)", "#F8FAFC", PRIMARY_DARK)
    box(draw,(1340,650,1720,870),"Output","Status session\nNotifikasi follow-up\nNotifikasi pembayaran", "#FFF1F2", ROSE)
    arrow(draw,(520,760),(700,760)); arrow(draw,(1160,760),(1340,760))
    image.save(ASSETS / "queue-scheduler.png")


def make_deployment():
    image, draw = canvas("Topologi Runtime Full Docker", "Baseline lokal dan target VPS menggunakan komponen terpisah", (1800, 1080))
    box(draw,(70,180,410,390),"Client","Browser desktop\nTablet / mobile", "#F8FAFC", PRIMARY)
    box(draw,(560,180,920,390),"Web","Nginx 1.27\nport aplikasi", "#EAF4FB", PRIMARY)
    box(draw,(1070,180,1430,390),"App","PHP 8.4 FPM\nLaravel 12", "#ECFDF5", GREEN)
    arrow(draw,(410,285),(560,285)); arrow(draw,(920,285),(1070,285))
    services=[
        (90,650,390,880,"MySQL 8.4","data utama\nvolume persisten",AMBER),
        (470,650,770,880,"Redis 7.4","cache / queue\nAOF",ROSE),
        (850,650,1150,880,"Worker","queue:work\nretry",PURPLE),
        (1230,650,1530,880,"Scheduler","schedule:work\nreminder/sync",PRIMARY_DARK),
        (1550,650,1780,880,"WAHA","GOWS\nsession volume",GREEN),
    ]
    for x1,y1,x2,y2,t,l,c in services: box(draw,(x1,y1,x2,y2),t,l,"#F8FAFC",c,25,19)
    for x in (240,620,1000,1380,1665): arrow(draw,(1250,390),(x,650),color="#94A3B8",width=3)
    image.save(ASSETS / "topologi-runtime.png")


def add_figure(doc, name, caption, width=6.15):
    b.add_figure(doc, ASSETS / name, caption, width)


def table(doc, caption, headers, rows, widths):
    b.add_table(doc, headers, rows, widths, caption)


def add_steps(doc, items):
    table(doc, "Urutan proses", ["Langkah", "Uraian"], [[str(i), item] for i,item in enumerate(items,1)], [0.75,5.75])


def new_page(doc):
    doc.add_page_break()


def build_document():
    for make in (make_cover, make_context, make_layers, make_components, make_module_map, make_web_request,
                 make_business_flow, make_api_sequence, make_async, make_deployment):
        make()

    doc=Document()
    section=doc.sections[0]
    section.page_width=Inches(8.5); section.page_height=Inches(11)
    section.top_margin=Inches(0.85); section.bottom_margin=Inches(0.85)
    section.left_margin=Inches(1); section.right_margin=Inches(1)
    section.header_distance=Inches(0.4); section.footer_distance=Inches(0.4)
    b.configure_styles(doc)
    section.header.paragraphs[0].text="SIVMI CRM  |  BAB VI: ARSITEKTUR APLIKASI"
    for run in section.header.paragraphs[0].runs: b.set_run_font(run,8.5,True,b.PRIMARY)
    b.add_page_number(section.footer.paragraphs[0])

    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(0)
    pic=p.add_run().add_picture(str(COVER),width=Inches(6.1)); pic._inline.docPr.set("descr","Sampul BAB VI Arsitektur Aplikasi SIVMI CRM")
    new_page(doc)

    b.add_heading(doc,"KENDALI DOKUMEN",1)
    table(doc,"Tabel 0-1. Kendali dan baseline dokumen",["Atribut","Nilai"],[
        ["Judul","BAB VI: Arsitektur Aplikasi SIVMI CRM"],["Versi","1.0"],["Tanggal baseline","24 Agustus 2026"],
        ["Status","Dokumen kerja terpisah - siap validasi teknis dan manajemen"],
        ["Ruang lingkup","Konteks, prinsip, layer, komponen, modul, request, transaksi, data, API, queue, file, UI, error, logging, test, dan gap"],
        ["Basis implementasi","PHP 8.3+, Laravel 12, Filament 5, Livewire 4, Tailwind CSS 4, Vite 8, MySQL 8.4, Redis 7.4, WAHA CORE/GOWS"],
        ["Runtime","Full Docker: Nginx, PHP-FPM, MySQL, Redis, queue worker, scheduler, Node dev server, dan WAHA"],
        ["Sumber kebenaran","Codebase, routes, provider, resource/page, service, model, command, queue job, config, Docker Compose, migration, dan test terbaru"],
    ],[1.55,4.95])
    b.add_callout(doc,"Kedudukan dokumen","Dokumen ini merekam arsitektur aplikasi as-built. Perbedaan terhadap rancangan lama harus diselesaikan dengan menjadikan kode, migration, konfigurasi, dan test terbaru sebagai baseline; target yang belum diimplementasikan ditandai sebagai gap.")
    b.add_heading(doc,"DAFTAR ISI",1); b.add_toc(doc); new_page(doc)

    b.add_heading(doc,"RINGKASAN EKSEKUTIF",1)
    b.add_body(doc,"SIVMI CRM adalah aplikasi web modular monolith berbasis Laravel 12. Filament 5 dan Livewire 4 membentuk presentation layer administratif, Eloquent menyediakan persistence dan branch scope, sedangkan service, observer, command, serta queued job mengorkestrasi aturan bisnis. Integrasi WhatsApp memakai WAHA CORE dengan engine GOWS melalui HTTP API dan webhook HMAC.")
    b.add_body(doc,"Boundary aplikasi dipusatkan pada role dan cabang. Superadmin dan Direksi dapat membaca lintas cabang sesuai fungsi, sedangkan Admin Cabang dan Staff dibatasi oleh query server-side. Proses berat atau eksternal dipisahkan ke queue; reminder dan sinkronisasi session dijalankan scheduler. Audit dan notifikasi diproduksi melalui observer dan service terpusat.")
    b.add_callout(doc,"Keputusan arsitektur utama","Modular monolith dipertahankan karena domain masih erat, tim dan operasi belum membutuhkan distribusi layanan, serta transaksi CRM perlu konsistensi kuat. Integrasi eksternal ditempatkan di adapter/service agar dapat diganti tanpa merombak resource dan model.",b.LIGHT_GREEN)
    table(doc,"Tabel 0-2. Ringkasan arsitektur",["Area","As-built","Kontrol"],[
        ["Presentation","Filament Resource/Page, Livewire, Blade, Tailwind","Auth, CSRF, validation, responsive UI"],
        ["Application","Controller, command, observer, notifier, import/export","Orkestrasi dan audit"],
        ["Domain","Enum, service, model event, state transition","Business rule dan branch scope"],
        ["Data","Eloquent, Query Builder, transaction, MySQL","FK, unique, index, visibleTo"],
        ["Integration","WahaClient, webhook controller, queued job","API key, HMAC, retry, idempotensi"],
        ["Operations","Docker, Nginx, PHP-FPM, Redis, queue, scheduler","Health check, volume, log, secret gate"],
    ],[1.2,3.0,2.3])

    b.add_heading(doc,"BAB VI: ARSITEKTUR APLIKASI",1)
    b.add_heading(doc,"6.1 Pengertian",2)
    b.add_body(doc,"Arsitektur aplikasi adalah susunan komponen perangkat lunak, batas tanggung jawab, pola interaksi, dependensi, dan kontrol operasional yang mengubah kebutuhan bisnis menjadi layanan sistem. Arsitektur ini tidak hanya menjelaskan teknologi, tetapi juga cara modul menjaga konsistensi, keamanan akses, keterlacakan, dan kemampuan berkembang.")
    table(doc,"Tabel 6-1. Sudut pandang arsitektur",["View","Pertanyaan","Artefak"],[
        ["Context","Siapa memakai dan apa yang terhubung?","Aktor, kanal, WAHA, datastore"],
        ["Logical","Bagaimana fungsi dibagi?","Layer, komponen, modul"],
        ["Process","Bagaimana request dan event mengalir?","Sequence web, webhook, transaksi"],
        ["Development","Bagaimana kode diorganisasi?","Namespace app, resource, service, model"],
        ["Deployment","Di mana komponen berjalan?","Container, volume, network, health check"],
    ],[1.1,2.75,2.65])

    b.add_heading(doc,"6.2 Tujuan",2)
    b.add_bullets(doc,[
        "Menyediakan batas tanggung jawab yang konsisten antara UI, orkestrasi, domain, persistence, dan infrastructure.",
        "Menjamin pembatasan role, cabang, PIC, serta kepemilikan session diterapkan di server, bukan hanya menu.",
        "Menjaga proses lead, follow-up, pendaftaran, pembayaran, KPI, dan WhatsApp dapat ditelusuri.",
        "Memisahkan pekerjaan lambat dan event eksternal dari request interaktif melalui queue dan scheduler.",
        "Menjadi baseline pengujian, debugging, deployment lokal full Docker, dan migrasi ke VPS.",
    ])
    table(doc,"Tabel 6-2. Sasaran kualitas",["Atribut","Sasaran","Mekanisme as-built"],[
        ["Security","Akses minimum dan input tervalidasi","Auth, CSRF, RBAC, scope, HMAC"],
        ["Reliability","Event tidak hilang/duplikat","DB transaction, payload hash, retry"],
        ["Performance","Dashboard/query responsif","Aggregate query, index, cache, queue"],
        ["Maintainability","Perubahan terlokalisasi","Resource/schema/service/enum"],
        ["Operability","Mudah dijalankan dan ditelusuri","Docker, health endpoint, log, failed_jobs"],
    ],[1.2,2.1,3.2])

    b.add_heading(doc,"6.3 Ruang Lingkup",2)
    table(doc,"Tabel 6-3. In-scope dan out-of-scope",["Kategori","Cakupan"],[
        ["In-scope","Panel admin, dashboard, master, pipeline CRM, transaksi, KPI/laporan, audit/notifikasi, import/export, akun dan inbox WhatsApp, queue, scheduler, file, runtime Docker."],
        ["Out-of-scope saat ini","Portal publik siswa, aplikasi mobile native, payment gateway otomatis, LMS penuh, data warehouse, microservice terpisah, public API partner lengkap, HA multi-node."],
        ["Target berikutnya","Hardening VPS, observability, backup teruji, object storage opsional, kontrak API/OpenAPI, CI/CD dan UAT lintas perangkat."],
    ],[1.6,4.9])

    new_page(doc); b.add_heading(doc,"6.4 Konteks Aplikasi",2)
    add_figure(doc,"konteks-aplikasi.png","Gambar 6-1. Diagram konteks aplikasi SIVMI CRM")
    table(doc,"Tabel 6-4. Antarmuka konteks",["Pihak/komponen","Pertukaran","Protokol/kontrol"],[
        ["Browser pengguna","Form, tabel, dashboard, chat, notifikasi","HTTP(S), session, CSRF, auth"],
        ["WAHA/GOWS","Session, QR, sendText, webhook event/ack","HTTP API key, HMAC SHA-512"],
        ["MySQL","Master, transaksi, audit, message, jobs","PDO, FK, transaction"],
        ["Redis","Cache/queue sesuai environment","Password produksi, AOF"],
        ["Filesystem","Build asset, CSV, storage privat/publik, log","Flysystem, storage link"],
        ["Operator","Worker, scheduler, migration, backup, health check","Artisan, Docker Compose"],
    ],[1.35,2.7,2.45])

    b.add_heading(doc,"6.5 Prinsip Arsitektur Aplikasi",2)
    principles=[
        ["P-01","Server-side authorization","Menu bukan boundary; resource/query memvalidasi akses."],
        ["P-02","Cabang sebagai data boundary","Setiap query operasional memakai visibleTo atau scope ekuivalen."],
        ["P-03","Single source of truth","Status dan relasi berasal dari database, bukan angka dummy."],
        ["P-04","Thin integration adapter","Detail WAHA berada di WahaClient/WahaSessionService."],
        ["P-05","Async by default untuk event eksternal","Webhook diterima cepat lalu diproses queue."],
        ["P-06","Idempotent processing","payload_hash dan unique external message mencegah duplikasi."],
        ["P-07","Auditability","Perubahan domain, login/logout, dan chat penting memiliki jejak."],
        ["P-08","Configuration over hard-code","Opsi program, service, queue, storage, dan runtime dari config/env."],
        ["P-09","Fail closed","HMAC kosong/invalid ditolak; produksi memerlukan secret dan HTTPS."],
        ["P-10","Progressive deployment","Satu codebase untuk lokal full Docker dan target VPS."],
    ]
    table(doc,"Tabel 6-5. Prinsip dan penerapan",["ID","Prinsip","Penerapan"],principles,[0.65,2.0,3.85])

    b.add_heading(doc,"6.6 Gaya Arsitektur",2)
    b.add_body(doc,"Gaya utama adalah modular monolith berlapis dengan UI server-driven. Semua domain berada dalam satu deployment Laravel dan satu database transaksi, tetapi dipisahkan melalui module/resource, service, enum, model, serta integrasi adapter. Livewire menambahkan interaksi reaktif tanpa memindahkan aturan bisnis ke browser.")
    table(doc,"Tabel 6-6. Evaluasi gaya arsitektur",["Gaya","Status","Alasan"],[
        ["Modular monolith","Dipakai","Transaksi erat, operasi sederhana, perubahan lintas domain masih sering."],
        ["Layered architecture","Dipakai","Memisahkan UI, application, domain, persistence, infrastructure."],
        ["Event-driven internal","Sebagian","Observer, notification, queue job, webhook event."],
        ["Microservices","Tidak dipakai","Belum ada kebutuhan independensi deploy dan overhead operasional belum layak."],
        ["SPA/API-first","Tidak dipakai","Filament/Livewire memenuhi kebutuhan panel internal."],
    ],[1.55,1.1,3.85])

    new_page(doc); b.add_heading(doc,"6.7 Layer dan Tanggung Jawab",2)
    add_figure(doc,"layer-aplikasi.png","Gambar 6-2. Layer logis dan dependensi aplikasi")
    table(doc,"Tabel 6-7. Kontrak tiap layer",["Layer","Komponen","Tanggung jawab","Larangan"],[
        ["Presentation","Filament, Livewire, Blade","Input, view model, feedback, navigation","Tidak menentukan scope hanya di UI"],
        ["Application","Controller, command, observer, notifier","Orkestrasi use case dan side effect","Tidak menggandakan aturan domain"],
        ["Domain service","Matcher, school creator, chat/session service","Aturan reusable dan integrasi use case","Tidak bergantung pada markup"],
        ["Persistence","Model, scope, relationship, query builder","Konsistensi, relasi, transaksi, agregasi","Tidak mengekspos data tanpa scope"],
        ["Infrastructure","DB, Redis, filesystem, WAHA, Docker","Runtime, I/O, external service","Tidak membawa keputusan role"],
    ],[1.1,1.65,2.35,1.4])

    b.add_heading(doc,"6.8 Struktur Komponen",2)
    add_figure(doc,"struktur-komponen.png","Gambar 6-3. Struktur komponen as-built")
    table(doc,"Tabel 6-8. Katalog komponen inti",["Komponen","Lokasi","Peran"],[
        ["AdminPanelProvider","app/Providers/Filament","Panel /admin, theme, menu, middleware, database notification"],
        ["Filament Resources","app/Filament/Resources","CRUD, form/table, policy praktis, scope query"],
        ["Dashboard/Laporan","app/Filament/Pages","Agregasi KPI, filter global, export laporan"],
        ["Model + Enum","app/Models; app/Enums","State, relasi, cast, visibleTo, transition"],
        ["WhatsApp services","app/Services/WhatsApp","Client HTTP, lifecycle session, chat, matching"],
        ["Webhook pipeline","Controller + Middleware + Job","Verifikasi, idempotensi, async processing"],
        ["CrmNotifier/Observers","app/Support; app/Observers","Notifikasi dedupe dan audit perubahan"],
        ["Commands","app/Console/Commands","Reminder dan sinkronisasi session"],
    ],[1.5,1.75,3.25])

    new_page(doc); b.add_heading(doc,"6.9 Peta Modul Aplikasi",2)
    add_figure(doc,"peta-modul.png","Gambar 6-4. Peta modul dan layanan bersama")
    table(doc,"Tabel 6-9. Peta modul fungsional",["Modul","Kapabilitas","Entitas/layanan","Role dominan"],[
        ["Dashboard","Filter, KPI, tren, ranking, top sekolah/kota","Query Builder, target","Semua sesuai scope"],
        ["Leads","CRUD, status, sumber, sekolah inline, import/export","Lead, School creator","Admin/Staff"],
        ["Follow-up","Histori, hasil, jadwal berikutnya, reminder","FollowupRiwayat","Admin/Staff"],
        ["Inbox WhatsApp","Conversation, unread, chat, match lead","WA services/models","Admin/Staff"],
        ["Pendaftaran","Konversi lead dan program","Pendaftaran, Lead","Admin/Staff"],
        ["Pembayaran","Pencatatan dan verifikasi","Pembayaran","Admin/Superadmin"],
        ["Program","Klasifikasi, nama otomatis, harga/diskon","ProgramBimbel, settings","Superadmin"],
        ["Target/Laporan","Target cabang/user, progres, export","Target models, queries","Manajemen"],
        ["Master/Kontrol","Cabang, sekolah, pengguna, settings, audit, akun WA","Master + control models","Superadmin/Admin"],
    ],[1.15,2.05,2.0,1.3])

    b.add_heading(doc,"6.10 Alur Request Web",2)
    add_figure(doc,"alur-request-web.png","Gambar 6-5. Alur request web dan pembaruan Livewire")
    add_steps(doc,[
        "Browser mengakses /admin atau mengirim action Livewire; Nginx hanya melayani public asset dan meneruskan PHP.",
        "Laravel memuat route, cookie terenkripsi, session, CSRF, binding, dan middleware autentikasi Filament.",
        "Resource/Page membangun schema, memvalidasi input, serta memanggil authorization dan query scope.",
        "Service atau model menjalankan aturan domain; operasi majemuk menggunakan transaction.",
        "MySQL/Redis/filesystem memproses I/O; observer dapat menghasilkan audit atau notifikasi.",
        "Respons HTML/payload Livewire dikirim; error validasi kembali ke form dan exception dicatat.",
    ])
    table(doc,"Tabel 6-10. Kontrol request",["Tahap","Kontrol","Kegagalan"],[
        ["Edge/web","try_files, public root, body 20 MB, security headers","404/413/502"],
        ["Framework","Session, CSRF, auth, binding","419/401/403/404"],
        ["Form/action","Required, enum, FK, role, branch","Validation error"],
        ["Domain/data","Transaction, unique/FK, idempotensi","409/422/500 sesuai konteks"],
        ["External","Timeout 15s, retry 2 x 300ms","WAHA unavailable/error"],
    ],[1.2,3.2,2.1])

    new_page(doc); b.add_heading(doc,"6.11 Alur Follow-up dan Transaksi",2)
    add_figure(doc,"alur-followup-transaksi.png","Gambar 6-6. Alur lead, follow-up, pendaftaran, pembayaran, dan KPI")
    table(doc,"Tabel 6-11. Transisi dan side effect",["Peristiwa","Precondition","Perubahan","Side effect"],[
        ["Lead dibuat","Cabang/inputter valid","Status awal Cold","Audit, assignment notification"],
        ["Follow-up pertama","Lead visible","Cold dapat menjadi Warm","Histori, reminder berikutnya"],
        ["Minat meningkat","Validasi petugas","Warm/Hot","Dashboard dan ranking"],
        ["Pendaftaran dibuat","Lead + program valid","Lead menjadi Daftar","Siswa aktif dan target"],
        ["Pembayaran dicatat","Pendaftaran visible","Status Menunggu/Lunas/Terverifikasi","Notifikasi pemeriksaan, omzet"],
        ["Tidak berminat","Alasan operasional tersedia","Status Tidak Berminat","Tetap mempertahankan histori"],
    ],[1.25,1.7,1.6,1.95])

    b.add_heading(doc,"6.12 Arsitektur Data pada Aplikasi",2)
    b.add_body(doc,"Detail model relasional dibahas pada BAB V. Pada layer aplikasi, data diakses melalui Eloquent untuk relasi dan rule per record, Query Builder untuk agregasi dashboard/laporan, serta DB transaction untuk perubahan majemuk. visibleTo menjadi pola baku pembatasan cabang, PIC, dan kepemilikan session.")
    table(doc,"Tabel 6-12. Pola akses data",["Pola","Penggunaan","Kontrol/performa"],[
        ["Eloquent Resource","CRUD master dan transaksi","Relationship eager load, visibleTo"],
        ["Query Builder","Dashboard, laporan, ranking, tren","Aggregate SQL, group, limit, indeks tanggal"],
        ["Transaction","Webhook, pengaturan, pengiriman chat","Atomicity dan konsistensi side effect"],
        ["firstOrCreate","Webhook event, conversation, message","Unique key dan idempotensi"],
        ["chunkById(100)","Reminder follow-up/pembayaran","Batas memori dan query stabil"],
        ["Cache forever + forget","crm_settings","Read cepat, invalidasi eksplisit"],
    ],[1.45,2.25,2.8])
    b.add_callout(doc,"Aturan query","Filter global pada dashboard mempersempit data yang sudah di-scope; filter dari browser tidak boleh memperluas scope pengguna. Query produksi tidak menyimpan angka dummy dan empty state harus mencerminkan hasil database.",b.LIGHT_GREEN)

    b.add_heading(doc,"6.13 Arsitektur API",2)
    add_figure(doc,"arsitektur-api-waha.png","Gambar 6-7. Sequence API dan webhook WAHA")
    table(doc,"Tabel 6-13. Endpoint dan kontrak integrasi",["Arah","Endpoint","Auth/kontrol","Hasil"],[
        ["Inbound","POST /api/webhooks/waha","HMAC SHA-512, validation, payload_hash","202 accepted + duplicate flag"],
        ["Internal web","GET /integrations/waha/accounts/{id}/qr","Session auth + visibility","image/png QR"],
        ["Outbound","GET/POST/PUT /api/sessions/*","X-Api-Key, timeout, retry","Lifecycle session"],
        ["Outbound","POST /api/sendText","X-Api-Key, working account","External message ID/status"],
        ["Health","GET /up; WAHA /health","Internal health check","Status runtime"],
    ],[0.9,2.15,2.05,1.4])
    table(doc,"Tabel 6-14. Event webhook",["Event","Pemrosesan","Status data"],[
        ["session.status","Mapping WORKING/STARTING/FAILED","Account status, connected/seen/error"],
        ["message / message.any","Filter group/broadcast, normalize, match, persist","Conversation, message, lead, unread"],
        ["message.ack","Mapping ERROR/PENDING/SERVER/DEVICE/READ/PLAYED","Message processing_status"],
        ["Lainnya","Ignored aman","Event Ignored"],
    ],[1.45,3.35,1.7])
    b.add_heading(doc,"6.13.1 Keamanan dan Idempotensi API",3)
    b.add_bullets(doc,[
        "Webhook ditolak 503 bila secret belum dikonfigurasi dan 401 bila signature hilang atau tidak valid.",
        "Body mentah dihitung SHA-256 sebagai payload_hash unik sebelum job didispatch.",
        "Message unik per account + external_message_id dan conversation unik per account + chat_id.",
        "API key dan HMAC secret wajib berada di environment/secret store, bukan source control.",
        "Public API partner dan OpenAPI formal belum tersedia; perlu dirancang sebelum integrasi pihak ketiga.",
    ])

    b.add_heading(doc,"6.14 Scheduler, Queue, dan Reminder",2)
    add_figure(doc,"queue-scheduler.png","Gambar 6-8. Arsitektur pekerjaan asynchronous dan terjadwal")
    table(doc,"Tabel 6-15. Katalog proses latar",["Proses","Trigger","Eksekusi","Ketahanan"],[
        ["ProcessWahaWebhook","Event diterima","Queue worker","tries 5; backoff 10/30/60/180; failed_jobs"],
        ["crm:sync-whatsapp-sessions","Setiap 5 menit","Scheduler; tiap akun aktif","withoutOverlapping; report error"],
        ["crm:send-reminders","Setiap jam","Scheduler; chunkById 100","Notifikasi dedupe"],
        ["Import/export Filament","Aksi pengguna","Queue/framework tables","Status, failed rows, owner"],
    ],[1.65,1.35,1.7,1.8])
    b.add_callout(doc,"Kesenjangan operasional","Compose lokal menempatkan queue dan scheduler pada profile workers. Keduanya harus benar-benar diaktifkan; tanpa worker, webhook diterima tetapi inbox tidak terisi. Produksi menjalankan worker dan scheduler sebagai service tetap.",b.LIGHT_AMBER)

    b.add_heading(doc,"6.15 Penyimpanan File",2)
    table(doc,"Tabel 6-16. Kelas penyimpanan",["Kelas","Lokasi/disk","Contoh","Kontrol"],[
        ["Private","storage/app/private","File proses/import sementara","Tidak diekspos langsung"],
        ["Public","storage/app/public","Asset unggahan yang diizinkan","storage:link, URL /storage"],
        ["Build asset","public/build","CSS/JS Vite terkompilasi","Immutable image build produksi"],
        ["Log","storage/logs","Laravel/exception log","Rotasi dan redaksi data"],
        ["Container volume","storage_data, waha_sessions","Storage Laravel dan auth WhatsApp","Persisten di luar container"],
        ["Object storage","S3 config tersedia","Target untuk skala/backup","Belum baseline aktif"],
    ],[1.25,1.55,1.8,1.9])
    b.add_bullets(doc,[
        "Nama file dan MIME type harus divalidasi; file CSV diolah sebagai data, bukan dieksekusi.",
        "Formula spreadsheet dinetralkan pada export untuk mencegah CSV injection.",
        "Backup database, storage, dan session WAHA harus memiliki jadwal, retensi, enkripsi, serta uji restore.",
        "File sementara dan failed import row perlu kebijakan pembersihan berkala.",
    ])

    b.add_heading(doc,"6.16 Presentasi dan UI",2)
    b.add_body(doc,"UI memakai light theme modern dalam bahasa Indonesia. Filament menyediakan form, table, notification, infolist, dan navigation; Tailwind theme mengadaptasi warna SIVMI. Sidebar biru mendukung collapse desktop dan drawer mobile, dengan status collapse disimpan di localStorage. Dashboard dan chat menggunakan Blade/Livewire khusus untuk kebutuhan yang tidak cukup diwakili CRUD standar.")
    table(doc,"Tabel 6-17. Sistem desain dan perilaku",["Area","Standar","Implementasi"],[
        ["Warna","Primary #1476B8; dark #075985; success/warning/danger/purple","Panel colors + theme CSS"],
        ["Surface","Background #F6F8FC, kartu putih, border tipis, shadow lembut","Resource/page/dashboard"],
        ["Bentuk","Radius 12-18 px dan transisi halus","Card, button, input"],
        ["Responsive","Mobile, tablet, laptop, desktop","Grid adaptif, drawer, scroll table"],
        ["Loading","Skeleton untuk query dashboard/chat","State loading Livewire"],
        ["Feedback","Empty, error, validation, database notification","Filament notification/state"],
        ["Aksesibilitas","Label, fokus, kontras, alt text, keyboard","Sebagian; perlu audit WCAG formal"],
    ],[1.2,2.65,2.65])

    b.add_heading(doc,"6.17 Error Handling dan Logging",2)
    table(doc,"Tabel 6-18. Strategi error",["Sumber","Penanganan","Bukti/keluaran"],[
        ["Form/Livewire","Validation rules dan form errors","Pesan Indonesia di field"],
        ["Authorization","403/forbidden, scoped query","Record tidak terlihat/tidak dapat dimutasi"],
        ["API webhook","JSON untuk api/*; 401/503/422","Respons machine-readable"],
        ["WAHA outbound","throw; retry 2; record failed message","last_error dan notification danger"],
        ["Queued job","Status Failed; rethrow ke failed_jobs","attempts, last_error, stack log"],
        ["Scheduler","Catch per account dan report(exception)","Command exit failure + log"],
        ["Model mutation","Activity observer dan redaksi field sensitif","activity_logs"],
    ],[1.25,2.75,2.5])
    b.add_heading(doc,"6.17.1 Logging Aman",3)
    b.add_bullets(doc,[
        "Jangan mencatat API key, HMAC secret, password, raw credential, atau body pesan penuh pada log umum.",
        "Observer mengecualikan password, email, nomor telepon, alamat, catatan, body, raw_payload, dan field sensitif lain dari properties.",
        "Gunakan correlation identifier event/message/job untuk investigasi tanpa menyalin payload personal.",
        "Produksi wajib APP_DEBUG=false, HTTPS, secret unik, rotasi log, dan akses log berbasis peran.",
    ])

    new_page(doc); b.add_heading(doc,"6.18 Pengujian Arsitektur Aplikasi",2)
    table(doc,"Tabel 6-19. Lapisan pengujian",["Lapisan","Cakupan aktual","Risiko yang ditutup"],[
        ["Unit","Contoh dasar, isolasi database","Test harness dan state leakage"],
        ["Feature access","MasterDataAccess, LeadManagement, RegistrationPayment, ProgramTarget","RBAC, branch scope, mutation"],
        ["Feature UI","Dashboard, Settings, Report","Render, filter, menu, Livewire form"],
        ["CSV","LeadCsv, SchoolCsv","Mapping, validation, formula injection"],
        ["Integration","WahaIntegration, chat/webhook flow","HMAC, queue, idempotensi, status"],
        ["Audit/notification","AuditNotification","Observer, dedupe, reminder"],
        ["Build/runtime","npm build, route/config checks, Docker health","Asset dan deploy readiness"],
        ["Visual UAT","Browser QA lintas viewport","Layout, skeleton, empty/error, chat"],
    ],[1.15,2.75,2.6])
    add_steps(doc,[
        "Jalankan test suite pada database test terisolasi dan pastikan tidak memakai data produksi.",
        "Jalankan npm run build dan periksa tidak ada asset atau Tailwind class yang hilang.",
        "Naikkan stack full Docker; cek /up, login, queue, scheduler, database, Redis, dan WAHA health.",
        "Uji alur end-to-end: lead masuk, follow-up, pendaftaran, pembayaran, KPI, webhook, inbox, balasan chat, audit, notifikasi.",
        "Lakukan visual UAT mobile/tablet/laptop/desktop dan catat screenshot serta hasil console/network.",
        "Simpan bukti versi image, migration, test result, dan rollback point sebelum rilis.",
    ])
    b.add_callout(doc,"Kriteria lulus minimum","Seluruh test otomatis lulus; build sukses; migration idempoten; endpoint health sehat; tidak ada error console kritis; akses lintas cabang ditolak; webhook valid masuk satu kali; worker mengisi inbox; dan smoke test transaksi menghasilkan KPI yang benar.",b.LIGHT_GREEN)

    b.add_heading(doc,"6.19 Temuan dan Gap Arsitektur Aplikasi",2)
    gaps=[
        ["G-01","High","Operasional worker/scheduler mudah terlewat pada profile lokal","Sediakan preflight/status check dan start-full sebagai jalur baku."],
        ["G-02","High","Backup/restore dan RPO/RTO belum terbukti otomatis","Otomasi backup terenkripsi dan drill restore berkala."],
        ["G-03","High","VPS belum memiliki rancangan TLS, secret store, firewall, dan monitoring final","Buat baseline hardening dan deployment checklist."],
        ["G-04","Medium","Public API dan OpenAPI belum tersedia","Tentukan use case partner, versioning, auth, rate limit, schema."],
        ["G-05","Medium","Observability masih log/health dasar","Tambah structured log, metric queue, alert, correlation ID."],
        ["G-06","Medium","Policy tersebar pada resource/model helper","Konsolidasikan policy/gate untuk audit akses lebih mudah."],
        ["G-07","Medium","Retensi file, event webhook, chat, dan audit belum formal","Tentukan klasifikasi, retensi, purge, legal hold."],
        ["G-08","Medium","UAT aksesibilitas dan lintas perangkat belum menjadi gate otomatis","Susun matriks viewport, keyboard, kontras, screen reader."],
        ["G-09","Low","S3 tersedia tetapi belum diputuskan","Aktifkan hanya bila kebutuhan skala/backup membenarkan."],
        ["G-10","Low","Modular boundary belum diuji secara statis","Tambahkan aturan dependency/architecture test bila codebase tumbuh."],
    ]
    table(doc,"Tabel 6-20. Register gap dan rekomendasi",["ID","Prioritas","Temuan","Rekomendasi"],gaps,[0.55,0.75,2.5,2.7])
    b.add_heading(doc,"6.19.1 Roadmap Teknis",3)
    table(doc,"Tabel 6-21. Tahapan penguatan",["Tahap","Fokus","Exit criteria"],[
        ["1. Stabilkan lokal","Test, build, worker/scheduler, WAHA, visual UAT","Alur utama hijau dan dapat direproduksi"],
        ["2. Siapkan VPS","Image immutable, TLS, secret, backup, monitoring","Staging menyerupai produksi"],
        ["3. Hardening","Rate limit, policy, retention, alert, restore drill","Risiko high ditutup"],
        ["4. Scale terukur","Redis queue, object storage, tuning DB, horizontal web","Berdasarkan metric, bukan asumsi"],
        ["5. Integrasi lanjut","API versioning/OpenAPI dan webhook partner","Kontrak, auth, SLA, audit tersedia"],
    ],[1.25,2.55,2.7])
    b.add_heading(doc,"6.19.2 Keputusan yang Perlu Disahkan",3)
    b.add_bullets(doc,[
        "RPO/RTO, jadwal backup, pemilik restore, dan masa retensi data/chat/audit.",
        "Topologi VPS, domain/TLS, pemisahan staging-production, dan pengelolaan secret.",
        "SLA WAHA, prosedur reconnect, kepemilikan nomor, dan eskalasi pesan gagal.",
        "Kebijakan API pihak ketiga, rate limit, data yang boleh dibagikan, dan persetujuan keamanan.",
        "Definition of Done yang mewajibkan test, dokumentasi, migration, logging, serta visual UAT.",
    ])

    new_page(doc); b.add_heading(doc,"LAMPIRAN A - TOPOLOGI RUNTIME",1)
    add_figure(doc,"topologi-runtime.png","Gambar A-1. Topologi runtime full Docker untuk lokal dan target VPS")
    table(doc,"Tabel A-1. Service dan tanggung jawab",["Service","Image/runtime","State","Health/operasi"],[
        ["web","Nginx 1.27 Alpine","Stateless","GET /up; proxy ke app"],
        ["app","PHP 8.4 FPM","Storage/cache lokal","Socket 9000; migration terkontrol"],
        ["database","MySQL 8.4","mysql_data","mysqladmin ping"],
        ["redis","Redis 7.4","redis_data/AOF","redis-cli ping"],
        ["queue","Image app","jobs/failed_jobs","queue:work; restart policy"],
        ["scheduler","Image app","Schedule state","schedule:work"],
        ["waha","devlikeapro/waha GOWS","waha_sessions","GET /health"],
        ["node (dev)","Node 22 + Vite","node_modules","HMR; tidak perlu produksi"],
    ],[1.05,1.75,1.55,2.15])
    b.add_callout(doc,"Catatan deployment","Compose produksi membangun image PHP dan Nginx immutable, mematikan dashboard WAHA, mewajibkan secret, APP_DEBUG=false, dan HTTPS. Migrasi produksi harus dijalankan sebagai langkah release terkontrol dengan backup dan rollback point.")

    b.add_heading(doc,"LAMPIRAN B - TRACEABILITY KOMPONEN",1)
    table(doc,"Tabel B-1. Traceability kebutuhan ke komponen dan test",["Kapabilitas","Komponen utama","Kontrol","Test terkait"],[
        ["RBAC/cabang","User, visibleTo, Resource query","Role enum, canManage, scope","MasterDataAccess, LeadManagement"],
        ["Dashboard KPI","Dashboard Page + Query Builder","Filter tervalidasi, aggregate DB","DashboardTest"],
        ["Lead/CSV","Lead Resource, Importer/Exporter","Enum, normalisasi, scope","LeadManagement, LeadCsv"],
        ["Sekolah inline","CreateSchoolFromLead, SchoolNamePrefix","Format nama, NPSN opsional","CreateSchoolFromLead, SchoolCsv"],
        ["Konversi/transaksi","Pendaftaran/Pembayaran Resource","Status dan kepemilikan","RegistrationPayment"],
        ["Program/target","Program/Target Resource, CrmSetting","Nama otomatis, unique period","ProgramTarget, Settings"],
        ["WhatsApp","WahaClient, webhook, job, chat service","HMAC, idempotensi, queue","WahaIntegration"],
        ["Audit/notifikasi","Observers, CrmNotifier, commands","Redaksi, dedupe, reminder","AuditNotification"],
        ["Laporan","Laporan Page","Scope, export setting","ReportTest"],
    ],[1.35,2.25,1.65,1.25])

    b.add_heading(doc,"PENUTUP",1)
    b.add_body(doc,"Arsitektur aplikasi SIVMI CRM telah membentuk fondasi modular monolith yang sesuai untuk operasi CRM bimbel saat ini: UI terpusat, data relasional konsisten, kontrol role/cabang, integrasi WhatsApp asynchronous, audit, notifikasi, serta runtime Docker yang dapat dipindahkan ke VPS. Fokus tahap berikutnya bukan memecah aplikasi menjadi microservice, melainkan memperkuat operability, security, observability, backup/restore, dan quality gate.")
    b.add_callout(doc,"Baseline final BAB VI","Setiap perubahan arsitektur setelah 24 Agustus 2026 harus memperbarui dokumen ini, diagram, migration/API contract, runbook deployment, dan test traceability secara bersamaan.",b.LIGHT_GREEN)

    OUTPUT.parent.mkdir(parents=True,exist_ok=True)
    doc.save(OUTPUT)
    print(f"Created: {OUTPUT}")


if __name__ == "__main__":
    build_document()
