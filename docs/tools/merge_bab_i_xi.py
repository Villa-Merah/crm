from __future__ import annotations

from copy import deepcopy
from io import BytesIO
from pathlib import Path
import sys


PROJECT_ROOT = Path(__file__).resolve().parents[2]
DOC_TOOLS = PROJECT_ROOT / "storage" / "app" / "doc-tools"
sys.path.insert(0, str(DOC_TOOLS))

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Pt, RGBColor


SOURCES = [
    PROJECT_ROOT / "docs" / "BAB I - Pendahuluan SIVMI CRM v1.0.docx",
    PROJECT_ROOT / "docs" / "BAB II - Landasan Teori SIVMI CRM v1.0.docx",
    PROJECT_ROOT / "docs" / "BAB III - Analisa Kebutuhan SIVMI CRM v1.0.docx",
    PROJECT_ROOT / "docs" / "BAB IV - Arsitektur Bisnis SIVMI CRM v1.0.docx",
    PROJECT_ROOT / "docs" / "BAB V - Arsitektur Data SIVMI CRM v1.1.docx",
    PROJECT_ROOT / "docs" / "BAB VI - Arsitektur Aplikasi SIVMI CRM v1.0.docx",
    PROJECT_ROOT / "docs" / "BAB VII - Arsitektur Teknologi SIVMI CRM v1.0.docx",
    PROJECT_ROOT / "docs" / "BAB VIII - Manajemen Debugging SIVMI CRM v1.0.docx",
    PROJECT_ROOT / "docs" / "BAB IX - Deployment, Migrasi, dan Pemeliharaan Sistem SIVMI CRM v1.0.docx",
    PROJECT_ROOT / "docs" / "BAB X - Panduan Penggunaan SIVMI CRM v1.0.docx",
    PROJECT_ROOT / "docs" / "BAB XI - Panduan Administrator SIVMI CRM v1.0.docx",
]

OUTPUT = PROJECT_ROOT / "docs" / "Dokumentasi SIVMI CRM BAB I-XI v1.0 - Tanpa Cover dan Daftar Isi.docx"

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def element_text(element) -> str:
    return "".join(node.text or "" for node in element.iter(qn("w:t"))).strip()


def has_page_break(element) -> bool:
    return bool(element.xpath(".//w:br[@w:type='page']"))


def body_without_cover_and_toc(document: Document):
    elements = list(document.element.body.iterchildren())
    content = []
    cover_removed = False
    skipping_toc = False

    for element in elements:
        if element.tag == qn("w:sectPr"):
            continue

        if not cover_removed:
            if has_page_break(element):
                cover_removed = True
            continue

        text = element_text(element).upper()
        if text == "DAFTAR ISI":
            skipping_toc = True
            continue

        if skipping_toc:
            if has_page_break(element):
                skipping_toc = False
            continue

        content.append(element)

    if not cover_removed:
        raise RuntimeError(f"Pemisah sampul tidak ditemukan pada {document.core_properties.title or 'dokumen sumber'}")

    return content


def merge_missing_styles(destination: Document, source: Document) -> None:
    destination_styles = destination.styles.element
    existing = {style.get(qn("w:styleId")) for style in destination_styles.findall(qn("w:style"))}
    for style in source.styles.element.findall(qn("w:style")):
        style_id = style.get(qn("w:styleId"))
        if style_id not in existing:
            destination_styles.append(deepcopy(style))
            existing.add(style_id)


def copy_relationships(element, source: Document, destination: Document) -> None:
    relationship_attributes = [qn("r:embed"), qn("r:id"), qn("r:link")]
    relationship_map: dict[str, str] = {}

    for node in element.iter():
        for attribute in relationship_attributes:
            source_rid = node.get(attribute)
            if not source_rid or source_rid not in source.part.rels:
                continue

            if source_rid not in relationship_map:
                relationship = source.part.rels[source_rid]
                if relationship.reltype == RT.IMAGE:
                    new_rid, _ = destination.part.get_or_add_image(BytesIO(relationship.target_part.blob))
                elif relationship.is_external:
                    new_rid = destination.part.relate_to(
                        relationship.target_ref,
                        relationship.reltype,
                        is_external=True,
                    )
                else:
                    raise RuntimeError(
                        f"Relasi internal yang belum didukung: {relationship.reltype} ({source_rid})"
                    )
                relationship_map[source_rid] = new_rid

            node.set(attribute, relationship_map[source_rid])


def append_page_break(document: Document) -> None:
    paragraph = OxmlElement("w:p")
    run = OxmlElement("w:r")
    page_break = OxmlElement("w:br")
    page_break.set(qn("w:type"), "page")
    run.append(page_break)
    paragraph.append(run)
    document.element.body.insert(-1, paragraph)


def neutralize_page_furniture(document: Document) -> None:
    section = document.sections[0]
    header = section.header
    header.is_linked_to_previous = False
    paragraph = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
    paragraph.clear()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("SIVMI CRM  |  DOKUMENTASI SISTEM BAB I-XI")
    run.font.name = "Arial"
    run.font.size = Pt(8)
    run.font.bold = True
    run.font.color.rgb = RGBColor(100, 116, 139)


def audit(document: Document) -> None:
    headings = [p.text.strip().upper() for p in document.paragraphs if p.text.strip()]
    toc_headings = [text for text in headings if text == "DAFTAR ISI"]
    if toc_headings:
        raise RuntimeError("Audit gagal: heading DAFTAR ISI masih ditemukan.")

    expected = [
        "BAB I: PENDAHULUAN",
        "BAB II: LANDASAN TEORI",
        "BAB III: ANALISA KEBUTUHAN",
        "BAB IV: ARSITEKTUR BISNIS",
        "BAB V: ARSITEKTUR DATA",
        "BAB VI: ARSITEKTUR APLIKASI",
        "BAB VII: ARSITEKTUR TEKNOLOGI",
        "BAB VIII: MANAJEMEN DEBUGGING",
        "BAB IX: DEPLOYMENT, MIGRASI, DAN PEMELIHARAAN SISTEM",
        "BAB X: PANDUAN PENGGUNAAN",
        "BAB XI: PANDUAN ADMINISTRATOR",
    ]
    cursor = -1
    for title in expected:
        try:
            cursor = headings.index(title, cursor + 1)
        except ValueError as exception:
            raise RuntimeError(f"Audit gagal: heading '{title}' tidak ditemukan atau tidak berurutan.") from exception

    if len(document.inline_shapes) < 20:
        raise RuntimeError("Audit gagal: jumlah diagram/gambar hasil gabungan terlalu sedikit.")


def main() -> None:
    missing = [str(path) for path in SOURCES if not path.exists()]
    if missing:
        raise FileNotFoundError("Dokumen sumber tidak ditemukan:\n" + "\n".join(missing))

    destination = Document(SOURCES[0])
    destination_body = destination.element.body

    for element in list(destination_body.iterchildren()):
        if element.tag != qn("w:sectPr"):
            destination_body.remove(element)

    for index, source_path in enumerate(SOURCES):
        source = Document(source_path)
        merge_missing_styles(destination, source)
        if index > 0:
            append_page_break(destination)

        for source_element in body_without_cover_and_toc(source):
            cloned = deepcopy(source_element)
            copy_relationships(cloned, source, destination)
            destination_body.insert(-1, cloned)

    neutralize_page_furniture(destination)
    destination.core_properties.title = "Dokumentasi SIVMI CRM BAB I-XI"
    destination.core_properties.subject = "Gabungan BAB I sampai BAB XI tanpa sampul dan daftar isi"
    destination.core_properties.author = "SIVMI"
    destination.core_properties.keywords = "SIVMI, CRM, dokumentasi, arsitektur, operasional"

    audit(destination)
    destination.save(OUTPUT)

    verification = Document(OUTPUT)
    audit(verification)
    print(f"Selesai: {OUTPUT}")
    print(f"Paragraf: {len(verification.paragraphs)}")
    print(f"Tabel: {len(verification.tables)}")
    print(f"Gambar: {len(verification.inline_shapes)}")


if __name__ == "__main__":
    main()
