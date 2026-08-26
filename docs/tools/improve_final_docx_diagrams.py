from __future__ import annotations

from copy import deepcopy
from io import BytesIO
from pathlib import Path

from PIL import Image
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[2]
DOCX = ROOT / "docs" / "DOKUMENTASI-FINAL-PROJECT-SIVMI-CRM-v2.0.docx"
TEMP = ROOT / "docs" / "DOKUMENTASI-FINAL-PROJECT-SIVMI-CRM-v2.0.tmp.docx"
ASSETS = ROOT / "docs" / "_diagram_final_assets"

# Ekspor resolusi tinggi dari Diagram Gabungan v1.5 yang telah disunting pengguna.
MANUAL_REPLACEMENTS = {
    "Gambar 4-1. Proses bisnis end-to-end SIVMI CRM": "proses-bisnis-v15.png",
    "Gambar 4-2. Flowmap proses akuisisi hingga monitoring": "flowmap-v15.png",
    "Gambar 4-3. Flowchart keputusan lead hingga pembayaran": "flowchart-v15.png",
    "Gambar 4-4. Arsitektur layer SIVMI CRM": "arsitektur-layer-v15.png",
    "Gambar 4-5. Sitemap panel administrasi SIVMI CRM": "sitemap-v15.png",
    "Gambar 5-1. DFD Level 0 - konteks pertukaran data SIVMI CRM": "dfd-level-0-v15.png",
    "Gambar 5-5. DFD Level 1 - dekomposisi proses CRM": "dfd-level-1-v15.png",
    "Gambar 5-6. DFD Level 1 - session, webhook, pesan, dan pencocokan lead": "dfd-whatsapp-v15.png",
    "Gambar 5-7. Peta domain entitas inti SIVMI CRM": "peta-entitas-v15.png",
}

ERD_REPLACEMENTS = {
    "Gambar 5-8. ERD operasional lead-to-payment": "erd-master-lead-v11.png",
    "Gambar 5-9. ERD target kinerja, pengaturan, notifikasi, dan audit": "erd-kinerja-kontrol-v11.png",
    "Gambar 5-10. ERD akun WhatsApp, conversation, message, event, dan lead": "erd-whatsapp-v11.png",
}

EMU_PER_INCH = 914400
MAX_WIDTH_IN = 6.38
MAX_HEIGHT_IN = 7.55


def paragraph_description(paragraph) -> str | None:
    props = paragraph._p.xpath(".//wp:docPr")
    return props[0].get("descr") if props else None


def image_relationship(paragraph):
    blips = paragraph._p.xpath(".//a:blip")
    if not blips:
        raise ValueError("Paragraf gambar tidak memiliki a:blip")
    rid = blips[0].get(qn("r:embed"))
    return paragraph.part.rels[rid]


def set_drawing_size(paragraph, image_bytes: bytes, max_width=MAX_WIDTH_IN, max_height=MAX_HEIGHT_IN):
    with Image.open(BytesIO(image_bytes)) as image:
        width_px, height_px = image.size
    ratio = width_px / height_px
    width = max_width
    height = width / ratio
    if height > max_height:
        height = max_height
        width = height * ratio
    cx = int(width * EMU_PER_INCH)
    cy = int(height * EMU_PER_INCH)
    for extent in paragraph._p.xpath(".//wp:extent"):
        extent.set("cx", str(cx))
        extent.set("cy", str(cy))
    for extent in paragraph._p.xpath(".//a:xfrm/a:ext"):
        extent.set("cx", str(cx))
        extent.set("cy", str(cy))


def set_alt_text(paragraph, value: str):
    props = paragraph._p.xpath(".//wp:docPr")
    if props:
        props[0].set("descr", value)


def replace_picture(paragraph, asset: Path, alt_text: str | None = None):
    blob = asset.read_bytes()
    relationship = image_relationship(paragraph)
    relationship.target_part._blob = blob
    set_drawing_size(paragraph, blob)
    if alt_text:
        set_alt_text(paragraph, alt_text)


def format_figure_paragraph(paragraph, page_break=False):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(8)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.keep_together = True
    paragraph.paragraph_format.page_break_before = page_break


def find_caption(document: Document, text: str):
    return next((paragraph for paragraph in document.paragraphs if paragraph.text.strip() == text), None)


def set_paragraph_text_preserve_first_run(paragraph, value: str):
    if not paragraph.runs:
        paragraph.add_run(value)
        return
    paragraph.runs[0].text = value
    for run in paragraph.runs[1:]:
        run.text = ""


def insert_picture_after(document: Document, anchor_caption, asset: Path, alt_text: str, caption_text: str):
    image_paragraph = document.add_paragraph()
    image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = image_paragraph.add_run()
    run.add_picture(str(asset), width=Inches(MAX_WIDTH_IN))
    set_alt_text(image_paragraph, alt_text)
    set_drawing_size(image_paragraph, asset.read_bytes())
    format_figure_paragraph(image_paragraph, page_break=True)

    caption = document.add_paragraph(style=anchor_caption.style)
    if anchor_caption.runs:
        template_run = anchor_caption.runs[0]
        caption_run = caption.add_run(caption_text)
        caption_run.bold = template_run.bold
        caption_run.italic = template_run.italic
        caption_run.font.name = template_run.font.name
        caption_run.font.size = template_run.font.size
        caption_run.font.color.rgb = template_run.font.color.rgb
    else:
        caption.add_run(caption_text)
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(2)
    caption.paragraph_format.space_after = Pt(10)
    caption.paragraph_format.keep_together = True

    anchor_caption._p.addnext(image_paragraph._p)
    image_paragraph._p.addnext(caption._p)
    return image_paragraph, caption


def heading_page_break(document: Document, heading: str):
    paragraph = next((p for p in document.paragraphs if p.text.strip() == heading), None)
    if paragraph:
        paragraph.paragraph_format.page_break_before = True
        paragraph.paragraph_format.keep_with_next = True


def main():
    document = Document(DOCX)
    figures = {
        paragraph_description(paragraph): paragraph
        for paragraph in document.paragraphs
        if paragraph_description(paragraph)
    }

    replaced = []
    for description, filename in {**MANUAL_REPLACEMENTS, **ERD_REPLACEMENTS}.items():
        paragraph = figures.get(description)
        if paragraph is None:
            raise RuntimeError(f"Gambar tidak ditemukan: {description}")
        replace_picture(paragraph, ASSETS / filename)
        format_figure_paragraph(paragraph)
        replaced.append(description)

    # Pecah ERD operasional agar kardinalitas tidak bertumpuk pada satu gambar padat.
    old_caption = find_caption(document, "Gambar 5-8. ERD operasional lead-to-payment")
    if old_caption is None:
        raise RuntimeError("Caption Gambar 5-8 tidak ditemukan")
    set_paragraph_text_preserve_first_run(
        old_caption,
        "Gambar 5-8a. ERD master data, pengguna, lead, dan follow-up",
    )
    old_caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    old_figure = figures["Gambar 5-8. ERD operasional lead-to-payment"]
    set_alt_text(old_figure, "Gambar 5-8a. ERD master data, pengguna, lead, dan follow-up")

    insert_picture_after(
        document,
        old_caption,
        ASSETS / "erd-pendaftaran-pembayaran-v11.png",
        "Gambar 5-8b. ERD pendaftaran dan pembayaran",
        "Gambar 5-8b. ERD pendaftaran dan pembayaran",
    )

    # Tambahkan ERD tabel input/output dan runtime untuk melengkapi susunan database.
    wa_caption = find_caption(document, "Gambar 5-10. ERD akun WhatsApp, conversation, message, event, dan lead")
    if wa_caption is None:
        raise RuntimeError("Caption Gambar 5-10 tidak ditemukan")
    insert_picture_after(
        document,
        wa_caption,
        ASSETS / "erd-runtime-import-export-v11.png",
        "Gambar 5-10b. ERD import, export, session, cache, dan queue runtime",
        "Gambar 5-10b. ERD import, export, session, cache, dan queue runtime",
    )

    # Ritme dan posisi seluruh gambar dibuat konsisten, tanpa mengubah isi bab lain.
    for paragraph in document.paragraphs:
        if paragraph._p.xpath(".//w:drawing"):
            format_figure_paragraph(paragraph, paragraph.paragraph_format.page_break_before)
        text = paragraph.text.strip()
        if text.startswith("Gambar "):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_before = Pt(2)
            paragraph.paragraph_format.space_after = Pt(10)
            paragraph.paragraph_format.keep_together = True

    heading_page_break(document, "5.8.2 ERD Operasional")
    heading_page_break(document, "5.8.3 ERD Kinerja dan Kontrol")
    heading_page_break(document, "5.8.4 ERD Integrasi WhatsApp")

    document.core_properties.title = "Dokumentasi Final Project SIVMI CRM v2.0 - Diagram Diperbaiki"
    document.core_properties.subject = "Diagram resolusi tinggi dan ERD terpisah per domain"
    document.save(TEMP)
    TEMP.replace(DOCX)
    print(f"output={DOCX}")
    print(f"replaced={len(replaced)} inserted=2")


if __name__ == "__main__":
    main()
