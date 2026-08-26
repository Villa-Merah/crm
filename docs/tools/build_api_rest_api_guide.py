from __future__ import annotations

from pathlib import Path
import sys

ROOT = Path(__file__).resolve().parents[2]
DOC_TOOLS = ROOT / "storage" / "app" / "doc-tools"
sys.path.insert(0, str(DOC_TOOLS))

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = ROOT / "docs" / "API-DAN-REST-API-SIVMI-CRM-v1.1.docx"
ASSETS = ROOT / "docs" / "_api_rest_assets"

# compact_reference_guide tokens
PAGE_WIDTH = Inches(8.5)
PAGE_HEIGHT = Inches(11)
MARGIN = Inches(1)
HEADER_DISTANCE = Inches(0.492)
FOOTER_DISTANCE = Inches(0.492)
CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_TOP = 80
CELL_BOTTOM = 80
CELL_START = 120
CELL_END = 120
BODY_FONT = "Calibri"
BODY_SIZE = 11
BODY_AFTER = 6
BODY_LINE = 1.25
H1_SIZE, H1_BEFORE, H1_AFTER = 16, 18, 10
H2_SIZE, H2_BEFORE, H2_AFTER = 13, 14, 7
H3_SIZE, H3_BEFORE, H3_AFTER = 12, 10, 5

PRIMARY = "1476B8"
PRIMARY_DARK = "075985"
INK = "0B2545"
MUTED = "64748B"
LIGHT = "F6F8FC"
BLUE_GRAY = "E8EEF5"
PALE_BLUE = "EAF4FB"
PALE_GREEN = "ECFDF5"
PALE_AMBER = "FFFBEB"
PALE_ROSE = "FFF1F2"
GREEN = "10B981"
AMBER = "F59E0B"
ROSE = "F43F5E"
PURPLE = "A855F7"
WHITE = "FFFFFF"
BORDER = "CBD5E1"


def rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


def set_run(run, *, name=BODY_FONT, size=BODY_SIZE, color=INK, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = rgb(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tag = OxmlElement("w:tblHeader")
    tag.set(qn("w:val"), "true")
    tr_pr.append(tag)


def set_cant_split(row):
    row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=BORDER, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=CELL_TOP, start=CELL_START, bottom=CELL_BOTTOM, end=CELL_END):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa: list[int]):
    if sum(widths_dxa) != CONTENT_WIDTH_DXA:
        raise ValueError(f"Lebar tabel harus {CONTENT_WIDTH_DXA} DXA, diterima {sum(widths_dxa)}")
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(CONTENT_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")


def keep_with_next(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    if p_pr.find(qn("w:keepNext")) is None:
        p_pr.append(OxmlElement("w:keepNext"))


def paragraph(doc, text="", *, size=BODY_SIZE, color=INK, bold=False, italic=False,
              align=WD_ALIGN_PARAGRAPH.LEFT, before=0, after=BODY_AFTER, line=BODY_LINE):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    set_run(p.add_run(text), size=size, color=color, bold=bold, italic=italic)
    return p


def heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    keep_with_next(p)
    run = p.add_run(text)
    if level == 1:
        size, before, after, color = H1_SIZE, H1_BEFORE, H1_AFTER, PRIMARY
    elif level == 2:
        size, before, after, color = H2_SIZE, H2_BEFORE, H2_AFTER, PRIMARY
    else:
        size, before, after, color = H3_SIZE, H3_BEFORE, H3_AFTER, PRIMARY_DARK
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.0
    set_run(run, size=size, color=color, bold=True)
    return p


def bullet_list(doc, items: list[str]):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.375)
        p.paragraph_format.first_line_indent = Inches(-0.188)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = BODY_LINE
        set_run(p.add_run(item))


def numbered_list(doc, items: list[str]):
    # Buat instance numbering baru untuk setiap kelompok agar langkah selalu
    # dimulai dari 1. Word dapat meneruskan nomor antarbagian bila hanya
    # mengandalkan style "List Number" yang sama.
    style = doc.styles["List Number"]
    base_num_id = int(style.element.pPr.numPr.numId.val)
    numbering = doc.part.numbering_part.element
    base_num = numbering.num_having_numId(base_num_id)
    abstract_num_id = int(base_num.abstractNumId.val)
    fresh_num = numbering.add_num(abstract_num_id)
    fresh_num.add_lvlOverride(ilvl=0).add_startOverride(1)

    for item in items:
        p = doc.add_paragraph(style="List Number")
        num_pr = p._p.get_or_add_pPr().get_or_add_numPr()
        num_pr.get_or_add_ilvl().val = 0
        num_pr.get_or_add_numId().val = fresh_num.numId
        p.paragraph_format.left_indent = Inches(0.375)
        p.paragraph_format.first_line_indent = Inches(-0.188)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = BODY_LINE
        set_run(p.add_run(item))


def callout(doc, title: str, body: str, *, fill=PALE_BLUE, accent=PRIMARY):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_repeat_header(table.rows[0])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_border(cell, accent, "10")
    set_cell_margins(cell, 130, 150, 130, 150)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.1
    set_run(p.add_run(title + "\n"), size=10.5, bold=True, color=accent)
    set_run(p.add_run(body), size=10.2, color=INK)
    set_table_geometry(table, [CONTENT_WIDTH_DXA])
    paragraph(doc, "", after=0)


def data_table(doc, headers: list[str], rows: list[list[str]], widths_dxa: list[int], *, compact=False):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    header = table.rows[0]
    set_repeat_header(header)
    set_cant_split(header)
    for idx, text in enumerate(headers):
        cell = header.cells[idx]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, PRIMARY_DARK)
        set_cell_border(cell, PRIMARY_DARK)
        set_cell_margins(cell, top=55, bottom=55) if compact else set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        set_run(p.add_run(text), size=8.3 if compact else 9, color=WHITE, bold=True)
    for row_index, values in enumerate(rows):
        row = table.add_row()
        set_cant_split(row)
        for idx, value in enumerate(values):
            cell = row.cells[idx]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            set_cell_shading(cell, WHITE if row_index % 2 == 0 else LIGHT)
            set_cell_border(cell)
            set_cell_margins(cell, top=55, bottom=55) if compact else set_cell_margins(cell)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            set_run(p.add_run(str(value)), size=7.8 if compact else 9.2)
    set_table_geometry(table, widths_dxa)
    paragraph(doc, "", after=0)
    return table


def code_block(doc, lines: list[str]):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_repeat_header(table.rows[0])
    cell = table.cell(0, 0)
    set_cell_shading(cell, "0F172A")
    set_cell_border(cell, "0F172A")
    set_cell_margins(cell, 130, 150, 130, 150)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    for idx, line in enumerate(lines):
        set_run(p.add_run(line + ("\n" if idx < len(lines) - 1 else "")), name="Cascadia Mono", size=7.7, color="E2E8F0")
    set_table_geometry(table, [CONTENT_WIDTH_DXA])
    paragraph(doc, "", after=0)


def note(doc, label: str, text: str):
    p = paragraph(doc, "", after=6)
    set_run(p.add_run(label + ": "), size=10.5, color=PRIMARY_DARK, bold=True)
    set_run(p.add_run(text), size=10.5, color=INK)


def font(size, bold=False):
    candidates = [
        Path("C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf"),
        Path("C:/Windows/Fonts/calibrib.ttf" if bold else "C:/Windows/Fonts/calibri.ttf"),
    ]
    for candidate in candidates:
        if candidate.exists():
            return ImageFont.truetype(str(candidate), size)
    return ImageFont.load_default()


def rounded_box(draw, xy, title, lines, fill, outline, title_color=INK):
    draw.rounded_rectangle(xy, radius=24, fill=fill, outline=outline, width=4)
    x1, y1, x2, y2 = xy
    resolved_title_color = title_color if title_color.startswith("#") else "#" + title_color
    draw.text((x1 + 28, y1 + 24), title, font=font(28, True), fill=resolved_title_color)
    y = y1 + 78
    for line in lines:
        draw.text((x1 + 28, y), line, font=font(21), fill="#" + INK)
        y += 34


def arrow(draw, start, end, color="64748B", width=5):
    fill = "#" + color
    draw.line([start, end], fill=fill, width=width)
    x2, y2 = end
    x1, y1 = start
    import math
    angle = math.atan2(y2 - y1, x2 - x1)
    length = 18
    for offset in (2.55, -2.55):
        draw.line([end, (x2 + length * math.cos(angle + offset), y2 + length * math.sin(angle + offset))], fill=fill, width=width)


def architecture_diagram():
    ASSETS.mkdir(parents=True, exist_ok=True)
    image = Image.new("RGB", (1800, 1040), "#F6F8FC")
    draw = ImageDraw.Draw(image)
    draw.text((70, 45), "Konteks API SIVMI CRM Saat Ini", font=font(40, True), fill="#075985")
    draw.text((70, 98), "Dua arah integrasi WAHA, ditambah route internal berbasis session", font=font(24), fill="#64748B")
    rounded_box(draw, (80, 250, 440, 590), "Pengguna CRM", ["Browser", "Filament + Livewire", "Session + RBAC"], "#FFFFFF", "#1476B8")
    rounded_box(draw, (660, 190, 1140, 660), "SIVMI CRM", ["Web routes internal", "REST client WAHA", "Webhook controller", "Queue + database"], "#EAF4FB", "#075985")
    rounded_box(draw, (1360, 250, 1720, 590), "WAHA GOWS", ["REST endpoints", "Session WhatsApp", "Webhook publisher"], "#ECFDF5", "#10B981")
    arrow(draw, (440, 340), (660, 340), "1476B8")
    draw.text((475, 290), "HTTPS/web", font=font(20, True), fill="#1476B8")
    arrow(draw, (1140, 310), (1360, 310), "A855F7")
    draw.text((1165, 255), "REST + X-Api-Key", font=font(20, True), fill="#7E22CE")
    arrow(draw, (1360, 515), (1140, 515), "F59E0B")
    draw.text((1160, 545), "Webhook + HMAC", font=font(20, True), fill="#B45309")
    rounded_box(draw, (420, 785, 1380, 960), "MySQL + Queue Worker + Private Storage", ["event idempotency | conversation/message | lead matching | media lokal"], "#FFFFFF", "#64748B")
    arrow(draw, (900, 660), (900, 785), "64748B")
    image.save(ASSETS / "konteks-api.png")


def webhook_diagram():
    image = Image.new("RGB", (1800, 1220), "#F6F8FC")
    draw = ImageDraw.Draw(image)
    draw.text((70, 45), "Alur Webhook Pesan Masuk", font=font(40, True), fill="#075985")
    draw.text((70, 98), "Penerimaan cepat, aman, idempotent, lalu diproses asynchronous", font=font(24), fill="#64748B")
    stages = [
        (80, 190, 500, 350, "1. WAHA POST", ["/api/webhooks/waha", "JSON + HMAC SHA-512"], "#EAF4FB", "#1476B8"),
        (690, 190, 1110, 350, "2. Middleware", ["Verifikasi signature", "401/503 bila gagal"], "#FFF1F2", "#F43F5E"),
        (1300, 190, 1720, 350, "3. Controller", ["Validasi payload", "Hash raw body"], "#FFFBEB", "#F59E0B"),
        (1300, 500, 1720, 680, "4. Event Store", ["firstOrCreate payload_hash", "Status Pending"], "#FFFFFF", "#64748B"),
        (690, 500, 1110, 680, "5. HTTP 202", ["accepted: true", "duplicate: true/false"], "#ECFDF5", "#10B981"),
        (80, 500, 500, 680, "6. Queue Job", ["5 tries", "backoff 10/30/60/180"], "#EDE9FE", "#A855F7"),
        (80, 850, 500, 1060, "7. Resolve", ["Session -> account", "Event type routing"], "#FFFFFF", "#1476B8"),
        (690, 850, 1110, 1060, "8. Persist", ["Conversation + message", "Media private"], "#FFFFFF", "#075985"),
        (1300, 850, 1720, 1060, "9. CRM Action", ["Match/create lead", "Notification + ACK"], "#ECFDF5", "#10B981"),
    ]
    for x1, y1, x2, y2, title, lines, fill, outline in stages:
        rounded_box(draw, (x1, y1, x2, y2), title, lines, fill, outline)
    for start, end in [((500,270),(690,270)),((1110,270),(1300,270)),((1510,350),(1510,500)),((1300,590),(1110,590)),((690,590),(500,590)),((290,680),(290,850)),((500,955),(690,955)),((1110,955),(1300,955))]:
        arrow(draw, start, end)
    image.save(ASSETS / "alur-webhook.png")


def public_api_diagram():
    image = Image.new("RGB", (1800, 1120), "#F6F8FC")
    draw = ImageDraw.Draw(image)
    draw.text((70, 45), "API Publik v1 SIVMI CRM", font=font(40, True), fill="#075985")
    draw.text((70, 98), "As-built: autentikasi client, ability, scope cabang, idempotensi, dan resource Leads", font=font(24), fill="#64748B")
    rounded_box(draw, (70, 240, 390, 540), "Consumer", ["Mobile / mitra", "Bearer API key", "JSON + HTTPS"], "#FFFFFF", "#1476B8")
    rounded_box(draw, (520, 180, 900, 600), "Middleware API", ["Token hash SHA-256", "Ability read/write", "Role + scope cabang", "Rate limit", "X-Request-ID"], "#EAF4FB", "#1476B8")
    rounded_box(draw, (1030, 180, 1410, 600), "Lead API v1", ["GET list/detail", "POST create", "FormRequest", "API Resource", "Eager loading"], "#ECFDF5", "#10B981")
    rounded_box(draw, (1510, 240, 1730, 540), "MySQL", ["leads", "api_clients", "idempotency"], "#FFFFFF", "#075985")
    arrow(draw, (390, 340), (520, 340), "1476B8")
    draw.text((405, 285), "/api/v1", font=font(20, True), fill="#1476B8")
    arrow(draw, (900, 340), (1030, 340), "10B981")
    arrow(draw, (1410, 340), (1510, 340), "075985")
    rounded_box(draw, (250, 780, 730, 1000), "Menerima data", ["GET /api/v1/leads", "GET /api/v1/leads/{id}", "Pagination + filter allowlist"], "#FFFFFF", "#A855F7")
    rounded_box(draw, (1070, 780, 1550, 1000), "Mengirim data", ["POST /api/v1/leads", "Idempotency-Key wajib", "201 atau replay tersimpan"], "#FFFFFF", "#F59E0B")
    arrow(draw, (710, 600), (520, 780), "A855F7")
    arrow(draw, (1220, 600), (1310, 780), "F59E0B")
    image.save(ASSETS / "api-publik-v1.png")


def configure_document():
    doc = Document()
    section = doc.sections[0]
    section.page_width = PAGE_WIDTH
    section.page_height = PAGE_HEIGHT
    section.top_margin = MARGIN
    section.bottom_margin = MARGIN
    section.left_margin = MARGIN
    section.right_margin = MARGIN
    section.header_distance = HEADER_DISTANCE
    section.footer_distance = FOOTER_DISTANCE

    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), BODY_FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), BODY_FONT)
    normal.font.size = Pt(BODY_SIZE)
    normal.font.color.rgb = rgb(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(BODY_AFTER)
    normal.paragraph_format.line_spacing = BODY_LINE

    for level, size, before, after, color in (
        (1, H1_SIZE, H1_BEFORE, H1_AFTER, PRIMARY),
        (2, H2_SIZE, H2_BEFORE, H2_AFTER, PRIMARY),
        (3, H3_SIZE, H3_BEFORE, H3_AFTER, PRIMARY_DARK),
    ):
        style = doc.styles[f"Heading {level}"]
        style.font.name = BODY_FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), BODY_FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), BODY_FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.0

    for list_style_name in ("List Bullet", "List Number"):
        style = doc.styles[list_style_name]
        style.font.name = BODY_FONT
        style.font.size = Pt(BODY_SIZE)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = BODY_LINE

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run(header.add_run("SIVMI CRM  /  REFERENSI API  /  v1.1"), size=8, color=MUTED, bold=True)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run(footer.add_run("Dokumen internal  |  Halaman "), size=8, color=MUTED)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    return doc


def editorial_cover(doc):
    paragraph(doc, "PANDUAN TEKNIS", size=11, bold=True, color=PRIMARY, align=WD_ALIGN_PARAGRAPH.CENTER, before=68, after=18)
    paragraph(doc, "API DAN REST API", size=29, bold=True, color=PRIMARY_DARK, align=WD_ALIGN_PARAGRAPH.CENTER, after=8, line=1.0)
    paragraph(doc, "Konsep, Arsitektur, Keamanan, dan Penerapan\npada SIVMI CRM", size=16, color=INK, align=WD_ALIGN_PARAGRAPH.CENTER, after=26, line=1.15)
    paragraph(doc, "Dokumentasi as-built integrasi WhatsApp melalui WAHA GOWS dan API publik Leads v1 untuk pertukaran data yang aman, terukur, dan dapat diuji.", size=11.5, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, after=55, line=1.25)
    data_table(doc, ["Atribut", "Nilai"], [
        ["Versi dokumen", "1.1"],
        ["Baseline", "25 Agustus 2026"],
        ["Framework", "Laravel 12.66.0, Filament 5.7.6, Livewire 4.4.0"],
        ["Integrasi utama", "WAHA GOWS + Public Leads API /api/v1"],
        ["Status", "As-built WAHA + API publik Leads v1; target deployment VPS"],
    ], [2700, 6660], compact=True)
    callout(doc, "Pernyataan penting", "API publik v1 untuk resource Leads kini tersedia sebagai template as-built. Pendaftaran, Pembayaran, Program, dan Laporan masih melalui Filament/Livewire dan belum menjadi REST API publik.", fill=PALE_AMBER, accent=AMBER)


def add_image(doc, path: Path, caption: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run()
    run.add_picture(str(path), width=Inches(6.45))
    drawing = run._r.find(qn("w:drawing"))
    if drawing is not None:
        doc_pr = drawing.find(".//" + qn("wp:docPr"))
        if doc_pr is not None:
            doc_pr.set("title", caption)
            doc_pr.set("descr", caption)
    cp = paragraph(doc, caption, size=9, color=MUTED, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=8, line=1.0)
    keep_with_next(p)
    return cp


def build():
    architecture_diagram()
    webhook_diagram()
    public_api_diagram()
    doc = configure_document()
    editorial_cover(doc)

    heading(doc, "1. Ringkasan Eksekutif", 1)
    paragraph(doc, "API adalah kontrak komunikasi antarsistem. REST API adalah bentuk API berbasis HTTP yang mengorganisasi fungsi sebagai resource, menggunakan metode HTTP dan status code secara konsisten, serta tidak menyimpan state percakapan pada server antar-request. Dalam SIVMI CRM, API digunakan untuk integrasi WAHA GOWS dan pertukaran data Leads melalui API publik versi 1.")
    callout(doc, "As-built SIVMI CRM", "SIVMI menjadi consumer REST WAHA, provider webhook WAHA, dan provider API publik Leads v1. Route QR/media tetap berupa web endpoint internal berbasis session/RBAC, sedangkan /api/v1 memakai API client, ability, role, dan scope cabang.", fill=PALE_GREEN, accent=GREEN)
    bullet_list(doc, [
        "Outbound REST: membuat, memperbarui, memulai, menghentikan, dan membaca session WAHA; mengirim teks, gambar, video; mengunduh media; mengambil QR.",
        "Inbound webhook: menerima event message, message.any, message.ack, dan session.status pada POST /api/webhooks/waha.",
        "Public REST: menerima daftar/detail Leads melalui GET dan menerima data lead baru melalui POST /api/v1/leads.",
        "Asynchronous processing: controller mengembalikan HTTP 202, sedangkan pemrosesan bisnis dilakukan oleh queue job.",
        "Data integrity: payload hash, unique external message ID, serta unique account-chat menjaga idempotensi.",
        "Security: API key WAHA, HMAC SHA-512 webhook, session auth QR/media, serta API client bertoken hash dengan ability dan scope cabang untuk /api/v1.",
    ])

    heading(doc, "2. Pengertian API", 1)
    heading(doc, "2.1 Definisi", 2)
    paragraph(doc, "Application Programming Interface (API) adalah batas resmi yang memungkinkan satu perangkat lunak meminta data atau tindakan dari perangkat lunak lain. Batas tersebut berisi kontrak: alamat yang dipanggil, metode, header, format request, aturan validasi, format response, error, autentikasi, dan batas penggunaan.")
    data_table(doc, ["Istilah", "Makna praktis"], [
        ["Provider/server", "Sistem yang menyediakan fungsi atau data melalui API."],
        ["Consumer/client", "Sistem yang memanggil API provider."],
        ["Endpoint", "Kombinasi method dan URL yang mewakili operasi tertentu."],
        ["Request", "Permintaan client: method, URL, header, body, query, dan credential."],
        ["Response", "Jawaban server: status code, header, body, dan metadata."],
        ["Payload", "Data request/response, pada proyek ini terutama JSON atau binary media."],
        ["Contract", "Kesepakatan struktur dan perilaku yang harus stabil bagi kedua sistem."],
        ["SDK/client wrapper", "Lapisan kode yang membungkus detail HTTP; pada proyek ini WahaClient."],
    ], [2400, 6960], compact=True)
    heading(doc, "2.2 API tidak selalu REST", 2)
    paragraph(doc, "API dapat berbentuk library API, operating-system API, RPC, SOAP, GraphQL, event stream, atau REST. Webhook juga merupakan pola API berbasis event: provider memanggil URL consumer ketika suatu kejadian terjadi. Karena itu, istilah API lebih luas daripada REST API.")

    heading(doc, "3. Pengertian REST API", 1)
    paragraph(doc, "REST (Representational State Transfer) adalah gaya arsitektur. Dalam implementasi web, resource diidentifikasi melalui URI, representasi biasanya JSON, dan operasi diekspresikan menggunakan semantik HTTP. REST bukan sekadar endpoint yang mengembalikan JSON; konsistensi resource, method, status, caching, statelessness, dan kontrak error sama pentingnya.")
    heading(doc, "3.1 Prinsip utama", 2)
    data_table(doc, ["Prinsip", "Penjelasan"], [
        ["Client-server", "UI dan penyimpanan/proses server dipisahkan sehingga dapat berkembang mandiri."],
        ["Stateless", "Setiap request membawa konteks autentikasi dan data yang dibutuhkan; server tidak bergantung pada urutan request sebelumnya."],
        ["Uniform interface", "Resource, method, media type, status code, dan hyperlink digunakan konsisten."],
        ["Cacheable", "Response menyatakan apakah boleh disimpan cache; data sensitif biasanya private/no-store."],
        ["Layered system", "Client tidak perlu mengetahui apakah request melewati proxy, gateway, load balancer, atau service lain."],
        ["Code on demand", "Opsional; server dapat mengirim kode yang dieksekusi client, tetapi jarang relevan untuk API JSON."],
    ], [2500, 6860], compact=True)
    heading(doc, "3.2 Resource dan URI", 2)
    paragraph(doc, "URI REST memakai kata benda plural, bukan nama aksi. Implementasi saat ini menggunakan /api/v1/leads dan /api/v1/leads/{lead}. Kandidat berikutnya adalah /api/v1/leads/{id}/follow-ups dan /api/v1/registrations. Command non-CRUD dapat memakai subresource, misalnya /api/v1/whatsapp-accounts/{id}/actions/start.")

    heading(doc, "4. Metode HTTP dan Idempotensi", 1)
    data_table(doc, ["Method", "Tujuan", "Safe", "Idempotent", "Contoh"], [
        ["GET", "Membaca resource", "Ya", "Ya", "GET /api/sessions/{name}"],
        ["POST", "Membuat resource atau menjalankan command", "Tidak", "Tidak secara default", "POST /api/sendText"],
        ["PUT", "Mengganti/update representasi pada URI tertentu", "Tidak", "Ya", "PUT /api/sessions/{name}"],
        ["PATCH", "Mengubah sebagian field", "Tidak", "Bergantung kontrak", "PATCH /api/v1/leads/{id}"],
        ["DELETE", "Menghapus resource", "Tidak", "Ya", "DELETE /api/v1/leads/{id}"],
    ], [1000, 2300, 900, 1750, 3410], compact=True)
    note(doc, "Safe", "Tidak dimaksudkan mengubah state bisnis. Idempotent berarti request yang sama diulang menghasilkan state akhir yang sama, bukan berarti response harus identik.")

    heading(doc, "5. Status Code, Header, dan Format Data", 1)
    heading(doc, "5.1 Status code", 2)
    data_table(doc, ["Kelompok", "Arti", "Penggunaan SIVMI / rekomendasi"], [
        ["2xx", "Berhasil/diterima", "200 untuk GET; 201 untuk lead baru; 202 untuk webhook yang masuk antrean."],
        ["400", "Request salah", "Idempotency-Key hilang/tidak valid atau request tidak dapat diproses."],
        ["401", "Tidak terautentikasi", "API token tidak valid, client nonaktif, atau signature webhook salah."],
        ["403", "Tidak berwenang", "Ability, role, cabang, PIC, QR, atau account berada di luar scope."],
        ["404", "Resource tidak ada/tidak terlihat", "Lead/media di luar scope dikembalikan sebagai tidak ditemukan."],
        ["409", "Konflik state", "Idempotency-Key dipakai untuk payload berbeda atau request masih berjalan."],
        ["422", "Validasi gagal", "Field payload atau parameter filter tidak memenuhi kontrak."],
        ["429", "Terlalu banyak request", "Rate limit per route API publik terlampaui."],
        ["503", "Service tidak tersedia", "API publik nonaktif, secret webhook kosong, atau QR belum siap."],
    ], [1300, 2200, 5860], compact=True)
    heading(doc, "5.2 Header penting", 2)
    bullet_list(doc, [
        "Accept: application/json menyatakan format response yang diharapkan oleh WahaClient.",
        "Content-Type: application/json menjelaskan format payload webhook dan REST request.",
        "X-Api-Key membawa credential dari SIVMI ke WAHA; nilainya tidak boleh dicatat pada log.",
        "X-Webhook-Hmac membawa signature raw body; X-Webhook-Hmac-Algorithm harus sha512.",
        "Authorization: Bearer atau X-API-Key membawa token API client publik; database hanya menyimpan hash SHA-256.",
        "Idempotency-Key wajib pada POST /api/v1/leads; replay payload identik menerima respons tersimpan.",
        "X-Request-ID menghubungkan respons API dengan log dan audit operasional.",
        "Cache-Control: no-store/private digunakan pada QR dan media agar data tidak tersimpan di cache publik.",
        "X-Content-Type-Options: nosniff mencegah browser menebak tipe media yang berpotensi berbahaya.",
    ])

    heading(doc, "6. Peta Penerapan API di SIVMI CRM", 1)
    add_image(doc, ASSETS / "konteks-api.png", "Gambar 1. SIVMI sebagai consumer REST WAHA dan provider webhook.")
    add_image(doc, ASSETS / "api-publik-v1.png", "Gambar 2. API publik Leads v1 sebagai jalur pertukaran data eksternal.")
    data_table(doc, ["Arah", "Provider", "Consumer", "Mekanisme"], [
        ["Outbound", "WAHA GOWS", "SIVMI CRM", "REST HTTP + X-Api-Key; JSON/binary"],
        ["Inbound", "SIVMI CRM", "WAHA GOWS", "POST webhook + HMAC SHA-512; response 202"],
        ["Internal", "SIVMI CRM", "Browser user", "Web route + session auth + RBAC; QR/media binary"],
        ["Publik v1", "SIVMI CRM", "Mobile/mitra terdaftar", "Bearer/X-API-Key + ability + scope cabang; JSON"],
    ], [1800, 2400, 2400, 2760], compact=True)
    callout(doc, "Batas implementasi saat ini", "API publik baru mencakup GET list/detail dan POST create untuk Leads. Modul Pendaftaran, Pembayaran, Program, Follow-up, dan Laporan belum diekspos melalui /api/v1.", fill=PALE_AMBER, accent=AMBER)

    heading(doc, "7. Inventaris Endpoint As-Built", 1)
    heading(doc, "7.1 Endpoint yang dimiliki SIVMI", 2)
    data_table(doc, ["Method + path", "Proteksi", "Response", "Fungsi"], [
        ["POST /api/webhooks/waha", "HMAC SHA-512", "202 JSON", "Menerima event WAHA dan memasukkannya ke queue."],
        ["GET /api/v1/leads", "API client + leads:read + rate limit", "200 JSON", "Daftar lead paginated sesuai role dan scope cabang."],
        ["GET /api/v1/leads/{lead}", "API client + leads:read + rate limit", "200/404 JSON", "Detail lead tanpa membocorkan resource di luar scope."],
        ["POST /api/v1/leads", "API client + leads:write + idempotensi", "201/409/422 JSON", "Membuat lead dengan actor audit dan cabang terkontrol."],
        ["GET /integrations/waha/accounts/{account}/qr", "Session auth + operational permission + scope", "200 PNG atau 503 HTML", "Menampilkan QR session WAHA kepada user berhak."],
        ["GET /integrations/waha/messages/{message}/media", "Session auth + conversation visibleTo", "200 binary atau 404", "Mengirim media privat secara inline."],
    ], [3500, 2500, 1300, 2060], compact=True)
    heading(doc, "7.2 Endpoint WAHA yang dipanggil SIVMI", 2)
    data_table(doc, ["Method", "Path WAHA", "Method WahaClient", "Tujuan"], [
        ["GET", "/api/sessions/{name}", "getSession", "Membaca status session."],
        ["POST", "/api/sessions", "createSession", "Membuat session tanpa auto-start."],
        ["PUT", "/api/sessions/{name}", "updateSession", "Memperbarui konfigurasi webhook."],
        ["POST", "/api/sessions/{name}/start", "startSession", "Memulai session."],
        ["POST", "/api/sessions/{name}/stop", "stopSession", "Menghentikan session."],
        ["POST", "/api/sessions/{name}/restart", "restartSession", "Memulihkan session gagal."],
        ["POST", "/api/sendText", "sendText", "Mengirim pesan teks."],
        ["POST", "/api/sendImage", "sendImage", "Mengirim gambar Base64."],
        ["POST", "/api/sendVideo", "sendVideo", "Mengirim video Base64 dengan convert=true."],
        ["GET", "path media dari payload", "downloadMedia", "Mengunduh media inbound dari host WAHA."],
        ["GET", "/api/{session}/auth/qr", "qrCode", "Mendapatkan QR image/png."],
    ], [900, 3220, 2100, 3140], compact=True)

    heading(doc, "8. Kontrak Outbound SIVMI ke WAHA", 1)
    heading(doc, "8.1 Konfigurasi HTTP client", 2)
    bullet_list(doc, [
        "Base URL berasal dari services.waha.base_url; default container adalah http://waha:3000.",
        "Semua request menerima JSON dan mengirim X-Api-Key dari WAHA_API_KEY.",
        "Timeout default 15 detik, dapat diubah melalui WAHA_TIMEOUT.",
        "Retry dikonfigurasi dua kali dengan jeda 300 ms dan throw=false pada level retry; setiap method tetap memanggil throw() pada response akhir.",
        "Nama session dan path parameter di-encode dengan rawurlencode untuk mengurangi risiko path injection.",
        "Download media hanya menerima path absolut dari URL terurai; client tidak langsung mengikuti host eksternal payload.",
    ])
    heading(doc, "8.2 Contoh request teks", 2)
    code_block(doc, [
        "POST /api/sendText HTTP/1.1",
        "Host: waha:3000",
        "Accept: application/json",
        "Content-Type: application/json",
        "X-Api-Key: <WAHA_API_KEY>",
        "",
        "{",
        "  \"session\": \"sivmi-user-3\",",
        "  \"chatId\": \"628123456789@c.us\",",
        "  \"text\": \"Halo, jadwal konsultasi tersedia.\",",
        "  \"linkPreview\": true",
        "}",
    ])
    heading(doc, "8.3 Media", 2)
    paragraph(doc, "Gambar atau video dibaca dari upload, diverifikasi MIME type, lalu dikirim sebagai objek file berisi mimetype, filename, dan data Base64. Format yang diterima aplikasi saat ini adalah JPEG, PNG, WebP, MP4, dan MOV. Setelah WAHA menerima media, salinan disimpan pada disk local di whatsapp-media/YYYY/MM dengan nama UUID.")

    heading(doc, "9. Kontrak Inbound Webhook WAHA", 1)
    heading(doc, "9.1 URL dan event", 2)
    paragraph(doc, "WAHA dikonfigurasi mengirim event ke WAHA_WEBHOOK_URL. Event yang didaftarkan adalah message, message.any, message.ack, dan session.status. Kebijakan retry WAHA menggunakan linear delay 2 detik dengan maksimum 8 percobaan.")
    heading(doc, "9.2 Contoh payload", 2)
    code_block(doc, [
        "POST /api/webhooks/waha HTTP/1.1",
        "Content-Type: application/json",
        "X-Webhook-Hmac-Algorithm: sha512",
        "X-Webhook-Hmac: <HMAC_RAW_BODY>",
        "",
        "{",
        "  \"event\": \"message\",",
        "  \"session\": \"sivmi-user-3\",",
        "  \"engine\": \"GOWS\",",
        "  \"payload\": {",
        "    \"id\": \"message-id-example\",",
        "    \"timestamp\": 1787600000,",
        "    \"from\": \"628123456789@c.us\",",
        "    \"fromMe\": false,",
        "    \"body\": \"Halo, saya ingin bertanya program bimbel.\",",
        "    \"pushName\": \"Calon Siswa\",",
        "    \"hasMedia\": false",
        "  }",
        "}",
    ])
    heading(doc, "9.3 Response", 2)
    code_block(doc, [
        "HTTP/1.1 202 Accepted",
        "Content-Type: application/json",
        "",
        "{ \"accepted\": true, \"duplicate\": false }",
    ])
    note(doc, "Makna 202", "Request telah diverifikasi dan disimpan untuk diproses, bukan berarti seluruh proses pencocokan lead dan penyimpanan pesan telah selesai.")

    heading(doc, "10. Keamanan Webhook HMAC", 1)
    paragraph(doc, "HMAC memastikan payload berasal dari pihak yang mengetahui shared secret dan body tidak berubah selama transit. Middleware menghitung hash_hmac('sha512', rawBody, secret), lalu membandingkannya dengan header menggunakan hash_equals agar comparison tahan timing attack.")
    numbered_list(doc, [
        "WAHA membentuk raw JSON body yang akan dikirim.",
        "WAHA menghitung HMAC SHA-512 menggunakan WAHA_WEBHOOK_HMAC_SECRET.",
        "Signature dikirim pada X-Webhook-Hmac dan algoritma pada X-Webhook-Hmac-Algorithm.",
        "SIVMI menghitung ulang HMAC dari request->getContent(), bukan dari JSON yang sudah diparsing.",
        "Jika secret kosong, SIVMI mengembalikan 503. Jika signature hilang/salah, SIVMI mengembalikan 401.",
        "Hanya request valid yang diteruskan ke controller dan database event.",
    ])
    callout(doc, "Aturan secret", "Gunakan nilai acak panjang, berbeda dari WAHA API key. Jangan commit docker/.env, jangan tampilkan secret pada screenshot/log, dan rotasi jika pernah terekspos.", fill=PALE_ROSE, accent=ROSE)

    heading(doc, "11. Alur Pemrosesan Webhook", 1)
    add_image(doc, ASSETS / "alur-webhook.png", "Gambar 3. Pipeline webhook dari penerimaan sampai aksi CRM.")
    heading(doc, "11.1 Controller: cepat dan idempotent", 2)
    paragraph(doc, "Controller memvalidasi event, session, dan payload; membentuk SHA-256 dari raw body; lalu menggunakan firstOrCreate pada payload_hash. Event baru atau event berstatus Failed dikirim ke ProcessWahaWebhook. Controller segera merespons 202 agar WAHA tidak menunggu proses bisnis yang lebih lama.")
    heading(doc, "11.2 Queue job", 2)
    paragraph(doc, "ProcessWahaWebhook mengimplementasikan ShouldQueue, memiliki 5 tries, dan backoff 10, 30, 60, 180 detik. Status event bergerak dari Pending ke Processing, kemudian Processed/Ignored, atau Failed. Error dibatasi 2.000 karakter agar tetap operasional dan tidak membuat kolom tak terkendali.")

    heading(doc, "12. Routing Event dan Status", 1)
    data_table(doc, ["Event", "Proses", "Hasil utama"], [
        ["session.status", "Map status WAHA", "Working/Connecting/Failed/Disconnected, me ID, connected_at, last_seen_at."],
        ["message.ack", "Cari outbound message berdasarkan account + external ID", "Pending, Sent, Delivered, Read, Played, atau Failed; notifikasi bila ERROR."],
        ["message / message.any", "Validasi chat, nomor, media, conversation dan message", "Match/create lead, unread count, notifikasi inbound, timestamp WhatsApp."],
        ["Event lain", "Tidak diproses", "Event menjadi Ignored."],
    ], [1900, 3700, 3760], compact=True)
    heading(doc, "12.1 Filter chat", 2)
    paragraph(doc, "Pesan tanpa chatId, grup @g.us, newsletter, dan status@broadcast diabaikan. Untuk identitas GOWS @lid, sistem memakai SenderAlt atau RecipientAlt agar nomor telepon dapat dinormalisasi dan dicocokkan.")

    heading(doc, "13. Pencocokan dan Pembuatan Lead", 1)
    data_table(doc, ["Kondisi", "match_status", "Tindakan"], [
        ["Satu lead dengan nomor sama dan cabang sama", "Matched", "Hubungkan conversation ke lead yang ada."],
        ["Ada lebih dari satu match atau match beda cabang", "Conflict", "Jangan auto-link; butuh review manual."],
        ["Tidak ada match dan pesan inbound", "NewLead", "Buat lead Cold, kanal WhatsApp, PIC pemilik akun WA."],
        ["Tidak ada match dan pesan outbound", "Unresolved", "Jangan membuat lead otomatis."],
    ], [3600, 1800, 3960], compact=True)
    heading(doc, "13.1 Normalisasi nomor", 2)
    paragraph(doc, "PhoneNormalizer menghapus karakter non-digit, mengubah awalan 0 menjadi 62, awalan 8 menjadi 62, menerima prefix internasional 00, lalu memvalidasi 8-15 digit E.164 tanpa tanda plus. Normalisasi ini menjadi dasar deduplikasi no_wa_normalized.")

    heading(doc, "14. Model Data Integrasi", 1)
    data_table(doc, ["Tabel", "Kunci/constraint", "Peran"], [
        ["user_whatsapp_accounts", "user_id, phone_e164, session_name unique", "Nomor dan session WAHA per admin/staff, cabang, status, error, metadata."],
        ["whatsapp_conversations", "unique account + chat_id", "Percakapan per akun; optional lead; match status; unread count."],
        ["whatsapp_messages", "unique account + external_message_id", "Pesan inbound/outbound, ACK status, body, media, payload aman."],
        ["waha_webhook_events", "payload_hash unique", "Inbox event idempotent, attempts, status proses, last error."],
        ["leads", "index no_wa_normalized dan whatsapp_last_seen_at", "Kanal masuk, jejak pertama/terakhir WhatsApp, deduplikasi."],
        ["api_clients", "key_hash unique; user_id; optional cabang_id", "Identitas machine-to-machine, ability, actor audit, scope cabang, status aktif."],
        ["api_idempotency_keys", "unique api_client_id + key", "Hash request dan respons replay untuk mencegah POST lead ganda."],
    ], [2500, 3100, 3760], compact=True)
    heading(doc, "14.1 Relasi utama", 2)
    bullet_list(doc, [
        "Satu user memiliki maksimum satu user_whatsapp_account; account selalu terkait satu cabang.",
        "Satu account memiliki banyak conversation dan message.",
        "Satu conversation memiliki banyak message dan dapat terhubung ke satu lead atau null.",
        "Satu lead dapat dihubungkan ke banyak conversation/message sepanjang kebijakan bisnis mengizinkan.",
        "Webhook event berdiri sebagai inbox integrasi dan tidak wajib memiliki foreign key langsung ke message.",
        "Satu api_client terkait satu user audit dan opsional satu cabang; client memiliki banyak api_idempotency_keys.",
    ])

    heading(doc, "15. Pengiriman Chat dari CRM", 1)
    heading(doc, "15.1 Validasi dan otorisasi", 2)
    bullet_list(doc, [
        "Teks wajib 1-4.000 karakter; caption media maksimum 4.000 karakter.",
        "Conversation harus terlihat oleh actor melalui scope visibleTo.",
        "Akun WhatsApp harus aktif dan berstatus Working.",
        "Staff hanya dapat memakai akun miliknya; admin cabang dibatasi cabang; role global mengikuti kebijakan aplikasi.",
        "Media disimpan privat dan hanya dilayani jika conversation masih terlihat oleh user yang login.",
    ])
    heading(doc, "15.2 Transaksi dan audit", 2)
    paragraph(doc, "Setelah WAHA menerima pesan, aplikasi membuat whatsapp_messages, memperbarui timestamp conversation, dan menulis activity_logs. Jika pengiriman gagal, sistem membuat record message berstatus Failed, menyimpan error terbatas, dan mengirim notifikasi kepada actor.")

    heading(doc, "16. Idempotensi, Retry, dan Konsistensi", 1)
    data_table(doc, ["Lapisan", "Mekanisme", "Tujuan"], [
        ["Webhook HTTP", "payload_hash SHA-256 unique", "Request event identik tidak membuat job/event ganda."],
        ["Public POST", "Idempotency-Key + request_hash + response replay", "POST lead identik tidak membuat record ganda; payload berbeda menjadi 409."],
        ["Conversation", "unique account + chat_id", "Satu thread WAHA tidak menjadi banyak conversation."],
        ["Message", "unique account + external_message_id", "Redelivery tidak menggandakan pesan."],
        ["Outbound HTTP", "timeout 15 s; retry 2 x 300 ms", "Mengatasi kegagalan jaringan singkat."],
        ["WAHA webhook", "linear, delay 2 s, attempts 8", "Mengirim ulang bila CRM sementara tidak dapat dijangkau."],
        ["Queue", "tries 5; backoff 10/30/60/180 s", "Memproses ulang error sementara secara terkendali."],
        ["Database", "DB transaction pada message flow", "Conversation, message, lead, dan event berubah secara konsisten."],
    ], [2100, 3900, 3360], compact=True)
    note(doc, "Perhatian", "POST /api/v1/leads telah mendukung Idempotency-Key. Endpoint mutasi berikutnya, terutama pembayaran dan pesan, wajib memakai pola yang sama sebelum dibuka ke consumer.")

    heading(doc, "17. Error Handling dan Observability", 1)
    data_table(doc, ["Sumber error", "Perilaku saat ini", "Bukti operasional"], [
        ["WAHA HTTP", "Response akhir throw; service mencatat failure dan notifikasi", "last_error, message Failed, Laravel log"],
        ["Webhook auth", "401/503 JSON sebelum controller", "HTTP access log tanpa secret"],
        ["Webhook validation", "422 JSON Laravel", "validation error"],
        ["Queue processing", "Event Failed, attempts++, exception dilempar ulang", "failed_jobs/log/event status"],
        ["QR", "503 recovery view, bukan 500", "account status dan last_error"],
        ["Media access", "404 untuk tidak ada atau tidak berhak", "Tidak membocorkan keberadaan resource"],
        ["Public API auth", "401/403/503 envelope sebelum controller", "error.code, request_id, X-Request-ID"],
        ["Public API validation", "422 field errors; 409 konflik idempotensi", "response tersimpan per client + key"],
    ], [2400, 4200, 2760], compact=True)
    bullet_list(doc, [
        "Jangan mencatat X-Api-Key, HMAC secret, session cookie, QR body, atau isi pesan penuh pada log umum.",
        "Public API mengirim X-Request-ID dan request_id pada envelope; perlu diteruskan ke structured log agar penelusuran lintas layanan lengkap.",
        "Pantau failed_jobs, event Failed, account status Failed, waktu last_seen_at, dan webhook processing latency.",
        "Gunakan alert untuk lonjakan 401, 422, 429, 5xx, retry, serta antrean yang tidak berkurang.",
    ])

    heading(doc, "18. Konfigurasi Local, Docker, dan Production", 1)
    data_table(doc, ["Variable", "Fungsi", "Contoh aman"], [
        ["WAHA_BASE_URL", "Alamat REST WAHA bagi Laravel", "http://waha:3000 di Docker"],
        ["WAHA_API_KEY", "Credential raw bagi WahaClient", "<secret-raw>"],
        ["WAHA_API_KEY_HASH", "Hash credential pada container WAHA production", "sha512:<128-hex>"],
        ["WAHA_WEBHOOK_URL", "URL publik/cross-container untuk callback", "https://crm.example.com/api/webhooks/waha"],
        ["WAHA_WEBHOOK_HMAC_SECRET", "Shared secret signature webhook", "<random-secret>"],
        ["WAHA_TIMEOUT", "Timeout HTTP client dalam detik", "15"],
        ["PUBLIC_API_ENABLED", "Feature gate API publik", "false secara default; true setelah client siap"],
        ["PUBLIC_API_RATE_LIMIT", "Batas request per menit/route", "60"],
        ["PUBLIC_API_MAX_PER_PAGE", "Batas maksimum pagination", "100"],
        ["PUBLIC_API_IDEMPOTENCY_TTL_HOURS", "Masa simpan respons replay", "24"],
    ], [2800, 3600, 2960], compact=True)
    callout(doc, "Hybrid lokal", "Dalam Docker lokal, Laravel memanggil http://waha:3000 dan WAHA memanggil http://web/api/webhooks/waha. Browser tetap mengakses CRM di http://localhost:8000 dan dashboard WAHA di http://localhost:3000.", fill=PALE_BLUE, accent=PRIMARY)
    heading(doc, "18.1 Aktivasi API publik", 2)
    numbered_list(doc, [
        "Jalankan php artisan migrate untuk membuat api_clients dan api_idempotency_keys.",
        "Buat token dengan php artisan api-client:create, tentukan user audit, cabang, dan ability minimum.",
        "Simpan raw token sekali pada secret manager; database hanya menyimpan key_hash SHA-256 dan prefix.",
        "Ubah PUBLIC_API_ENABLED=true, lalu jalankan php artisan config:clear.",
        "Uji GET/POST dengan file docs/api/SIVMI-Public-API-v1.http sebelum diberikan kepada consumer.",
    ])
    heading(doc, "18.2 Production", 2)
    bullet_list(doc, [
        "Gunakan HTTPS untuk URL publik dan jangan expose port WAHA langsung ke internet tanpa firewall/reverse proxy.",
        "Matikan dashboard dan Swagger WAHA bila tidak diperlukan; compose production sudah mengarah ke konfigurasi tersebut.",
        "Simpan secret pada secret manager/environment, bukan repository atau dokumentasi.",
        "Batasi ingress webhook hanya pada jalur yang diperlukan dan pertimbangkan IP allowlist jika sumber stabil.",
        "Rotasi API key/HMAC secret secara terencana dan uji sinkronisasi webhook setelah rotasi.",
    ])

    heading(doc, "19. Pengujian API yang Sudah Ada", 1)
    paragraph(doc, "WahaIntegrationTest dan PublicLeadApiTest memakai database SQLite in-memory. Integrasi WAHA menggunakan mock client/fake storage, sedangkan API Leads menguji token hash, ability, role/cabang, idempotensi, dan envelope JSON tanpa ketergantungan internet.")
    data_table(doc, ["Skenario", "Yang dibuktikan"], [
        ["Invalid HMAC", "Webhook menolak request dengan 401."],
        ["Incoming message", "Lead baru, conversation, message, unread count, dan notifikasi terbentuk."],
        ["Duplicate payload", "Event/pesan tidak digandakan; response duplicate=true."],
        ["GOWS LID", "SenderAlt dipakai untuk menemukan nomor telepon."],
        ["Duplicate/cross-branch lead", "Conversation menjadi Conflict dan tidak auto-link."],
        ["RBAC conversation", "Staff hanya melihat conversation akun miliknya."],
        ["Send text/media", "Outbound record, private file, dan authorization media."],
        ["Failed outbound", "Message Failed dan notifikasi dikirim."],
        ["ACK", "Status berubah hingga Read/Failed sesuai ackName."],
        ["Webhook config", "URL, event, HMAC, dan retry dikirim ke WAHA."],
        ["QR recovery", "Session Failed direstart; unavailable menjadi 503 terkontrol."],
    ], [3500, 5860], compact=True)
    heading(doc, "19.1 Pengujian API publik Leads", 2)
    data_table(doc, ["Skenario", "Yang dibuktikan"], [
        ["Tanpa token", "Request ditolak 401 dengan error code dan X-Request-ID."],
        ["Scope cabang", "Client hanya menerima lead yang diizinkan user audit dan cabangnya."],
        ["Ability", "Client read-only tidak dapat melakukan POST."],
        ["Create + replay", "POST pertama membuat lead; replay key/payload sama mengembalikan data yang sama tanpa duplikasi."],
        ["Konflik idempotensi", "Key sama dengan payload berbeda dikembalikan sebagai 409."],
        ["Cross branch", "Client cabang tidak dapat menulis ke cabang lain."],
    ], [3500, 5860], compact=True)
    callout(doc, "Baseline pengujian 25 Agustus 2026", "PublicLeadApiTest: 6 test, 22 assertion. Seluruh suite proyek: 80 test, 472 assertion, seluruhnya lulus.", fill=PALE_GREEN, accent=GREEN)
    heading(doc, "19.2 Perintah test", 2)
    code_block(doc, [
        "php artisan test tests/Feature/WahaIntegrationTest.php",
        "php artisan test tests/Feature/PublicLeadApiTest.php",
        "php artisan test",
        "docker compose --env-file docker\\.env -f docker\\compose.yml exec -T app \\",
        "  php artisan test",
        "php artisan route:list --path=api/v1",
        "php artisan route:list --path=integrations/waha",
    ])

    heading(doc, "20. Gap dan Batasan As-Built", 1)
    data_table(doc, ["Area", "Kondisi saat ini", "Dampak"], [
        ["Public CRUD API", "GET list/detail dan POST Leads tersedia", "Update/delete, follow-up, pendaftaran, pembayaran, program, dan laporan belum publik."],
        ["Versioning", "/api/v1 tersedia; webhook tetap tanpa versi", "Perubahan webhook masih harus dikelola hati-hati."],
        ["OpenAPI", "Spesifikasi 3.1 untuk Leads tersedia", "Belum ada validation/compatibility test otomatis terhadap schema."],
        ["Rate limiting", "Throttle route publik tersedia", "Produksi multi-instance perlu shared limiter dan monitoring 429."],
        ["Replay window", "HMAC valid tanpa timestamp/nonce", "Payload valid lama dapat dikirim ulang, meski dedupe hash menahan duplikasi data."],
        ["Lifecycle token", "Token hash, prefix, active flag, dan CLI tersedia", "Belum ada expiry token, rotasi otomatis, OAuth2, atau portal client."],
        ["Outbound partner", "Belum ada push/webhook publik untuk mitra", "Consumer menerima data dengan pola pull GET, bukan subscription event."],
    ], [2300, 3000, 4060], compact=True)

    heading(doc, "21. API Publik Leads v1 As-Built", 1)
    callout(doc, "Status implementasi", "Template API publik telah aktif pada source tetapi dinonaktifkan secara konfigurasi. API hanya dapat digunakan setelah migration, pembuatan client, penyimpanan token yang aman, dan PUBLIC_API_ENABLED=true.", fill=PALE_GREEN, accent=GREEN)
    heading(doc, "21.1 Endpoint yang tersedia", 2)
    data_table(doc, ["Method", "Endpoint", "Tujuan"], [
        ["GET", "/api/v1/leads", "Daftar lead paginated; filter status, sumber, search, cabang global, dan per_page."],
        ["GET", "/api/v1/leads/{lead}", "Detail lead bila terlihat oleh role, ownership Staff, dan scope cabang."],
        ["POST", "/api/v1/leads", "Membuat lead dengan actor audit, normalisasi nomor, dan Idempotency-Key wajib."],
    ], [900, 4100, 4360], compact=True)
    heading(doc, "21.2 Autentikasi dan otorisasi", 2)
    bullet_list(doc, [
        "api_clients menyimpan name, key_prefix, key_hash SHA-256, abilities, user audit, optional cabang, active, dan last_used_at.",
        "Token diterima melalui Authorization: Bearer atau X-API-Key; raw token hanya ditampilkan sekali oleh api-client:create.",
        "Ability yang tersedia adalah leads:read dan leads:write; wildcard hanya untuk client yang benar-benar dipercaya.",
        "Role Direksi tidak dapat diberi operasi tulis; Staff hanya membaca lead miliknya dan otomatis menjadi PIC lead yang dibuat.",
        "Client bercabang tidak dapat menimpa cabang request; client global wajib mengirim cabang_id aktif untuk POST.",
        "FormRequest mengendalikan validasi dan API Resource membatasi field respons serta eager loading relasi.",
    ])
    heading(doc, "21.3 Mengirim dan menerima data", 2)
    numbered_list(doc, [
        "Consumer membuat API client dan menyimpan token di secret manager.",
        "Untuk menerima data, consumer memanggil GET /api/v1/leads atau detail dengan ability leads:read.",
        "Untuk mengirim data, consumer memanggil POST /api/v1/leads dengan ability leads:write dan Idempotency-Key unik.",
        "Middleware memverifikasi feature gate, token, client aktif, user audit, ability, role, rate limit, dan idempotensi.",
        "Controller menerapkan scope cabang/PIC, menyimpan lead dalam transaksi, lalu mengembalikan 201 dan X-Request-ID.",
        "Replay payload identik mengembalikan respons tersimpan; key sama dengan payload berbeda menghasilkan 409.",
    ])
    code_block(doc, [
        "POST /api/v1/leads HTTP/1.1",
        "Authorization: Bearer <SIVMI_API_TOKEN>",
        "Accept: application/json",
        "Content-Type: application/json",
        "Idempotency-Key: lead-import-20260825-0001",
        "",
        "{",
        "  \"nama\": \"Calon Siswa Contoh\",",
        "  \"no_wa\": \"081234567890\",",
        "  \"status_lead\": \"Cold\",",
        "  \"sumber_lead\": \"Instagram\",",
        "  \"tanggal_masuk\": \"2026-08-25\"",
        "}",
    ])
    heading(doc, "21.4 Kontrak dan berkas pendukung", 2)
    bullet_list(doc, [
        "OpenAPI 3.1: docs/api/openapi-public-v1.yaml.",
        "Request siap pakai: docs/api/SIVMI-Public-API-v1.http.",
        "Panduan aktivasi dan contoh PowerShell/cURL: docs/TEMPLATE-API-PUBLIK-SIVMI.md.",
        "Roadmap berikutnya: PATCH lead, follow-up, program read-only, pendaftaran, pembayaran, laporan, expiry/rotasi token, dan OAuth2 bila partner bertambah.",
    ])

    heading(doc, "22. Contoh Desain Response API v1", 1)
    heading(doc, "22.1 Create success", 2)
    code_block(doc, [
        "HTTP/1.1 201 Created",
        "X-Request-ID: <uuid>",
        "{",
        "  \"data\": {",
        "    \"id\": 141,",
        "    \"nama\": \"Calon Siswa Contoh\",",
        "    \"no_wa\": \"081234567890\",",
        "    \"status\": \"Cold\",",
        "    \"sumber\": \"Instagram\",",
        "    \"cabang\": { \"id\": 1, \"kode\": \"BDG\", \"nama\": \"Bandung\" }",
        "  },",
        "  \"meta\": { \"request_id\": \"<uuid>\" }",
        "}",
    ])
    heading(doc, "22.2 Validation error", 2)
    code_block(doc, [
        "HTTP/1.1 422 Unprocessable Entity",
        "{",
        "  \"error\": {",
        "    \"code\": \"validation_failed\",",
        "    \"message\": \"Data yang diberikan tidak valid.\",",
        "    \"fields\": { \"no_wa\": [\"Nomor WhatsApp tidak valid.\"] },",
        "    \"request_id\": \"<uuid>\"",
        "  }",
        "}",
    ])

    heading(doc, "23. Checklist Implementasi dan Operasional", 1)
    heading(doc, "23.1 Sebelum menambah endpoint", 2)
    bullet_list(doc, [
        "Tetapkan consumer, use case, owner data, scope role/cabang, dan klasifikasi sensitivitas.",
        "Definisikan resource, method, schema, status, error, idempotensi, rate limit, dan versioning.",
        "Pastikan query menggunakan eager loading, pagination, indeks, dan tidak membuka mass assignment berbahaya.",
        "Tambahkan policy tests, validation tests, happy path, duplicate/retry, cross-branch, dan unauthorized tests.",
    ])
    heading(doc, "23.2 Operasional API publik", 2)
    bullet_list(doc, [
        "Pastikan PUBLIC_API_ENABLED hanya true pada environment yang sudah menjalankan migration dan security review.",
        "Berikan ability minimum, gunakan client terpisah per integrasi/cabang, dan nonaktifkan client yang tidak digunakan.",
        "Pantau last_used_at, 401/403/409/422/429/5xx, pola replay, latency, dan pertumbuhan api_idempotency_keys.",
        "Bersihkan idempotency key yang melewati expires_at melalui scheduler sebelum volume membesar.",
        "Uji scope lintas cabang dan role Staff setiap kali endpoint atau policy diperluas.",
    ])
    heading(doc, "23.3 Operasional WAHA", 2)
    bullet_list(doc, [
        "Pantau account Working, last_seen_at, webhook event Failed, failed_jobs, dan message Failed.",
        "Setelah mengubah URL/secret, jalankan Sinkronkan/Aktifkan chat agar config session WAHA diperbarui.",
        "Uji pesan inbound, outbound, ACK Read, image/video, deduplikasi nomor, dan konflik lintas cabang.",
        "Backup database dan storage whatsapp-media; uji restore dan hak akses file.",
        "Jangan menggunakan payload chat nyata sebagai contoh dokumentasi atau fixture publik.",
    ])

    heading(doc, "24. Glosarium", 1)
    data_table(doc, ["Istilah", "Definisi singkat"], [
        ["API", "Kontrak komunikasi antarkomponen atau antarsistem."],
        ["REST", "Gaya arsitektur resource-oriented di atas uniform interface."],
        ["Webhook", "Callback HTTP berbasis event dari provider ke consumer."],
        ["HMAC", "Message authentication code berbasis hash dan shared secret."],
        ["Idempotensi", "Pengulangan request tidak mengubah state akhir lebih dari sekali."],
        ["E.164", "Format nomor telepon internasional yang dinormalisasi."],
        ["ACK", "Acknowledgement status pengiriman/pembacaan pesan."],
        ["RBAC", "Pembatasan akses berdasarkan peran."],
        ["Scope cabang", "Pembatasan data berdasarkan cabang pengguna."],
        ["OpenAPI", "Spesifikasi machine-readable untuk kontrak HTTP API."],
        ["API client", "Identitas machine-to-machine dengan token, ability, user audit, dan scope cabang."],
        ["Ability", "Izin granular seperti leads:read atau leads:write yang melekat pada client."],
    ], [2400, 6960], compact=True)

    paragraph(doc, "Dokumentasi ini menjelaskan kondisi source per 25 Agustus 2026. Secret, nomor aktual, isi percakapan, dan kredensial tidak disertakan.", size=9.5, color=ROSE, bold=True, before=8, after=4)

    doc.add_page_break()
    heading(doc, "Lampiran A — Pemetaan Kode", 1)
    data_table(doc, ["Komponen", "Lokasi", "Komponen", "Lokasi"], [
        ["Route API + webhook", "routes/api.php", "Lead API controller", "app/Http/Controllers/Api/V1/LeadController.php"],
        ["Auth API middleware", "app/Http/Middleware/AuthenticatePublicApi.php", "Idempotency middleware", "app/Http/Middleware/EnsureApiIdempotency.php"],
        ["API requests", "app/Http/Requests/Api/V1", "Lead API Resource", "app/Http/Resources/Api/V1/LeadResource.php"],
        ["API client model", "app/Models/ApiClient.php", "Idempotency model", "app/Models/ApiIdempotencyKey.php"],
        ["API client command", "app/Console/Commands/CreateApiClient.php", "Public API config", "config/public_api.php"],
        ["OpenAPI + contoh", "docs/api/openapi-public-v1.yaml; SIVMI-Public-API-v1.http", "API feature tests", "tests/Feature/PublicLeadApiTest.php"],
        ["HTTP client WAHA", "app/Services/WhatsApp/WahaClient.php", "Webhook controller", "app/Http/Controllers/Webhooks/WahaWebhookController.php"],
        ["HMAC middleware", "app/Http/Middleware/VerifyWahaWebhookHmac.php", "Queue processor", "app/Jobs/ProcessWahaWebhook.php"],
        ["Chat outbound", "app/Services/WhatsApp/WhatsAppChatService.php", "Lead matching", "app/Services/WhatsApp/LeadMatcher.php"],
        ["Migration API", "database/migrations/2026_08_25_000001_create_public_api_tables.php", "WAHA feature tests", "tests/Feature/WahaIntegrationTest.php"],
    ], [1500, 3180, 1500, 3180], compact=True)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
