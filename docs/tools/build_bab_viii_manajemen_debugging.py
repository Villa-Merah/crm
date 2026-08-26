from __future__ import annotations

from pathlib import Path
import sys

PROJECT_ROOT = Path(__file__).resolve().parents[2]
DOC_TOOLS = PROJECT_ROOT / "storage" / "app" / "doc-tools"
sys.path.insert(0, str(DOC_TOOLS))

from PIL import Image, ImageDraw
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

import build_bab_iv_arsitektur_bisnis as b
import build_bab_vii_arsitektur_teknologi as v7


OUTPUT = PROJECT_ROOT / "docs" / "BAB VIII - Manajemen Debugging SIVMI CRM v1.0.docx"
ASSETS = PROJECT_ROOT / "docs" / "_bab_viii_assets"
COVER = ASSETS / "cover-bab-viii.png"

PRIMARY = "#1476B8"
PRIMARY_DARK = "#075985"
INK = "#0B2545"
MUTED = "#64748B"
BG = "#F6F8FC"
GREEN = "#10B981"
AMBER = "#F59E0B"
ROSE = "#F43F5E"
PURPLE = "#A855F7"


def font(name: str, size: int):
    return v7.font(name, size)


def canvas(title: str, subtitle: str, size=(1800, 1050)):
    return v7.canvas(title, subtitle, size)


def box(draw, xy, title, lines, fill="#F8FAFC", outline=PRIMARY,
        title_size=27, body_size=21):
    return v7.box(draw, xy, title, lines, fill, outline, title_size, body_size)


def arrow(draw, start, end, color="#64748B", width=5):
    return v7.arrow(draw, start, end, color, width)


def make_cover() -> None:
    ASSETS.mkdir(parents=True, exist_ok=True)
    image = Image.new("RGB", (1600, 2200), BG)
    draw = ImageDraw.Draw(image)
    for y in range(2200):
        ratio = y / 2200
        start = (20, 118, 184)
        end = (7, 89, 133)
        draw.line((0, y, 1600, y), fill=tuple(round(start[i] * (1 - ratio) + end[i] * ratio) for i in range(3)))
    draw.ellipse((1010, -180, 1770, 580), fill=(27, 139, 199))
    draw.ellipse((1140, 140, 1640, 640), fill=(8, 103, 151))
    draw.rounded_rectangle((120, 250, 1480, 1910), radius=44, fill="white")
    draw.text((190, 350), "SIVMI CRM", font=font("arialbd.ttf", 34), fill=PRIMARY)
    draw.text((190, 600), "BAB VIII", font=font("arialbd.ttf", 48), fill=PRIMARY_DARK)
    draw.multiline_text((190, 690), "MANAJEMEN\nDEBUGGING", font=font("arialbd.ttf", 86), fill=INK, spacing=18)
    draw.text((190, 1045), "Deteksi, reproduksi, analisis akar masalah, perbaikan,", font=font("arial.ttf", 32), fill="#475569")
    draw.text((190, 1097), "pengujian ulang, triage insiden, dan bukti QA", font=font("arial.ttf", 32), fill="#475569")
    draw.rounded_rectangle((190, 1245, 805, 1325), radius=40, fill="#EAF4FB")
    draw.text((230, 1267), "BASELINE VERIFIKASI 25 AGUSTUS 2026", font=font("arial.ttf", 28), fill=PRIMARY_DARK)
    draw.text((190, 1610), "Versi dokumen 1.0", font=font("arial.ttf", 28), fill=MUTED)
    draw.text((190, 1665), "Status: pedoman operasional debugging dan QA", font=font("arial.ttf", 28), fill=MUTED)
    draw.text((190, 1775), "SIVMI CRM - Catat, Kelola, Closing", font=font("arialbd.ttf", 34), fill=PRIMARY)
    image.save(COVER)


def make_debug_cycle() -> None:
    image, draw = canvas("Siklus Debugging dan QA", "Setiap perubahan bergerak dari bukti masalah menuju verifikasi dan pembelajaran", (1800, 1120))
    nodes = [
        (90, 220, 380, 420, "1. Deteksi", "Alarm, laporan user,\ntest, atau log", PRIMARY),
        (520, 120, 810, 320, "2. Reproduksi", "Lingkungan, data,\nlangkah, hasil", PURPLE),
        (980, 120, 1270, 320, "3. Analisis", "Timeline, hipotesis,\nroot cause", AMBER),
        (1410, 220, 1700, 420, "4. Perbaikan", "Patch minimum,\nreview, refactor", GREEN),
        (1410, 690, 1700, 890, "5. Verifikasi", "Test regresi, build,\nsmoke, visual QA", PRIMARY_DARK),
        (980, 790, 1270, 990, "6. Release", "Gate, rollback,\nobservasi", ROSE),
        (520, 790, 810, 990, "7. Penutupan", "Bukti, approval,\ncatatan bug", GREEN),
        (90, 690, 380, 890, "8. Pencegahan", "Test baru, runbook,\nmonitoring", PURPLE),
    ]
    for x1, y1, x2, y2, title, lines, color in nodes:
        box(draw, (x1, y1, x2, y2), title, lines, "#F8FAFC", color, 25, 19)
    path = [(380, 320, 520, 220), (810, 220, 980, 220), (1270, 220, 1410, 320),
            (1555, 420, 1555, 690), (1410, 790, 1270, 890), (980, 890, 810, 890),
            (520, 890, 380, 790), (235, 690, 235, 420)]
    for x1, y1, x2, y2 in path:
        arrow(draw, (x1, y1), (x2, y2), color="#94A3B8", width=4)
    draw.rounded_rectangle((665, 430, 1135, 700), radius=26, fill="#EAF4FB", outline=PRIMARY, width=4)
    draw.multiline_text((900, 565), "Bukti yang dapat diulang\n+ pemilik keputusan\n+ jejak audit", anchor="mm", align="center", spacing=8,
                        font=font("arialbd.ttf", 27), fill=INK)
    image.save(ASSETS / "siklus-debugging-qa.png")


def make_detection() -> None:
    image, draw = canvas("Deteksi dan Identifikasi", "Sinyal dikorelasikan sebelum dinyatakan sebagai bug aplikasi", (1800, 1080))
    sources = [
        (70, 180, 390, 360, "Pengguna", "Keluhan, screenshot,\nwaktu kejadian", PRIMARY),
        (70, 420, 390, 600, "Automated QA", "PHPUnit, build,\nsmoke check", GREEN),
        (70, 660, 390, 840, "Runtime", "Laravel/Nginx/WAHA,\nqueue, database", AMBER),
    ]
    for x1, y1, x2, y2, title, lines, color in sources:
        box(draw, (x1, y1, x2, y2), title, lines, "#F8FAFC", color, 26, 19)
        arrow(draw, (x2, (y1 + y2) // 2), (600, 510), color="#94A3B8", width=4)
    box(draw, (600, 300, 1010, 720), "Korelasi awal", "Timestamp & timezone\nActor/role/cabang\nURL/request/job\nRelease/config\nError + dampak\nData sensitif disamarkan", "#EAF4FB", PRIMARY_DARK, 30, 22)
    box(draw, (1220, 180, 1690, 400), "Insiden valid", "Dapat direproduksi\natau ada bukti kuat", "#ECFDF5", GREEN, 27, 20)
    box(draw, (1220, 470, 1690, 690), "Perlu observasi", "Bukti belum cukup;\ntambah telemetry", "#FFFBEB", AMBER, 27, 20)
    box(draw, (1220, 760, 1690, 960), "Bukan bug", "Konfigurasi, data,\natau penggunaan", "#FFF1F2", ROSE, 27, 20)
    for y in (290, 580, 860):
        arrow(draw, (1010, 510), (1220, y), color="#94A3B8", width=4)
    image.save(ASSETS / "deteksi-identifikasi.png")


def make_reproduction() -> None:
    image, draw = canvas("Kontrak Reproduksi Masalah", "Kasus dianggap dapat ditangani bila konteks dan hasilnya dapat diulang", (1800, 1080))
    items = [
        (80, 180, 500, 410, "Lingkungan", "local / test / staging\nimage tag, PHP, browser", PRIMARY),
        (690, 180, 1110, 410, "Identitas & scope", "role, user, cabang\ntanpa kredensial", PURPLE),
        (1300, 180, 1720, 410, "Data awal", "fixture minimum\nID anonim & state", AMBER),
        (80, 640, 500, 870, "Langkah", "aksi bernomor\nrequest/event yang sama", PRIMARY_DARK),
        (690, 640, 1110, 870, "Expected vs actual", "hasil yang benar\ndan hasil aktual", GREEN),
        (1300, 640, 1720, 870, "Bukti", "log tersanitasi, test,\nscreenshot, correlation ID", ROSE),
    ]
    for x1, y1, x2, y2, title, lines, color in items:
        box(draw, (x1, y1, x2, y2), title, lines, "#F8FAFC", color, 28, 21)
    for x in (290, 900, 1510):
        arrow(draw, (x, 410), (x, 640), color="#94A3B8", width=4)
    draw.rounded_rectangle((540, 450, 1260, 600), radius=24, fill="#EAF4FB", outline=PRIMARY, width=3)
    draw.text((900, 525), "Kasus reproduksi minimum dan tidak merusak data nyata", anchor="mm",
              font=font("arialbd.ttf", 25), fill=INK)
    image.save(ASSETS / "kontrak-reproduksi.png")


def make_rca() -> None:
    image, draw = canvas("Analisis Akar Masalah", "Hipotesis diuji lintas lapisan; gejala tidak disamakan dengan penyebab", (1800, 1120))
    box(draw, (690, 410, 1110, 710), "Kejadian", "Gejala tervalidasi\n+ timeline\n+ dampak", "#FFF1F2", ROSE, 32, 23)
    causes = [
        (80, 150, 500, 350, "Kode", "validation, query,\nstate, concurrency", PRIMARY),
        (80, 780, 500, 980, "Data", "constraint, migrasi,\nduplikasi, format", AMBER),
        (1300, 150, 1720, 350, "Konfigurasi", "env, cache, route,\npermission, secret", PURPLE),
        (1300, 780, 1720, 980, "Infrastruktur", "network, disk, DB,\nqueue, container", GREEN),
        (690, 90, 1110, 300, "Integrasi", "WAHA/API, timeout,\nretry, webhook", PRIMARY_DARK),
        (690, 820, 1110, 1030, "Proses", "gap test, release,\nmonitoring, runbook", ROSE),
    ]
    for x1, y1, x2, y2, title, lines, color in causes:
        box(draw, (x1, y1, x2, y2), title, lines, "#F8FAFC", color, 27, 20)
        arrow(draw, ((x1 + x2) // 2, (y1 + y2) // 2), (900, 560), color="#94A3B8", width=4)
    image.save(ASSETS / "analisis-akar-masalah.png")


def make_fix_pipeline() -> None:
    image, draw = canvas("Perbaikan, Refactoring, dan Release Gate", "Patch harus kecil, dapat ditinjau, dapat diuji, dan dapat dibatalkan", (1800, 1050))
    stages = [
        (60, 270, 330, 620, "Patch", "Perubahan minimum\nfeature flag bila perlu", PRIMARY),
        (410, 270, 680, 620, "Review", "scope, security,\ndata, backward compat", PURPLE),
        (760, 270, 1030, 620, "Verify", "targeted test\nfull regression\nbuild", GREEN),
        (1110, 270, 1380, 620, "Release", "backup, migration,\nsmoke, approval", AMBER),
        (1460, 270, 1740, 620, "Observe", "log, metric, queue,\nuser confirmation", ROSE),
    ]
    for x1, y1, x2, y2, title, lines, color in stages:
        box(draw, (x1, y1, x2, y2), title, lines, "#F8FAFC", color, 29, 20)
    for x in (330, 680, 1030, 1380):
        arrow(draw, (x, 445), (x + 80, 445))
    draw.rounded_rectangle((430, 750, 1370, 910), radius=22, fill="#EAF4FB", outline=PRIMARY, width=3)
    draw.multiline_text((900, 830), "Refactoring terpisah dari hotfix kecuali dibutuhkan untuk menghilangkan akar masalah.\nRollback tersedia sebelum perubahan menyentuh produksi.",
                        anchor="mm", align="center", spacing=7, font=font("arialbd.ttf", 22), fill=INK)
    image.save(ASSETS / "pipeline-perbaikan.png")


def make_test_layers() -> None:
    image, draw = canvas("Lapisan Pengujian Ulang", "Bukti bergerak dari test spesifik menuju alur bisnis dan UAT visual", (1800, 1080))
    layers = [
        (560, 760, 1240, 940, "Full regression", "72 test / 448 assertion - PHPUnit", PRIMARY_DARK),
        (650, 580, 1150, 730, "Feature & policy", "RBAC, cabang, form, CSV, WAHA", PRIMARY),
        (740, 410, 1060, 550, "Unit", "normalizer, isolation, pure rules", GREEN),
        (805, 250, 995, 380, "Targeted", "test bug baru", PURPLE),
    ]
    for x1, y1, x2, y2, title, lines, color in layers:
        box(draw, (x1, y1, x2, y2), title, lines, "#F8FAFC", color, 25, 18)
    box(draw, (70, 300, 430, 560), "Build", "Vite production\nmanifest + asset", "#FFFBEB", AMBER, 28, 20)
    box(draw, (1370, 300, 1730, 560), "Smoke/UAT", "login, dashboard,\ntransaksi, chat, mobile", "#FFF1F2", ROSE, 28, 20)
    arrow(draw, (430, 430), (740, 480)); arrow(draw, (1060, 480), (1370, 430))
    draw.text((900, 1000), "Testing database: SQLite :memory: | cache/session: array | queue: sync", anchor="mm",
              font=font("arialbd.ttf", 22), fill=MUTED)
    image.save(ASSETS / "lapisan-pengujian.png")


def make_triage() -> None:
    image, draw = canvas("Triage Insiden dan Eskalasi", "Severity ditetapkan dari dampak, luas pengguna, kehilangan data, dan workaround", (1800, 1080))
    levels = [
        (70, 180, 430, 850, "SEV-1", "Kritis", "Aplikasi/kanal utama berhenti\nrisiko data/security\ntanpa workaround\n\nRespons awal: 15 menit\nWar room + pimpinan", ROSE),
        (500, 180, 860, 850, "SEV-2", "Tinggi", "Fungsi utama terganggu\nbanyak user/cabang\nworkaround terbatas\n\nRespons awal: 30 menit\nOwner teknis + bisnis", AMBER),
        (930, 180, 1290, 850, "SEV-3", "Sedang", "Fungsi terbatas\ndampak terkendali\nworkaround tersedia\n\nRespons awal: 4 jam kerja\nBacklog prioritas", PRIMARY),
        (1360, 180, 1730, 850, "SEV-4", "Rendah", "Kosmetik/peningkatan\ntidak menghambat operasi\n\nRespons awal: 1 hari kerja\nSprint terencana", GREEN),
    ]
    for x1, y1, x2, y2, code, title, lines, color in levels:
        box(draw, (x1, y1, x2, y2), f"{code} - {title}", lines, "#F8FAFC", color, 29, 20)
    draw.text((900, 950), "Target respons adalah usulan operasional sampai SLA disahkan Direksi", anchor="mm",
              font=font("arialbd.ttf", 23), fill=MUTED)
    image.save(ASSETS / "triage-insiden.png")


def make_evidence_chain() -> None:
    image, draw = canvas("Rantai Bukti Penutupan Bug", "Status Closed hanya diberikan bila bukti teknis dan penerimaan operasional lengkap", (1800, 1050))
    stages = [
        (70, 260, 360, 610, "Laporan", "ID, waktu, dampak,\nreporter", PRIMARY),
        (440, 260, 730, 610, "Reproduksi", "data minimum,\nexpected/actual", PURPLE),
        (810, 260, 1100, 610, "Perbaikan", "commit/change, RCA,\ntest baru", GREEN),
        (1180, 260, 1470, 610, "Verifikasi", "test/build/smoke,\nlingkungan", AMBER),
        (1540, 260, 1770, 610, "Tutup", "approval, release,\nobservasi", ROSE),
    ]
    for x1, y1, x2, y2, title, lines, color in stages:
        box(draw, (x1, y1, x2, y2), title, lines, "#F8FAFC", color, 27, 19)
    for x in (360, 730, 1100, 1470):
        arrow(draw, (x, 435), (x + 80, 435))
    draw.rounded_rectangle((360, 750, 1440, 920), radius=24, fill="#ECFDF5", outline=GREEN, width=3)
    draw.multiline_text((900, 835), "Bukti disimpan tanpa secret/PII mentah dan dapat ditelusuri ke release.\nJika observasi gagal, status kembali Open/Reopened.",
                        anchor="mm", align="center", spacing=7, font=font("arialbd.ttf", 23), fill=INK)
    image.save(ASSETS / "rantai-bukti.png")


def add_figure(doc: Document, name: str, caption: str, width=6.15) -> None:
    b.add_figure(doc, ASSETS / name, caption, width)


def table(doc: Document, caption: str, headers: list[str], rows: list[list[str]], widths: list[float]) -> None:
    b.add_table(doc, headers, rows, widths, caption)


def callout(doc: Document, title: str, body: str, fill: str = b.LIGHT_BLUE) -> None:
    b.add_callout(doc, title, body, fill)
    row = doc.tables[-1].rows[0]
    row._tr.get_or_add_trPr().append(b.OxmlElement("w:cantSplit"))


def page(doc: Document) -> None:
    doc.add_page_break()


def steps(doc: Document, caption: str, items: list[str]) -> None:
    table(doc, caption, ["Langkah", "Uraian"], [[str(i), item] for i, item in enumerate(items, 1)], [0.75, 5.75])


def build_document() -> None:
    for fn in (make_cover, make_debug_cycle, make_detection, make_reproduction, make_rca,
               make_fix_pipeline, make_test_layers, make_triage, make_evidence_chain):
        fn()

    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = Inches(0.85)
    sec.bottom_margin = Inches(0.85)
    sec.left_margin = Inches(1)
    sec.right_margin = Inches(1)
    sec.header_distance = Inches(0.4)
    sec.footer_distance = Inches(0.4)
    b.configure_styles(doc)
    sec.header.paragraphs[0].text = "SIVMI CRM  |  BAB VIII: MANAJEMEN DEBUGGING"
    for run in sec.header.paragraphs[0].runs:
        b.set_run_font(run, 8.5, True, b.PRIMARY)
    b.add_page_number(sec.footer.paragraphs[0])

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    pic = p.add_run().add_picture(str(COVER), width=Inches(6.1))
    pic._inline.docPr.set("descr", "Sampul BAB VIII Manajemen Debugging SIVMI CRM")
    page(doc)

    b.add_heading(doc, "KENDALI DOKUMEN", 1)
    table(doc, "Tabel 0-1. Kendali dan baseline dokumen", ["Atribut", "Nilai"], [
        ["Judul", "BAB VIII: Manajemen Debugging SIVMI CRM"],
        ["Versi", "1.0"],
        ["Tanggal baseline", "25 Agustus 2026"],
        ["Status", "Pedoman operasional terpisah - siap validasi QA dan operasi"],
        ["Ruang lingkup", "Deteksi, reproduksi, RCA, perbaikan, pengujian ulang, perintah diagnostik, logging aman, triage, pemulihan, bukti, dan penutupan bug"],
        ["Baseline aplikasi", "Laravel 12, Filament 5, Livewire 4, PHPUnit 12.5, Vite 8, full Docker, MySQL, Redis, queue, scheduler, dan WAHA/GOWS"],
        ["Bukti terbaru", "72 test dan 448 assertion lulus; build Vite produksi lulus pada 25 Agustus 2026"],
        ["Sumber kebenaran", "Codebase, konfigurasi PHPUnit/logging, test suite, Docker scripts, deployment runbook, migration, route, dan artefak build"],
    ], [1.55, 4.95])
    callout(doc, "Kedudukan dokumen", "BAB VIII mengatur cara masalah ditemukan, dibuktikan, diperbaiki, diuji, dirilis, dan ditutup. Ia tidak menggantikan incident response organisasi, tetapi menyediakan baseline teknis dan bukti yang harus dipakai bersama owner bisnis.")
    b.add_heading(doc, "DAFTAR ISI", 1)
    b.add_toc(doc)
    page(doc)

    b.add_heading(doc, "RINGKASAN EKSEKUTIF", 1)
    b.add_body(doc, "Manajemen debugging SIVMI CRM menggunakan pendekatan evidence-first: gejala harus dikaitkan dengan lingkungan, actor, cabang, data awal, release, dan timestamp sebelum akar masalah ditetapkan. Perbaikan dibuat sekecil mungkin, diuji pada database terisolasi, dilanjutkan full regression dan build, lalu dikonfirmasi melalui smoke test atau UAT sesuai risiko.")
    b.add_body(doc, "Baseline 25 Agustus 2026 mencatat 72 test dengan 448 assertion lulus dalam 102,631 detik serta build Vite produksi berhasil. Test mencakup RBAC, pembatasan cabang, form Filament/Livewire, dashboard, CSV, sekolah dinamis, pendaftaran, pembayaran, pengaturan, audit/notifikasi, laporan, serta integrasi WAHA/GOWS. Angka ini adalah snapshot, bukan jaminan tanpa cacat; setiap perubahan tetap harus menambah atau menyesuaikan test yang relevan.")
    callout(doc, "Keputusan utama", "Bug tidak ditutup hanya karena kode berubah. Status Closed membutuhkan reproduksi terdokumentasi, akar masalah, bukti test, build/smoke sesuai scope, dampak data dinilai, release/rollback diketahui, dan penerimaan owner dicatat.", b.LIGHT_GREEN)
    table(doc, "Tabel 0-2. Baseline QA ringkas", ["Area", "As-built", "Target penguatan"], [
        ["Automated test", "72 test / 448 assertion lulus", "CI wajib dan trend hasil"],
        ["Isolation", "SQLite :memory:, array cache/session, queue sync", "Contract test MySQL/Redis terjadwal"],
        ["Build", "Vite 8.2.1 production build lulus", "Build image + dependency/security scan"],
        ["Logging", "Laravel stack/single/daily; report exception", "Structured log, correlation ID, central retention"],
        ["Runtime", "Health endpoint, Docker status/log scripts", "Queue/scheduler heartbeat + alert"],
        ["Visual QA", "Dilakukan per perubahan utama", "Matriks browser/device dan bukti UAT tetap"],
    ], [1.3, 2.55, 2.65])

    b.add_heading(doc, "BAB VIII: MANAJEMEN DEBUGGING", 1)
    b.add_heading(doc, "8.1 Pengertian", 2)
    b.add_body(doc, "Debugging adalah proses sistematis untuk mengubah gejala menjadi penjelasan sebab-akibat yang dapat dibuktikan, lalu menghilangkan akar masalah tanpa menambah regresi. Dalam SIVMI CRM, objek debugging dapat berupa request web, state Livewire, query Eloquent, scope RBAC/cabang, migration, queue job, scheduler, file CSV, notifikasi, atau integrasi WAHA.")
    table(doc, "Tabel 8-1. Istilah utama", ["Istilah", "Definisi operasional", "Bukti minimum"], [
        ["Bug", "Perilaku aktual menyimpang dari requirement yang sah", "Expected vs actual"],
        ["Incident", "Gangguan layanan yang memerlukan respons terkoordinasi", "Dampak, severity, timeline"],
        ["Root cause", "Kondisi dasar yang bila dihilangkan mencegah kejadian berulang", "Hipotesis teruji"],
        ["Regression", "Fungsi yang sebelumnya benar menjadi rusak setelah perubahan", "Test sebelum/sesudah"],
        ["Workaround", "Cara sementara mengurangi dampak tanpa menghapus penyebab", "Instruksi dan batasan"],
        ["Corrective action", "Perbaikan untuk kejadian saat ini", "Patch dan verifikasi"],
        ["Preventive action", "Kontrol agar masalah tidak berulang atau lebih cepat terdeteksi", "Test, monitor, runbook"],
    ], [1.2, 3.2, 2.1])

    b.add_heading(doc, "8.2 Tujuan", 2)
    b.add_bullets(doc, [
        "Memulihkan fungsi dan melindungi integritas data secepat mungkin sesuai severity.",
        "Memisahkan gejala, faktor pemicu, dan akar masalah berdasarkan bukti yang dapat diulang.",
        "Mencegah regresi melalui targeted test, full regression, build, smoke test, dan UAT proporsional.",
        "Menjaga keamanan dengan tidak mengekspos credential, API key, cookie, isi chat pribadi, atau data siswa pada log dan tiket.",
        "Menyediakan jejak keputusan yang dapat diaudit dari laporan sampai release dan penutupan.",
        "Mengubah insiden menjadi perbaikan permanen pada test, monitoring, dokumentasi, dan proses release.",
    ])
    table(doc, "Tabel 8-2. Sasaran mutu", ["Sasaran", "Indikator", "Kriteria awal"], [
        ["Reproducibility", "Kasus memiliki langkah dan fixture minimum", "Dapat diulang atau ada bukti runtime kuat"],
        ["Correctness", "Test target dan regression", "Semua yang relevan lulus"],
        ["Safety", "Review scope/security/data", "Tidak ada data leak atau destructive step"],
        ["Traceability", "Bug ID ke change dan release", "Rantai bukti lengkap"],
        ["Recovery", "Rollback/workaround", "Dapat dieksekusi sebelum release"],
        ["Learning", "Preventive action", "Test/runbook/monitor diperbarui"],
    ], [1.25, 2.75, 2.5])

    page(doc)
    b.add_heading(doc, "8.3 Prinsip Debugging", 2)
    table(doc, "Tabel 8-3. Prinsip dan penerapan", ["Prinsip", "Penerapan pada SIVMI CRM"], [
        ["Evidence first", "Gunakan timestamp, request/job, release, actor, scope cabang, log, test, dan data minimum."],
        ["Reproduce safely", "Pakai testing/staging dan fixture; jangan bereksperimen pada data produksi."],
        ["One hypothesis at a time", "Ubah satu variabel atau tambahkan instrumentasi sebelum menyimpulkan."],
        ["Fix root cause", "Workaround dapat dipakai untuk pemulihan, tetapi corrective action tetap dicatat."],
        ["Small reversible change", "Patch sempit, migration direview, backup/rollback tersedia."],
        ["Least privilege", "Diagnosis mengikuti RBAC; akses data lintas cabang hanya bila berwenang."],
        ["No secret in evidence", "Redact token, password, cookie, API key, HMAC, raw PII, dan body chat."],
        ["Automate recurrence", "Tambahkan regression test atau monitor untuk pola kegagalan yang ditemukan."],
    ], [1.65, 4.85])

    b.add_heading(doc, "8.4 Siklus Debugging dan QA", 2)
    add_figure(doc, "siklus-debugging-qa.png", "Gambar 8-1. Siklus debugging, QA, release, dan pencegahan")
    steps(doc, "Tabel 8-4. Tahapan siklus", [
        "Terima sinyal dan buat identitas kasus, severity awal, owner, serta waktu kejadian.",
        "Kumpulkan konteks minimum lalu reproduksi pada lingkungan aman.",
        "Bangun timeline, hipotesis lintas lapisan, dan identifikasi akar masalah.",
        "Rancang patch minimum, dampak data, compatibility, keamanan, dan rollback.",
        "Jalankan targeted test, full regression, build, smoke/UAT, serta pemeriksaan log.",
        "Rilis melalui gate yang sesuai severity dan amati periode stabilisasi.",
        "Tutup kasus hanya setelah bukti, penerimaan owner, dan preventive action lengkap.",
    ])

    b.add_heading(doc, "8.5 Deteksi dan Identifikasi", 2)
    add_figure(doc, "deteksi-identifikasi.png", "Gambar 8-2. Alur deteksi, korelasi, dan klasifikasi sinyal")
    table(doc, "Tabel 8-5. Sumber deteksi", ["Sumber", "Contoh", "Tindakan pertama"], [
        ["User/UAT", "Tampilan, alur form, hak akses", "Simpan screenshot, URL, role, device, waktu"],
        ["PHPUnit", "Assertion atau exception", "Jalankan test target dengan output verbose"],
        ["Build", "Vite/asset/dependency error", "Catat versi runtime dan log build"],
        ["Laravel", "Exception, 500, validation", "Korelasi request, user, route, stack trace"],
        ["Docker/Nginx", "502, unhealthy, refused", "Periksa ps, health, logs, port, DNS service"],
        ["Database", "Connection, constraint, slow query", "Cek config efektif, migration, query plan"],
        ["Queue/scheduler", "Job tertunda/gagal, reminder tidak jalan", "Cek worker, failed jobs, schedule state"],
        ["WAHA/GOWS", "Session failed, QR, webhook/chat kosong", "Cek service, session Working, event, HMAC, queue"],
    ], [1.25, 2.45, 2.8])
    callout(doc, "Aturan identifikasi", "Jangan menyebut masalah sebagai bug kode sebelum perbedaan requirement, konfigurasi, data, jaringan, dependency, dan penggunaan telah dinilai. Status awal boleh Investigating sampai bukti cukup.", b.LIGHT_AMBER)

    b.add_heading(doc, "8.6 Reproduksi Masalah", 2)
    add_figure(doc, "kontrak-reproduksi.png", "Gambar 8-3. Kontrak informasi untuk reproduksi yang aman")
    table(doc, "Tabel 8-6. Template reproduksi", ["Elemen", "Isi wajib", "Contoh SIVMI"], [
        ["Lingkungan", "local/test/staging, image/tag, browser", "Full Docker; Chrome; build hash"],
        ["Actor", "role dan scope tanpa credential", "Staff Bandung"],
        ["Data awal", "fixture/ID anonim dan status", "Lead Cold belum punya pendaftaran"],
        ["Langkah", "urutan aksi presisi", "Buka Leads > Buat > isi > simpan"],
        ["Expected", "hasil sesuai rule", "Lead tersimpan pada cabang sendiri"],
        ["Actual", "hasil dan pesan", "Validation error atau record tidak muncul"],
        ["Frequency", "selalu/intermiten dan jumlah percobaan", "3 dari 3 percobaan"],
        ["Evidence", "test/log/screenshot/correlation ID", "Nama test dan timestamp"],
    ], [1.15, 2.6, 2.75])
    b.add_heading(doc, "8.6.1 Isolasi Lingkungan", 3)
    b.add_bullets(doc, [
        "Automated test memakai APP_ENV=testing, SQLite :memory:, CACHE_STORE=array, SESSION_DRIVER=array, QUEUE_CONNECTION=sync, dan MAIL_MAILER=array.",
        "Untuk perilaku khusus MySQL, Redis, queue asynchronous, Nginx, atau WAHA, gunakan integration/smoke test di full Docker atau staging terisolasi.",
        "Jangan menjalankan test dengan konfigurasi yang dapat mengarah ke database development; DatabaseIsolationTest mempertahankan guard ini.",
        "Data produksi hanya boleh direproduksi setelah dianonimkan dan disetujui sesuai kebutuhan akses.",
    ])

    b.add_heading(doc, "8.7 Analisis Akar Masalah", 2)
    add_figure(doc, "analisis-akar-masalah.png", "Gambar 8-4. Domain hipotesis analisis akar masalah")
    steps(doc, "Tabel 8-7. Metode RCA", [
        "Susun timeline dari first bad event, perubahan terakhir, dan saat deteksi.",
        "Bedakan gejala, trigger, contributing factor, dan root cause.",
        "Buat hipotesis pada kode, data, konfigurasi, integrasi, infrastruktur, dan proses.",
        "Uji hipotesis dengan perubahan satu variabel, instrumentation, atau test terkecil.",
        "Gunakan Five Whys secukupnya; hentikan ketika sebab dapat dikontrol dan dibuktikan.",
        "Catat faktor yang tidak terbukti agar penyelidikan tidak berulang.",
        "Tentukan corrective action, preventive action, owner, dan bukti keberhasilan.",
    ])
    table(doc, "Tabel 8-8. Contoh pemetaan gejala", ["Gejala", "Kemungkinan lapisan", "Bukti pembeda"], [
        ["502 Bad Gateway", "Nginx, PHP-FPM, container DNS/startup", "web/app health dan error log"],
        ["Redirect loop", "auth/session/URL/proxy", "route, cookie, APP_URL, headers"],
        ["Skeleton terus tampil", "Livewire error/query/state JS", "network, console, Laravel log, component test"],
        ["Data lintas cabang", "scope/policy/filter tampering", "user role, SQL, visibleTo, policy test"],
        ["Inbox WA kosong", "webhook, HMAC, queue, event parsing", "WAHA event, webhook_event status, failed job"],
        ["QR/session gagal", "WAHA engine/session state", "health, status STARTING/FAILED/WORKING, logs"],
    ], [1.6, 2.25, 2.65])

    b.add_heading(doc, "8.8 Perbaikan Kode dan Refactoring", 2)
    add_figure(doc, "pipeline-perbaikan.png", "Gambar 8-5. Pipeline patch, review, verifikasi, dan observasi")
    table(doc, "Tabel 8-9. Checklist desain perbaikan", ["Aspek", "Pertanyaan wajib"], [
        ["Scope", "Apakah patch hanya menyentuh sebab dan consumer yang relevan?"],
        ["Data", "Apakah migration, constraint, idempotensi, duplicate, atau rollback terdampak?"],
        ["RBAC", "Apakah role dan cabang tetap dibatasi pada query, action, dan UI?"],
        ["Livewire/Filament", "Apakah state path, validation, option query, hydration, dan operation create/edit konsisten?"],
        ["Integration", "Apakah timeout, retry, HMAC, event duplicate, dan failure state aman?"],
        ["Compatibility", "Apakah schema/API/image lama masih dapat berjalan selama rollout?"],
        ["Observability", "Apakah failure baru dapat dilihat tanpa merekam data sensitif?"],
        ["Rollback", "Bisakah release atau data dikembalikan secara aman?"],
    ], [1.55, 4.95])
    b.add_heading(doc, "8.8.1 Kebijakan Refactoring", 3)
    b.add_bullets(doc, [
        "Hotfix memprioritaskan pemulihan dan patch minimum; refactoring luas dibuat sebagai change terpisah kecuali dibutuhkan untuk memperbaiki akar masalah.",
        "Ekstraksi service, normalizer, scope, atau enum dilakukan bila mengurangi duplikasi rule yang terbukti menimbulkan bug.",
        "Perubahan query harus mempertahankan eager loading, aggregate efisien, index/constraint, dan batas role/cabang.",
        "Perubahan migration harus backward-compatible bila release tidak atomik dan wajib memiliki backup/restore path.",
        "Jangan menyamarkan kegagalan dengan catch kosong; report exception, simpan status aman, dan berikan error state yang dapat ditindaklanjuti.",
    ])

    b.add_heading(doc, "8.9 Pengujian Ulang", 2)
    add_figure(doc, "lapisan-pengujian.png", "Gambar 8-6. Lapisan bukti pengujian ulang SIVMI CRM")
    table(doc, "Tabel 8-10. Urutan pengujian", ["Lapisan", "Tujuan", "Gate"], [
        ["Targeted test", "Membuktikan bug gagal sebelum patch dan lulus sesudahnya", "Wajib untuk rule yang dapat diautomasi"],
        ["Related feature", "Menguji consumer dan policy sekitar", "Semua lulus"],
        ["Full regression", "Mendeteksi dampak lintas modul", "72 test/448 assertion baseline; angka mengikuti suite terbaru"],
        ["Build", "Memastikan asset produksi dapat dibangun", "Vite build lulus; warning non-blocking dicatat"],
        ["Integration", "MySQL/Redis/queue/Nginx/WAHA nyata", "Scope risiko terpenuhi"],
        ["Smoke", "Alur kritis setelah start/deploy", "Login, dashboard, transaksi, chat, health"],
        ["Visual UAT", "Responsif dan usability", "Mobile, tablet, laptop, desktop disetujui"],
    ], [1.25, 3.2, 2.05])
    callout(doc, "Snapshot saat dokumen dibuat", "PHPUnit: 72 test dan 448 assertion lulus dalam 102,631 detik. Vite 8.2.1: build lulus dalam 2,12 detik. Terdapat warning opsional tentang paket fontaine untuk optimized font fallbacks; warning tidak menggagalkan build dan harus diperlakukan sebagai improvement terpisah.", b.LIGHT_GREEN)

    b.add_heading(doc, "8.10 Perintah Diagnostik", 2)
    table(doc, "Tabel 8-11. Perintah lokal dan aplikasi", ["Tujuan", "Perintah", "Catatan aman"], [
        ["Full test", "php artisan test", "Memakai phpunit.xml; pastikan isolation guard"],
        ["Docker test", ".\\docker\\test-full.ps1", "Direkomendasikan untuk full Docker Windows"],
        ["Test target", "php artisan test --filter=NamaTest", "Mulai dari kasus terkecil"],
        ["Route", "php artisan route:list", "Verifikasi URI, middleware, action"],
        ["Migration", "php artisan migrate:status", "Read-only terhadap data"],
        ["Config", "php artisan config:show database", "Jangan menyalin secret ke tiket"],
        ["Scheduler", "php artisan schedule:list", "Cek interval/next run"],
        ["Failed jobs", "php artisan queue:failed", "Review sebelum retry/forget"],
        ["Build", "npm run build", "Catat warning dan manifest"],
    ], [1.15, 2.75, 2.6])
    table(doc, "Tabel 8-12. Perintah Docker dan service", ["Tujuan", "Perintah", "Bukti"], [
        ["Status", ".\\docker\\status-full.ps1", "Service, health, port"],
        ["Log app", ".\\docker\\logs-full.ps1 -Service app -Tail 200", "Exception tersanitasi"],
        ["Log queue", ".\\docker\\logs-full.ps1 -Service queue -Tail 200", "Job/retry/failure"],
        ["Log WAHA", ".\\docker\\logs-full.ps1 -Service waha -Tail 200", "Engine/session/webhook"],
        ["Compose config", "docker compose --env-file docker/.env -f docker/compose.yml config", "Config efektif tanpa publikasi secret"],
        ["Health app", "Invoke-WebRequest -UseBasicParsing http://localhost:8000/up", "HTTP status"],
        ["Container exec", "docker compose ... exec app php artisan about", "Runtime efektif"],
    ], [1.2, 3.55, 1.75])
    callout(doc, "Larangan diagnostik", "Jangan menjalankan migrate:fresh, db:wipe, down -v, truncate, reset --hard, atau penghapusan volume sebagai langkah diagnosis tanpa scope dan persetujuan eksplisit. Jangan menempelkan docker/.env, cookie, token, atau payload WhatsApp mentah ke tiket.", b.LIGHT_AMBER)

    b.add_heading(doc, "8.11 Logging Aman", 2)
    table(doc, "Tabel 8-13. Kebijakan logging", ["Boleh dicatat", "Wajib disamarkan/dihindari"], [
        ["Timestamp, level, environment, release/tag", "Password, APP_KEY, WAHA API key/HMAC secret"],
        ["Route/action/job class dan status", "Session cookie, CSRF token, Authorization header"],
        ["User ID internal, role, cabang ID bila perlu", "Nama/telepon/email lengkap bila tidak diperlukan"],
        ["Correlation/request/event ID", "Isi chat WhatsApp dan raw webhook payload umum"],
        ["Exception class dan stack trace terbatas", "SQL binding yang membawa PII/credential"],
        ["Durasi, retry, queue age, HTTP status", "File CSV/import yang berisi data calon siswa"],
    ], [3.25, 3.25])
    b.add_heading(doc, "8.11.1 Implementasi As-built", 3)
    b.add_bullets(doc, [
        "Default channel adalah stack; channel single, daily, monthly, stderr, syslog, Slack, dan Papertrail tersedia melalui konfigurasi.",
        "Channel daily memiliki retensi default 14 file; production baseline menggunakan level yang harus ditetapkan sesuai kebutuhan operasional.",
        "Exception pada dashboard, laporan, audit, WAHA QR/chat/session, dan webhook dilaporkan melalui report() atau disimpan sebagai last_error yang dibatasi.",
        "ActivityLog mencatat event bisnis dan audit; observer menghindari field sensitif serta menerapkan scope role/cabang pada tampilan audit.",
        "Target berikutnya adalah structured logging, correlation ID, redaction processor, central shipping, alert, dan retention yang disahkan.",
    ])

    page(doc)
    b.add_heading(doc, "8.12 Triage Insiden dan Pemulihan", 2)
    add_figure(doc, "triage-insiden.png", "Gambar 8-7. Usulan severity dan target respons awal")
    b.add_body(doc, "Severity dan target respons pada gambar merupakan baseline usulan sampai SLA disahkan. Penetapan mempertimbangkan kehilangan/korupsi data, keamanan, luas role/cabang, kanal WhatsApp, kemampuan transaksi, durasi, serta ketersediaan workaround.")
    table(doc, "Tabel 8-14. Alur triage", ["Fase", "Aktivitas", "Output"], [
        ["Acknowledge", "Catat ID, waktu, reporter, dampak, owner", "Kasus terdaftar"],
        ["Classify", "Tentukan severity, scope, security/data risk", "Prioritas dan eskalasi"],
        ["Contain", "Maintenance, feature off, queue pause, revoke secret bila perlu", "Dampak berhenti meluas"],
        ["Recover", "Rollback, restart terkontrol, restore, reconnect WAHA", "Layanan minimum pulih"],
        ["Correct", "RCA dan patch permanen", "Akar masalah ditangani"],
        ["Verify", "Test, smoke, data reconciliation, observasi", "Bukti stabil"],
        ["Review", "Post-incident dan preventive action", "Pembelajaran tercatat"],
    ], [1.1, 3.25, 2.15])
    b.add_heading(doc, "8.12.1 Runbook Pemulihan Ringkas", 3)
    steps(doc, "Tabel 8-15. Tahapan pemulihan", [
        "Bekukan perubahan non-esensial dan tunjuk incident commander serta pencatat timeline.",
        "Lindungi data dan credential; isolasi komponen terdampak tanpa menghapus bukti.",
        "Pilih containment/workaround yang paling kecil risikonya dan komunikasikan dampak.",
        "Rollback image/config atau restore hanya dari backup terverifikasi dan sesuai approval.",
        "Validasi migration, health, login, RBAC, data, queue, scheduler, dan WAHA sesuai insiden.",
        "Pantau periode stabilisasi, rekonsiliasi data, lalu turunkan severity bila bukti cukup.",
        "Lakukan post-incident review untuk SEV-1/SEV-2 dan kasus data/security.",
    ])

    b.add_heading(doc, "8.13 Temuan Debugging yang Telah Ditangani", 2)
    table(doc, "Tabel 8-16. Register temuan historis", ["ID", "Gejala", "Akar/faktor utama", "Perbaikan dan bukti"], [
        ["DBG-01", "500 dan redirect berulang pada /admin", "Konfigurasi route/auth/session belum konsisten selama bootstrap panel", "Route/login diperbaiki; DashboardTest memverifikasi redirect satu kali dan shell"],
        ["DBG-02", "Login Docker tanpa styling / 502", "Asset/runtime/upstream belum siap atau konfigurasi full Docker belum lengkap", "Entrypoint, Nginx, start/health, dan build diperkuat; full Docker menjadi baseline"],
        ["DBG-03", "Widget dashboard berhenti pada skeleton", "Error/state widget dan data/filter belum selesai dimuat", "Error/empty/loading state dan aggregate diperbaiki; DashboardTest mencakup filter/periode/KPI"],
        ["DBG-04", "Lima test form Filament/Livewire gagal", "Test masih memakai struktur/state path form lama", "Test disesuaikan dengan schema form terbaru; suite 72/448 lulus"],
        ["DBG-05", "Trend leads/pendaftaran kosong", "Pembentukan periode dan dataset grafik tidak sinkron dengan filter semua bulan", "Range periode dan aggregate diperbaiki; test bulan dan semua bulan tersedia"],
        ["DBG-06", "WAHA tidak terhubung/QR gagal", "Service, API key, engine/session state, dan recovery response", "GOWS hybrid, health/session sync, QR recovery; WahaIntegrationTest mencakup failed/restart/recovery"],
        ["DBG-07", "Chat masuk tidak muncul", "Webhook/event GOWS LID, HMAC, queue, atau matching belum lengkap", "SenderAlt/LID parsing, event/job, dedupe dan inbox diuji end-to-end"],
        ["DBG-08", "Data lead/sekolah duplikat atau format berbeda", "Normalisasi nomor/nama dan unique matching tidak terpusat", "PhoneNormalizer, LeadMatcher, create-school rules, CSV validation, dan conflict state"],
        ["DBG-09", "Status/program tidak sesuai rule terbaru", "Enum dan penamaan tersebar", "Enum lead diperbarui; nama program tiga unsur; setting dan test consistency"],
        ["DBG-10", "Risiko test menyentuh DB development", "Environment container dapat tetap local", "phpunit.xml force testing + DatabaseIsolationTest + test-full.ps1"],
    ], [0.65, 1.6, 2.0, 2.25])
    callout(doc, "Batas interpretasi", "Register ini merangkum temuan yang terlihat dari riwayat pengembangan dan bukti test saat ini. Ia bukan daftar seluruh insiden produksi; aplikasi masih lokal dan belum memiliki telemetry produksi terpusat.")

    b.add_heading(doc, "8.14 Baseline Verifikasi Tercatat", 2)
    table(doc, "Tabel 8-17. Snapshot verifikasi 25 Agustus 2026", ["Pemeriksaan", "Hasil", "Bukti/ruang lingkup"], [
        ["PHPUnit", "PASS", "72 test, 448 assertion, 102,631 detik"],
        ["Database isolation", "PASS", "SQLite :memory:; guard unit test"],
        ["Vite build", "PASS", "Vite 8.2.1; manifest dan asset dibuat; 2,12 detik"],
        ["Build warning", "NON-BLOCKING", "Paket fontaine opsional belum dipasang untuk optimized fallback"],
        ["RBAC/cabang", "PASS automated", "Superadmin, Admin Cabang, Staff, Direksi; scope data"],
        ["CRM core", "PASS automated", "Lead, sekolah, program, target, pendaftaran, pembayaran"],
        ["CSV/report", "PASS automated", "Import/export, formula neutralization, permission, metric"],
        ["WAHA/GOWS", "PASS mocked/feature", "HMAC, incoming/outbound, LID, QR, ACK, visibility"],
        ["Production runtime", "BELUM DIBUKTIKAN", "Staging/VPS, TLS, alert, backup/restore drill masih gate"],
    ], [1.45, 1.25, 3.8])
    table(doc, "Tabel 8-18. Cakupan test feature", ["Test file", "Domain utama"], [
        ["DashboardTest", "Auth, filter tampering, periode, KPI"],
        ["LeadManagementTest / LeadCsvTest", "Scope, status, import/export, formula safety"],
        ["CreateSchoolFromLeadTest / SchoolCsvTest", "Prefix sekolah, alumni/gapyear, duplicate, grade"],
        ["ProgramTargetTest / SettingsTest", "Nama program, target, setting, permission"],
        ["RegistrationPaymentTest", "Pendaftaran, pembayaran, scope dan status"],
        ["AuditNotificationTest / ReportTest", "Audit, reminder, notifikasi, laporan"],
        ["WahaIntegrationTest", "Webhook, matching, chat, QR, account, ACK, failure"],
        ["MasterDataAccessTest", "Hak akses master data dan user cabang"],
    ], [2.7, 3.8])

    b.add_heading(doc, "8.15 Format Catatan Debugging", 2)
    table(doc, "Tabel 8-19. Field catatan bug/insiden", ["Field", "Isi"], [
        ["Bug/Incident ID", "Kode unik dan judul singkat berbasis gejala"],
        ["Status/Severity", "New, Investigating, Reproduced, Fixing, Verifying, Monitoring, Closed/Reopened; SEV-1..4"],
        ["Reporter/Owner", "Pelapor, owner teknis, owner bisnis, incident commander bila ada"],
        ["Waktu", "Detected, started, contained, recovered, fixed, closed; timezone WIB/UTC jelas"],
        ["Environment/Release", "local/test/staging/production, tag/commit/image"],
        ["Dampak", "User, role, cabang, modul, data, security, workaround"],
        ["Reproduksi", "Data awal, langkah, expected, actual, frekuensi"],
        ["Evidence", "Log tersanitasi, test, screenshot, request/event ID"],
        ["RCA", "Gejala, trigger, contributing factor, root cause"],
        ["Resolution", "Patch/config/data repair, migration, rollback"],
        ["Verification", "Test target, regression, build, smoke/UAT, hasil"],
        ["Prevention/Approval", "Test/monitor/runbook baru dan pihak yang menyetujui"],
    ], [1.55, 4.95])
    b.add_heading(doc, "8.15.1 Status Workflow", 3)
    table(doc, "Tabel 8-20. Definisi status", ["Status", "Makna", "Syarat pindah"], [
        ["New", "Laporan diterima", "Triage awal"],
        ["Investigating", "Bukti sedang dikumpulkan", "Reproduced/Not a Bug/Need Info"],
        ["Reproduced", "Kasus dapat diulang atau bukti kuat", "RCA dan fix plan"],
        ["Fixing", "Patch sedang dikerjakan", "Review dan test target"],
        ["Verifying", "Patch tersedia", "Regression/build/smoke"],
        ["Monitoring", "Sudah dirilis/dipulihkan", "Periode stabilisasi"],
        ["Closed", "Kriteria lengkap dan disetujui", "Dapat Reopened bila berulang"],
        ["Deferred", "Ditunda dengan alasan/risiko", "Owner dan tanggal review"],
    ], [1.25, 2.65, 2.6])

    page(doc)
    b.add_heading(doc, "8.16 Kriteria Penutupan Bug", 2)
    add_figure(doc, "rantai-bukti.png", "Gambar 8-8. Rantai bukti dari laporan hingga penutupan")
    table(doc, "Tabel 8-21. Definition of Done bug", ["Kriteria", "Wajib", "Bukti"], [
        ["Requirement/expected disepakati", "Ya", "Referensi rule atau approval owner"],
        ["Reproduksi atau bukti runtime cukup", "Ya", "Langkah/test/timeline"],
        ["Root cause dan dampak data dinilai", "Ya", "RCA singkat"],
        ["Patch direview dan aman", "Ya", "Change/commit/review"],
        ["Regression test ditambah/disesuaikan", "Jika dapat diautomasi", "Nama test"],
        ["Full test dan build lulus", "Ya untuk code change", "Ringkasan hasil"],
        ["Smoke/UAT lulus", "Sesuai risiko", "Checklist/screenshot/approval"],
        ["Migration/repair/rollback tervalidasi", "Jika terdampak", "Runbook dan hasil"],
        ["Logging tidak membocorkan data", "Ya", "Review bukti"],
        ["Dokumentasi/diagram/runbook diperbarui", "Jika perubahan arsitektur/proses", "Versi artefak"],
        ["Observasi pascarilis selesai", "Untuk release runtime", "Waktu dan metric/log"],
        ["Owner menyetujui penutupan", "Ya", "Nama/peran/tanggal"],
    ], [2.8, 1.25, 2.45])
    callout(doc, "Reopened", "Kasus dibuka kembali bila gejala yang sama muncul, corrective action tidak lengkap, data belum konsisten, atau bukti monitoring menunjukkan regresi. Reopened mempertahankan ID/riwayat lama dan menambah timeline baru.", b.LIGHT_AMBER)

    b.add_heading(doc, "8.17 Rekomendasi Debugging dan QA", 2)
    table(doc, "Tabel 8-22. Roadmap peningkatan", ["Prioritas", "Rekomendasi", "Kriteria selesai"], [
        ["P0", "CI menjalankan test terisolasi dan build pada setiap merge", "Artifact hasil, exit code, branch protection"],
        ["P0", "Staging full Docker terpisah dengan data sintetis", "Smoke/UAT tanpa shared secret/data"],
        ["P0", "Queue/scheduler heartbeat dan failed-job alert", "Simulasi gagal menghasilkan alert"],
        ["P0", "Central log dengan redaction dan correlation ID", "Request-job-webhook dapat ditelusuri aman"],
        ["P1", "Integration test berkala pada MySQL/Redis/queue", "Perbedaan SQLite/MySQL terdeteksi"],
        ["P1", "Browser QA lintas breakpoint dan console/network capture", "Matriks mobile-tablet-laptop-desktop"],
        ["P1", "Error tracking dan release tagging", "Exception terkelompok ke versi release"],
        ["P1", "Dependency, secret, dan image vulnerability scan", "Critical finding diblokir/diterima formal"],
        ["P1", "Backup/restore dan rollback rehearsal", "RPO/RTO aktual tercatat"],
        ["P2", "Performance baseline dan slow-query profiling", "SLO/threshold berbasis data"],
        ["P2", "Chaos/game day terbatas untuk WAHA/Redis/queue", "Runbook dan recovery time diuji"],
    ], [0.75, 3.25, 2.5])
    b.add_heading(doc, "8.17.1 Gate Sebelum Deployment VPS", 3)
    steps(doc, "Tabel 8-23. Gate QA dan debugging", [
        "Pastikan working tree/release candidate teridentifikasi dan migration direview.",
        "Jalankan test terisolasi, build production, dependency/security scan, dan simpan hasil.",
        "Build image immutable lalu deploy ke staging dengan konfigurasi production-like.",
        "Jalankan smoke, RBAC/cabang, transaksi, CSV, queue/scheduler, WAHA, serta visual UAT.",
        "Validasi log redaction, alert, backup, restore, rollback, dan owner on-call.",
        "Tutup seluruh P0/P1 blocker atau catat risk acceptance berwenang.",
        "Rilis, observasi, rekonsiliasi data, dan dokumentasikan baseline baru.",
    ])

    b.add_heading(doc, "LAMPIRAN A - TEMPLATE CATATAN DEBUGGING", 1)
    table(doc, "Tabel A-1. Template siap pakai", ["Label", "Isi yang dicatat"], [
        ["ID / Judul", ""],
        ["Status / Severity", ""],
        ["Reporter / Owner", ""],
        ["Environment / Release", ""],
        ["Detected / Started / Timezone", ""],
        ["Dampak / Scope / Workaround", ""],
        ["Data awal / Langkah", ""],
        ["Expected / Actual", ""],
        ["Evidence tersanitasi", ""],
        ["Timeline / RCA", ""],
        ["Corrective / Preventive Action", ""],
        ["Test / Build / Smoke / UAT", ""],
        ["Release / Rollback / Observasi", ""],
        ["Approval / Tanggal tutup", ""],
    ], [2.2, 4.3])

    b.add_heading(doc, "LAMPIRAN B - REGISTER RISIKO DEBUGGING", 1)
    table(doc, "Tabel B-1. Risiko dan kontrol", ["ID", "Risiko", "Dampak", "Kontrol"], [
        ["D-01", "Diagnosis pada data produksi", "Perubahan/kehilangan data", "Reproduce di test/staging; approval untuk produksi"],
        ["D-02", "Secret/PII masuk log atau tiket", "Kebocoran data", "Redaction, least access, retention"],
        ["D-03", "Hotfix terlalu luas", "Regresi baru", "Patch minimum, review, feature flag"],
        ["D-04", "SQLite menyembunyikan beda MySQL", "Gagal saat deploy", "Integration test MySQL terjadwal"],
        ["D-05", "Container hidup tetapi queue macet", "Webhook/reminder tertunda", "Heartbeat, queue age, alert"],
        ["D-06", "Bug ditutup tanpa bukti", "Masalah berulang", "Definition of Done dan approval"],
        ["D-07", "Log terlalu sedikit", "RCA lambat", "Correlation ID dan structured logging"],
        ["D-08", "Log terlalu banyak", "Noise, biaya, privacy", "Level, sampling, retention, redaction"],
    ], [0.6, 2.0, 1.7, 2.2])

    b.add_heading(doc, "PENUTUP", 1)
    b.add_body(doc, "Manajemen debugging SIVMI CRM telah memiliki baseline automated test dan build yang kuat untuk fase lokal. Kesiapan deployment membutuhkan kesinambungan bukti dari CI, staging, integration/smoke test, observability, backup/rollback, dan penerimaan operasional. Pedoman ini harus diperbarui ketika test suite, arsitektur, environment, logging, severity, atau release gate berubah.")
    callout(doc, "Baseline final BAB VIII", "Perubahan yang menyelesaikan bug wajib meninggalkan bukti yang lebih kuat daripada sebelum bug ditemukan: test baru atau diperbarui, log/monitor yang memadai, runbook yang lebih jelas, dan keputusan penutupan yang dapat diaudit.", b.LIGHT_GREEN)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(f"Created: {OUTPUT}")


if __name__ == "__main__":
    build_document()
