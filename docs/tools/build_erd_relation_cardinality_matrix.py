from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
OUTPUT = ROOT / "docs" / "Matriks ERD Relasi dan Kardinalitas SIVMI CRM v1.0.docx"

PRIMARY = "1476B8"
PRIMARY_DARK = "075985"
INK = "0B2545"
MUTED = "64748B"
LIGHT = "F6F8FC"
LIGHT_BLUE = "EAF4FB"
GREEN = "10B981"
LIGHT_GREEN = "ECFDF5"
AMBER = "F59E0B"
LIGHT_AMBER = "FFFBEB"
ROSE = "F43F5E"
PURPLE = "A855F7"
BORDER = "CBD5E1"
WHITE = "FFFFFF"

# Landscape matrix override for compact_reference_guide:
# Letter landscape, margin 0.55", usable 9.9", table 9.8" + indent 120 DXA.
TABLE_WIDTH_DXA = 14112
TABLE_INDENT_DXA = 120
SUMMARY_WIDTHS = [2200, 1800, 2200, 1400, 6512]
DETAIL_WIDTHS = [500, 1800, 2300, 1700, 1300, 2200, 4312]


def rel(parent, parent_key, relation, child, child_key, cardinality, optionality, action, note, physical=True):
    return {
        "parent": parent,
        "parent_key": parent_key,
        "relation": relation,
        "child": child,
        "child_key": child_key,
        "cardinality": cardinality,
        "optionality": optionality,
        "action": action,
        "note": note,
        "physical": physical,
    }


CORE = [
    rel("cabang", "cabang_id", "memiliki", "users", "cabang_id", "1 : 0..N", "FK anak nullable", "SET NULL", "Satu cabang dapat menaungi banyak pengguna. Superadmin/direksi boleh tidak memiliki cabang."),
    rel("sekolah", "sekolah_id", "menjadi asal", "leads", "sekolah_id", "1 : 0..N", "FK anak nullable", "SET NULL", "Lead boleh belum memiliki sekolah atau memilih Alumni/Gapyear."),
    rel("program_bimbel", "program_bimbel_id", "diminati oleh", "leads", "minat_program_id", "1 : 0..N", "FK anak nullable", "SET NULL", "Minat program pada lead bersifat opsional dan dapat ditetapkan kemudian."),
    rel("cabang", "cabang_id", "memiliki", "leads", "cabang_id", "1 : 0..N", "FK anak wajib", "RESTRICT/NO ACTION", "Setiap lead wajib berada pada satu cabang; penghapusan cabang ditolak saat masih direferensikan."),
    rel("users", "id", "menginput", "leads", "diinput_oleh_user_id", "1 : 0..N", "FK anak wajib", "RESTRICT/NO ACTION", "Mencatat pengguna pembuat lead dan mempertahankan jejak asal data."),
    rel("users", "id", "menjadi PIC", "leads", "diserahkan_ke_user_id", "1 : 0..N", "FK anak nullable", "SET NULL", "Lead boleh belum ditugaskan; PIC dilepas jika akun pengguna dihapus."),
    rel("leads", "lead_id", "memiliki", "followup_riwayat", "lead_id", "1 : 0..N", "FK anak wajib", "CASCADE", "Riwayat follow-up melekat pada lead dan ikut dihapus bersama lead."),
    rel("users", "id", "mencatat", "followup_riwayat", "user_id", "1 : 0..N", "FK anak wajib", "RESTRICT/NO ACTION", "Setiap follow-up memiliki pencatat yang valid."),
    rel("leads", "lead_id", "dikonversi menjadi", "pendaftaran", "lead_id (UNIQUE)", "1 : 0..1", "FK anak wajib + unik", "RESTRICT/NO ACTION", "Satu lead paling banyak menghasilkan satu pendaftaran."),
    rel("program_bimbel", "program_bimbel_id", "dipilih pada", "pendaftaran", "program_bimbel_id", "1 : 0..N", "FK anak wajib", "RESTRICT/NO ACTION", "Program dapat dipilih oleh banyak pendaftaran."),
    rel("cabang", "cabang_id", "memproses", "pendaftaran", "cabang_id", "1 : 0..N", "FK anak wajib", "RESTRICT/NO ACTION", "Pendaftaran wajib berada pada satu cabang operasional."),
    rel("pendaftaran", "pendaftaran_id", "memiliki", "pembayaran", "pendaftaran_id", "1 : 0..N", "FK anak wajib", "CASCADE", "Satu pendaftaran dapat dibayar dalam satu atau beberapa transaksi."),
]

PERFORMANCE = [
    rel("cabang", "cabang_id", "memiliki target", "target_kinerja_cabang", "cabang_id", "1 : 0..N", "FK anak wajib", "RESTRICT/NO ACTION", "Kombinasi cabang, tahun, dan bulan bersifat unik."),
    rel("users", "id", "membuat target", "target_kinerja_cabang", "created_by", "1 : 0..N", "FK anak nullable", "SET NULL", "Identitas pembuat target boleh menjadi null saat akun dihapus."),
    rel("target_kinerja_cabang", "target_kinerja_cabang_id", "membagi target", "target_kinerja_user", "target_kinerja_cabang_id", "1 : 0..N", "FK anak wajib", "CASCADE", "Target user merupakan rincian target cabang."),
    rel("users", "id", "menerima target", "target_kinerja_user", "user_id", "1 : 0..N", "FK anak wajib", "CASCADE", "Pasangan target cabang dan user bersifat unik."),
    rel("users", "id", "memperbarui", "crm_settings", "updated_by", "1 : 0..N", "FK anak nullable", "SET NULL", "Nilai pengaturan tetap tersimpan saat akun editor dihapus."),
    rel("users", "id", "menghasilkan", "activity_logs", "user_id", "1 : 0..N", "FK anak nullable", "SET NULL", "Log aktivitas dipertahankan walaupun pengguna dihapus."),
    rel("cabang", "cabang_id", "membatasi", "activity_logs", "cabang_id", "1 : 0..N", "FK anak nullable", "SET NULL", "Cabang mendukung pembatasan visibilitas audit."),
    rel("subject polymorphic", "type + id", "dicatat sebagai", "activity_logs", "subject_type + subject_id", "1 : 0..N", "Referensi nullable", "Tanpa FK fisik", "Subjek dapat berupa berbagai model Eloquent; integritas dijaga aplikasi.", False),
    rel("notifiable polymorphic", "type + id", "menerima", "notifications", "notifiable_type + notifiable_id", "1 : 0..N", "Referensi wajib", "Tanpa FK fisik", "Saat ini terutama menunjuk users; notifikasi database memakai relasi polymorphic.", False),
]

WHATSAPP = [
    rel("users", "id", "memiliki akun WA", "user_whatsapp_accounts", "user_id (UNIQUE)", "1 : 0..1", "FK anak wajib + unik", "CASCADE", "Setiap admin/staff paling banyak memiliki satu akun WhatsApp aktif."),
    rel("cabang", "cabang_id", "menaungi akun WA", "user_whatsapp_accounts", "cabang_id", "1 : 0..N", "FK anak wajib", "CASCADE", "Akun WhatsApp selalu mengikuti cabang pengguna."),
    rel("user_whatsapp_accounts", "id", "memiliki", "whatsapp_conversations", "whatsapp_account_id", "1 : 0..N", "FK anak wajib", "CASCADE", "Kombinasi akun dan chat_id bersifat unik."),
    rel("leads", "lead_id", "dicocokkan dengan", "whatsapp_conversations", "lead_id", "1 : 0..N", "FK anak nullable", "SET NULL", "Percakapan dapat belum cocok, cocok, duplikat, atau konflik lintas cabang."),
    rel("whatsapp_conversations", "id", "berisi", "whatsapp_messages", "conversation_id", "1 : 0..N", "FK anak wajib", "CASCADE", "Semua pesan berada dalam satu percakapan."),
    rel("user_whatsapp_accounts", "id", "mengirim/menerima", "whatsapp_messages", "whatsapp_account_id", "1 : 0..N", "FK anak wajib", "CASCADE", "Kombinasi akun dan external_message_id mencegah duplikasi pesan."),
    rel("leads", "lead_id", "terkait dengan", "whatsapp_messages", "lead_id", "1 : 0..N", "FK anak nullable", "SET NULL", "Pesan dapat tersimpan sebelum proses pencocokan lead selesai."),
]

DATA_EXCHANGE = [
    rel("users", "id", "menjalankan", "imports", "user_id", "1 : 0..N", "FK anak wajib", "CASCADE", "Setiap pekerjaan import dimiliki pengguna."),
    rel("imports", "id", "memiliki kegagalan", "failed_import_rows", "import_id", "1 : 0..N", "FK anak wajib", "CASCADE", "Baris gagal ikut dihapus saat riwayat import dihapus."),
    rel("users", "id", "menjalankan", "exports", "user_id", "1 : 0..N", "FK anak wajib", "CASCADE", "Setiap pekerjaan export dimiliki pengguna."),
    rel("users", "id", "menjadi owner", "api_clients", "user_id", "1 : 0..N", "FK anak wajib", "CASCADE", "API client mewarisi konteks pengguna pembuat/owner."),
    rel("cabang", "cabang_id", "membatasi", "api_clients", "cabang_id", "1 : 0..N", "FK anak nullable", "SET NULL", "Null berarti client dapat memakai scope yang ditentukan role/ability; cabang membatasi data operasional."),
    rel("api_clients", "id", "memiliki kunci idempoten", "api_idempotency_keys", "api_client_id", "1 : 0..N", "FK anak wajib", "CASCADE", "Kombinasi api_client_id dan key bersifat unik."),
]

LOGICAL = [
    rel("users", "id", "memiliki sesi", "sessions", "user_id", "1 : 0..N", "Kolom nullable", "Tanpa FK fisik", "Sesi anonim tidak memiliki user_id; indeks tersedia tetapi constraint FK tidak dibuat.", False),
    rel("users", "email", "memiliki token reset", "password_reset_tokens", "email (PK)", "1 : 0..1", "Relasi berdasarkan email", "Tanpa FK fisik", "Token menggunakan email sebagai primary key dan dijaga oleh alur autentikasi.", False),
]

INDEPENDENT = [
    ("waha_webhook_events", "Event store idempoten berdasarkan payload_hash; diproses queue tanpa FK langsung."),
    ("cache", "Penyimpanan cache framework; tidak berelasi dengan entitas bisnis."),
    ("cache_locks", "Lock cache framework; tabel mandiri."),
    ("jobs", "Antrian database; payload menyimpan data job terenkapsulasi tanpa FK bisnis."),
    ("job_batches", "Metadata batch job; tabel mandiri."),
    ("failed_jobs", "Arsip kegagalan queue berdasarkan UUID; tabel mandiri."),
]


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def set_run(run, size=9, bold=False, color=INK, italic=False, name="Calibri"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = rgb(color)


def shade(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def cell_margins(cell, top=55, start=105, bottom=55, end=105):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    assert sum(widths) == TABLE_WIDTH_DXA
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    for tag, value, kind in (("tblW", TABLE_WIDTH_DXA, "dxa"), ("tblInd", TABLE_INDENT_DXA, "dxa")):
        node = tbl_pr.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tbl_pr.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), kind)
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        ET = OxmlElement("w:gridCol")
        ET.set(qn("w:w"), str(width))
        grid.append(ET)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[index]))
            tc_w.set(qn("w:type"), "dxa")
            cell_margins(cell)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    tr_pr.append(node)


def prevent_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))
    row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST


def add_page_field(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, value, end])


def configure_document(doc: Document):
    section = doc.sections[0]
    doc.settings.odd_and_even_pages_header_footer = False
    section.different_first_page_header_footer = False
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11)
    section.page_height = Inches(8.5)
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)
    section.header_distance = Inches(0.28)
    section.footer_distance = Inches(0.28)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = rgb(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.15

    for style_name, size, color_value, before, after in (
        ("Heading 1", 16, PRIMARY, 14, 7),
        ("Heading 2", 12.5, PRIMARY_DARK, 10, 5),
        ("Heading 3", 11, PRIMARY_DARK, 7, 4),
    ):
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(color_value)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hp.paragraph_format.space_after = Pt(0)
    set_run(hp.add_run("SIVMI CRM  |  Matriks ERD Relasi dan Kardinalitas"), 8.5, True, MUTED)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fp.paragraph_format.space_before = Pt(0)
    set_run(fp.add_run("Baseline 26 Agustus 2026  |  Halaman "), 8, False, MUTED)
    add_page_field(fp)


def title_block(doc: Document):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(3)
    set_run(p.add_run("MATRIKS ERD RELASI DAN KARDINALITAS"), 23, True, PRIMARY_DARK)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(12)
    set_run(p.add_run("SIVMI CRM - Skema operasional, kinerja, WhatsApp, import/export, dan REST API"), 11.5, False, MUTED)

    metadata = [
        ("Versi", "1.0"),
        ("Baseline skema", "Migration sampai 2026_08_25_000001_create_public_api_tables.php"),
        ("Sumber verifikasi", "database/migrations dan app/Models"),
        ("Notasi", "1 : 0..N = satu induk dapat memiliki nol sampai banyak anak; 1 : 0..1 = opsional tunggal"),
    ]
    for label, value in metadata:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        set_run(p.add_run(f"{label}: "), 9.5, True, INK)
        set_run(p.add_run(value), 9.5, False, INK)

    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    values = [
        ("36", "relasi terpetakan", LIGHT_BLUE, PRIMARY),
        ("32", "relasi FK fisik", LIGHT_GREEN, GREEN),
        ("4", "relasi logis/polymorphic", LIGHT_AMBER, AMBER),
        ("6", "tabel mandiri", "FFF1F2", ROSE),
    ]
    for index, (number, label, fill, accent) in enumerate(values):
        cell = table.cell(0, index)
        shade(cell, fill)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        set_run(p.add_run(number + "\n"), 16, True, accent)
        set_run(p.add_run(label), 8.5, True, INK)
    set_table_geometry(table, [3528, 3528, 3528, 3528])


def add_note(doc: Document, title: str, text: str, fill=LIGHT_BLUE):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    shade(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    set_run(p.add_run(title + "  "), 9.5, True, PRIMARY_DARK)
    set_run(p.add_run(text), 9.2, False, INK)
    set_table_geometry(table, [TABLE_WIDTH_DXA])


def add_heading(doc: Document, text: str, level=1, new_page=False):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.page_break_before = new_page
    return p


def fill_header(row, labels):
    repeat_header(row)
    prevent_split(row)
    for cell, label in zip(row.cells, labels):
        shade(cell, PRIMARY)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        set_run(p.add_run(label), 8.2, True, WHITE)


def relation_text(cell, item, compact=False):
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    size = 8.0 if compact else 7.6
    # Nama tabel sudah tersedia pada kolom induk dan anak. Hanya tampilkan
    # pasangan key di kolom relasi agar identifier panjang tidak memaksa baris
    # tabel menjadi terlalu tinggi.
    set_run(p.add_run(item["parent_key"]), size, True, INK, name="Consolas")
    set_run(p.add_run("  ke  "), size, False, MUTED, name="Consolas")
    set_run(p.add_run(item["child_key"]), size, True, PURPLE, name="Consolas")


def add_summary_table(doc: Document, rows):
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    fill_header(table.rows[0], ["Entitas induk", "Relasi", "Entitas anak", "Kardinalitas", "Keterangan ringkas"])
    for index, item in enumerate(rows):
        row = table.add_row()
        prevent_split(row)
        if index % 2:
            for cell in row.cells:
                shade(cell, LIGHT)
        values = [item["parent"], item["relation"], item["child"], item["cardinality"], item["note"]]
        for col, (cell, value) in enumerate(zip(row.cells, values)):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if col == 3 else WD_ALIGN_PARAGRAPH.LEFT
            set_run(p.add_run(value), 8.2, col in (0, 2), INK)
    set_table_geometry(table, SUMMARY_WIDTHS)


def add_detail_table(doc: Document, rows, start_number):
    # Pisahkan heading dari tabel secara eksplisit. Tanpa paragraf jangkar ini,
    # Word sesekali menempatkan repeated header table pada koordinat heading
    # saat dokumen dirender ke PDF.
    anchor = doc.add_paragraph()
    anchor.paragraph_format.space_before = Pt(0)
    anchor.paragraph_format.space_after = Pt(0)
    anchor.paragraph_format.line_spacing = Pt(1)
    set_run(anchor.add_run("\u00a0"), 1, False, WHITE)
    table = doc.add_table(rows=1, cols=7)
    table.style = "Table Grid"
    fill_header(table.rows[0], ["No.", "Entitas induk / PK", "Relasi FK", "Entitas anak", "Kardinalitas", "Opsionalitas / delete", "Keterangan"])
    for offset, item in enumerate(rows):
        row = table.add_row()
        prevent_split(row)
        if offset % 2:
            for cell in row.cells:
                shade(cell, LIGHT)
        number = start_number + offset
        values = [
            str(number),
            f"{item['parent']}\nPK/ref: {item['parent_key']}",
            None,
            item["child"],
            item["cardinality"],
            f"{item['optionality']}\n{item['action']}",
            item["note"],
        ]
        for col, (cell, value) in enumerate(zip(row.cells, values)):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            if col == 2:
                relation_text(cell, item)
                continue
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if col in (0, 4) else WD_ALIGN_PARAGRAPH.LEFT
            set_run(p.add_run(value), 7.6, col in (1, 3, 4), INK)
        if not item["physical"]:
            shade(row.cells[5], LIGHT_AMBER)
    set_table_geometry(table, DETAIL_WIDTHS)
    return start_number + len(rows)


def add_independent_table(doc: Document):
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    fill_header(table.rows[0], ["No.", "Tabel mandiri", "Fungsi dan alasan tidak memiliki relasi FK bisnis"])
    widths = [650, 3000, 10462]
    for index, (name, note) in enumerate(INDEPENDENT, start=1):
        row = table.add_row()
        prevent_split(row)
        if index % 2 == 0:
            for cell in row.cells:
                shade(cell, LIGHT)
        for col, value in enumerate((str(index), name, note)):
            p = row.cells[col].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if col == 0 else WD_ALIGN_PARAGRAPH.LEFT
            set_run(p.add_run(value), 8.4, col == 1, INK)
    set_table_geometry(table, widths)


def page_break(doc: Document):
    doc.add_page_break()


def build():
    doc = Document()
    configure_document(doc)
    title_block(doc)

    add_heading(doc, "1. Notasi dan Aturan Pembacaan", 1)
    add_note(
        doc,
        "Cara membaca:",
        "Kardinalitas ditulis dari perspektif entitas induk ke entitas anak. FK nullable membuat sisi anak opsional; UNIQUE membatasi satu anak; CASCADE menghapus anak; SET NULL mempertahankan anak; RESTRICT/NO ACTION menolak penghapusan induk yang masih direferensikan.",
    )

    add_heading(doc, "Cakupan Matriks", 2)
    scope = doc.add_table(rows=1, cols=5)
    scope.style = "Table Grid"
    scope_values = [
        ("12", "operasional"), ("9", "kinerja & kontrol"),
        ("7", "WhatsApp"), ("6", "pertukaran data"), ("2", "relasi logis"),
    ]
    for index, (number, label) in enumerate(scope_values):
        cell = scope.cell(0, index)
        shade(cell, LIGHT)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        set_run(p.add_run(number + "\n"), 13, True, PRIMARY)
        set_run(p.add_run(label), 8.2, True, INK)
    set_table_geometry(scope, [2822, 2822, 2822, 2822, 2824])

    add_heading(doc, "2. Matriks Ringkas ERD Operasional", 1, new_page=True)
    p = doc.add_paragraph("Ringkasan berikut mengikuti alur utama lead sampai pembayaran seperti ERD operasional.")
    p.paragraph_format.keep_with_next = True
    add_summary_table(doc, CORE)

    add_heading(doc, "3. Matriks Detail Relasi dan Kardinalitas", 1, new_page=True)
    add_note(doc, "Baseline:", "Relasi fisik diverifikasi dari migration; nama method relasi dikonfirmasi melalui model Eloquent. Nomor berlanjut antar-domain.", LIGHT_GREEN)

    number = 1
    add_heading(doc, "3.1 Operasional: Lead sampai Pembayaran", 2)
    number = add_detail_table(doc, CORE, number)

    add_heading(doc, "3.2 Kinerja, Pengaturan, Audit, dan Notifikasi", 2, new_page=True)
    number = add_detail_table(doc, PERFORMANCE, number)

    # Page break eksplisit mencegah heading dan baris pertama tabel bertumpuk
    # pada renderer Microsoft Word ketika tabel sebelumnya cukup panjang.
    page_break(doc)
    add_heading(doc, "3.3 Integrasi WhatsApp dan WAHA", 2)
    number = add_detail_table(doc, WHATSAPP, number)

    # Gunakan page break eksplisit pada batas domain ini. Word dapat menumpuk
    # repeated table header dengan heading yang memakai page_break_before tepat
    # setelah tabel panjang, sehingga judul dan kepala tabel saling menimpa.
    page_break(doc)
    add_heading(doc, "3.4 Import, Export, dan REST API Publik", 2)
    number = add_detail_table(doc, DATA_EXCHANGE, number)

    add_heading(doc, "3.5 Relasi Logis Framework", 2)
    number = add_detail_table(doc, LOGICAL, number)

    add_heading(doc, "4. Tabel Mandiri", 1, new_page=True)
    p = doc.add_paragraph("Tabel berikut tetap menjadi bagian skema tetapi tidak mempunyai foreign key langsung ke entitas bisnis utama.")
    p.paragraph_format.keep_with_next = True
    add_independent_table(doc)

    add_heading(doc, "5. Catatan Integritas dan Kardinalitas", 1)
    notes = [
        "Relasi leads ke pendaftaran adalah 1 : 0..1 karena pendaftaran.lead_id memiliki UNIQUE constraint.",
        "Relasi users ke user_whatsapp_accounts adalah 1 : 0..1 karena user_id memiliki UNIQUE constraint.",
        "activity_logs.subject dan notifications.notifiable bersifat polymorphic; database tidak dapat menerapkan FK ke satu tabel induk tertentu.",
        "sessions.user_id dan password_reset_tokens.email adalah relasi logis framework tanpa constraint FK fisik.",
        "waha_webhook_events sengaja mandiri sebagai event store idempoten; session_name dan payload_hash dipakai untuk korelasi proses.",
        "Scope role dan cabang diterapkan pada query aplikasi dan tidak mengubah kardinalitas fisik database.",
    ]
    for text in notes:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.28)
        p.paragraph_format.first_line_indent = Inches(-0.18)
        p.paragraph_format.space_after = Pt(3)
        set_run(p.add_run(text), 9, False, INK)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(f"Created: {OUTPUT}")
    print(f"Relations: {number - 1}; independent tables: {len(INDEPENDENT)}")


if __name__ == "__main__":
    build()
