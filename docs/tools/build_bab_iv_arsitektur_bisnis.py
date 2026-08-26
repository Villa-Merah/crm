from __future__ import annotations

from datetime import datetime
from pathlib import Path
import sys

PROJECT_ROOT = Path(__file__).resolve().parents[2]
DOC_TOOLS = PROJECT_ROOT / "storage" / "app" / "doc-tools"
sys.path.insert(0, str(DOC_TOOLS))

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = PROJECT_ROOT / "docs" / "BAB IV - Arsitektur Bisnis SIVMI CRM v1.0.docx"
ASSETS = PROJECT_ROOT / "docs" / "_bab_iv_assets"
COVER = ASSETS / "cover-bab-iv.png"

PRIMARY = "1476B8"
PRIMARY_DARK = "075985"
INK = "0B2545"
MUTED = "64748B"
LIGHT = "F6F8FC"
LIGHT_BLUE = "EAF4FB"
LIGHT_GREEN = "ECFDF5"
LIGHT_AMBER = "FFFBEB"
LIGHT_ROSE = "FFF1F2"
WHITE = "FFFFFF"
BORDER = "D8E1EC"


def rgb(hex_value: str) -> RGBColor:
    return RGBColor.from_string(hex_value)


def set_run_font(run, size: float | None = None, bold: bool | None = None,
                 color: str | None = None, italic: bool | None = None,
                 name: str = "Calibri") -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = rgb(color)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 80, start: int = 120,
                     bottom: int = 80, end: int = 120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        element = tc_mar.find(qn(f"w:{tag}"))
        if element is None:
            element = OxmlElement(f"w:{tag}")
            tc_mar.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_geometry(table, widths_inches: list[float]) -> None:
    widths_dxa = [round(width * 1440) for width in widths_inches]
    total = sum(widths_dxa)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for value in widths_dxa:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(value))
        grid.append(grid_col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[index]))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(widths_inches[index])
            set_cell_margins(cell)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Halaman ")
    set_run_font(run, 8.5, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, value, end])


def configure_styles(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = rgb(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    heading_tokens = {
        "Heading 1": (16, PRIMARY, 16, 8),
        "Heading 2": (13, PRIMARY, 12, 6),
        "Heading 3": (12, PRIMARY_DARK, 8, 4),
    }
    for name, (size, color, before, after) in heading_tokens.items():
        style = document.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for list_name in ("List Bullet", "List Number"):
        style = document.styles[list_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.5)
        style.paragraph_format.first_line_indent = Inches(-0.25)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing = 1.167

    caption = document.styles["Caption"]
    caption.font.name = "Calibri"
    caption._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    caption._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    caption.font.size = Pt(9)
    caption.font.italic = True
    caption.font.color.rgb = rgb(MUTED)
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(8)
    caption.paragraph_format.keep_with_next = True

    if "Table Text" not in [style.name for style in document.styles]:
        table_text = document.styles.add_style("Table Text", WD_STYLE_TYPE.PARAGRAPH)
    else:
        table_text = document.styles["Table Text"]
    table_text.font.name = "Calibri"
    table_text._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    table_text._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    table_text.font.size = Pt(8.5)
    table_text.font.color.rgb = rgb(INK)
    table_text.paragraph_format.space_before = Pt(0)
    table_text.paragraph_format.space_after = Pt(2)
    table_text.paragraph_format.line_spacing = 1.0


def add_body(document: Document, text: str, bold_lead: str | None = None) -> None:
    paragraph = document.add_paragraph(style="Normal")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    if bold_lead and text.startswith(bold_lead):
        lead = paragraph.add_run(bold_lead)
        set_run_font(lead, 11, True, INK)
        rest = paragraph.add_run(text[len(bold_lead):])
        set_run_font(rest, 11, False, INK)
    else:
        run = paragraph.add_run(text)
        set_run_font(run, 11, False, INK)


def add_bullets(document: Document, items: list[str]) -> None:
    for item in items:
        paragraph = document.add_paragraph(style="List Bullet")
        run = paragraph.add_run(item)
        set_run_font(run, 11, False, INK)


def add_numbered(document: Document, items: list[str]) -> None:
    # Nomor ditulis eksplisit agar setiap kelompok selalu dimulai dari 1.
    # Word cenderung menggabungkan kembali sequence OOXML terpisah ketika field
    # dan style diperbarui melalui COM, sehingga hasil render dapat meloncat.
    for index, item in enumerate(items, start=1):
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.35)
        paragraph.paragraph_format.first_line_indent = Inches(-0.25)
        paragraph.paragraph_format.space_after = Pt(2)
        run = paragraph.add_run(f"{index}.  {item}")
        set_run_font(run, 11, False, INK)


def add_heading(document: Document, text: str, level: int) -> None:
    paragraph = document.add_heading(text, level=level)
    paragraph.paragraph_format.keep_with_next = True


def add_callout(document: Document, title: str, body: str, fill: str = LIGHT_BLUE) -> None:
    table = document.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    set_table_geometry(table, [6.5])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    paragraph = cell.paragraphs[0]
    paragraph.style = document.styles["Normal"]
    title_run = paragraph.add_run(title + "\n")
    set_run_font(title_run, 10.5, True, PRIMARY_DARK)
    body_run = paragraph.add_run(body)
    set_run_font(body_run, 10, False, INK)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def add_table(document: Document, headers: list[str], rows: list[list[str]],
              widths: list[float], caption: str) -> None:
    caption_paragraph = document.add_paragraph(style="Caption")
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    caption_paragraph.add_run(caption)
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_repeat_table_header(table.rows[0])
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        set_cell_shading(cell, PRIMARY)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.style = document.styles["Table Text"]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(header)
        set_run_font(run, 8.5, True, WHITE)
    for row_index, values in enumerate(rows):
        row = table.add_row()
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        cells = row.cells
        for column_index, value in enumerate(values):
            cell = cells[column_index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            if row_index % 2:
                set_cell_shading(cell, LIGHT)
            paragraph = cell.paragraphs[0]
            paragraph.style = document.styles["Table Text"]
            run = paragraph.add_run(str(value))
            set_run_font(run, 8.5, False, INK)
    set_table_geometry(table, widths)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def add_figure(document: Document, path: Path, caption: str,
               width: float = 6.2) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    picture = paragraph.add_run().add_picture(str(path), width=Inches(width))
    picture._inline.docPr.set("descr", caption)
    caption_paragraph = document.add_paragraph(style="Caption")
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_paragraph.add_run(caption)


def add_toc(document: Document) -> None:
    paragraph = document.add_paragraph()
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = 'TOC \\o "1-3" \\h \\z \\u'
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "Daftar isi akan diperbarui otomatis saat dokumen dibuka."
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, placeholder, end])


def make_cover() -> None:
    width, height = 1600, 2200
    image = Image.new("RGB", (width, height), "#F6F8FC")
    draw = ImageDraw.Draw(image)
    for y in range(height):
        ratio = y / height
        start = (20, 118, 184)
        end = (7, 89, 133)
        color = tuple(round(start[i] * (1 - ratio) + end[i] * ratio) for i in range(3))
        draw.line((0, y, width, y), fill=color)
    draw.ellipse((1010, -180, 1770, 580), fill=(27, 139, 199))
    draw.ellipse((1140, 140, 1640, 640), fill=(8, 103, 151))
    draw.rounded_rectangle((120, 250, 1480, 1910), radius=44, fill=(255, 255, 255))

    font_regular = Path("C:/Windows/Fonts/arial.ttf")
    font_bold = Path("C:/Windows/Fonts/arialbd.ttf")
    title_font = ImageFont.truetype(str(font_bold), 102)
    chapter_font = ImageFont.truetype(str(font_bold), 48)
    subtitle_font = ImageFont.truetype(str(font_regular), 34)
    meta_font = ImageFont.truetype(str(font_regular), 28)
    brand_font = ImageFont.truetype(str(font_bold), 34)

    draw.text((190, 350), "SIVMI CRM", font=brand_font, fill="#1476B8")
    draw.text((190, 600), "BAB IV", font=chapter_font, fill="#075985")
    draw.multiline_text((190, 690), "ARSITEKTUR\nBISNIS", font=title_font, fill="#0B2545", spacing=18)
    draw.text((190, 1040), "Model operasi, proses end-to-end, aktor, KPI,", font=subtitle_font, fill="#475569")
    draw.text((190, 1092), "enterprise architecture, sasaran, dan aturan bisnis", font=subtitle_font, fill="#475569")
    draw.rounded_rectangle((190, 1245, 760, 1325), radius=40, fill="#EAF4FB")
    draw.text((230, 1267), "BASELINE AS-BUILT 24 AGUSTUS 2026", font=meta_font, fill="#075985")
    draw.text((190, 1610), "Versi dokumen 1.0", font=meta_font, fill="#64748B")
    draw.text((190, 1665), "Status: dokumen kerja untuk validasi manajemen", font=meta_font, fill="#64748B")
    draw.text((190, 1775), "SIVMI CRM - Catat, Kelola, Closing", font=brand_font, fill="#1476B8")
    image.save(COVER)


def build_document() -> None:
    make_cover()
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    configure_styles(document)

    header = section.header
    header_paragraph = header.paragraphs[0]
    header_paragraph.text = "SIVMI CRM  |  BAB IV: ARSITEKTUR BISNIS"
    for run in header_paragraph.runs:
        set_run_font(run, 8.5, True, PRIMARY)
    footer = section.footer
    add_page_number(footer.paragraphs[0])

    cover_paragraph = document.add_paragraph()
    cover_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cover_paragraph.paragraph_format.space_after = Pt(0)
    cover_picture = cover_paragraph.add_run().add_picture(str(COVER), width=Inches(6.1))
    cover_picture._inline.docPr.set("descr", "Sampul BAB IV Arsitektur Bisnis SIVMI CRM")
    document.add_page_break()

    add_heading(document, "KENDALI DOKUMEN", 1)
    add_table(document, ["Atribut", "Nilai"], [
        ["Judul", "BAB IV: Arsitektur Bisnis SIVMI CRM"],
        ["Versi", "1.0"],
        ["Tanggal baseline", "24 Agustus 2026"],
        ["Status", "Dokumen kerja terpisah - siap validasi manajemen"],
        ["Ruang lingkup", "Organisasi, stakeholder, proses bisnis, KPI, enterprise architecture, sasaran, dan aturan bisnis"],
        ["Baseline teknologi terkait", "Laravel 12, Filament 5, Livewire 4, MySQL, Redis, full Docker, WAHA 2026.8.1 CORE/GOWS"],
        ["Baseline pengujian", "72 test lulus dengan 448 assertion; build Vite produksi berhasil"],
        ["Sumber", "Codebase as-built, migration, seeder, test, diagram gabungan v1.4, dan dokumentasi induk v1.6"],
    ], [1.6, 4.9], "Tabel 0-1. Kendali dan baseline dokumen")
    add_callout(document, "Kedudukan dokumen", "Dokumen ini merupakan BAB IV versi terpisah dan lebih detail. Apabila terdapat perbedaan dengan dokumen induk, codebase as-built dan keputusan manajemen terbaru menjadi sumber kebenaran utama.")

    add_heading(document, "DAFTAR ISI", 1)
    add_toc(document)
    document.add_page_break()

    add_heading(document, "RINGKASAN EKSEKUTIF", 1)
    add_body(document, "Arsitektur bisnis SIVMI CRM mendefinisikan cara organisasi mengubah minat calon siswa menjadi layanan bimbingan belajar yang terukur. Model operasi menghubungkan tiga cabang awal - Bandung, JakSel, dan JakPus - dengan empat peran akses: Superadmin, Admin Cabang, Staff, dan Direksi. Seluruh proses menggunakan satu sumber data untuk lead, tindak lanjut, pendaftaran, pembayaran, program, target, laporan, komunikasi WhatsApp, notifikasi, dan audit.")
    add_body(document, "Rantai nilai utama terdiri dari akuisisi, validasi dan distribusi lead, tindak lanjut, konversi, transaksi, serta monitoring. Lead dapat masuk secara manual, CSV, atau pesan WhatsApp melalui WAHA/GOWS. Nomor telepon dinormalisasi dan dibandingkan dengan data lead untuk mencegah duplikasi. Setiap aktivitas dibatasi berdasarkan role dan cabang; perubahan penting meninggalkan jejak audit.")
    add_callout(document, "Keputusan arsitektural utama", "Cabang adalah batas kepemilikan data operasional. Pengguna global dapat melihat lintas cabang sesuai role, sedangkan Admin Cabang dan Staff hanya memproses data cabangnya. Dashboard, laporan, KPI, inbox WhatsApp, dan audit menerapkan batas yang sama.", LIGHT_GREEN)
    add_table(document, ["Dimensi", "Kondisi terbaru", "Dampak bisnis"], [
        ["Organisasi", "4 role; 3 cabang awal", "Tanggung jawab dan visibilitas dapat dibedakan"],
        ["Akuisisi", "Manual, CSV, WhatsApp; sumber lead terklasifikasi", "Asal lead dapat diukur dan ditelusuri"],
        ["Konversi", "Cold, Warm, Hot, Daftar, Tidak Berminat", "Pipeline lebih sederhana dan konsisten"],
        ["Program", "Nama otomatis 3 komponen: kategori + kelas + metode", "Penamaan produk seragam"],
        ["Komunikasi", "Akun WA per admin/staff, inbox dan chat dua arah", "Percakapan menjadi bukti operasional"],
        ["Pengendalian", "Branch scope, RBAC, notifikasi, audit", "Risiko akses silang dan kehilangan jejak berkurang"],
        ["Pengukuran", "Filter global, target cabang/user, KPI multidimensi", "Direksi memperoleh indikator yang konsisten"],
    ], [1.3, 2.25, 2.95], "Tabel 0-2. Ringkasan model operasi as-built")

    add_heading(document, "BAB IV: ARSITEKTUR BISNIS", 1)
    add_body(document, "Bab ini menjelaskan organisasi, aktor, kemampuan bisnis, proses, indikator, sasaran, dan aturan yang menjadi landasan SIVMI CRM. Fokusnya adalah menjawab siapa melakukan apa, data apa yang digunakan, keputusan apa yang dihasilkan, serta kontrol apa yang wajib dipenuhi.")

    add_heading(document, "4.1 Profil dan Visi Misi Perusahaan", 2)
    add_heading(document, "4.1.1 Profil Operasional SIVMI", 3)
    add_body(document, "SIVMI pada baseline sistem diposisikan sebagai penyelenggara layanan bimbingan belajar multi-cabang. Aktivitas utamanya meliputi pemasaran dan akuisisi calon siswa, konsultasi program, follow-up, pendaftaran, penerimaan pembayaran, pengelolaan target, serta monitoring performa cabang dan pengguna.")
    add_body(document, "Codebase belum memuat nama badan hukum, sejarah pendirian, visi, misi, nilai perusahaan, struktur jabatan formal, SLA layanan, atau kebijakan retensi yang telah disahkan. Oleh karena itu, bagian strategis di bawah diperlakukan sebagai rumusan operasional untuk divalidasi, bukan pernyataan korporat resmi.")
    add_table(document, ["Elemen", "Baseline as-built"], [
        ["Model layanan", "CRM bimbingan belajar berbasis web dengan operasi multi-cabang"],
        ["Cabang awal", "Bandung (BDG), JakSel (JKS), dan JakPus (JKP)"],
        ["Pelanggan eksternal", "Calon siswa, siswa, orang tua/wali, dan pihak sekolah"],
        ["Pengguna internal", "Superadmin, Admin Cabang, Staff, dan Direksi"],
        ["Objek utama", "Lead, sekolah, program, follow-up, pendaftaran, pembayaran, target, akun WhatsApp"],
        ["Saluran akuisisi", "Input manual, impor CSV, dan pesan WhatsApp melalui WAHA/GOWS"],
        ["Sumber lead", "Iklan, Instagram, Kunjungan, Saudara, Facebook, dan Tiktok"],
        ["Nilai yang dihasilkan", "Pipeline terpantau, tanggung jawab jelas, konversi tercatat, omzet terukur, komunikasi terlacak"],
    ], [1.7, 4.8], "Tabel 4-1. Profil operasional SIVMI berdasarkan sistem")

    add_heading(document, "4.1.2 Rumusan Visi dan Misi Operasional", 3)
    add_callout(document, "Perlu pengesahan manajemen", "Rumusan berikut disusun sebagai working statement yang konsisten dengan fungsi sistem. Redaksi tidak boleh dipublikasikan sebagai visi dan misi resmi sebelum disetujui Direksi.", LIGHT_AMBER)
    add_body(document, "Visi operasional usulan: membangun pengelolaan hubungan calon siswa dan layanan bimbingan belajar yang terukur, responsif, transparan, dan konsisten di seluruh cabang.", "Visi operasional usulan:")
    add_body(document, "Misi operasional usulan:")
    add_bullets(document, [
        "Menyatukan data calon siswa, komunikasi, pendaftaran, pembayaran, dan kinerja dalam satu platform.",
        "Memastikan setiap lead memiliki kepemilikan, status, tindak lanjut, dan histori yang dapat ditelusuri.",
        "Mendorong keputusan berbasis data melalui KPI yang memiliki definisi, periode, sumber, dan penanggung jawab.",
        "Menjaga kerahasiaan data melalui RBAC, pembatasan cabang, audit, dan pengendalian integrasi.",
        "Meningkatkan kecepatan respons melalui inbox WhatsApp terpusat, notifikasi, reminder, dan quick action.",
    ])

    add_heading(document, "4.1.3 Proposisi Nilai Bisnis", 3)
    add_table(document, ["Penerima nilai", "Masalah yang diatasi", "Nilai dari SIVMI CRM"], [
        ["Calon siswa/orang tua", "Respons lambat dan informasi terfragmentasi", "Kontak tercatat, tindak lanjut terjadwal, program dan pembayaran lebih jelas"],
        ["Staff", "Lead tersebar dan histori komunikasi hilang", "Inbox, kepemilikan lead, follow-up, dan status berada dalam satu alur"],
        ["Admin Cabang", "Sulit mengendalikan aktivitas dan target tim", "Monitoring cabang, distribusi PIC, pending action, ranking, dan audit"],
        ["Direksi", "Laporan lintas cabang tidak seragam", "KPI dan laporan terfilter dengan definisi agregasi yang konsisten"],
        ["Superadmin", "Konfigurasi dan akses tidak terkendali", "Master data, role, cabang, pengaturan, akun WA, serta audit terpusat"],
    ], [1.25, 2.15, 3.1], "Tabel 4-2. Proposisi nilai per stakeholder")

    add_heading(document, "4.1.4 Ruang Lingkup Bisnis", 3)
    add_bullets(document, [
        "Termasuk: akuisisi lead, master sekolah, master program, distribusi PIC, follow-up, pendaftaran, pembayaran, target, dashboard, laporan, WhatsApp, notifikasi, pengaturan, dan audit.",
        "Termasuk kontrol: autentikasi, RBAC, branch scope, validasi, deduplikasi nomor, impor/ekspor CSV, logging, empty/error/loading state, dan quality gate.",
        "Belum dinyatakan sebagai ruang lingkup as-built: LMS pembelajaran, jadwal pengajar, absensi akademik, payroll, akuntansi buku besar, dan portal siswa mandiri.",
    ])

    add_heading(document, "4.2 Struktur Organisasi dan Stakeholder", 2)
    add_heading(document, "4.2.1 Model Organisasi", 3)
    add_body(document, "Model organisasi sistem memisahkan pengendalian global dan eksekusi lokal. Direksi mengonsumsi informasi lintas cabang untuk keputusan strategis. Superadmin menjaga konfigurasi serta integritas sistem. Admin Cabang mengendalikan operasi cabang. Staff menangani lead, follow-up, pendaftaran, pembayaran, dan percakapan sesuai penugasan.")
    add_table(document, ["Role", "Cakupan data", "Mandat utama", "Batas utama"], [
        ["Superadmin", "Global", "Konfigurasi, master data, pengguna, integrasi, dan pengawasan", "Tidak menggantikan keputusan bisnis Direksi tanpa mandat"],
        ["Direksi", "Global, terutama baca/monitor", "Menilai KPI, laporan, pencapaian target, dan risiko", "Aksi transaksi dibatasi sesuai kebijakan resource"],
        ["Admin Cabang", "Cabang sendiri", "Mengelola operasi, tim, distribusi lead, target, dan kontrol transaksi", "Tidak melihat atau mengubah data cabang lain"],
        ["Staff", "Cabang sendiri dan data yang relevan", "Menangani lead, follow-up, chat, pendaftaran, serta pencatatan pembayaran", "Tidak mengelola konfigurasi global"],
    ], [1.05, 1.15, 2.4, 1.9], "Tabel 4-3. Struktur role dan batas kewenangan")

    add_heading(document, "4.2.2 Peta Stakeholder", 3)
    add_table(document, ["Stakeholder", "Kepentingan", "Informasi yang dibutuhkan", "Pengaruh"], [
        ["Direksi", "Pertumbuhan dan akuntabilitas", "Omzet, siswa, funnel, cabang, program, sekolah/kota, risiko", "Tinggi"],
        ["Superadmin", "Konsistensi dan keamanan sistem", "Role, user, master, pengaturan, log, integrasi, health", "Tinggi"],
        ["Admin Cabang", "Target dan produktivitas cabang", "Lead/PIC, follow-up, pendaftaran, pembayaran, pending action", "Tinggi"],
        ["Staff", "Kecepatan layanan dan closing", "Kontak, histori, jadwal follow-up, program, chat, status pembayaran", "Sedang"],
        ["Calon siswa/orang tua", "Informasi dan layanan responsif", "Program, biaya, jadwal komunikasi, status pendaftaran", "Sedang"],
        ["Sekolah", "Hubungan sumber calon siswa", "Identitas dan klasifikasi sekolah yang akurat", "Rendah-sedang"],
        ["Penyedia WAHA/WhatsApp", "Konektivitas komunikasi", "Session, webhook, API key, status pesan", "Teknis tinggi"],
    ], [1.25, 1.8, 2.65, 0.8], "Tabel 4-4. Peta kebutuhan stakeholder")

    add_heading(document, "4.2.3 Matriks RACI Proses Utama", 3)
    add_body(document, "R = Responsible, A = Accountable, C = Consulted, I = Informed. Matriks berikut adalah baseline operasional; penetapan akhir tetap memerlukan pengesahan organisasi.")
    add_table(document, ["Aktivitas", "Direksi", "Superadmin", "Admin Cabang", "Staff"], [
        ["Menetapkan target cabang", "A", "C", "R", "I"],
        ["Mengelola konfigurasi dan role", "I", "A/R", "C", "I"],
        ["Menerima dan memvalidasi lead", "I", "C", "A", "R"],
        ["Menetapkan/menyesuaikan PIC", "I", "C", "A/R", "R"],
        ["Melakukan follow-up", "I", "I", "A", "R"],
        ["Membuat pendaftaran", "I", "I", "A", "R"],
        ["Mencatat/verifikasi pembayaran", "I", "C", "A/R", "R"],
        ["Mengelola akun dan session WhatsApp", "I", "A/R", "R", "C"],
        ["Meninjau laporan dan KPI", "A/R", "C", "R", "I"],
        ["Menangani konflik data lintas cabang", "I", "A/R", "C", "I"],
    ], [2.45, 0.75, 1.1, 1.25, 0.95], "Tabel 4-5. Matriks RACI operasional")

    add_heading(document, "4.2.4 Prinsip Eskalasi", 3)
    add_numbered(document, [
        "Staff menyelesaikan aktivitas rutin pada lead dan percakapan yang menjadi cakupannya.",
        "Admin Cabang menangani redistribusi PIC, tunggakan tindak lanjut, koreksi transaksi, dan konflik dalam cabang.",
        "Superadmin menangani konflik lintas cabang, akses, konfigurasi, integrasi, dan kegagalan teknis yang berdampak luas.",
        "Direksi memutuskan kebijakan target, retensi, SLA, pengecualian bisnis material, dan prioritas pengembangan.",
    ])

    add_heading(document, "4.3 Proses dan Alur Kerja Bisnis", 2)
    add_heading(document, "4.3.1 Rantai Nilai End-to-End", 3)
    add_body(document, "Rantai nilai mengubah sinyal minat menjadi hubungan, konversi, transaksi, dan informasi kinerja. Setiap tahap menghasilkan status dan bukti yang menjadi input tahap berikutnya.")
    add_figure(document, ASSETS / "proses-bisnis.png", "Gambar 4-1. Proses bisnis end-to-end SIVMI CRM")
    add_table(document, ["Tahap", "Input", "Aktivitas", "Output"], [
        ["Akuisisi", "Identitas/minat/pesan", "Catat sumber, sekolah, cabang, program", "Lead tervalidasi"],
        ["Distribusi", "Lead tervalidasi", "Tetapkan PIC sesuai cabang", "Kepemilikan jelas"],
        ["Tindak lanjut", "Lead + PIC", "Kontak, catat hasil, jadwalkan berikutnya", "Status dan histori"],
        ["Konversi", "Lead siap", "Pilih program dan buat pendaftaran", "Siswa/pendaftaran aktif"],
        ["Transaksi", "Pendaftaran", "Catat nominal, metode, referensi, verifikasi", "Pembayaran tercatat"],
        ["Monitoring", "Data operasional", "Agregasi KPI, target, laporan, audit", "Keputusan dan tindakan"],
    ], [1.0, 1.4, 2.55, 1.55], "Tabel 4-6. Input-output rantai nilai")

    add_heading(document, "4.3.2 Akuisisi dan Validasi Lead", 3)
    add_body(document, "Lead dapat dibuat melalui form, diimpor melalui CSV, atau terbentuk dari pesan WhatsApp masuk. Kanal masuk disimpan terpisah dari sumber pemasaran: kanal menjelaskan bagaimana data masuk (Manual, CSV, WhatsApp), sedangkan sumber menjelaskan asal minat (Iklan, Instagram, Kunjungan, Saudara, Facebook, Tiktok).")
    add_numbered(document, [
        "Terima nama, nomor WhatsApp, sekolah, minat program, sumber, cabang, PIC, status awal, dan tanggal masuk sesuai kanal.",
        "Normalisasi nomor telepon ke bentuk yang dapat dibandingkan; cegah duplikasi tanpa menghapus histori.",
        "Pastikan cabang aktif dan PIC berasal dari cabang yang sama untuk pengguna non-global.",
        "Tentukan hasil pencocokan WhatsApp: NewLead, Matched, Conflict, Unresolved, atau Ignored.",
        "Simpan lead dengan status awal Cold apabila berasal dari pesan baru dan belum ada penilaian lebih lanjut.",
        "Catat aktor, waktu, kanal, dan perubahan untuk audit.",
    ])
    add_callout(document, "Aturan duplikasi WhatsApp", "Nomor yang cocok tepat pada cabang yang sama ditautkan ke lead. Nomor yang cocok lebih dari satu atau berada pada cabang lain ditandai Conflict dan memerlukan resolusi. Sistem tidak boleh menggabungkan data lintas cabang secara otomatis.", LIGHT_AMBER)

    add_heading(document, "4.3.3 Pengelolaan Sekolah saat Akuisisi", 3)
    add_body(document, "Jika sekolah calon siswa belum tersedia, pengguna dapat membuat sekolah dari alur lead tanpa meninggalkan proses. NPSN bersifat opsional. Penamaan harus memakai prefix yang disediakan agar pencarian dan laporan konsisten.")
    add_table(document, ["Prefix", "Makna", "Status sekolah"], [
        ["SMAN", "Sekolah Menengah Atas Negeri", "Negeri"],
        ["SMAS", "Sekolah Menengah Atas Swasta", "Swasta"],
        ["SMKN", "Sekolah Menengah Kejuruan Negeri", "Negeri"],
        ["SMKS", "Sekolah Menengah Kejuruan Swasta", "Swasta"],
        ["MAN", "Madrasah Aliyah Negeri", "Negeri"],
        ["ALUMNI/GAPYEAR", "Calon siswa yang sudah lulus/tidak terikat sekolah aktif", "Kategori khusus"],
    ], [1.35, 3.3, 1.85], "Tabel 4-7. Prefix nama sekolah dan kategori calon siswa")
    add_body(document, "Kluster sekolah dibatasi Grade A sampai Grade E. Definisi yang telah diberikan untuk Grade A-C menjadi acuan bisnis; Grade D-E tetap tersedia di enum tetapi kriteria resminya masih perlu diputuskan manajemen.")
    add_table(document, ["Grade", "Definisi kerja"], [
        ["Grade A", "Sering menyumbang siswa terbanyak walaupun bukan sekolah favorit menurut catatan kementerian"],
        ["Grade B", "Sedikit/jarang menyumbang siswa dan merupakan sekolah favorit menurut catatan kementerian"],
        ["Grade C", "Tidak pernah menyumbang siswa dan merupakan sekolah favorit menurut catatan kementerian"],
        ["Grade D", "Tersedia pada sistem; definisi bisnis belum disahkan"],
        ["Grade E", "Tersedia pada sistem; definisi bisnis belum disahkan"],
    ], [1.0, 5.5], "Tabel 4-8. Klasifikasi grade sekolah")

    add_heading(document, "4.3.4 Distribusi dan Follow-up", 3)
    add_body(document, "Setelah validasi, lead memiliki PIC. Follow-up menyimpan waktu, metode, status kontak, jadwal berikutnya, dan catatan hasil. Metode baseline mencakup WhatsApp, telepon, dan tatap muka; status kontak membedakan terhubung dan tidak terhubung.")
    add_numbered(document, [
        "PIC memprioritaskan lead berdasarkan status, usia lead, jadwal follow-up, dan arahan cabang.",
        "Setiap upaya kontak dicatat sebagai riwayat baru; histori lama tidak ditimpa.",
        "Jika belum menghasilkan keputusan, PIC menetapkan tanggal follow-up berikutnya.",
        "Status lead diperbarui menjadi Cold, Warm, Hot, Daftar, atau Tidak Berminat berdasarkan hasil terbaru.",
        "Follow-up jatuh tempo muncul pada panel tindakan tertunda dan menjadi bahan monitoring Admin Cabang.",
    ])

    add_heading(document, "4.3.5 Komunikasi WhatsApp, Chat, dan Notifikasi", 3)
    add_body(document, "Setiap Admin Cabang atau Staff dapat memiliki nomor WhatsApp dan session WAHA tersendiri. WAHA 2026.8.1 CORE menggunakan engine GOWS. Pesan masuk diterima melalui webhook bertanda tangan HMAC, diproses asynchronous, ditautkan ke akun/cabang, dinormalisasi, dicocokkan ke lead, lalu disimpan sebagai conversation dan message.")
    add_table(document, ["Peristiwa", "Respons sistem", "Notifikasi/tindakan"], [
        ["Pesan inbound baru", "Simpan event, conversation, message, unread count", "Badge inbox dan notifikasi penerima"],
        ["Nomor belum dikenal", "Buat lead Cold atau tandai NewLead sesuai hasil matcher", "PIC/cabang dapat meninjau"],
        ["Nomor cocok", "Tautkan conversation dengan lead", "Histori chat berada pada lead"],
        ["Konflik", "Simpan tanpa merge otomatis", "Admin/Superadmin menyelesaikan"],
        ["Pesan outbound", "Kirim lewat session pemilik dan simpan status", "Sent/Delivered/Read/Failed ditampilkan"],
        ["Kegagalan kirim", "Status Failed dan error tercatat", "Notifikasi sistem ke pengguna"],
    ], [1.45, 3.25, 1.8], "Tabel 4-9. Peristiwa bisnis chat dan notifikasi")

    add_heading(document, "4.3.6 Pendaftaran dan Program Bimbel", 3)
    add_body(document, "Lead yang siap dikonversi memilih program aktif dan dibuatkan pendaftaran. Nama program dibuat otomatis dari tepat tiga komponen: Kategori + Kelas Program + Metode Belajar, misalnya SR GOLD ONLINE. Tipe program, jenjang, dan tingkat kelas tetap disimpan sebagai atribut klasifikasi tetapi tidak menjadi bagian nama program.")
    add_bullets(document, [
        "Kategori, kelas program, tipe program, metode belajar, dan jenjang wajib dipilih sesuai konfigurasi aktif.",
        "Tingkat kelas dapat menyesuaikan jenjang dan tidak selalu diperlukan untuk Gap Year.",
        "Biaya pendaftaran, biaya pendidikan, diskon, jumlah pertemuan, fasilitas, dan status aktif menjadi atribut layanan.",
        "Pendaftaran mengikat lead, program, cabang, tanggal daftar, kelas/angkatan, status, dan catatan.",
    ])

    add_heading(document, "4.3.7 Pembayaran dan Verifikasi", 3)
    add_body(document, "Pembayaran hanya dicatat untuk pendaftaran yang sah dan berada dalam cakupan cabang pengguna. Informasi minimum meliputi tanggal, nominal, metode, status, nomor referensi bila ada, dan catatan. Omzet dashboard/laporan dihitung dari status pembayaran Lunas atau Terverifikasi.")
    add_callout(document, "Pemisahan pencatatan dan kebijakan finansial", "Sistem saat ini mencatat transaksi operasional CRM. Rekonsiliasi bank, jurnal akuntansi, pajak, refund, dan approval finansial berlapis belum dinyatakan sebagai modul akuntansi penuh.", LIGHT_AMBER)

    add_heading(document, "4.3.8 Monitoring, Laporan, dan Audit", 3)
    add_body(document, "Dashboard dan laporan menggunakan filter cabang, tahun, bulan - termasuk Semua bulan - dan pengguna. Widget memakai periode dan scope yang sama sehingga total lead, pendaftaran, siswa aktif, omzet, tren, target, ranking, sekolah/kota teratas, aktivitas, dan tindakan tertunda dapat dibandingkan secara konsisten.")
    add_figure(document, ASSETS / "flowmap.png", "Gambar 4-2. Flowmap proses akuisisi hingga monitoring")
    add_figure(document, ASSETS / "flowchart.png", "Gambar 4-3. Flowchart keputusan lead hingga pembayaran", width=4.35)

    add_heading(document, "4.3.9 SLA dan Titik Kontrol", 3)
    add_table(document, ["Titik kontrol", "Kontrol as-built", "SLA/parameter yang perlu disahkan"], [
        ["Lead baru", "Tanggal masuk, status, PIC, pending action", "Batas waktu respons pertama"],
        ["Follow-up", "Jadwal berikutnya dan overdue", "Frekuensi per status dan batas overdue"],
        ["Konflik WA", "Status Conflict/Unresolved", "Waktu resolusi dan pemilik eskalasi"],
        ["Pendaftaran", "Lead, program, cabang, status", "Dokumen wajib dan masa berlaku"],
        ["Pembayaran", "Nominal, metode, status, referensi", "Batas verifikasi dan approval koreksi"],
        ["Audit", "Aktor, event, cabang, waktu, metadata", "Retensi log dan proses review"],
    ], [1.4, 2.45, 2.65], "Tabel 4-10. Titik kontrol dan parameter SLA")

    add_heading(document, "4.4 Pemodelan Aktor dan Key Performance Indicator (KPI)", 2)
    add_heading(document, "4.4.1 Aktor dan Outcome", 3)
    add_table(document, ["Aktor", "Outcome yang diharapkan", "Fitur utama"], [
        ["Direksi", "Keputusan pertumbuhan berbasis data", "Dashboard, laporan, target, KPI lintas cabang"],
        ["Superadmin", "Sistem konsisten, aman, dan tersedia", "Master, pengaturan, user, WAHA, audit"],
        ["Admin Cabang", "Target cabang tercapai dan antrean terkendali", "Distribusi PIC, target, ranking, pending action"],
        ["Staff", "Lead ditangani cepat dan terdokumentasi", "Leads, follow-up, inbox/chat, pendaftaran, pembayaran"],
        ["Calon siswa/orang tua", "Mendapat respons dan informasi yang jelas", "Interaksi WA/telepon/tatap muka dan proses pendaftaran"],
    ], [1.35, 2.65, 2.5], "Tabel 4-11. Aktor, outcome, dan fitur pendukung")

    add_heading(document, "4.4.2 Definisi KPI Inti", 3)
    add_table(document, ["KPI", "Definisi/formula", "Sumber", "Frekuensi"], [
        ["Total leads", "Jumlah lead pada periode/filter", "leads.tanggal_masuk", "Real-time/filter"],
        ["Follow-up hari ini", "Jumlah aktivitas follow-up pada tanggal berjalan", "followup_riwayat", "Harian"],
        ["Pendaftaran", "Jumlah pendaftaran pada periode/filter", "pendaftaran.tanggal_daftar", "Real-time/filter"],
        ["Siswa aktif", "Jumlah pendaftaran berstatus aktif", "pendaftaran.status_pendaftaran", "Real-time/filter"],
        ["Omzet", "Jumlah nominal pembayaran Lunas/Terverifikasi", "pembayaran + pendaftaran", "Real-time/filter"],
        ["Pencapaian target omzet", "Omzet aktual / target omzet x 100%", "pembayaran + target_kinerja", "Bulanan/tahunan"],
        ["Pencapaian target siswa", "Siswa aktual / target siswa x 100%", "pendaftaran + target_kinerja", "Bulanan/tahunan"],
        ["Ranking pengguna", "Jumlah lead yang menjadi tanggung jawab pengguna", "users + leads", "Periode terpilih"],
        ["Lead overdue", "Follow-up berikutnya <= waktu sekarang dan belum ditangani", "followup_riwayat", "Harian"],
        ["Konversi lead", "Pendaftaran / total lead x 100%", "leads + pendaftaran", "Periode terpilih"],
    ], [1.25, 2.65, 1.55, 1.05], "Tabel 4-12. Kamus KPI inti")

    add_heading(document, "4.4.3 KPI Dimensional", 3)
    add_body(document, "Analisis jumlah lead tidak berhenti pada total. Dashboard menyediakan dimensi program dan pasar untuk membantu keputusan produk serta akuisisi.")
    add_bullets(document, [
        "Program yang diminati: jumlah lead per nama program tiga komponen.",
        "Kategori program, tipe program, metode belajar, kelas program, dan jenjang.",
        "Lima sekolah penyumbang lead terbanyak dan lima wilayah/kota terbanyak.",
        "Sumber lead untuk membandingkan kontribusi Iklan, Instagram, Kunjungan, Saudara, Facebook, dan Tiktok.",
        "Cabang dan PIC untuk mengukur distribusi beban serta performa operasional.",
    ])

    add_heading(document, "4.4.4 Aturan Filter dan Agregasi", 3)
    add_numbered(document, [
        "Semua widget memakai filter global cabang, tahun, bulan, dan pengguna.",
        "Pilihan Semua bulan mengagregasi seluruh bulan pada tahun yang dipilih, bukan seluruh tahun lintas periode.",
        "Pengguna non-global tidak dapat memperluas filter di luar cabangnya.",
        "Filter pengguna harus konsisten terhadap PIC atau aktor yang relevan pada masing-masing dataset.",
        "Nilai kosong ditampilkan sebagai empty state, bukan angka dummy produksi.",
        "Kegagalan query harus menghasilkan error state yang dapat dipahami dan tercatat.",
    ])

    add_heading(document, "4.4.5 Kualitas dan Kepemilikan KPI", 3)
    add_table(document, ["Kontrol", "Tujuan", "Pemilik utama"], [
        ["Definisi tunggal", "Mencegah perbedaan rumus dashboard dan laporan", "Superadmin + Direksi"],
        ["Tanggal acuan", "Memastikan periode memakai kolom tanggal yang tepat", "Pemilik proses"],
        ["Status transaksi", "Mencegah pembayaran belum sah masuk omzet", "Admin Cabang"],
        ["Branch scope", "Mencegah agregasi data di luar kewenangan", "Superadmin"],
        ["Rekonsiliasi", "Membandingkan agregat dengan record transaksi", "Admin Cabang"],
        ["Audit perubahan", "Menjelaskan perubahan target/data sensitif", "Superadmin"],
    ], [1.45, 3.35, 1.7], "Tabel 4-13. Kontrol kualitas KPI")

    add_heading(document, "4.5 Perancangan Enterprise Architecture (Overview)", 2)
    add_heading(document, "4.5.1 Konteks Enterprise Architecture", 3)
    add_body(document, "Enterprise architecture menghubungkan tujuan bisnis dengan kemampuan, proses, aplikasi, data, dan teknologi. SIVMI CRM menggunakan pendekatan berlapis: presentasi Filament/Tailwind, interaksi Livewire, aplikasi Laravel, data Eloquent/MySQL, serta platform Docker dengan Redis, queue, scheduler, Nginx, dan WAHA.")
    add_figure(document, ASSETS / "architecture-layer.png", "Gambar 4-4. Arsitektur layer SIVMI CRM")

    add_heading(document, "4.5.2 Prinsip Arsitektur Bisnis", 3)
    add_table(document, ["Prinsip", "Konsekuensi desain"], [
        ["Satu sumber kebenaran", "Lead, transaksi, target, dan komunikasi tidak dikelola sebagai spreadsheet terpisah setelah masuk sistem"],
        ["Cabang sebagai batas data", "Semua query operasional dan agregasi menerapkan branch scope"],
        ["Histori tidak ditimpa", "Follow-up, pesan, event webhook, dan audit disimpan sebagai rekaman berurutan"],
        ["Konfigurasi sebelum hard-code", "Klasifikasi program dan pengaturan bisnis dikelola melalui menu Pengaturan bila tersedia"],
        ["Otomasi dengan kontrol", "Deduplikasi dan pembuatan lead otomatis tidak boleh menggabungkan konflik secara diam-diam"],
        ["Observabilitas", "Error, retry, failed job, session WA, notifikasi, dan health harus dapat diperiksa"],
        ["Reproducible deployment", "Full Docker menjadi baseline pengembangan menuju VPS"],
    ], [2.0, 4.5], "Tabel 4-14. Prinsip enterprise architecture")

    add_heading(document, "4.5.3 Peta Kapabilitas Bisnis", 3)
    add_table(document, ["Domain", "Kapabilitas inti", "Modul pendukung"], [
        ["Strategi & kinerja", "Target cabang/user, KPI, analisis, ranking", "Dashboard, Target Kinerja, Laporan"],
        ["Akuisisi", "Capture, import/export, sumber, sekolah, minat", "Leads, Sekolah, Program Bimbel"],
        ["Engagement", "Follow-up, reminder, chat dua arah, notifikasi", "Follow-up, Inbox WhatsApp, Notifikasi"],
        ["Konversi", "Pendaftaran, program, kelas/angkatan", "Pendaftaran, Program Bimbel"],
        ["Revenue", "Pencatatan dan verifikasi pembayaran", "Pembayaran, Dashboard, Laporan"],
        ["Governance", "RBAC, cabang, pengguna, audit, pengaturan", "Cabang, Pengguna, Pengaturan, Audit"],
        ["Platform", "Session WA, webhook, queue, scheduler, health", "Koneksi WhatsApp, Docker services"],
    ], [1.45, 2.85, 2.2], "Tabel 4-15. Peta kapabilitas bisnis dan modul")

    add_heading(document, "4.5.4 Peta Modul Aplikasi", 3)
    add_figure(document, ASSETS / "sitemap.png", "Gambar 4-5. Sitemap panel administrasi SIVMI CRM")
    add_body(document, "Menu operasional SIVMI CRM memuat Dashboard, Leads, Follow-up, Inbox WhatsApp, Target Kinerja, dan Laporan. Kelompok Master Data memuat Cabang, Sekolah, Pengguna, Pendaftaran, Pembayaran, Program Bimbel, Pengaturan, Audit Aktivitas, serta Koneksi WhatsApp sesuai kewenangan role.")

    add_heading(document, "4.5.5 Traceability Sasaran ke Teknologi", 3)
    add_table(document, ["Sasaran bisnis", "Kapabilitas", "Data", "Komponen"], [
        ["Respons lead lebih cepat", "Inbox, follow-up, reminder", "conversation, message, followup", "WAHA, queue, Livewire"],
        ["Konversi terukur", "Pipeline dan pendaftaran", "lead, pendaftaran, program", "Filament resources, Eloquent"],
        ["Omzet akurat", "Pembayaran dan laporan", "pembayaran, pendaftaran", "MySQL, dashboard query"],
        ["Data cabang aman", "RBAC dan branch scope", "user, cabang, seluruh FK", "Authorization, scoped query"],
        ["Operasi dapat diaudit", "Audit dan event", "activity_logs, webhook_events", "Observer, log, queue"],
        ["Deployment konsisten", "Container orchestration", "volume/config", "Docker Compose, Nginx, Redis"],
    ], [1.45, 1.6, 1.75, 1.7], "Tabel 4-16. Traceability bisnis-aplikasi-teknologi")

    add_heading(document, "4.5.6 Kondisi As-Built dan Target", 3)
    add_table(document, ["Area", "As-built", "Target pengembangan"], [
        ["Deployment", "Full Docker lokal; siap diarahkan ke VPS", "CI/CD, secret manager, HTTPS, backup dan monitoring terkelola"],
        ["WhatsApp", "WAHA CORE/GOWS, session per user, webhook, chat", "Hardening endpoint, observability, SLA, pengelolaan kapasitas session"],
        ["KPI", "Dashboard/laporan terfilter dan KPI dimensional", "Kamus KPI disahkan, rekonsiliasi berkala, target historis terkunci"],
        ["Notifikasi", "Database notification untuk peristiwa operasional", "Preferensi, escalation matrix, kanal tambahan bila diperlukan"],
        ["Audit", "Activity log dan webhook event", "Retensi, review terjadwal, ekspor insiden, alert anomali"],
        ["Data sekolah", "NPSN opsional, prefix, Grade A-E", "Referensi wilayah baku dan definisi resmi Grade D-E"],
    ], [1.25, 2.55, 2.7], "Tabel 4-17. Baseline dan target arsitektur bisnis")

    add_heading(document, "4.6 Matriks Sasaran, Target Terukur (Objective), dan Indikator Kinerja (KPI)", 2)
    add_heading(document, "4.6.1 Prinsip Penetapan Sasaran", 3)
    add_body(document, "Setiap sasaran harus memiliki baseline, target, periode, pemilik, sumber data, frekuensi review, dan ambang eskalasi. Nilai target demo dari seeder tidak boleh dianggap target produksi. Target produksi ditetapkan per cabang/periode dan dapat dialokasikan ke pengguna.")
    add_table(document, ["Sasaran", "Objective terukur", "KPI", "Pemilik", "Review"], [
        ["Percepat respons", "Lead baru ditangani dalam SLA", "Median waktu respons; lead belum ditangani", "Admin Cabang", "Harian"],
        ["Tingkatkan engagement", "Follow-up tepat waktu", "% follow-up on-time; overdue", "Admin Cabang", "Harian/mingguan"],
        ["Naikkan konversi", "Lebih banyak lead menjadi pendaftar", "Conversion rate; pendaftaran", "Admin Cabang", "Mingguan/bulanan"],
        ["Capai target siswa", "Aktual memenuhi target periode", "% target siswa", "Admin Cabang", "Bulanan"],
        ["Capai target omzet", "Pembayaran sah memenuhi target", "% target omzet; omzet", "Direksi/Admin", "Bulanan"],
        ["Optimalkan kanal", "Fokus pada sumber berkualitas", "Lead dan konversi per sumber", "Direksi/Marketing", "Bulanan"],
        ["Optimalkan produk", "Program sesuai permintaan", "Lead/daftar/omzet per klasifikasi", "Direksi", "Bulanan/kuartal"],
        ["Perkuat pasar sekolah", "Prioritaskan sekolah/wilayah potensial", "Top school/city; konversi per sekolah", "Admin Cabang", "Bulanan"],
        ["Jaga kualitas data", "Minimalkan konflik/duplikasi", "% conflict/unresolved; data wajib lengkap", "Superadmin", "Mingguan"],
        ["Jaga keandalan WA", "Session dan webhook tersedia", "Working session; failed event/message", "Superadmin", "Harian"],
    ], [1.25, 1.65, 1.75, 1.15, 0.7], "Tabel 4-18. Matriks sasaran, objective, dan KPI")

    add_heading(document, "4.6.2 Contoh Kontrak KPI", 3)
    add_table(document, ["Elemen", "Contoh: pencapaian target omzet"], [
        ["Definisi", "Persentase omzet sah terhadap target omzet pada scope dan periode yang sama"],
        ["Formula", "SUM pembayaran nominal berstatus Lunas/Terverifikasi / SUM target omzet x 100%"],
        ["Dimensi", "Cabang, tahun, bulan/Semua bulan, pengguna bila dialokasikan"],
        ["Sumber", "pembayaran, pendaftaran, target_kinerja_cabang, target_kinerja_user"],
        ["Pemilik", "Admin Cabang untuk operasional; Direksi untuk target dan evaluasi"],
        ["Frekuensi", "Real-time pada dashboard; review formal bulanan"],
        ["Pengecualian", "Pembayaran Pending/Gagal/Tidak terverifikasi tidak dihitung"],
    ], [1.45, 5.05], "Tabel 4-19. Contoh kontrak definisi KPI")

    add_heading(document, "4.6.3 Siklus Review Kinerja", 3)
    add_numbered(document, [
        "Awal periode: Direksi/Admin menetapkan target cabang dan alokasi pengguna.",
        "Harian: Staff dan Admin memantau lead baru, follow-up jatuh tempo, chat belum dibaca, serta pembayaran tertunda.",
        "Mingguan: Admin meninjau funnel, sumber lead, ranking, konflik data, dan tindakan korektif.",
        "Bulanan: Direksi merekonsiliasi siswa/omzet, membandingkan cabang dan program, lalu mengesahkan evaluasi.",
        "Setelah review: perubahan target/kebijakan dicatat, dikomunikasikan, dan diaudit.",
    ])

    add_heading(document, "4.7 Aturan Bisnis", 2)
    add_heading(document, "4.7.1 Akses, Role, dan Cabang", 3)
    add_table(document, ["ID", "Aturan bisnis"], [
        ["BR-001", "Pengguna harus aktif dan memiliki role yang valid untuk mengakses panel."],
        ["BR-002", "Superadmin dan Direksi memiliki akses lintas cabang sesuai kebijakan resource; Admin Cabang dan Staff dibatasi ke cabangnya."],
        ["BR-003", "Pengguna non-global tidak boleh memilih cabang lain melalui filter, form, import, URL, atau request yang dimanipulasi."],
        ["BR-004", "PIC lead harus berasal dari cabang yang sama dengan lead, kecuali proses migrasi resmi yang diaudit."],
        ["BR-005", "Aksi create, edit, delete, import, export, dan konfigurasi mengikuti otorisasi resource serta role."],
    ], [0.9, 5.6], "Tabel 4-20. Aturan akses dan cabang")

    add_heading(document, "4.7.2 Lead dan Follow-up", 3)
    add_table(document, ["ID", "Aturan bisnis"], [
        ["BR-006", "Nama dan nomor WhatsApp merupakan identitas operasional utama lead; nomor dinormalisasi sebelum pencocokan."],
        ["BR-007", "Status lead hanya Cold, Warm, Hot, Daftar, atau Tidak Berminat."],
        ["BR-008", "Kanal masuk hanya Manual, CSV, atau WhatsApp; sumber lead hanya Iklan, Instagram, Kunjungan, Saudara, Facebook, atau Tiktok."],
        ["BR-009", "Satu lead dapat memiliki banyak riwayat follow-up; riwayat tidak diubah menjadi satu catatan agregat."],
        ["BR-010", "Follow-up berikutnya yang jatuh tempo harus muncul sebagai tindakan tertunda sesuai scope."],
        ["BR-011", "Lead Tidak Berminat tidak dihapus otomatis; retensinya mengikuti kebijakan manajemen."],
        ["BR-012", "Impor lead memvalidasi cabang, PIC, status, sumber, sekolah/program, dan melaporkan kegagalan baris."],
    ], [0.9, 5.6], "Tabel 4-21. Aturan lead dan follow-up")

    add_heading(document, "4.7.3 Sekolah dan Program", 3)
    add_table(document, ["ID", "Aturan bisnis"], [
        ["BR-013", "NPSN sekolah bersifat opsional; jika diisi harus diperlakukan sebagai identitas referensi dan diperiksa keunikannya."],
        ["BR-014", "Sekolah baru dari form lead memakai prefix SMAN, SMAS, SMKN, SMKS, MAN, atau ALUMNI/GAPYEAR."],
        ["BR-015", "Kluster sekolah hanya Grade A, B, C, D, atau E; perubahan definisi grade harus disahkan dan terdokumentasi."],
        ["BR-016", "Nama program dibuat otomatis dari Kategori + Kelas Program + Metode Belajar dan ditulis uppercase."],
        ["BR-017", "Tipe program, jenjang, dan tingkat kelas tidak ditambahkan ke nama program tiga komponen."],
        ["BR-018", "Hanya program aktif yang dapat dipilih untuk transaksi baru."],
    ], [0.9, 5.6], "Tabel 4-22. Aturan sekolah dan program")

    add_heading(document, "4.7.4 Pendaftaran dan Pembayaran", 3)
    add_table(document, ["ID", "Aturan bisnis"], [
        ["BR-019", "Satu lead memiliki paling banyak satu pendaftaran aktif pada baseline model saat ini."],
        ["BR-020", "Pendaftaran harus mengikat lead, program, cabang, dan tanggal daftar yang valid."],
        ["BR-021", "Pembayaran harus mengacu pada pendaftaran yang berada pada cabang pengguna."],
        ["BR-022", "Omzet hanya menjumlahkan pembayaran berstatus Lunas atau Terverifikasi."],
        ["BR-023", "Nominal pembayaran tidak boleh negatif; referensi transaksi digunakan bila tersedia."],
        ["BR-024", "Perubahan transaksi sensitif harus menghasilkan audit dan mengikuti role."],
    ], [0.9, 5.6], "Tabel 4-23. Aturan pendaftaran dan pembayaran")

    add_heading(document, "4.7.5 WhatsApp, Chat, dan Notifikasi", 3)
    add_table(document, ["ID", "Aturan bisnis"], [
        ["BR-025", "Setiap akun WA terkait satu pengguna, satu cabang, dan session name yang unik."],
        ["BR-026", "Pesan hanya dapat dikirim ketika akun aktif dan session berstatus Working."],
        ["BR-027", "Webhook WAHA harus divalidasi dengan HMAC dan event idempotent sebelum diproses."],
        ["BR-028", "Conversation dan message mengikuti scope akun/cabang; Staff tidak boleh membaca conversation pengguna lain di luar kebijakan visibilitas."],
        ["BR-029", "Nomor cocok cabang sama dapat ditautkan; konflik lintas cabang tidak digabung otomatis."],
        ["BR-030", "Status pesan outbound mengikuti Sent, Delivered, Read/Played, atau Failed; kegagalan menghasilkan notifikasi."],
        ["BR-031", "Secret API, HMAC, dan kredensial tidak boleh disimpan pada activity log atau metadata terbuka."],
    ], [0.9, 5.6], "Tabel 4-24. Aturan WhatsApp dan notifikasi")

    add_heading(document, "4.7.6 KPI, Target, dan Laporan", 3)
    add_table(document, ["ID", "Aturan bisnis"], [
        ["BR-032", "Dashboard dan laporan memakai filter cabang, tahun, bulan, dan pengguna yang konsisten."],
        ["BR-033", "Semua bulan berarti agregasi Januari-Desember pada tahun terpilih."],
        ["BR-034", "Target disimpan per cabang/periode dan dapat dialokasikan ke pengguna."],
        ["BR-035", "Target demo/seeder tidak boleh dipakai sebagai target produksi tanpa keputusan manajemen."],
        ["BR-036", "Widget tidak boleh memakai angka dummy produksi; kondisi kosong dan error harus ditampilkan eksplisit."],
        ["BR-037", "Query agregasi wajib menerapkan branch scope dan meminimalkan N+1 melalui join/eager loading yang sesuai."],
    ], [0.9, 5.6], "Tabel 4-25. Aturan KPI, target, dan laporan")

    add_heading(document, "4.7.7 Audit, Retensi, dan Perubahan", 3)
    add_table(document, ["ID", "Aturan bisnis"], [
        ["BR-038", "Perubahan data penting mencatat aktor, event, cabang, waktu, subjek, serta metadata yang aman."],
        ["BR-039", "Log dan event tidak boleh memuat password, API key, secret webhook, atau payload sensitif yang tidak diperlukan."],
        ["BR-040", "Penghapusan atau koreksi material harus dapat ditelusuri dan mengikuti prosedur persetujuan."],
        ["BR-041", "Perubahan enum, rumus KPI, role, dan aturan cabang harus disertai migration/test serta pembaruan dokumentasi."],
        ["BR-042", "Seeder demo dilewati pada environment production dan tidak boleh menimpa data operasional."],
    ], [0.9, 5.6], "Tabel 4-26. Aturan audit dan perubahan")

    add_heading(document, "4.7.8 Keputusan Bisnis yang Masih Diperlukan", 3)
    add_table(document, ["Keputusan", "Alasan", "Pemilik rekomendasi"], [
        ["Visi, misi, dan nilai resmi", "Tidak tersedia pada codebase", "Direksi"],
        ["Definisi Grade D dan E", "Enum tersedia tetapi kriteria belum disahkan", "Direksi/Operasional"],
        ["SLA respons/follow-up", "Sistem mendukung pengukuran tetapi ambang belum resmi", "Direksi/Admin Cabang"],
        ["Retensi lead/chat/audit", "Diperlukan untuk privasi, kapasitas, dan kepatuhan", "Direksi/Superadmin"],
        ["Approval koreksi/refund", "Belum menjadi workflow finansial berlapis", "Direksi/Keuangan"],
        ["Definisi siswa aktif", "Saat ini mengacu status pendaftaran; perlu sinkron dengan kebijakan akademik", "Direksi/Operasional"],
        ["Kamus KPI resmi", "Rumus as-built perlu disahkan agar menjadi kontrak organisasi", "Direksi"],
        ["Kebijakan kepemilikan conversation", "Perlu aturan saat pergantian PIC/user", "Superadmin/Admin Cabang"],
    ], [2.0, 2.9, 1.6], "Tabel 4-27. Daftar keputusan untuk validasi manajemen")

    add_heading(document, "PENUTUP BAB IV", 1)
    add_body(document, "Arsitektur bisnis SIVMI CRM telah memiliki fondasi as-built yang jelas: role dan cabang, rantai nilai lead-to-revenue, master sekolah dan program, komunikasi WhatsApp, KPI, target, laporan, notifikasi, dan audit. Kesiapan menuju VPS terutama bergantung pada pengesahan kebijakan organisasi, SLA, kamus KPI, retensi, serta hardening operasional yang telah diidentifikasi.")
    add_callout(document, "Kriteria BAB IV dinyatakan disahkan", "Direksi menyetujui profil/visi/misi, struktur kewenangan, RACI, definisi KPI, target, SLA, aturan retensi, definisi Grade D-E, serta pengecualian proses finansial; hasil keputusan kemudian diterapkan pada konfigurasi, SOP, test, dan dokumentasi.", LIGHT_GREEN)

    add_heading(document, "SUMBER VERIFIKASI INTERNAL", 1)
    add_bullets(document, [
        "Codebase SIVMI CRM pada baseline 24 Agustus 2026: app/Enums, app/Models, app/Filament, app/Services, database/migrations, database/seeders, dan tests/Feature.",
        "Dokumentasi dan Pengembangan Sistem CRM SIVMI v1.6 - Sementara.docx.",
        "Diagram Gabungan SIVMI CRM v1.4.drawio, khususnya arsitektur sistem, sitemap, proses bisnis, flowmap, flowchart, WAHA, full Docker, chat/notifikasi, dan quality gate.",
        "Panduan instalasi pertama, konfigurasi WAHA, deployment, serta skrip full Docker pada folder docs/ dan docker/.",
        "Baseline quality gate: 72 test lulus dengan 448 assertion dan build produksi Vite berhasil.",
    ])

    document.core_properties.title = "BAB IV - Arsitektur Bisnis SIVMI CRM"
    document.core_properties.subject = "Arsitektur bisnis, proses, stakeholder, KPI, enterprise architecture, sasaran, dan aturan bisnis"
    document.core_properties.author = "Tim Pengembangan SIVMI CRM"
    document.core_properties.keywords = "SIVMI CRM, arsitektur bisnis, KPI, proses bisnis, RBAC, WAHA, Docker"
    document.core_properties.comments = "Dibuat berdasarkan baseline as-built 24 Agustus 2026."
    document.core_properties.created = datetime(2026, 8, 24)
    document.core_properties.modified = datetime(2026, 8, 24)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
