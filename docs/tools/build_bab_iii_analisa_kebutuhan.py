from __future__ import annotations

from pathlib import Path
import sys

PROJECT_ROOT = Path(__file__).resolve().parents[2]
DOC_TOOLS = PROJECT_ROOT / "storage" / "app" / "doc-tools"
sys.path.insert(0, str(DOC_TOOLS))

from PIL import Image, ImageDraw
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

import build_bab_iv_arsitektur_bisnis as b
import build_bab_vii_arsitektur_teknologi as v7


OUTPUT = PROJECT_ROOT / "docs" / "BAB III - Analisa Kebutuhan SIVMI CRM v1.0.docx"
ASSETS = PROJECT_ROOT / "docs" / "_bab_iii_assets"
COVER = ASSETS / "cover-bab-iii.png"

PRIMARY = "#1476B8"
PRIMARY_DARK = "#075985"
INK = "#0B2545"
MUTED = "#64748B"
BG = "#F6F8FC"
GREEN = "#10B981"
AMBER = "#F59E0B"
ROSE = "#F43F5E"
PURPLE = "#A855F7"


def font(name: str, size: int):
    return v7.font(name, size)


def canvas(title: str, subtitle: str, size=(1800, 1080)):
    return v7.canvas(title, subtitle, size)


def box(draw, xy, title, lines, fill="#F8FAFC", outline=PRIMARY, title_size=27, body_size=20):
    return v7.box(draw, xy, title, lines, fill, outline, title_size, body_size)


def arrow(draw, start, end, color="#64748B", width=5):
    return v7.arrow(draw, start, end, color, width)


def make_cover() -> None:
    ASSETS.mkdir(parents=True, exist_ok=True)
    image = Image.new("RGB", (1600, 2200), BG)
    draw = ImageDraw.Draw(image)
    for y in range(2200):
        ratio = y / 2200
        start, end = (20, 118, 184), (7, 89, 133)
        draw.line((0, y, 1600, y), fill=tuple(round(start[i] * (1 - ratio) + end[i] * ratio) for i in range(3)))
    draw.ellipse((1010, -180, 1770, 580), fill=(27, 139, 199))
    draw.ellipse((1140, 140, 1640, 640), fill=(8, 103, 151))
    draw.rounded_rectangle((120, 250, 1480, 1910), radius=44, fill="white")
    draw.text((190, 350), "SIVMI CRM", font=font("arialbd.ttf", 34), fill=PRIMARY)
    draw.text((190, 610), "BAB III", font=font("arialbd.ttf", 48), fill=PRIMARY_DARK)
    draw.multiline_text((190, 700), "ANALISA\nKEBUTUHAN", font=font("arialbd.ttf", 88), fill=INK, spacing=18)
    draw.text((190, 1045), "Kebutuhan fungsional dan nonfungsional", font=font("arial.ttf", 31), fill="#475569")
    draw.text((190, 1097), "berdasarkan implementasi terbaru SIVMI CRM", font=font("arial.ttf", 31), fill="#475569")
    draw.rounded_rectangle((190, 1245, 805, 1325), radius=40, fill="#EAF4FB")
    draw.text((230, 1267), "BASELINE APLIKASI 25 AGUSTUS 2026", font=font("arial.ttf", 28), fill=PRIMARY_DARK)
    draw.text((190, 1610), "Versi dokumen 1.0", font=font("arial.ttf", 28), fill=MUTED)
    draw.text((190, 1665), "Sasaran: pemilik bisnis, analis, pengembang, QA, dan administrator", font=font("arial.ttf", 25), fill=MUTED)
    draw.text((190, 1775), "SIVMI CRM - Catat, Kelola, Closing", font=font("arialbd.ttf", 34), fill=PRIMARY)
    image.save(COVER)


def make_context() -> None:
    image, draw = canvas("Konteks Kebutuhan SIVMI CRM", "Satu ruang kerja CRM menghubungkan aktor, data operasional, dan layanan eksternal")
    box(draw, (650, 360, 1150, 700), "SIVMI CRM", "Dashboard | Lead | Follow-up\nWhatsApp | Pendaftaran | Pembayaran\nTarget | Laporan | Master Data", "#EAF4FB", PRIMARY_DARK, 31, 21)
    actors = [
        (60, 150, 430, 420, "TIM OPERASIONAL", "Admin Cabang\nStaff", PRIMARY),
        (60, 650, 430, 920, "MANAJEMEN", "Superadmin\nDireksi", PURPLE),
        (1370, 150, 1740, 420, "CALON SISWA", "WhatsApp\nData sekolah\nMinat program", GREEN),
        (1370, 650, 1740, 920, "LAYANAN", "MySQL | Redis\nWAHA/GOWS\nEmail/notification", AMBER),
    ]
    for x1, y1, x2, y2, title, lines, color in actors:
        box(draw, (x1, y1, x2, y2), title, lines, "#F8FAFC", color, 26, 20)
    arrow(draw, (430, 285), (650, 455)); arrow(draw, (430, 785), (650, 605))
    arrow(draw, (1150, 455), (1370, 285)); arrow(draw, (1150, 605), (1370, 785))
    image.save(ASSETS / "konteks-kebutuhan.png")


def make_function_map() -> None:
    image, draw = canvas("Dekomposisi Kebutuhan Fungsional", "Fungsi dikelompokkan berdasarkan alur kerja, data master, analitik, dan kontrol")
    box(draw, (650, 80, 1150, 290), "KEBUTUHAN FUNGSIONAL", "Akses | Operasi | Analitik | Kontrol", "#EAF4FB", PRIMARY_DARK, 29, 21)
    groups = [
        (45, 480, 385, 900, "AKSES", "Login & session\nRBAC\nScope cabang\nProfil & notifikasi", PRIMARY),
        (430, 480, 770, 900, "OPERASI", "Lead & CSV\nFollow-up\nWhatsApp\nDaftar & bayar", GREEN),
        (815, 480, 1155, 900, "MASTER DATA", "Cabang\nSekolah\nPengguna\nProgram", PURPLE),
        (1200, 480, 1540, 900, "ANALITIK", "Dashboard\nTarget\nLaporan\nKPI program", AMBER),
        (1585, 480, 1770, 900, "KONTROL", "Setting\nAudit\nError/empty\nImport/export", ROSE),
    ]
    for x1, y1, x2, y2, title, lines, color in groups:
        box(draw, (x1, y1, x2, y2), title, lines, "#F8FAFC", color, 24, 18)
        arrow(draw, (900, 290), ((x1 + x2) // 2, 480), "#94A3B8", 4)
    image.save(ASSETS / "dekomposisi-fungsional.png")


def make_rbac() -> None:
    image, draw = canvas("Kebutuhan RBAC dan Scope Cabang", "Hak menu, aksi, query, widget, import/export, dan notifikasi harus konsisten")
    roles = [
        (60, 190, 430, 820, "SUPERADMIN", "Semua cabang\nKelola sistem & master\nAudit global\nKoneksi WA global", PRIMARY),
        (500, 190, 860, 820, "ADMIN CABANG", "Cabang sendiri\nKelola tim/operasi\nTarget cabang\nWA cabang", PURPLE),
        (930, 190, 1290, 820, "STAFF", "Cabang sendiri\nAssignment sendiri\nFollow-up & chat\nOperasi sesuai izin", GREEN),
        (1360, 190, 1730, 820, "DIREKSI", "Semua cabang\nDashboard/laporan\nMonitoring\nBaca saja operasional", AMBER),
    ]
    for x1, y1, x2, y2, title, lines, color in roles:
        box(draw, (x1, y1, x2, y2), title, lines, "#F8FAFC", color, 27, 20)
    draw.text((900, 950), "Cabang baseline: Bandung | JakSel | JakPus", anchor="mm", font=font("arialbd.ttf", 25), fill=MUTED)
    image.save(ASSETS / "rbac-scope.png")


def make_lead_flow() -> None:
    image, draw = canvas("Alur Kebutuhan Lead sampai Omzet", "Setiap perpindahan tahap mempertahankan cabang, PIC, sumber, program, waktu, dan audit")
    nodes = [
        (50, 240, 330, 610, "MASUK", "Manual | CSV | WA\nnomor dinormalisasi", PRIMARY),
        (390, 240, 670, 610, "LEAD", "Cold | Warm | Hot\nTidak Berminat", AMBER),
        (730, 240, 1010, 610, "FOLLOW-UP", "Metode | hasil\njadwal berikutnya", PURPLE),
        (1070, 240, 1350, 610, "DAFTAR", "Program | cabang\nstatus siswa", GREEN),
        (1410, 240, 1750, 610, "BAYAR & KPI", "Lunas/Terverifikasi\nomzet & laporan", PRIMARY_DARK),
    ]
    for x1, y1, x2, y2, title, lines, color in nodes:
        box(draw, (x1, y1, x2, y2), title, lines, "#F8FAFC", color, 25, 19)
    for x in (330, 670, 1010, 1350): arrow(draw, (x, 425), (x + 60, 425))
    draw.rounded_rectangle((310, 760, 1490, 950), radius=22, fill="#FFF7ED", outline=AMBER, width=3)
    draw.multiline_text((900, 855), "Duplikat nomor dan konflik lintas cabang tidak boleh membuat record baru otomatis.\nStatus Daftar menghubungkan lead ke pendaftaran; omzet berasal dari pembayaran valid.", anchor="mm", align="center", spacing=8, font=font("arialbd.ttf", 22), fill=INK)
    image.save(ASSETS / "alur-lead-omzet.png")


def make_waha() -> None:
    image, draw = canvas("Kebutuhan Integrasi WhatsApp WAHA/GOWS", "Webhook disimpan dahulu, diverifikasi HMAC, lalu diproses queue secara idempotent")
    nodes = [
        (55, 260, 335, 630, "PENGIRIM", "Pesan personal\nnomor / @lid", PRIMARY),
        (405, 260, 685, 630, "WAHA/GOWS", "Session Working\nAPI key\nwebhook", PURPLE),
        (755, 260, 1035, 630, "INGEST", "HMAC\nsimpan event\nrespons cepat", AMBER),
        (1105, 260, 1385, 630, "QUEUE", "normalisasi\ndedupe\nmatching cabang", GREEN),
        (1455, 260, 1745, 630, "HASIL", "Inbox/chat/ACK\nNewLead/Matched\nConflict", PRIMARY_DARK),
    ]
    for x1, y1, x2, y2, title, lines, color in nodes:
        box(draw, (x1, y1, x2, y2), title, lines, "#F8FAFC", color, 24, 18)
    for x in (335, 685, 1035, 1385): arrow(draw, (x, 445), (x + 70, 445))
    draw.text((900, 850), "Satu admin/staff = satu nomor = satu session aktif; group/newsletter/status diabaikan.", anchor="mm", font=font("arialbd.ttf", 23), fill=MUTED)
    image.save(ASSETS / "kebutuhan-waha.png")


def make_quality() -> None:
    image, draw = canvas("Model Kebutuhan Nonfungsional", "Kualitas harus dibuktikan melalui pengukuran, test, monitoring, dan prosedur pemulihan")
    box(draw, (650, 100, 1150, 310), "KUALITAS SISTEM", "Aman | Cepat | Andal | Mudah dirawat", "#EAF4FB", PRIMARY_DARK, 29, 21)
    items = [
        (55, 500, 370, 890, "KEAMANAN", "RBAC & scope\nHTTPS/secret\nCSRF/session\naudit", ROSE),
        (415, 500, 730, 890, "KINERJA", "query efisien\nindex/pagination\nqueue\ncache", PRIMARY),
        (775, 500, 1090, 890, "KEANDALAN", "health check\nretry/idempotency\nbackup/restore\nvolume", GREEN),
        (1135, 500, 1450, 890, "USABILITY", "responsive\nBahasa Indonesia\nskeleton/empty/error\naksesibilitas", PURPLE),
        (1495, 500, 1760, 890, "OPERABILITAS", "Docker\nlogging\nmonitoring\ntest/build", AMBER),
    ]
    for x1, y1, x2, y2, title, lines, color in items:
        box(draw, (x1, y1, x2, y2), title, lines, "#F8FAFC", color, 23, 18)
        arrow(draw, (900, 310), ((x1 + x2) // 2, 500), "#94A3B8", 4)
    image.save(ASSETS / "model-nonfungsional.png")


def make_traceability() -> None:
    image, draw = canvas("Traceability Kebutuhan", "Setiap kebutuhan harus mempunyai implementasi, bukti uji, dan pemilik penerimaan")
    nodes = [
        (70, 270, 360, 650, "KEBUTUHAN", "ID FR/NFR\naktor\naturan\nprioritas", PRIMARY),
        (450, 270, 740, 650, "DESAIN", "UI/form\npolicy/scope\nmodel/query\nintegrasi", PURPLE),
        (830, 270, 1120, 650, "IMPLEMENTASI", "Laravel/Filament\nDocker/WAHA\nmigration/config", AMBER),
        (1210, 270, 1500, 650, "VERIFIKASI", "unit/feature\nbuild\nsmoke\nvisual QA", GREEN),
        (1590, 270, 1760, 650, "UAT", "owner\nhasil\ngap", PRIMARY_DARK),
    ]
    for x1, y1, x2, y2, title, lines, color in nodes:
        box(draw, (x1, y1, x2, y2), title, lines, "#F8FAFC", color, 23, 18)
    for x in (360, 740, 1120, 1500): arrow(draw, (x, 460), (x + 90, 460))
    image.save(ASSETS / "traceability.png")


def add_figure(doc: Document, name: str, caption: str, width=6.15) -> None:
    b.add_figure(doc, ASSETS / name, caption, width)


def table(doc: Document, caption: str, headers: list[str], rows: list[list[str]], widths: list[float]) -> None:
    b.add_table(doc, headers, rows, widths, caption)


def callout(doc: Document, title: str, body: str, fill: str = b.LIGHT_BLUE) -> None:
    b.add_callout(doc, title, body, fill)
    doc.tables[-1].rows[0]._tr.get_or_add_trPr().append(b.OxmlElement("w:cantSplit"))


def page(doc: Document) -> None:
    doc.add_page_break()


def build_document() -> None:
    for fn in (make_cover, make_context, make_function_map, make_rbac, make_lead_flow, make_waha, make_quality, make_traceability):
        fn()

    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.5); sec.page_height = Inches(11)
    sec.top_margin = Inches(0.85); sec.bottom_margin = Inches(0.85)
    sec.left_margin = Inches(1); sec.right_margin = Inches(1)
    sec.header_distance = Inches(0.4); sec.footer_distance = Inches(0.4)
    b.configure_styles(doc)
    sec.header.paragraphs[0].text = "SIVMI CRM  |  BAB III: ANALISA KEBUTUHAN"
    for run in sec.header.paragraphs[0].runs:
        b.set_run_font(run, 8.5, True, b.PRIMARY)
    b.add_page_number(sec.footer.paragraphs[0])

    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(0)
    pic = p.add_run().add_picture(str(COVER), width=Inches(6.1))
    pic._inline.docPr.set("descr", "Sampul BAB III Analisa Kebutuhan SIVMI CRM")
    page(doc)

    b.add_heading(doc, "KENDALI DOKUMEN", 1)
    table(doc, "Tabel 0-1. Kendali dokumen", ["Atribut", "Nilai"], [
        ["Judul", "BAB III: Analisa Kebutuhan SIVMI CRM"],
        ["Versi", "1.0"],
        ["Tanggal baseline", "25 Agustus 2026"],
        ["Tujuan", "Mendefinisikan kebutuhan fungsional dan nonfungsional yang dapat ditelusuri ke implementasi, pengujian, dan penerimaan"],
        ["Pemangku kepentingan", "Direksi, Superadmin, Admin Cabang, Staff, administrator, analis, pengembang, QA, dan pemilik proses bisnis"],
        ["Lingkungan", "Full Docker sebagai baseline lokal; Docker Compose production sebagai target VPS; Laragon/hybrid legacy"],
        ["Sumber kebenaran", "Resource/Page Filament, enum, model/migration, policy/scope, service WAHA, config, Docker Compose, route, scheduler, seeder, dan test suite"],
        ["Quality baseline", "72 test dan 448 assertion serta build Vite production lulus pada baseline proyek"],
    ], [1.55, 4.95])
    callout(doc, "Tanpa daftar isi", "Sesuai permintaan, dokumen ini tidak memuat halaman daftar isi. Setelah kendali dokumen, pembahasan langsung masuk ke ringkasan analisis dan BAB III.")

    b.add_heading(doc, "RINGKASAN ANALISIS", 1)
    b.add_body(doc, "SIVMI CRM dibutuhkan untuk memusatkan pengelolaan calon siswa dari lead masuk sampai pendaftaran, pembayaran, dan pelaporan. Sistem juga menghubungkan nomor WhatsApp masing-masing admin/staff melalui WAHA/GOWS, mencegah duplikasi nomor, membatasi data menurut role dan cabang, serta menyediakan bukti aktivitas melalui notifikasi dan audit.")
    add_figure(doc, "konteks-kebutuhan.png", "Gambar 3-1. Konteks kebutuhan SIVMI CRM")
    table(doc, "Tabel 0-2. Pernyataan kebutuhan utama", ["Masalah", "Kebutuhan sistem", "Hasil yang diharapkan"], [
        ["Data lead tersebar", "Satu basis data dan alur status yang konsisten", "Tidak ada pencatatan ganda lintas kanal"],
        ["Follow-up tidak terukur", "Riwayat, hasil, dan jadwal tindak lanjut", "Tindakan tertunda dapat dipantau"],
        ["Chat terpisah dari CRM", "Session WA per pengguna, webhook, inbox, dan matching lead", "Percakapan dapat ditelusuri ke pemilik/cabang"],
        ["Kinerja sulit dibandingkan", "Filter global, KPI, target, tren, ranking, sekolah/wilayah", "Keputusan berbasis data sesuai scope"],
        ["Hak akses rawan tercampur", "RBAC dan branch scope pada menu, aksi, query, widget, dan export", "Tidak ada kebocoran lintas cabang"],
        ["Operasi sulit dipulihkan", "Docker, health, log, queue, backup/restore, test", "Lingkungan dapat diulang dan diaudit"],
    ], [1.65, 2.95, 1.9])
    callout(doc, "Status kebutuhan", "Label As-built berarti sudah terdapat pada implementasi yang diaudit. Label Target berarti persyaratan penerimaan/operasional yang harus diverifikasi saat UAT atau sebelum produksi; target angka perlu disahkan pemilik sistem.", b.LIGHT_AMBER)

    b.add_heading(doc, "BAB III: ANALISA KEBUTUHAN", 1)
    b.add_heading(doc, "3.1 Kebutuhan Fungsional", 2)
    b.add_body(doc, "Kebutuhan fungsional menjelaskan perilaku yang harus diberikan sistem kepada aktor. Setiap kebutuhan diberi ID agar dapat dipetakan ke desain, kode, test, dan UAT. Prioritas Must berarti fungsi wajib untuk operasi; Should berarti penting tetapi dapat dirilis bertahap; Could berarti peningkatan setelah baseline stabil.")
    add_figure(doc, "dekomposisi-fungsional.png", "Gambar 3-2. Dekomposisi kebutuhan fungsional")
    page(doc)

    b.add_heading(doc, "3.1.1 Aktor, Role, dan Scope", 3)
    add_figure(doc, "rbac-scope.png", "Gambar 3-3. Kebutuhan RBAC dan scope cabang")
    table(doc, "Tabel 3-1. Aktor dan tujuan", ["Aktor", "Tujuan", "Scope utama"], [
        ["Superadmin", "Mengelola sistem, master data, user, pengaturan, koneksi, dan audit", "Global"],
        ["Admin Cabang", "Mengelola tim dan operasi cabang", "Cabang sendiri"],
        ["Staff", "Mengelola lead, follow-up, chat, dan transaksi yang ditugaskan", "Cabang/assignment"],
        ["Direksi", "Memantau KPI, target, dan laporan", "Global baca-saja operasional"],
        ["Calon siswa", "Berinteraksi melalui WhatsApp dan proses pendaftaran", "Data pribadi sendiri"],
        ["Administrator", "Menjaga runtime, data, queue, WAHA, backup, dan keamanan", "Teknis sesuai otorisasi"],
    ], [1.25, 3.45, 1.8])
    table(doc, "Tabel 3-2. Kebutuhan autentikasi dan otorisasi", ["ID", "Kebutuhan", "Prioritas", "Status"], [
        ["FR-AUTH-01", "Sistem menyediakan login dengan email dan kata sandi bagi akun aktif.", "Must", "As-built"],
        ["FR-AUTH-02", "Session panel harus dipertahankan secara aman dan dapat diakhiri melalui logout.", "Must", "As-built"],
        ["FR-RBAC-01", "Menu, page, resource, aksi, query, widget, notifikasi, import/export mengikuti role.", "Must", "As-built/verify"],
        ["FR-RBAC-02", "Superadmin dan Direksi dapat membaca lintas cabang; Direksi tidak mengubah data operasional.", "Must", "As-built"],
        ["FR-RBAC-03", "Admin Cabang dan Staff hanya mengakses data cabang yang ditetapkan.", "Must", "As-built/verify"],
        ["FR-RBAC-04", "Staff dibatasi lagi berdasarkan assignment/pemilik ketika aturan modul mewajibkan.", "Must", "As-built/verify"],
        ["FR-RBAC-05", "Cabang baseline yang tersedia adalah Bandung, JakSel, dan JakPus.", "Must", "As-built"],
    ], [1.15, 3.85, 0.75, 0.75])

    b.add_heading(doc, "3.1.2 Dashboard, Filter, KPI, dan Laporan", 3)
    table(doc, "Tabel 3-3. Kebutuhan dashboard dan laporan", ["ID", "Kebutuhan", "Prioritas", "Status"], [
        ["FR-DASH-01", "Dashboard menampilkan hero, ringkasan hari ini, KPI, target, tren, aktivitas, quick action, ranking, dan tindakan tertunda.", "Must", "As-built"],
        ["FR-DASH-02", "Filter Cabang, Tahun, Bulan, dan Pengguna berlaku global ke seluruh widget yang relevan.", "Must", "As-built"],
        ["FR-DASH-03", "Filter bulan menyediakan Semua bulan untuk agregasi satu tahun.", "Must", "As-built"],
        ["FR-DASH-04", "Tren leads dan pendaftaran menampilkan data sesuai periode dan scope, serta empty/error state bila tidak tersedia.", "Must", "As-built"],
        ["FR-DASH-05", "KPI lead dapat dianalisis per program, kategori, tipe, metode, kelas program, dan jenjang.", "Must", "As-built"],
        ["FR-DASH-06", "Dashboard menampilkan lima sekolah dan lima wilayah/kota dengan kontribusi lead terbanyak.", "Should", "As-built"],
        ["FR-REP-01", "Laporan menyajikan KPI, tren, funnel status, pembayaran, cabang, program, dan progress target.", "Must", "As-built"],
        ["FR-REP-02", "Export laporan mengikuti filter, role, cabang, dan pengaturan izin Staff.", "Must", "As-built/verify"],
    ], [1.15, 3.85, 0.75, 0.75])

    b.add_heading(doc, "3.1.3 Lead, Sekolah, dan Import/Export", 3)
    add_figure(doc, "alur-lead-omzet.png", "Gambar 3-4. Alur kebutuhan lead sampai omzet")
    table(doc, "Tabel 3-4. Kebutuhan lead", ["ID", "Kebutuhan", "Prioritas", "Status"], [
        ["FR-LEAD-01", "Pengguna berwenang dapat membuat, melihat, mengubah, mencari, memfilter, dan menugaskan lead dalam scope.", "Must", "As-built"],
        ["FR-LEAD-02", "Lead menyimpan nama, nomor WhatsApp, sekolah, program diminati, sumber, cabang, PIC, status, dan tanggal masuk.", "Must", "As-built"],
        ["FR-LEAD-03", "Status lead dibatasi pada Cold, Warm, Hot, Daftar, dan Tidak Berminat.", "Must", "As-built"],
        ["FR-LEAD-04", "Sumber lead dibatasi pada Iklan, Instagram, Kunjungan, Saudara, Facebook, dan Tiktok.", "Must", "As-built"],
        ["FR-LEAD-05", "Kanal masuk dicatat sebagai Manual, CSV, atau WhatsApp.", "Must", "As-built"],
        ["FR-LEAD-06", "Nomor WhatsApp dinormalisasi agar variasi 08..., 8..., +62..., dan 62... dapat dibandingkan.", "Must", "As-built"],
        ["FR-LEAD-07", "Nomor yang sama harus dideteksi sebelum membuat duplikat; konflik lintas cabang memerlukan review manual.", "Must", "As-built"],
        ["FR-LEAD-08", "Import dan export lead CSV mengikuti validasi enum, referensi, role, dan scope cabang.", "Must", "As-built"],
        ["FR-LEAD-09", "Import menyediakan contoh/template, laporan berhasil/gagal, dan tidak membuat data parsial yang tidak valid.", "Should", "As-built/verify"],
    ], [1.15, 3.85, 0.75, 0.75])
    table(doc, "Tabel 3-5. Kebutuhan sekolah", ["ID", "Kebutuhan", "Prioritas", "Status"], [
        ["FR-SCH-01", "Sekolah dapat dikelola melalui CRUD serta import/export CSV.", "Must", "As-built"],
        ["FR-SCH-02", "NPSN bersifat opsional dan tidak menghalangi penyimpanan sekolah.", "Must", "As-built"],
        ["FR-SCH-03", "Kluster sekolah menggunakan Grade A sampai Grade E.", "Must", "As-built"],
        ["FR-SCH-04", "Nama sekolah baru memakai prefix SMAN, SMAS, SMKN, SMKS, MAN, atau ALUMNI/GAPYEAR.", "Must", "As-built"],
        ["FR-SCH-05", "Prefix menentukan status Negeri, Swasta, atau Lainnya dan mencegah format nama yang tidak konsisten.", "Must", "As-built"],
        ["FR-SCH-06", "Sekolah dapat dibuat dari form Lead saat belum tersedia tanpa meninggalkan data lead.", "Must", "As-built"],
        ["FR-SCH-07", "Pilihan ALUMNI/GAPYEAR tidak mewajibkan nama sekolah.", "Must", "As-built"],
    ], [1.15, 3.85, 0.75, 0.75])

    b.add_heading(doc, "3.1.4 Follow-up, Pendaftaran, Pembayaran, dan Program", 3)
    table(doc, "Tabel 3-6. Kebutuhan proses operasional", ["ID", "Kebutuhan", "Prioritas", "Status"], [
        ["FR-FU-01", "Follow-up menyimpan lead, waktu, metode, status kontak, tanggal berikutnya, catatan, dan pengguna.", "Must", "As-built"],
        ["FR-FU-02", "Follow-up jatuh tempo harus dapat ditampilkan sebagai tindakan tertunda/reminder.", "Must", "As-built"],
        ["FR-REG-01", "Pendaftaran menghubungkan lead, program, cabang, tanggal, angkatan, status, dan catatan.", "Must", "As-built"],
        ["FR-REG-02", "Pendaftaran aktif dihitung sebagai siswa aktif dan dapat menjadi dasar target siswa.", "Must", "As-built"],
        ["FR-PAY-01", "Pembayaran menghubungkan pendaftaran dengan tanggal, nominal, metode, status, referensi, dan catatan.", "Must", "As-built"],
        ["FR-PAY-02", "Omzet hanya menghitung pembayaran Lunas atau Terverifikasi sesuai periode dan scope.", "Must", "As-built"],
        ["FR-PRG-01", "Program menyimpan kategori, kelas program, tipe, metode, jenjang, tingkat kelas, biaya, diskon, pertemuan, fasilitas, dan status aktif.", "Must", "As-built"],
        ["FR-PRG-02", "Nama program dibentuk otomatis dari Kategori + Kelas program + Metode belajar, misalnya SR GOLD ONLINE.", "Must", "As-built"],
        ["FR-PRG-03", "Opsi klasifikasi program dapat dikelola dari Pengaturan dan tidak dapat dihapus saat masih digunakan.", "Must", "As-built"],
    ], [1.15, 3.85, 0.75, 0.75])

    b.add_heading(doc, "3.1.5 Target Kinerja dan Master Data", 3)
    table(doc, "Tabel 3-7. Kebutuhan target dan master data", ["ID", "Kebutuhan", "Prioritas", "Status"], [
        ["FR-TGT-01", "Target cabang menyimpan cabang, tahun, bulan, status, target omzet, target siswa, pembuat, dan catatan.", "Must", "As-built"],
        ["FR-TGT-02", "Target pengguna mengalokasikan omzet dan siswa kepada user dalam target cabang.", "Must", "As-built"],
        ["FR-TGT-03", "Dashboard membandingkan realisasi terhadap target yang sesuai filter.", "Must", "As-built"],
        ["FR-MST-01", "Cabang, sekolah, pengguna, pendaftaran, pembayaran, program, pengaturan, audit, dan koneksi WA dikelompokkan sebagai MASTER DATA.", "Should", "As-built"],
        ["FR-MST-02", "Pengguna menyimpan identitas, email unik, role, cabang, status aktif, dan nomor WhatsApp operasional.", "Must", "As-built"],
        ["FR-MST-03", "Satu pengguna hanya memiliki satu koneksi WhatsApp aktif.", "Must", "As-built"],
    ], [1.15, 3.85, 0.75, 0.75])

    b.add_heading(doc, "3.1.6 Integrasi WhatsApp, Chat, dan Notifikasi", 3)
    add_figure(doc, "kebutuhan-waha.png", "Gambar 3-5. Kebutuhan integrasi WhatsApp WAHA/GOWS")
    table(doc, "Tabel 3-8. Kebutuhan WAHA dan chat", ["ID", "Kebutuhan", "Prioritas", "Status"], [
        ["FR-WA-01", "Superadmin/Admin Cabang dapat membuat koneksi WA untuk pengguna dalam scope dan menyiapkan session GOWS.", "Must", "As-built"],
        ["FR-WA-02", "Sistem menyediakan aksi Mulai, Sinkronkan, tampilkan QR, Aktifkan chat, dan status Disconnected/Connecting/Working/Failed.", "Must", "As-built"],
        ["FR-WA-03", "Webhook WAHA diverifikasi HMAC, disimpan sebagai event, lalu diproses queue.", "Must", "As-built"],
        ["FR-WA-04", "Payload GOWS @lid menggunakan SenderAlt untuk nomor dan PushName untuk nama kontak.", "Must", "As-built"],
        ["FR-WA-05", "Pesan personal menghasilkan conversation/message, unread, arah pesan, waktu, dan status ACK.", "Must", "As-built"],
        ["FR-WA-06", "Nomor baru membuat lead Cold melalui kanal WhatsApp; satu kecocokan menautkan lead otomatis.", "Must", "As-built"],
        ["FR-WA-07", "Nomor ganda/lintas cabang berstatus Conflict; group, newsletter, dan status berstatus Ignored.", "Must", "As-built"],
        ["FR-WA-08", "Pengguna dapat membaca dan mengirim chat dari layout responsif dalam scope pemilik session/cabang.", "Must", "As-built"],
        ["FR-NOT-01", "Sistem menyediakan database notification, polling, unread badge, dan notifikasi sukses/peringatan/error.", "Must", "As-built"],
        ["FR-NOT-02", "Reminder follow-up dan status session WA disinkronkan oleh scheduler tanpa overlap.", "Must", "As-built"],
    ], [1.15, 3.85, 0.75, 0.75])

    b.add_heading(doc, "3.1.7 Pengaturan dan Audit Aktivitas", 3)
    table(doc, "Tabel 3-9. Kebutuhan kontrol sistem", ["ID", "Kebutuhan", "Prioritas", "Status"], [
        ["FR-SET-01", "Superadmin dapat mengelola identitas aplikasi, dukungan, izin export Staff, dan opsi program.", "Must", "As-built"],
        ["FR-AUD-01", "Aktivitas create/update/delete/login/logout dan perubahan penting dicatat dengan actor, cabang, waktu, subject, event, dan properties aman.", "Must", "As-built"],
        ["FR-AUD-02", "Audit dapat dicari dan difilter berdasarkan cabang, pengguna, event, dan rentang tanggal sesuai scope.", "Must", "As-built"],
        ["FR-AUD-03", "Audit tidak menyimpan password, API key, HMAC secret, cookie, QR, atau isi chat sensitif berlebih.", "Must", "Target/verify"],
        ["FR-STATE-01", "Daftar/widget menyediakan skeleton loading, empty state, dan error state yang dapat dipahami.", "Should", "As-built/verify"],
    ], [1.15, 3.85, 0.75, 0.75])

    b.add_heading(doc, "3.1.8 Aturan Bisnis Fungsional", 3)
    table(doc, "Tabel 3-10. Aturan bisnis kunci", ["Kode", "Aturan", "Dampak"], [
        ["BR-01", "Role global hanya Superadmin dan Direksi; Direksi read-only operasional.", "Semua policy/scope/widget/export"],
        ["BR-02", "Cabang Admin/Staff wajib dan membatasi data.", "Form, query, import, notification"],
        ["BR-03", "Nomor WhatsApp dinormalisasi dan dibandingkan sebelum create.", "Dedupe lead dan conversation"],
        ["BR-04", "Lead WhatsApp baru berstatus Cold; konflik tidak membuat duplikat.", "Webhook processing"],
        ["BR-05", "Nama program hanya tiga unsur: kategori, kelas, metode.", "Create/update/import program"],
        ["BR-06", "NPSN opsional; prefix dan Grade mengendalikan kualitas master sekolah.", "Form/import sekolah"],
        ["BR-07", "Omzet berasal dari pembayaran Lunas/Terverifikasi.", "Dashboard dan laporan"],
        ["BR-08", "Seeder demo dilewati di production.", "Deployment dan integritas data"],
    ], [1.0, 3.8, 1.7])

    b.add_heading(doc, "3.2 Kebutuhan Nonfungsional", 2)
    b.add_body(doc, "Kebutuhan nonfungsional menentukan mutu dan batas operasional sistem. Nilai terukur yang belum tercatat sebagai as-built ditandai Target UAT/Produksi dan harus disahkan sebelum go-live. Pengukuran dilakukan pada beban, infrastruktur, dan volume data yang representatif.")
    add_figure(doc, "model-nonfungsional.png", "Gambar 3-6. Model kebutuhan nonfungsional")

    b.add_heading(doc, "3.2.1 Keamanan dan Privasi", 3)
    table(doc, "Tabel 3-11. Kebutuhan keamanan", ["ID", "Kebutuhan/ukuran", "Verifikasi", "Status"], [
        ["NFR-SEC-01", "Seluruh halaman panel memerlukan autentikasi dan CSRF/session protection.", "Feature test dan browser", "As-built"],
        ["NFR-SEC-02", "Pembatasan role/cabang diterapkan di server, bukan hanya menyembunyikan UI.", "Policy/scope test lintas role", "As-built/verify"],
        ["NFR-SEC-03", "Production memakai HTTPS, APP_DEBUG=false, secure cookie, session encryption, dan secret unik.", "Config/deployment audit", "Target produksi"],
        ["NFR-SEC-04", "WAHA API key dan HMAC secret berbeda; dashboard/Swagger tidak diekspos production.", "Config/network test", "Target produksi"],
        ["NFR-SEC-05", "Password, token, cookie, QR, dump, dan isi chat sensitif tidak masuk log/audit/export tak berwenang.", "Log/privacy review", "Target/verify"],
        ["NFR-SEC-06", "Database, Redis, SMTP, volume, dan backup memakai credential/access minimum.", "Permission/secret audit", "Target produksi"],
        ["NFR-SEC-07", "Formula spreadsheet pada CSV dinetralkan untuk mengurangi injection.", "CSV feature test", "As-built"],
    ], [1.15, 3.45, 1.15, 0.75])

    b.add_heading(doc, "3.2.2 Kinerja dan Efisiensi", 3)
    table(doc, "Tabel 3-12. Kebutuhan kinerja", ["ID", "Kebutuhan/ukuran", "Verifikasi", "Status"], [
        ["NFR-PERF-01", "Query dashboard/report memakai agregasi, eager loading, scope, dan index yang efisien; hindari N+1.", "Query log/profiling", "As-built/verify"],
        ["NFR-PERF-02", "Daftar data memakai pagination, pencarian, filter, dan pemilihan kolom.", "Browser/query review", "As-built"],
        ["NFR-PERF-03", "P95 halaman utama target <= 2 detik pada beban dan data UAT yang disetujui.", "Load/APM test", "Target UAT"],
        ["NFR-PERF-04", "Endpoint webhook mengakui request dengan cepat setelah validasi/simpan; pekerjaan berat melalui queue.", "Integration/load test", "As-built/verify"],
        ["NFR-PERF-05", "Import/export besar dijalankan asinkron dan memberikan status kepada pengguna.", "Batch test", "As-built/verify"],
        ["NFR-PERF-06", "Cache/session/queue production menggunakan Redis dan tidak menghilangkan scope data.", "Config dan functional test", "Target produksi"],
    ], [1.15, 3.45, 1.15, 0.75])

    b.add_heading(doc, "3.2.3 Keandalan, Integritas, dan Pemulihan", 3)
    table(doc, "Tabel 3-13. Kebutuhan keandalan", ["ID", "Kebutuhan/ukuran", "Verifikasi", "Status"], [
        ["NFR-REL-01", "MySQL, Redis, app, web, dan WAHA mempunyai health check serta restart policy.", "Compose/health test", "As-built"],
        ["NFR-REL-02", "Queue memakai retry/timeout; job webhook dan sinkronisasi harus idempotent terhadap duplikasi event.", "Failure/replay test", "As-built/verify"],
        ["NFR-REL-03", "Volume database, storage, Redis, dan session WAHA persisten sesuai lingkungan.", "Restart/recreate test", "As-built"],
        ["NFR-REL-04", "Migration production eksplisit, didahului backup, dan tidak menjalankan seeder demo.", "Release checklist", "Target produksi"],
        ["NFR-REL-05", "Backup mencakup database, storage, session WAHA, manifest, checksum, enkripsi, dan salinan off-host.", "Backup audit", "Target produksi"],
        ["NFR-REL-06", "RPO dan RTO ditetapkan pemilik bisnis; restore diuji dan direkonsiliasi sebelum go-live.", "Restore drill", "Target keputusan"],
        ["NFR-REL-07", "Relasi, unique constraint, foreign key, validasi enum, dan transaksi menjaga integritas data.", "Migration/model/test", "As-built/verify"],
    ], [1.15, 3.45, 1.15, 0.75])

    b.add_heading(doc, "3.2.4 Usability, Responsivitas, dan Aksesibilitas", 3)
    table(doc, "Tabel 3-14. Kebutuhan pengalaman pengguna", ["ID", "Kebutuhan/ukuran", "Verifikasi", "Status"], [
        ["NFR-UX-01", "Antarmuka menggunakan Bahasa Indonesia dan light theme modern sesuai identitas SIVMI.", "Visual/content review", "As-built"],
        ["NFR-UX-02", "Layout responsif pada mobile, tablet, laptop, dan desktop; sidebar menjadi drawer pada mobile.", "Cross-device UAT", "As-built/verify"],
        ["NFR-UX-03", "Sidebar desktop collapsible dan status collapse disimpan di localStorage.", "Browser test", "As-built"],
        ["NFR-UX-04", "Loading memakai skeleton; kondisi tanpa data dan gagal menampilkan empty/error state serta langkah pemulihan.", "State-based UI test", "As-built/verify"],
        ["NFR-UX-05", "Form memberikan label, required marker, helper, validasi, dan pesan error yang jelas.", "Form/UAT test", "As-built/verify"],
        ["NFR-UX-06", "Kontras, fokus keyboard, heading, alt text, dan header tabel memenuhi baseline aksesibilitas yang diaudit.", "A11y audit/manual", "Target UAT"],
        ["NFR-UX-07", "Layout chat membedakan pesan masuk/keluar, waktu, status ACK, composer, dan informasi kontak tanpa overlap.", "Responsive chat UAT", "As-built/verify"],
    ], [1.15, 3.45, 1.15, 0.75])

    b.add_heading(doc, "3.2.5 Maintainability, Portabilitas, dan Skalabilitas", 3)
    table(doc, "Tabel 3-15. Kebutuhan teknis", ["ID", "Kebutuhan/ukuran", "Verifikasi", "Status"], [
        ["NFR-MNT-01", "Kode mengikuti Laravel 12, Filament 5, Livewire 4, Eloquent, enum, service, policy/scope, dan migration yang terstruktur.", "Code review", "As-built"],
        ["NFR-MNT-02", "Konfigurasi dipisahkan melalui environment; secret tidak di-hardcode atau masuk Git.", "Repository/secret scan", "As-built/verify"],
        ["NFR-MNT-03", "Full Docker dapat menjalankan aplikasi tanpa Laragon dan menjadi baseline yang dapat diulang.", "Fresh install test", "As-built"],
        ["NFR-MNT-04", "Image production membangun dependency/aset tanpa source bind mount dan dapat diberi tag immutable.", "Image/build audit", "As-built"],
        ["NFR-MNT-05", "Queue dan scheduler berjalan sebagai service terpisah dan dapat diskalakan sesuai profil beban.", "Compose/scale test", "As-built/target"],
        ["NFR-MNT-06", "Target VPS mendukung TLS, firewall, backup off-host, monitoring, dan pemisahan database/Redis terkelola bila diperlukan.", "Architecture review", "Target"],
        ["NFR-MNT-07", "Perubahan model/form wajib memperbarui test, seeder/example, dokumentasi, dan diagram terkait.", "Definition of done", "Target proses"],
    ], [1.15, 3.45, 1.15, 0.75])

    b.add_heading(doc, "3.2.6 Observability, Audit, dan Pengujian", 3)
    table(doc, "Tabel 3-16. Kebutuhan verifikasi operasional", ["ID", "Kebutuhan/ukuran", "Verifikasi", "Status"], [
        ["NFR-OBS-01", "Endpoint /up dan WAHA /health tersedia tanpa membocorkan data sensitif.", "Health request", "As-built"],
        ["NFR-OBS-02", "Log terpisah/tertelusur untuk app, web, queue, scheduler, database, Redis, dan WAHA.", "Log review", "As-built/target central"],
        ["NFR-OBS-03", "Monitoring mencakup error, latency, failed job, restart, disk, CPU, RAM, database, Redis, dan session WA.", "Dashboard/alert test", "Target produksi"],
        ["NFR-OBS-04", "Audit aktivitas dapat menjawab siapa, melakukan apa, kapan, pada data/cabang mana tanpa secret.", "Audit sampling", "As-built/verify"],
        ["NFR-TST-01", "Test otomatis memakai environment terisolasi dan tidak menulis database operasional.", "Test config/review", "As-built"],
        ["NFR-TST-02", "Quality gate mencakup test, build, migration review, smoke, RBAC, WhatsApp, dan UAT visual lintas perangkat.", "Release checklist", "As-built/target"],
        ["NFR-TST-03", "Baseline tercatat adalah 72 test/448 assertion dan build Vite lulus; regresi tidak boleh menurunkan hasil tanpa keputusan.", "CI/local test", "As-built"],
    ], [1.15, 3.45, 1.15, 0.75])

    b.add_heading(doc, "3.3 Traceability dan Kriteria Penerimaan", 2)
    b.add_body(doc, "Traceability memastikan kebutuhan tidak berhenti sebagai narasi. ID kebutuhan harus muncul pada desain atau backlog, dihubungkan dengan komponen implementasi, memiliki bukti verifikasi, dan diterima oleh pemilik proses. Perubahan kebutuhan harus memperbarui seluruh hubungan tersebut.")
    add_figure(doc, "traceability.png", "Gambar 3-7. Alur traceability kebutuhan hingga UAT")
    table(doc, "Tabel 3-17. Matriks traceability ringkas", ["Kelompok", "Implementasi utama", "Bukti", "Penerima"], [
        ["AUTH/RBAC", "Panel auth, UserRole, policy/scope, resource/page", "Feature test lintas role/cabang", "Superadmin/Direksi"],
        ["Dashboard/Laporan", "Dashboard page, query agregasi, filter global", "DashboardTest/ReportTest dan visual UAT", "Direksi/Admin Cabang"],
        ["Lead/CSV/Sekolah", "Lead/Sekolah resource, importer/exporter, enum", "Lead/School CSV dan management tests", "Admin Cabang/Staff"],
        ["Transaksi/Target/Program", "Pendaftaran, Pembayaran, Program, Target resource", "RegistrationPayment/ProgramTarget tests", "Admin Cabang/Direksi"],
        ["WAHA/Chat", "Account, conversation, webhook, job, service GOWS", "WahaIntegrationTest dan pesan end-to-end", "Admin WA/Staff"],
        ["Audit/Notifikasi/Setting", "Activity, database notification, CrmSetting", "AuditNotification/Settings tests", "Superadmin"],
        ["NFR Platform", "Docker Compose, Nginx, Redis, worker, scheduler", "Fresh install, health, build, backup/restore drill", "Administrator"],
    ], [1.35, 2.55, 1.65, 0.95])
    table(doc, "Tabel 3-18. Kriteria penerimaan minimum", ["Area", "Kriteria lulus"], [
        ["Fungsi", "Seluruh kebutuhan Must yang masuk scope rilis memiliki hasil test/UAT lulus atau waiver tertulis"],
        ["RBAC", "Superadmin, Admin Cabang, Staff, dan Direksi diuji pada Bandung, JakSel, dan JakPus"],
        ["Data", "Dedupe nomor, enum, relasi, import/export, transaksi, target, serta KPI konsisten"],
        ["WhatsApp", "Session Working, webhook HMAC, queue, inbox/chat, ACK, NewLead/Matched/Conflict terbukti"],
        ["Visual", "Mobile, tablet, laptop, desktop; skeleton, empty, error, form, sidebar, dan chat disetujui"],
        ["Operasional", "Health, log, queue, scheduler, backup/restore, monitoring, dan runbook tersedia"],
        ["Keamanan", "HTTPS, debug false, secret unik, scope server-side, log aman, dan permission disetujui"],
        ["Rilis", "Automated test/build lulus, migration ditinjau, smoke test dan observasi pascarilis selesai"],
    ], [1.4, 5.1])

    b.add_heading(doc, "3.4 Asumsi, Batasan, dan Gap Keputusan", 2)
    table(doc, "Tabel 3-19. Asumsi dan batasan", ["Jenis", "Pernyataan", "Tindak lanjut"], [
        ["Asumsi", "Setiap admin/staff memiliki nomor WhatsApp operasional dan satu session aktif.", "Validasi daftar pemilik nomor"],
        ["Asumsi", "Cabang aktif baseline Bandung, JakSel, JakPus.", "Approval perubahan master cabang"],
        ["Batasan", "WAHA bergantung pada kebijakan dan kestabilan WhatsApp serta engine GOWS.", "Monitoring, retry, runbook penautan ulang"],
        ["Batasan", "Shared hosting tidak menjadi target utama karena worker, scheduler, Redis, dan WAHA memerlukan proses persisten.", "Gunakan VPS/full Docker"],
        ["Gap", "Definisi operasional Grade D dan E belum memiliki deskripsi bisnis.", "Direksi menetapkan kriteria tertulis"],
        ["Gap", "RPO, RTO, target availability, concurrency, dan volume maksimum belum disahkan.", "Capacity test dan keputusan SLA"],
        ["Gap", "Kebijakan retensi chat, lead, audit, export, dan backup perlu keputusan organisasi.", "Privacy/retention policy"],
        ["Gap", "API publik pihak ketiga belum menjadi baseline.", "Kontrak, auth, rate limit, versioning bila diperlukan"],
    ], [1.0, 3.75, 1.75])
    callout(doc, "Keputusan yang diperlukan", "Sebelum produksi, pemilik sistem perlu mengesahkan SLA/availability, RPO/RTO, kapasitas, retensi data, definisi Grade D-E, prosedur persetujuan akses global, dan batas penggunaan WhatsApp.", b.LIGHT_AMBER)

    b.add_heading(doc, "LAMPIRAN A - RINGKASAN PRIORITAS", 1)
    table(doc, "Tabel A-1. Prioritas implementasi", ["Prioritas", "Makna", "Perlakuan"], [
        ["Must", "Wajib untuk operasi, keamanan, integritas, atau kontrol scope", "Tidak boleh go-live tanpa lulus/waiver resmi"],
        ["Should", "Penting untuk efektivitas dan pengalaman", "Masuk rencana rilis terdekat dan diuji"],
        ["Could", "Peningkatan setelah baseline stabil", "Dikelola sebagai backlog terukur"],
    ], [1.2, 2.7, 2.6])

    b.add_heading(doc, "PENUTUP", 1)
    b.add_body(doc, "BAB III ini menjadi baseline kebutuhan SIVMI CRM per 25 Agustus 2026. Setiap perubahan status lead, sumber, program, sekolah, role, cabang, WhatsApp, dashboard, arsitektur Docker, atau kebijakan operasional harus diperbarui pada ID kebutuhan, traceability, test, diagram, dan dokumen bab terkait.")
    callout(doc, "Prinsip akhir", "Kebutuhan dinyatakan selesai hanya bila perilakunya tersedia, pembatasan aksesnya benar, datanya konsisten, kualitasnya terukur, dan bukti penerimaannya dapat ditelusuri.", b.LIGHT_GREEN)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(f"Created: {OUTPUT}")


if __name__ == "__main__":
    build_document()
