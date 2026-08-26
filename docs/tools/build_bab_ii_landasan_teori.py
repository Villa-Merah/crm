from __future__ import annotations

from pathlib import Path
import sys

PROJECT_ROOT = Path(__file__).resolve().parents[2]
DOC_TOOLS = PROJECT_ROOT / "storage" / "app" / "doc-tools"
sys.path.insert(0, str(DOC_TOOLS))

from PIL import Image, ImageDraw
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

import build_bab_iv_arsitektur_bisnis as b
import build_bab_vii_arsitektur_teknologi as v7


OUTPUT = PROJECT_ROOT / "docs" / "BAB II - Landasan Teori SIVMI CRM v1.0.docx"
ASSETS = PROJECT_ROOT / "docs" / "_bab_ii_assets"
COVER = ASSETS / "cover-bab-ii.png"

PRIMARY = "#1476B8"
PRIMARY_DARK = "#075985"
INK = "#0B2545"
MUTED = "#64748B"
BG = "#F6F8FC"
GREEN = "#10B981"
AMBER = "#F59E0B"
ROSE = "#F43F5E"
PURPLE = "#A855F7"


def font(name: str, size: int):
    return v7.font(name, size)


def canvas(title: str, subtitle: str, size=(1800, 1080)):
    return v7.canvas(title, subtitle, size)


def box(draw, xy, title, lines, fill="#F8FAFC", outline=PRIMARY, title_size=27, body_size=20):
    return v7.box(draw, xy, title, lines, fill, outline, title_size, body_size)


def arrow(draw, start, end, color="#64748B", width=5):
    return v7.arrow(draw, start, end, color, width)


def make_cover() -> None:
    ASSETS.mkdir(parents=True, exist_ok=True)
    image = Image.new("RGB", (1600, 2200), BG)
    draw = ImageDraw.Draw(image)
    for y in range(2200):
        ratio = y / 2200
        start, end = (20, 118, 184), (7, 89, 133)
        draw.line((0, y, 1600, y), fill=tuple(round(start[i] * (1 - ratio) + end[i] * ratio) for i in range(3)))
    draw.ellipse((1010, -180, 1770, 580), fill=(27, 139, 199))
    draw.ellipse((1140, 140, 1640, 640), fill=(8, 103, 151))
    draw.rounded_rectangle((120, 250, 1480, 1910), radius=44, fill="white")
    draw.text((190, 350), "SIVMI CRM", font=font("arialbd.ttf", 34), fill=PRIMARY)
    draw.text((190, 610), "BAB II", font=font("arialbd.ttf", 48), fill=PRIMARY_DARK)
    draw.multiline_text((190, 700), "LANDASAN\nTEORI", font=font("arialbd.ttf", 92), fill=INK, spacing=18)
    draw.text((190, 1045), "Konsep sistem informasi, arsitektur, pemodelan,", font=font("arial.ttf", 30), fill="#475569")
    draw.text((190, 1097), "teknologi web, pengujian, integrasi, dan AI", font=font("arial.ttf", 30), fill="#475569")
    draw.rounded_rectangle((190, 1245, 805, 1325), radius=40, fill="#EAF4FB")
    draw.text((230, 1267), "BASELINE APLIKASI 25 AGUSTUS 2026", font=font("arial.ttf", 28), fill=PRIMARY_DARK)
    draw.text((190, 1610), "Versi dokumen 1.0", font=font("arial.ttf", 28), fill=MUTED)
    draw.text((190, 1665), "Sasaran: pembaca akademik, analis, pengembang, dan penguji", font=font("arial.ttf", 25), fill=MUTED)
    draw.text((190, 1775), "SIVMI CRM - Catat, Kelola, Closing", font=font("arialbd.ttf", 34), fill=PRIMARY)
    image.save(COVER)


def make_information_system() -> None:
    image, draw = canvas("Model Sistem Informasi", "Nilai terbentuk ketika manusia, proses, data, dan teknologi bekerja sebagai satu sistem")
    box(draw, (650, 365, 1150, 700), "INFORMASI", "Data yang telah diolah\nrelevan | akurat | tepat waktu", "#EAF4FB", PRIMARY_DARK, 31, 21)
    items = [
        (60, 150, 430, 420, "MANUSIA", "aktor\nperan\nkeputusan", PRIMARY),
        (60, 650, 430, 920, "PROSES", "aturan\nalur kerja\nkontrol", PURPLE),
        (1370, 150, 1740, 420, "DATA", "input\npenyimpanan\nkualitas", GREEN),
        (1370, 650, 1740, 920, "TEKNOLOGI", "aplikasi\ndatabase\ninfrastruktur", AMBER),
    ]
    for x1, y1, x2, y2, title, lines, color in items:
        box(draw, (x1, y1, x2, y2), title, lines, "#F8FAFC", color, 27, 20)
    arrow(draw, (430, 285), (650, 455)); arrow(draw, (430, 785), (650, 610))
    arrow(draw, (1370, 285), (1150, 455)); arrow(draw, (1370, 785), (1150, 610))
    image.save(ASSETS / "model-sistem-informasi.png")


def make_web_architecture() -> None:
    image, draw = canvas("Arsitektur Aplikasi Web", "Request dan response melewati browser, web server, aplikasi, data, dan integrasi")
    nodes = [
        (60, 270, 340, 640, "CLIENT", "Browser\nHTML/CSS/JS\nresponsive UI", PRIMARY),
        (410, 270, 690, 640, "WEB", "HTTP/HTTPS\nNginx\nstatic assets", PRIMARY_DARK),
        (760, 270, 1040, 640, "APP", "PHP/Laravel\nFilament/Livewire\npolicy/service", PURPLE),
        (1110, 270, 1390, 640, "DATA", "MySQL\nRedis\nfile storage", GREEN),
        (1460, 270, 1740, 640, "INTEGRASI", "WAHA/GOWS\nwebhook/API\nemail", AMBER),
    ]
    for x1, y1, x2, y2, title, lines, color in nodes:
        box(draw, (x1, y1, x2, y2), title, lines, "#F8FAFC", color, 26, 19)
    for x in (340, 690, 1040, 1390): arrow(draw, (x, 455), (x + 70, 455))
    draw.text((900, 850), "As-built SIVMI: full Docker, Nginx, PHP-FPM, Laravel 12, MySQL 8.4, Redis 7.4, WAHA/GOWS", anchor="mm", font=font("arialbd.ttf", 22), fill=MUTED)
    image.save(ASSETS / "arsitektur-web.png")


def make_frontend_stack() -> None:
    image, draw = canvas("Lapisan Frontend Web", "Standar web menjadi fondasi; framework membantu produktivitas dan konsistensi")
    layers = [
        (120, 700, 1680, 930, "HTML", "Struktur semantik dokumen dan komponen antarmuka", PRIMARY),
        (220, 510, 1580, 700, "CSS + TAILWIND", "Presentasi, layout responsif, warna, spacing, dan state visual", PURPLE),
        (320, 320, 1480, 510, "JAVASCRIPT + LIVEWIRE", "Interaksi, event, validasi, pembaruan antarmuka, dan komunikasi server", GREEN),
        (500, 140, 1300, 320, "FILAMENT", "Komponen panel admin, form, tabel, action, widget, dan notifikasi", AMBER),
    ]
    for x1, y1, x2, y2, title, lines, color in layers:
        box(draw, (x1, y1, x2, y2), title, lines, "#F8FAFC", color, 27, 19)
    image.save(ASSETS / "lapisan-frontend.png")


def make_ea() -> None:
    image, draw = canvas("Lapisan Enterprise Architecture", "Arsitektur menyelaraskan tujuan bisnis, data, aplikasi, dan teknologi")
    layers = [
        (180, 120, 1620, 315, "ARSITEKTUR BISNIS", "aktor | kapabilitas | proses | aturan | KPI", PRIMARY),
        (180, 355, 1620, 550, "ARSITEKTUR DATA", "entitas | kualitas | relasi | keamanan | siklus hidup", PURPLE),
        (180, 590, 1620, 785, "ARSITEKTUR APLIKASI", "modul | layanan | integrasi | antarmuka | tanggung jawab", GREEN),
        (180, 825, 1620, 1020, "ARSITEKTUR TEKNOLOGI", "runtime | jaringan | container | database | operasi", AMBER),
    ]
    for x1, y1, x2, y2, title, lines, color in layers:
        box(draw, (x1, y1, x2, y2), title, lines, "#F8FAFC", color, 29, 20)
    image.save(ASSETS / "lapisan-ea.png")


def make_modeling_map() -> None:
    image, draw = canvas("Peta Teknik Pemodelan", "Setiap diagram menjawab pertanyaan yang berbeda dan saling melengkapi")
    box(draw, (650, 90, 1150, 300), "SISTEM SIVMI CRM", "Aktor | Proses | Data | Perilaku", "#EAF4FB", PRIMARY_DARK, 29, 21)
    items = [
        (50, 500, 370, 890, "USE CASE", "Siapa menggunakan\nfungsi apa?", PRIMARY),
        (410, 500, 730, 890, "DFD", "Data mengalir\ndari/ke mana?", PURPLE),
        (770, 500, 1090, 890, "ERD", "Data disimpan\ndan berelasi bagaimana?", GREEN),
        (1130, 500, 1450, 890, "FLOWCHART", "Urutan logika\ndan keputusan?", AMBER),
        (1490, 500, 1760, 890, "ACTIVITY", "Aktivitas, peran,\nparalel, state?", ROSE),
    ]
    for x1, y1, x2, y2, title, lines, color in items:
        box(draw, (x1, y1, x2, y2), title, lines, "#F8FAFC", color, 24, 19)
        arrow(draw, (900, 300), ((x1 + x2) // 2, 500), "#94A3B8", 4)
    image.save(ASSETS / "peta-pemodelan.png")


def make_laravel() -> None:
    image, draw = canvas("Arsitektur Laravel pada SIVMI", "Framework mengatur request, aturan, data, presentasi, dan pekerjaan asinkron")
    nodes = [
        (60, 260, 340, 660, "ROUTE & MIDDLEWARE", "web/api\nauth/session\nCSRF/binding", PRIMARY),
        (410, 260, 690, 660, "PRESENTASI", "Filament\nLivewire\nBlade/Tailwind", PURPLE),
        (760, 260, 1040, 660, "DOMAIN", "model/enum\npolicy/scope\nservice/job", AMBER),
        (1110, 260, 1390, 660, "DATA", "Eloquent\nmigration\nMySQL/Redis", GREEN),
        (1460, 260, 1740, 660, "OPERASI", "queue\nscheduler\nnotification/log", PRIMARY_DARK),
    ]
    for x1, y1, x2, y2, title, lines, color in nodes:
        box(draw, (x1, y1, x2, y2), title, lines, "#F8FAFC", color, 23, 18)
    for x in (340, 690, 1040, 1390): arrow(draw, (x, 460), (x + 70, 460))
    image.save(ASSETS / "arsitektur-laravel.png")


def make_integration() -> None:
    image, draw = canvas("API, Webhook, dan Event Asinkron", "API diminta client; webhook mendorong event dari provider ke aplikasi")
    box(draw, (70, 210, 500, 760), "API (PULL/COMMAND)", "Client mengirim request\nAutentikasi + validasi\nServer mengembalikan response\nContoh: lifecycle session WAHA", "#F8FAFC", PRIMARY, 29, 20)
    box(draw, (685, 210, 1115, 760), "SIVMI CRM", "Route/API controller\nHMAC/API key\nEvent storage\nQueue + idempotency", "#EAF4FB", PRIMARY_DARK, 29, 20)
    box(draw, (1300, 210, 1730, 760), "WEBHOOK (PUSH)", "Provider mengirim event\nRespons penerimaan cepat\nProses berat asinkron\nContoh: pesan masuk WAHA", "#F8FAFC", GREEN, 29, 20)
    arrow(draw, (500, 400), (685, 400)); arrow(draw, (685, 580), (500, 580))
    arrow(draw, (1300, 400), (1115, 400)); arrow(draw, (1115, 580), (1300, 580))
    image.save(ASSETS / "api-webhook.png")


def make_ai_cycle() -> None:
    image, draw = canvas("AI dalam Siklus Pengembangan", "AI mempercepat pekerjaan, tetapi keputusan dan akuntabilitas tetap pada manusia")
    nodes = [
        (55, 280, 330, 650, "ANALISIS", "ringkas kebutuhan\nidentifikasi gap\nstruktur backlog", PRIMARY),
        (400, 280, 675, 650, "DESAIN", "alternatif arsitektur\ndiagram awal\nreview UX", PURPLE),
        (745, 280, 1020, 650, "IMPLEMENTASI", "scaffold\nrefactoring\ndokumentasi", AMBER),
        (1090, 280, 1365, 650, "PENGUJIAN", "test case\ndebugging\nanalisis log", GREEN),
        (1435, 280, 1745, 650, "GOVERNANCE", "review manusia\nprivasi/IP\ntraceability\napproval", ROSE),
    ]
    for x1, y1, x2, y2, title, lines, color in nodes:
        box(draw, (x1, y1, x2, y2), title, lines, "#F8FAFC", color, 24, 18)
    for x in (330, 675, 1020, 1365): arrow(draw, (x, 465), (x + 70, 465))
    draw.text((900, 850), "Aturan: jangan mengirim secret/data pribadi; verifikasi fakta, kode, keamanan, lisensi, dan hasil test.", anchor="mm", font=font("arialbd.ttf", 23), fill=MUTED)
    image.save(ASSETS / "siklus-ai.png")


def add_figure(doc: Document, name: str, caption: str, width=6.15) -> None:
    b.add_figure(doc, ASSETS / name, caption, width)


def table(doc: Document, caption: str, headers: list[str], rows: list[list[str]], widths: list[float]) -> None:
    b.add_table(doc, headers, rows, widths, caption)


def callout(doc: Document, title: str, body: str, fill: str = b.LIGHT_BLUE) -> None:
    b.add_callout(doc, title, body, fill)
    doc.tables[-1].rows[0]._tr.get_or_add_trPr().append(b.OxmlElement("w:cantSplit"))


def bullets(doc: Document, items: list[str]) -> None:
    b.add_bullets(doc, items)


def page(doc: Document) -> None:
    doc.add_page_break()


def build_document() -> None:
    for fn in (make_cover, make_information_system, make_web_architecture, make_frontend_stack, make_ea, make_modeling_map, make_laravel, make_integration, make_ai_cycle):
        fn()

    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.5); sec.page_height = Inches(11)
    sec.top_margin = Inches(0.85); sec.bottom_margin = Inches(0.85)
    sec.left_margin = Inches(1); sec.right_margin = Inches(1)
    sec.header_distance = Inches(0.4); sec.footer_distance = Inches(0.4)
    b.configure_styles(doc)
    sec.header.paragraphs[0].text = "SIVMI CRM  |  BAB II: LANDASAN TEORI"
    for run in sec.header.paragraphs[0].runs:
        b.set_run_font(run, 8.5, True, b.PRIMARY)
    b.add_page_number(sec.footer.paragraphs[0])

    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(0)
    pic = p.add_run().add_picture(str(COVER), width=Inches(6.1))
    pic._inline.docPr.set("descr", "Sampul BAB II Landasan Teori SIVMI CRM")
    page(doc)

    b.add_heading(doc, "KENDALI DOKUMEN", 1)
    table(doc, "Tabel 0-1. Kendali dokumen", ["Atribut", "Nilai"], [
        ["Judul", "BAB II: Landasan Teori SIVMI CRM"],
        ["Versi", "1.0"],
        ["Tanggal baseline", "25 Agustus 2026"],
        ["Tujuan", "Memberikan dasar konseptual dan teknis untuk analisis, perancangan, implementasi, pengujian, serta operasi SIVMI CRM"],
        ["Ruang lingkup", "Sistem informasi, web, PHP, MySQL, XAMPP, frontend, enterprise architecture, pemodelan, Laravel, UAT, API/webhook, dan AI"],
        ["Teknologi as-built", "PHP 8.4 container, Laravel 12, Filament 5, Livewire 4, Tailwind CSS 4, Vite 8, Nginx, MySQL 8.4, Redis 7.4, Docker, dan WAHA/GOWS"],
        ["Catatan", "XAMPP dan Bootstrap dibahas sebagai landasan/alternatif; keduanya bukan runtime dan UI framework utama pada baseline aktual"],
    ], [1.55, 4.95])
    callout(doc, "Tanpa daftar isi", "Untuk konsistensi dengan BAB III yang baru dibuat, BAB II ini tidak memuat halaman daftar isi dan langsung memasuki ringkasan landasan teori.")

    b.add_heading(doc, "RINGKASAN LANDASAN TEORI", 1)
    b.add_body(doc, "SIVMI CRM adalah aplikasi sistem informasi berbasis web yang menyatukan manusia, proses, data, aplikasi, dan teknologi. Bab ini menjelaskan teori yang menjadi dasar pengelolaan lead, follow-up, WhatsApp, pendaftaran, pembayaran, target, laporan, RBAC, audit, serta deployment berbasis container.")
    add_figure(doc, "model-sistem-informasi.png", "Gambar 2-1. Komponen pembentuk sistem informasi")
    table(doc, "Tabel 0-2. Pemetaan teori ke proyek", ["Kelompok teori", "Penerapan pada SIVMI CRM"], [
        ["Sistem dan web", "Aplikasi multiaktor, request-response, UI responsif, dan penyimpanan terpusat"],
        ["Bahasa dan data", "PHP/Laravel, MySQL, Redis, HTML, Tailwind, JavaScript, Livewire"],
        ["Arsitektur", "Bisnis, data, aplikasi, teknologi, Docker, queue, scheduler, WAHA"],
        ["Pemodelan", "Use case, DFD, ERD, flowchart, activity diagram, dan draw.io"],
        ["Verifikasi", "Test otomatis, build, UAT visual, smoke test, audit, dan monitoring"],
        ["Integrasi dan AI", "Webhook/API WAHA serta AI sebagai alat bantu pengembangan yang diawasi"],
    ], [2.0, 4.5])

    b.add_heading(doc, "BAB II: LANDASAN TEORI", 1)
    b.add_heading(doc, "2.1 Sistem Informasi", 2)
    b.add_body(doc, "Sistem informasi adalah susunan komponen yang mengumpulkan, memproses, menyimpan, dan mendistribusikan informasi untuk mendukung operasi, pengendalian, koordinasi, dan pengambilan keputusan. Sistem tidak hanya terdiri atas perangkat lunak; manusia, prosedur, data, kontrol, dan infrastruktur mempunyai peran yang sama penting.")
    b.add_body(doc, "Data merupakan fakta mentah, sedangkan informasi adalah data yang telah diberi konteks dan makna. Informasi yang baik perlu relevan terhadap keputusan, akurat, lengkap secara memadai, tepat waktu, konsisten, dan dapat ditelusuri. Dalam CRM, mutu informasi sangat dipengaruhi oleh ketepatan status lead, cabang, PIC, waktu follow-up, program, dan transaksi.")
    table(doc, "Tabel 2-1. Komponen sistem informasi", ["Komponen", "Fungsi", "Contoh SIVMI"], [
        ["Manusia", "Menginput, memproses, mengawasi, memutuskan", "Staff, Admin Cabang, Direksi"],
        ["Proses", "Menetapkan urutan dan aturan kerja", "Lead, follow-up, closing"],
        ["Data", "Mewakili objek dan kejadian", "Lead, sekolah, chat, pembayaran"],
        ["Aplikasi", "Mengotomasi validasi dan layanan", "Laravel/Filament"],
        ["Teknologi", "Menjalankan dan menghubungkan layanan", "Docker, MySQL, Redis, WAHA"],
        ["Kontrol", "Menjaga akses, kualitas, dan keterlacakan", "RBAC, scope cabang, audit"],
    ], [1.25, 2.7, 2.55])
    callout(doc, "Relevansi", "SIVMI CRM mengubah data operasional menjadi informasi KPI, tren, target, ranking, sekolah/wilayah, dan tindakan tertunda. Karena itu, kualitas dashboard bergantung langsung pada kualitas pencatatan transaksi sumber.", b.LIGHT_GREEN)

    b.add_heading(doc, "2.2 Aplikasi Berbasis Web", 2)
    add_figure(doc, "arsitektur-web.png", "Gambar 2-2. Arsitektur konseptual aplikasi web SIVMI CRM")
    b.add_body(doc, "Aplikasi web menggunakan browser sebagai client dan protokol HTTP/HTTPS untuk berkomunikasi dengan server. Browser mengirim request, server melakukan autentikasi, validasi, aturan bisnis, akses data, dan integrasi, kemudian mengembalikan response berupa HTML, JSON, file, atau status operasi.")
    bullets(doc, [
        "Thin client: pengguna cukup memakai browser tanpa instalasi aplikasi khusus.",
        "Centralized update: perubahan aplikasi dilakukan pada server sehingga versi lebih mudah dikendalikan.",
        "Responsive access: antarmuka dapat menyesuaikan mobile, tablet, laptop, dan desktop.",
        "Security boundary: autentikasi, otorisasi, validasi, dan scope data wajib ditegakkan di server.",
        "State management: session, cookie, cache, dan token digunakan secara terbatas dan aman.",
    ])
    callout(doc, "Penerapan", "SIVMI memakai Nginx sebagai web server, PHP-FPM/Laravel sebagai application server, MySQL dan Redis sebagai data services, serta WAHA/GOWS sebagai integration service. Seluruhnya dapat dijalankan melalui Docker Compose.")

    b.add_heading(doc, "2.3 PHP (Hypertext Preprocessor)", 2)
    b.add_body(doc, "PHP adalah bahasa pemrograman server-side yang dirancang untuk pengembangan web dan juga dapat digunakan untuk command-line, worker, scheduler, serta integrasi. Kode PHP dieksekusi pada server; hasilnya dikirim kepada client sehingga source code dan credential tidak semestinya terekspos ke browser.")
    table(doc, "Tabel 2-2. Karakteristik PHP yang relevan", ["Karakteristik", "Manfaat", "Penerapan"], [
        ["Dynamically typed", "Cepat dikembangkan", "Type declaration dan enum tetap digunakan"],
        ["Object-oriented", "Modular dan dapat diuji", "Model, service, job, policy"],
        ["Composer ecosystem", "Dependency terkelola", "Laravel dan Filament"],
        ["CLI", "Otomasi operasional", "Artisan migration, queue, scheduler, test"],
        ["Extension database/cache", "Integrasi efisien", "PDO MySQL dan phpredis"],
    ], [1.45, 2.3, 2.75])
    b.add_body(doc, "Manifest proyek mensyaratkan PHP ^8.3, sedangkan image runtime yang didokumentasikan memakai PHP 8.4. Praktik yang diperlukan meliputi dependency locking, input validation, exception handling, secret management, automated testing, dan pembaruan keamanan secara berkala.")

    b.add_heading(doc, "2.4 MySQL", 2)
    b.add_body(doc, "MySQL adalah relational database management system (RDBMS) yang menyimpan data dalam tabel, baris, dan kolom. Hubungan antarentitas dinyatakan melalui primary key dan foreign key, sedangkan constraint, index, transaksi, dan isolation membantu menjaga integritas serta kinerja.")
    table(doc, "Tabel 2-3. Konsep basis data relasional", ["Konsep", "Definisi", "Contoh SIVMI"], [
        ["Primary key", "Identitas unik record", "lead_id, cabang_id"],
        ["Foreign key", "Referensi ke entitas lain", "leads.cabang_id"],
        ["Unique constraint", "Mencegah nilai ganda tertentu", "Email pengguna"],
        ["Index", "Mempercepat pencarian/filter/join", "Status, cabang, tanggal"],
        ["Transaction", "Sekumpulan operasi atomik", "Create data saling terkait"],
        ["Migration", "Versi perubahan skema", "File database/migrations"],
    ], [1.4, 2.45, 2.65])
    callout(doc, "As-built", "SIVMI menggunakan MySQL 8.4 pada full Docker. Eloquent ORM mempermudah query dan relasi, tetapi desain index, scope cabang, agregasi dashboard, eager loading, backup, serta restore tetap memerlukan pengawasan database.", b.LIGHT_GREEN)

    b.add_heading(doc, "2.5 XAMPP", 2)
    b.add_body(doc, "XAMPP adalah paket development yang menggabungkan Apache, MariaDB/MySQL, PHP, dan Perl untuk memudahkan penyediaan server lokal. Paket semacam ini berguna bagi pembelajaran atau proyek sederhana karena komponen dapat dijalankan melalui satu control panel.")
    table(doc, "Tabel 2-4. XAMPP/Laragon dibanding full Docker", ["Aspek", "Paket lokal", "Full Docker SIVMI"], [
        ["Isolasi", "Bergantung instalasi host", "Service dan dependency terisolasi"],
        ["Reproduksibilitas", "Dapat berbeda antar komputer", "Compose dan image berversi"],
        ["Portabilitas", "Terikat konfigurasi Windows", "Dapat diarahkan ke Linux/VPS"],
        ["Worker/service", "Sering dijalankan manual", "Queue/scheduler sebagai service"],
        ["WAHA/Redis", "Perlu komponen tambahan", "Masuk satu network Compose"],
        ["Target proyek", "Legacy/kompatibilitas", "Baseline utama"],
    ], [1.35, 2.45, 2.7])
    callout(doc, "Klarifikasi teknologi", "Walaupun XAMPP tercantum pada landasan teori awal, SIVMI CRM tidak menjadikannya runtime utama. Laragon/hybrid juga dipertahankan hanya sebagai referensi legacy; jalur pengembangan yang direkomendasikan adalah full Docker.", b.LIGHT_AMBER)

    b.add_heading(doc, "2.6 HTML, CSS, JavaScript, dan Bootstrap", 2)
    add_figure(doc, "lapisan-frontend.png", "Gambar 2-3. Lapisan teknologi frontend pada implementasi SIVMI")
    b.add_body(doc, "HTML membentuk struktur semantik halaman; CSS mengatur presentasi, layout, warna, spacing, dan responsivitas; JavaScript menangani interaksi dan perubahan state pada browser. Ketiganya merupakan teknologi dasar aplikasi web modern.")
    b.add_body(doc, "Bootstrap adalah framework CSS/JavaScript berbasis komponen dan grid yang mempercepat pembangunan antarmuka. Namun, framework bukan bagian wajib dari standar web. SIVMI memilih Tailwind CSS 4, yaitu pendekatan utility-first, bersama tema Filament dan Livewire. Karena itu, Bootstrap dipahami sebagai alternatif teori, bukan dependency aktual.")
    table(doc, "Tabel 2-5. Peran teknologi frontend", ["Teknologi", "Peran", "As-built SIVMI"], [
        ["HTML", "Struktur, form, tabel, aksesibilitas", "Blade/Filament menghasilkan HTML"],
        ["CSS", "Visual, layout, responsive states", "Tailwind CSS dan theme.css"],
        ["JavaScript", "Event dan interaksi client", "Livewire/Filament/Vite"],
        ["Bootstrap", "Framework komponen siap pakai", "Tidak digunakan sebagai dependency utama"],
        ["Vite", "Dev server dan production bundling", "Vite 8"],
    ], [1.25, 2.5, 2.75])

    b.add_heading(doc, "2.7 Enterprise Architecture (EA)", 2)
    add_figure(doc, "lapisan-ea.png", "Gambar 2-4. Lapisan enterprise architecture SIVMI CRM")
    b.add_body(doc, "Enterprise Architecture adalah disiplin untuk menyelaraskan strategi, proses bisnis, informasi, aplikasi, dan teknologi. EA membantu organisasi memahami baseline, target, gap, prinsip, dependensi, dan roadmap perubahan sehingga sistem tidak dikembangkan sebagai kumpulan fitur yang terpisah.")
    table(doc, "Tabel 2-6. Domain EA", ["Domain", "Pertanyaan", "Artefak SIVMI"], [
        ["Bisnis", "Siapa, untuk tujuan apa, proses dan aturan apa?", "Aktor, proses lead, KPI, SOP"],
        ["Data", "Data apa, kualitas dan relasinya bagaimana?", "ERD, kamus data, kardinalitas"],
        ["Aplikasi", "Modul dan layanan apa yang mendukung proses?", "Peta modul, request flow, API"],
        ["Teknologi", "Platform apa dan bagaimana dioperasikan?", "Docker, Nginx, MySQL, Redis, WAHA"],
    ], [1.25, 2.4, 2.85])
    callout(doc, "Prinsip EA", "Dokumen BAB IV sampai BAB VII memisahkan arsitektur bisnis, data, aplikasi, dan teknologi. Pemisahan tersebut memperjelas tanggung jawab, tetapi seluruh lapisan tetap harus konsisten terhadap kebutuhan BAB III.")

    b.add_heading(doc, "2.8 Use Case Diagram", 2)
    add_figure(doc, "peta-pemodelan.png", "Gambar 2-5. Peta hubungan teknik pemodelan sistem")
    b.add_body(doc, "Use case diagram merupakan bagian dari UML yang menggambarkan aktor eksternal, tujuan/interaksi aktor dengan sistem, serta batas sistem. Diagram ini menekankan apa yang disediakan sistem dari sudut pandang pengguna, bukan urutan teknis atau struktur database.")
    bullets(doc, [
        "Actor mewakili peran, bukan selalu satu orang tertentu.",
        "Use case menggambarkan tujuan bernilai bagi aktor, misalnya Kelola Lead atau Lihat Laporan.",
        "System boundary membedakan fungsi internal dari pihak eksternal.",
        "Association menunjukkan aktor berpartisipasi pada use case.",
        "Include/extend/generalization digunakan seperlunya dan tidak menggantikan narasi skenario.",
    ])
    callout(doc, "Penerapan", "Aktor utama SIVMI adalah Superadmin, Admin Cabang, Staff, dan Direksi. WAHA dapat dipandang sebagai sistem eksternal. Detail precondition, alur utama, alternatif, dan postcondition sebaiknya dicatat pada spesifikasi use case, bukan hanya gambar.")

    b.add_heading(doc, "2.9 Data Flow Diagram (DFD)", 2)
    b.add_body(doc, "DFD memodelkan bagaimana data bergerak antara entitas eksternal, proses, data store, dan aliran data. Context diagram atau DFD Level 0 menampilkan sistem sebagai satu proses besar, sedangkan Level 1 menguraikannya menjadi proses yang lebih rinci.")
    table(doc, "Tabel 2-7. Elemen DFD", ["Elemen", "Makna", "Contoh SIVMI"], [
        ["Entitas eksternal", "Sumber/penerima data di luar batas proses", "Pengguna, calon siswa, WAHA"],
        ["Proses", "Transformasi input menjadi output", "Kelola lead, proses webhook"],
        ["Data store", "Penyimpanan data", "Leads, conversation, pembayaran"],
        ["Data flow", "Data yang berpindah", "Form lead, event pesan, laporan"],
    ], [1.45, 2.5, 2.55])
    b.add_body(doc, "DFD harus menjaga balancing: aliran masuk dan keluar pada proses tingkat atas tetap konsisten ketika diuraikan. DFD tidak menunjukkan detail urutan waktu, struktur tabel, class, atau teknologi implementasi.")

    b.add_heading(doc, "2.10 Entity Relationship Diagram (ERD)", 2)
    b.add_body(doc, "ERD menggambarkan entitas data, atribut, identifier, relasi, serta kardinalitas. ERD konseptual berfokus pada makna bisnis; ERD logis menambahkan atribut dan normalisasi; ERD fisik mencerminkan tabel, tipe data, primary key, foreign key, index, dan constraint.")
    table(doc, "Tabel 2-8. Kardinalitas", ["Notasi", "Makna", "Contoh"], [
        ["1:1", "Satu record berhubungan dengan paling banyak satu record", "Relasi khusus sesuai aturan"],
        ["1:N", "Satu parent memiliki banyak child", "Cabang ke leads"],
        ["M:N", "Banyak record berhubungan dengan banyak record", "Dipecah melalui associative entity"],
        ["0..1", "Relasi opsional tunggal", "Lead mungkin belum memiliki sekolah"],
        ["0..N", "Relasi opsional jamak", "Lead mungkin belum memiliki follow-up"],
        ["Polymorphic", "Satu relasi menunjuk beberapa tipe subject", "Audit activity subject"],
    ], [1.15, 3.0, 2.35])
    callout(doc, "Penerapan", "Entitas inti SIVMI meliputi cabang, users, sekolah, program_bimbel, leads, followup_riwayat, pendaftaran, pembayaran, target, WhatsApp account/conversation/message/event, settings, notification, dan activity log.", b.LIGHT_GREEN)

    b.add_heading(doc, "2.11 Flowchart", 2)
    b.add_body(doc, "Flowchart adalah diagram alur yang menunjukkan urutan proses, keputusan, input/output, dan terminator menggunakan simbol standar. Flowchart efektif untuk menjelaskan logika prosedural dan percabangan, misalnya validasi duplikasi nomor atau pemulihan session WAHA.")
    table(doc, "Tabel 2-9. Simbol flowchart umum", ["Simbol", "Fungsi", "Contoh"], [
        ["Terminator", "Mulai atau selesai", "Pesan webhook diterima"],
        ["Process", "Aktivitas/pengolahan", "Normalisasi nomor"],
        ["Decision", "Percabangan kondisi", "Nomor sudah ada?"],
        ["Input/Output", "Data masuk atau keluar", "Payload dan response"],
        ["Database", "Penyimpanan", "Simpan event/conversation"],
        ["Connector", "Penghubung alur", "Lanjutan antarbagian/halaman"],
    ], [1.35, 2.55, 2.6])
    b.add_body(doc, "Flowchart sebaiknya memakai satu arah baca yang dominan, label keputusan Ya/Tidak, simbol yang konsisten, serta tingkat detail yang sesuai. Diagram tidak perlu menampilkan setiap baris kode.")

    b.add_heading(doc, "2.12 Activity Diagram", 2)
    b.add_body(doc, "Activity diagram adalah diagram perilaku UML untuk menggambarkan aliran aktivitas, control flow, object flow, decision, merge, fork, join, initial node, final node, dan partition/swimlane. Dibanding flowchart, activity diagram lebih kuat untuk menunjukkan tanggung jawab antaraktor dan aktivitas paralel.")
    table(doc, "Tabel 2-10. Activity diagram dan penerapannya", ["Elemen", "Makna", "Contoh SIVMI"], [
        ["Activity", "Pekerjaan yang dilakukan", "Mencatat follow-up"],
        ["Decision/Merge", "Memilih dan menggabungkan cabang alur", "Lead duplikat atau baru"],
        ["Fork/Join", "Aktivitas paralel dan sinkronisasi", "Simpan event dan dispatch job"],
        ["Swimlane", "Penanggung jawab aktivitas", "Staff, sistem, WAHA"],
        ["Object flow", "Data yang berpindah antaraktivitas", "Lead, pesan, pembayaran"],
    ], [1.45, 2.55, 2.5])
    callout(doc, "Perbedaan praktis", "Use case menjelaskan tujuan aktor; activity diagram menjelaskan aliran aktivitas; sequence diagram (bila digunakan) menjelaskan urutan pesan antarobjek. Ketiganya tidak saling menggantikan.")

    b.add_heading(doc, "2.13 Laravel", 2)
    add_figure(doc, "arsitektur-laravel.png", "Gambar 2-6. Lapisan Laravel pada implementasi SIVMI CRM")
    b.add_body(doc, "Laravel adalah framework aplikasi web PHP yang menyediakan routing, middleware, dependency injection, Eloquent ORM, migration, validation, authentication, authorization, queue, scheduler, notification, cache, logging, dan testing support. Konvensi framework membantu struktur kode tetap konsisten.")
    table(doc, "Tabel 2-11. Ekosistem Laravel SIVMI", ["Komponen", "Peran"], [
        ["Laravel 12", "Fondasi aplikasi, HTTP, data, queue, scheduler, notification, testing"],
        ["Filament 5", "Panel admin, resource, form, table, action, widget, notification"],
        ["Livewire 4", "Komponen reaktif berbasis server dengan interaksi browser"],
        ["Eloquent", "Model, relasi, scope, cast, dan query"],
        ["Artisan", "Migration, seeder, queue, scheduler, cache, test, maintenance"],
        ["Vite/Tailwind", "Build aset dan styling responsif"],
    ], [1.65, 4.85])
    callout(doc, "Batas framework", "Framework mempercepat pengembangan tetapi tidak otomatis menjamin keamanan, performa, atau arsitektur yang baik. Policy/scope server-side, query efisien, transaction, validation, secret management, test, dan review tetap wajib.", b.LIGHT_AMBER)

    b.add_heading(doc, "2.14 User Acceptance Testing (UAT)", 2)
    b.add_body(doc, "UAT adalah proses validasi oleh pengguna atau pemilik bisnis untuk memastikan sistem memenuhi kebutuhan dan dapat digunakan dalam konteks operasional. UAT berbeda dari unit test dan integration test: fokusnya adalah kesesuaian dengan proses nyata, data, kewenangan, dan hasil bisnis.")
    table(doc, "Tabel 2-12. Struktur skenario UAT", ["Bagian", "Isi"], [
        ["ID dan tujuan", "Kebutuhan atau proses yang divalidasi"],
        ["Aktor/scope", "Role, cabang, dan kepemilikan data"],
        ["Prasyarat", "Akun, data master, session, dan environment"],
        ["Langkah", "Tindakan pengguna yang dapat diulang"],
        ["Hasil diharapkan", "Output, perubahan data, notifikasi, dan audit"],
        ["Hasil aktual", "Bukti screenshot/log aman dan status"],
        ["Keputusan", "Pass, fail, blocked, waiver, owner, dan tindak lanjut"],
    ], [1.55, 4.95])
    bullets(doc, [
        "Uji empat role: Superadmin, Admin Cabang, Staff, dan Direksi.",
        "Uji scope Bandung, JakSel, dan JakPus termasuk percobaan akses negatif.",
        "Uji mobile, tablet, laptop, desktop, skeleton, empty, error, dan layout chat.",
        "Uji end-to-end lead, CSV, sekolah inline, follow-up, WAHA, daftar, bayar, target, laporan, notifikasi, dan audit.",
        "Pisahkan data UAT dari production dan jangan memakai credential demo saat go-live.",
    ])

    b.add_heading(doc, "2.15 Webhook dan API", 2)
    add_figure(doc, "api-webhook.png", "Gambar 2-7. Perbedaan API request dan webhook event")
    b.add_body(doc, "Application Programming Interface (API) adalah kontrak komunikasi antarperangkat lunak. API mendefinisikan endpoint, method, request, response, autentikasi, error, dan versi. Webhook adalah mekanisme event-driven ketika provider mengirim notifikasi HTTP ke endpoint consumer setelah suatu kejadian.")
    table(doc, "Tabel 2-13. Kontrol integrasi", ["Kontrol", "Tujuan", "Penerapan SIVMI-WAHA"], [
        ["Authentication", "Membuktikan client", "X-Api-Key"],
        ["Integrity", "Mendeteksi payload berubah/palsu", "HMAC signature secret"],
        ["Validation", "Menolak struktur/nilai salah", "Event/payload validation"],
        ["Idempotency", "Mencegah efek ganda", "Event/message dedupe"],
        ["Async processing", "Respons cepat dan tahan beban", "Simpan event lalu queue"],
        ["Retry/timeout", "Pulih dari gangguan sementara", "Worker dan HTTP client policy"],
        ["Observability", "Menelusuri kegagalan", "Log aman, status event, failed job"],
    ], [1.45, 2.35, 2.7])
    callout(doc, "Penerapan", "SIVMI menerima POST /api/webhooks/waha, memverifikasi HMAC, menyimpan event, lalu memproses pesan melalui queue. Laravel memanggil API WAHA untuk session, QR, status, dan pengiriman pesan menggunakan API key.", b.LIGHT_GREEN)

    b.add_heading(doc, "2.16 Artificial Intelligence (AI) dalam Pengembangan Perangkat Lunak", 2)
    add_figure(doc, "siklus-ai.png", "Gambar 2-8. Peran AI dalam siklus pengembangan perangkat lunak")
    b.add_body(doc, "Artificial Intelligence dalam pengembangan perangkat lunak mencakup penggunaan model untuk membantu analisis kebutuhan, pencarian pola, pembuatan rancangan awal, penulisan kode, refactoring, test generation, debugging, dokumentasi, dan review. AI berperan sebagai alat bantu, bukan pengganti akuntabilitas analis, developer, tester, administrator, atau pemilik bisnis.")
    table(doc, "Tabel 2-14. Manfaat dan risiko AI", ["Area", "Manfaat", "Risiko", "Kontrol"], [
        ["Analisis", "Merangkum dan mengelompokkan", "Konteks salah/terlewat", "Review stakeholder"],
        ["Kode", "Scaffold dan alternatif solusi", "Bug/security/licensing", "Code review dan test"],
        ["Test", "Variasi kasus dan edge case", "False confidence", "Coverage dan manual test"],
        ["Debug", "Hipotesis akar masalah", "Diagnosis keliru", "Log/bukti/reproduksi"],
        ["Dokumen", "Struktur dan konsistensi", "Fakta/versi tidak akurat", "Audit source/as-built"],
        ["Data", "Klasifikasi/pencarian", "Privasi dan kebocoran", "Minimasi/redaksi/izin"],
    ], [1.15, 2.05, 1.55, 1.75])
    callout(doc, "Prinsip penggunaan AI", "Jangan memasukkan secret, credential, dump database, isi chat pribadi, atau data siswa yang dapat diidentifikasi ke layanan AI tanpa dasar hukum, persetujuan, dan kontrol organisasi. Semua output AI harus diverifikasi sebelum digunakan.", b.LIGHT_AMBER)

    b.add_heading(doc, "LAMPIRAN A - POSISI TEKNOLOGI AS-BUILT", 1)
    table(doc, "Tabel A-1. Teori dan implementasi aktual", ["Topik teori", "Pilihan aktual SIVMI", "Status"], [
        ["Local stack", "Full Docker Compose", "Utama"],
        ["XAMPP/Laragon", "Referensi legacy/kompatibilitas", "Bukan runtime utama"],
        ["Web server", "Nginx 1.27", "As-built"],
        ["Backend", "PHP 8.4 container, Laravel 12", "As-built"],
        ["Admin UI", "Filament 5 dan Livewire 4", "As-built"],
        ["CSS framework", "Tailwind CSS 4", "As-built"],
        ["Bootstrap", "Alternatif teori", "Tidak menjadi dependency utama"],
        ["Asset build", "Vite 8", "As-built"],
        ["Database/cache", "MySQL 8.4 dan Redis 7.4", "As-built"],
        ["WhatsApp", "WAHA CORE/GOWS", "As-built"],
    ], [1.5, 3.3, 1.7])

    b.add_heading(doc, "LAMPIRAN B - REFERENSI LANDASAN", 1)
    table(doc, "Tabel B-1. Referensi utama", ["No.", "Referensi"], [
        ["1", "Laudon, K. C. & Laudon, J. P. Management Information Systems. Pearson."],
        ["2", "Pressman, R. S. & Maxim, B. R. Software Engineering: A Practitioner's Approach. McGraw-Hill."],
        ["3", "PHP Documentation. PHP Manual, php.net/docs.php."],
        ["4", "Oracle. MySQL 8.4 Reference Manual, dev.mysql.com/doc/refman/8.4/en/."],
        ["5", "Apache Friends. XAMPP Documentation, apachefriends.org."],
        ["6", "WHATWG. HTML Living Standard; W3C CSS specifications; ECMA-262 ECMAScript specification."],
        ["7", "The Open Group. TOGAF Standard, 10th Edition."],
        ["8", "Object Management Group. UML 2.5.1 Specification."],
        ["9", "Chen, P. P. The Entity-Relationship Model—Toward a Unified View of Data. ACM TODS, 1976."],
        ["10", "ISO 5807. Information processing—Documentation symbols and conventions for data, program and system flowcharts."],
        ["11", "Laravel Documentation 12.x; Filament Documentation 5.x; Livewire Documentation 4.x."],
        ["12", "ISO/IEC 25010. Systems and software Quality Requirements and Evaluation—Product quality model."],
        ["13", "OpenAPI Initiative. OpenAPI Specification; Standard Webhooks specification."],
        ["14", "NIST. Artificial Intelligence Risk Management Framework (AI RMF 1.0)."],
    ], [0.55, 5.95])
    callout(doc, "Catatan referensi", "Dokumentasi versi framework dan library harus dicocokkan kembali dengan composer.lock/package-lock.json ketika dependency proyek berubah. Referensi teoritis ini tidak menggantikan dokumentasi as-built dan hasil verifikasi kode.")

    b.add_heading(doc, "PENUTUP", 1)
    b.add_body(doc, "Landasan teori pada BAB II menyediakan istilah, konsep, dan batas pemodelan yang digunakan pada analisa kebutuhan, arsitektur, implementasi, pengujian, deployment, serta SOP SIVMI CRM. Teori harus diterapkan secara proporsional dan selalu diverifikasi terhadap kondisi sistem aktual.")
    callout(doc, "Prinsip akhir", "Diagram menjelaskan sudut pandang; kode dan konfigurasi menjelaskan implementasi; test dan UAT menjelaskan bukti; operasi dan audit menjelaskan keberlanjutan. Keempatnya harus konsisten.", b.LIGHT_GREEN)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(f"Created: {OUTPUT}")


if __name__ == "__main__":
    build_document()
