from __future__ import annotations

from copy import deepcopy
from datetime import date
from io import BytesIO
from pathlib import Path
import sys

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "storage" / "app" / "doc-tools"))

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Cm, Pt, RGBColor


CHAPTERS = [
    ROOT / "docs" / "BAB I - Pendahuluan SIVMI CRM v1.0.docx",
    ROOT / "docs" / "BAB II - Landasan Teori SIVMI CRM v1.0.docx",
    ROOT / "docs" / "BAB III - Analisa Kebutuhan SIVMI CRM v1.0.docx",
    ROOT / "docs" / "BAB IV - Arsitektur Bisnis SIVMI CRM v1.0.docx",
    ROOT / "docs" / "BAB V - Arsitektur Data SIVMI CRM v1.1.docx",
    ROOT / "docs" / "BAB VI - Arsitektur Aplikasi SIVMI CRM v1.0.docx",
    ROOT / "docs" / "BAB VII - Arsitektur Teknologi SIVMI CRM v1.0.docx",
    ROOT / "docs" / "BAB VIII - Manajemen Debugging SIVMI CRM v1.0.docx",
    ROOT / "docs" / "BAB IX - Deployment, Migrasi, dan Pemeliharaan Sistem SIVMI CRM v1.0.docx",
    ROOT / "docs" / "BAB X - Panduan Penggunaan SIVMI CRM v1.0.docx",
    ROOT / "docs" / "BAB XI - Panduan Administrator SIVMI CRM v1.0.docx",
]

APPENDICES = [
    ("LAMPIRAN A: API DAN REST API", ROOT / "docs" / "API-DAN-REST-API-SIVMI-CRM-v1.1.docx", "strip"),
    ("LAMPIRAN B: LAPORAN TESTING DAN DEBUGGING", ROOT / "docs" / "LAPORAN-TESTING-DAN-DEBUGGING-SIVMI-CRM-FINAL-2026-08-26.docx", "all"),
    ("LAMPIRAN C: MATRIKS ERD, RELASI, DAN KARDINALITAS", ROOT / "docs" / "Matriks ERD Relasi dan Kardinalitas SIVMI CRM v1.0.docx", "all"),
]

OUTPUT = ROOT / "docs" / "DOKUMENTASI-FINAL-PROJECT-SIVMI-CRM-v2.0.docx"

PRIMARY = "1476B8"
PRIMARY_DARK = "075985"
INK = "0B2545"
MUTED = "64748B"
LIGHT = "F6F8FC"
GREEN = "10B981"
AMBER = "F59E0B"
PURPLE = "A855F7"
WHITE = "FFFFFF"


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def text_of(element) -> str:
    return "".join(node.text or "" for node in element.iter(qn("w:t"))).strip()


def has_page_break(element) -> bool:
    return bool(element.xpath(".//w:br[@w:type='page']"))


def content_without_cover_and_toc(document: Document):
    content = []
    cover_removed = False
    skipping_toc = False
    for element in document.element.body.iterchildren():
        if element.tag == qn("w:sectPr"):
            continue
        if not cover_removed:
            if has_page_break(element):
                cover_removed = True
            continue
        value = text_of(element).upper()
        if value == "DAFTAR ISI":
            skipping_toc = True
            continue
        if skipping_toc:
            if has_page_break(element):
                skipping_toc = False
            continue
        content.append(element)
    if not cover_removed:
        return [e for e in document.element.body.iterchildren() if e.tag != qn("w:sectPr")]
    return content


def all_content(document: Document):
    return [e for e in document.element.body.iterchildren() if e.tag != qn("w:sectPr")]


def merge_styles(destination: Document, source: Document) -> None:
    target = destination.styles.element
    known = {node.get(qn("w:styleId")) for node in target.findall(qn("w:style"))}
    for node in source.styles.element.findall(qn("w:style")):
        style_id = node.get(qn("w:styleId"))
        if style_id not in known:
            target.append(deepcopy(node))
            known.add(style_id)


def copy_relationships(element, source: Document, destination: Document) -> None:
    mapping: dict[str, str] = {}
    for node in element.iter():
        for attribute in (qn("r:embed"), qn("r:id"), qn("r:link")):
            old = node.get(attribute)
            if not old or old not in source.part.rels:
                continue
            if old not in mapping:
                rel = source.part.rels[old]
                if rel.reltype == RT.IMAGE:
                    new, _ = destination.part.get_or_add_image(BytesIO(rel.target_part.blob))
                elif rel.is_external:
                    new = destination.part.relate_to(rel.target_ref, rel.reltype, is_external=True)
                else:
                    # Hubungan internal non-gambar (misalnya hyperlink internal)
                    # tidak dibutuhkan untuk isi statis hasil gabungan.
                    node.attrib.pop(attribute, None)
                    continue
                mapping[old] = new
            node.set(attribute, mapping[old])


def page_break(document: Document) -> None:
    p = document.add_paragraph()
    p.add_run().add_break()
    p.runs[-1]._r[-1].set(qn("w:type"), "page")


def set_run(run, size=10.5, bold=False, color=INK, font="Aptos"):
    run.font.name = font
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), font)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), font)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)


def shade(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def cover(document: Document) -> None:
    p = document.add_paragraph()
    p.paragraph_format.space_before = Pt(88)
    p.paragraph_format.space_after = Pt(10)
    set_run(p.add_run("SIVMI CRM"), 17, True, PRIMARY)

    p = document.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    set_run(p.add_run("DOKUMENTASI FINAL PROJECT"), 30, True, PRIMARY_DARK)

    p = document.add_paragraph()
    p.paragraph_format.space_after = Pt(28)
    set_run(p.add_run("Arsitektur bisnis, data, aplikasi, teknologi, operasi, API, pengujian, dan panduan penggunaan"), 14, False, MUTED)

    table = document.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    values = [
        ("11", "bab sistem", PRIMARY),
        ("REST", "API v1 leads", GREEN),
        ("WAHA", "WhatsApp GOWS", PURPLE),
        ("83", "test lulus", AMBER),
    ]
    for cell, (number, label, fill) in zip(table.rows[0].cells, values):
        shade(cell, fill)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(10)
        set_run(p.add_run(number + "\n"), 18, True, WHITE)
        set_run(p.add_run(label), 8.5, True, WHITE)

    p = document.add_paragraph()
    p.paragraph_format.space_before = Pt(58)
    p.paragraph_format.space_after = Pt(4)
    set_run(p.add_run("Baseline final"), 9, True, MUTED)
    p = document.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    set_run(p.add_run("26 Agustus 2026 • Asia/Jakarta"), 11, True, INK)
    p = document.add_paragraph()
    set_run(p.add_run("Laravel 12 • Filament 5 • Livewire 4 • Tailwind CSS 4 • MySQL 8.4 • Docker • WAHA GOWS"), 9.2, False, MUTED)
    page_break(document)


def front_matter(document: Document) -> None:
    p = document.add_paragraph()
    set_run(p.add_run("Ringkasan Eksekutif"), 20, True, PRIMARY_DARK)
    p = document.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    set_run(p.add_run(
        "SIVMI CRM adalah sistem CRM bimbel berbasis web untuk mengelola lead calon siswa dari masuk, follow-up, pendaftaran, pembayaran, hingga KPI dan pelaporan. Sistem menerapkan RBAC superadmin, admin cabang, staff, dan direksi; scope data cabang; REST API leads; serta integrasi WhatsApp melalui WAHA dengan engine GOWS."
    ), 10.5, False, INK)

    table = document.add_table(rows=6, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    rows = [
        ("Arsitektur", "Modular monolith Laravel dengan web admin, queue, scheduler, webhook, dan REST API"),
        ("Data", "MySQL; Eloquent; audit; notifikasi; idempotensi API dan webhook"),
        ("Operasional", "Full Docker untuk lokal; target VPS; Nginx, PHP-FPM, Redis, MySQL, WAHA"),
        ("API publik", "GET/POST /api/v1/leads dengan bearer token, ability, rate limit, scope cabang, dan request ID"),
        ("Kualitas", "83 test dan 499 assertion lulus; build Vite serta validasi Composer berhasil"),
        ("Dokumen pendamping", "Satu file draw.io gabungan dan PowerPoint ringkasan stakeholder"),
    ]
    for idx, (label, value) in enumerate(rows):
        left, right = table.rows[idx].cells
        shade(left, PRIMARY_DARK if idx % 2 == 0 else PRIMARY)
        shade(right, "FFFFFF" if idx % 2 == 0 else LIGHT)
        set_run(left.paragraphs[0].add_run(label), 9, True, WHITE)
        set_run(right.paragraphs[0].add_run(value), 9, False, INK)

    p = document.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    set_run(p.add_run("Struktur Dokumen"), 18, True, PRIMARY_DARK)
    entries = [
        "BAB I — Pendahuluan", "BAB II — Landasan Teori", "BAB III — Analisa Kebutuhan",
        "BAB IV — Arsitektur Bisnis", "BAB V — Arsitektur Data", "BAB VI — Arsitektur Aplikasi",
        "BAB VII — Arsitektur Teknologi", "BAB VIII — Manajemen Debugging",
        "BAB IX — Deployment, Migrasi, dan Pemeliharaan Sistem", "BAB X — Panduan Penggunaan",
        "BAB XI — Panduan Administrator", "Lampiran A — API dan REST API",
        "Lampiran B — Laporan Testing dan Debugging", "Lampiran C — Matriks ERD Relasi dan Kardinalitas",
    ]
    for entry in entries:
        p = document.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(1.5)
        set_run(p.add_run(entry), 9.6, False, INK)
    page_break(document)


def configure(document: Document):
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.7)
    section.bottom_margin = Cm(1.6)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)
    section.header_distance = Cm(0.7)
    section.footer_distance = Cm(0.7)
    document.settings.odd_and_even_pages_header_footer = False
    section.different_first_page_header_footer = False

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run(hp.add_run("SIVMI CRM  |  DOKUMENTASI FINAL PROJECT"), 8, True, MUTED)
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run(fp.add_run("Baseline 26 Agustus 2026  |  "), 8, False, MUTED)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    fp._p.append(field)


def append_source(destination: Document, path: Path, mode="strip"):
    source = Document(path)
    merge_styles(destination, source)
    elements = content_without_cover_and_toc(source) if mode == "strip" else all_content(source)
    body = destination.element.body
    for element in elements:
        cloned = deepcopy(element)
        copy_relationships(cloned, source, destination)
        body.insert(-1, cloned)


def audit(document: Document):
    text = "\n".join(p.text for p in document.paragraphs)
    required = ["BAB I", "BAB XI", "API DAN REST API", "TESTING DAN DEBUGGING", "MATRIKS ERD"]
    missing = [item for item in required if item not in text.upper()]
    if missing:
        raise RuntimeError("Bagian final tidak ditemukan: " + ", ".join(missing))
    if len(document.inline_shapes) < 40:
        raise RuntimeError(f"Jumlah gambar terlalu sedikit: {len(document.inline_shapes)}")


def main():
    required = CHAPTERS + [item[1] for item in APPENDICES]
    absent = [str(path) for path in required if not path.exists()]
    if absent:
        raise FileNotFoundError("Sumber tidak ditemukan:\n" + "\n".join(absent))

    document = Document(CHAPTERS[0])
    for element in list(document.element.body.iterchildren()):
        if element.tag != qn("w:sectPr"):
            document.element.body.remove(element)
    configure(document)
    cover(document)
    front_matter(document)

    for index, path in enumerate(CHAPTERS):
        if index:
            page_break(document)
        append_source(document, path, "strip")

    for title, path, mode in APPENDICES:
        page_break(document)
        p = document.add_paragraph()
        p.paragraph_format.space_after = Pt(12)
        set_run(p.add_run(title), 21, True, PRIMARY_DARK)
        append_source(document, path, mode)

    document.core_properties.title = "Dokumentasi Final Project SIVMI CRM"
    document.core_properties.subject = "BAB I-XI, API/REST API, testing/debugging, dan matriks ERD"
    document.core_properties.author = "SIVMI"
    document.core_properties.keywords = "SIVMI, CRM, Laravel, Filament, WAHA, REST API, testing, ERD"
    audit(document)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    check = Document(OUTPUT)
    audit(check)
    print(f"Created: {OUTPUT}")
    print(f"Paragraphs: {len(check.paragraphs)}")
    print(f"Tables: {len(check.tables)}")
    print(f"Images: {len(check.inline_shapes)}")


if __name__ == "__main__":
    main()
