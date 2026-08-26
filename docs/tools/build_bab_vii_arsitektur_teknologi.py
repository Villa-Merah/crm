from __future__ import annotations

from pathlib import Path
import sys

PROJECT_ROOT = Path(__file__).resolve().parents[2]
DOC_TOOLS = PROJECT_ROOT / "storage" / "app" / "doc-tools"
sys.path.insert(0, str(DOC_TOOLS))

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

import build_bab_iv_arsitektur_bisnis as b


OUTPUT = PROJECT_ROOT / "docs" / "BAB VII - Arsitektur Teknologi SIVMI CRM v1.0.docx"
ASSETS = PROJECT_ROOT / "docs" / "_bab_vii_assets"
COVER = ASSETS / "cover-bab-vii.png"

PRIMARY = "#1476B8"
PRIMARY_DARK = "#075985"
INK = "#0B2545"
MUTED = "#64748B"
BG = "#F6F8FC"
GREEN = "#10B981"
AMBER = "#F59E0B"
ROSE = "#F43F5E"
PURPLE = "#A855F7"


def font(name: str, size: int):
    return ImageFont.truetype(f"C:/Windows/Fonts/{name}", size)


def canvas(title: str, subtitle: str, size=(1800, 1050)):
    image = Image.new("RGB", size, "white")
    draw = ImageDraw.Draw(image)
    draw.text((size[0] // 2, 48), title, anchor="ma", font=font("arialbd.ttf", 44), fill="#14357A")
    draw.text((size[0] // 2, 105), subtitle, anchor="ma", font=font("arial.ttf", 24), fill=MUTED)
    return image, draw


def box(draw, xy, title, lines, fill="#EAF4FB", outline=PRIMARY, title_size=27, body_size=21):
    x1, y1, x2, y2 = xy
    draw.rounded_rectangle(xy, radius=22, fill=fill, outline=outline, width=4)
    draw.text(((x1 + x2) // 2, y1 + 30), title, anchor="ma", font=font("arialbd.ttf", title_size), fill=INK)
    draw.multiline_text(((x1 + x2) // 2, y1 + 80), lines, anchor="ma", align="center", spacing=8,
                        font=font("arial.ttf", body_size), fill="#334155")


def arrow(draw, start, end, color="#64748B", width=5):
    draw.line((start, end), fill=color, width=width)
    x, y = end; sx, sy = start
    if abs(x - sx) >= abs(y - sy):
        sign = 1 if x > sx else -1
        pts = [(x, y), (x-sign*18, y-11), (x-sign*18, y+11)]
    else:
        sign = 1 if y > sy else -1
        pts = [(x, y), (x-11, y-sign*18), (x+11, y-sign*18)]
    draw.polygon(pts, fill=color)


def make_cover():
    ASSETS.mkdir(parents=True, exist_ok=True)
    image = Image.new("RGB", (1600, 2200), BG); draw = ImageDraw.Draw(image)
    for y in range(2200):
        ratio=y/2200; a=(20,118,184); z=(7,89,133)
        draw.line((0,y,1600,y),fill=tuple(round(a[i]*(1-ratio)+z[i]*ratio) for i in range(3)))
    draw.ellipse((1010,-180,1770,580),fill=(27,139,199)); draw.ellipse((1140,140,1640,640),fill=(8,103,151))
    draw.rounded_rectangle((120,250,1480,1910),radius=44,fill="white")
    draw.text((190,350),"SIVMI CRM",font=font("arialbd.ttf",34),fill=PRIMARY)
    draw.text((190,600),"BAB VII",font=font("arialbd.ttf",48),fill=PRIMARY_DARK)
    draw.multiline_text((190,690),"ARSITEKTUR\nTEKNOLOGI",font=font("arialbd.ttf",90),fill=INK,spacing=18)
    draw.text((190,1045),"Baseline, topologi, lingkungan, state, jaringan,",font=font("arial.ttf",34),fill="#475569")
    draw.text((190,1097),"backup, monitoring, availability, dan skalabilitas",font=font("arial.ttf",34),fill="#475569")
    draw.rounded_rectangle((190,1245,760,1325),radius=40,fill="#EAF4FB")
    draw.text((230,1267),"BASELINE AS-BUILT 25 AGUSTUS 2026",font=font("arial.ttf",28),fill=PRIMARY_DARK)
    draw.text((190,1610),"Versi dokumen 1.0",font=font("arial.ttf",28),fill=MUTED)
    draw.text((190,1665),"Status: dokumen kerja untuk kesiapan VPS",font=font("arial.ttf",28),fill=MUTED)
    draw.text((190,1775),"SIVMI CRM - Catat, Kelola, Closing",font=font("arialbd.ttf",34),fill=PRIMARY)
    image.save(COVER)


def make_stack():
    image, draw = canvas("Baseline Teknologi", "Versi runtime dan framework utama pada codebase as-built", (1800, 1050))
    rows = [
        [("UI Admin","Filament 5\nLivewire 4",PRIMARY),("Styling","Tailwind CSS 4\nVite 8",PURPLE),("Runtime","PHP 8.4 FPM\nrequire >= 8.3",GREEN)],
        [("Framework","Laravel 12\nEloquent ORM",PRIMARY_DARK),("Database","MySQL 8.4\nPDO MySQL",AMBER),("State","Redis 7.4\nphpredis",ROSE)],
        [("Web","Nginx 1.27\nAlpine",PRIMARY),("Build","Node.js 22\nComposer 2",PURPLE),("WhatsApp","WAHA CORE\nGOWS",GREEN)],
    ]
    for ri,row in enumerate(rows):
        for ci,(title,body,color) in enumerate(row):
            x1=120+ci*560; y1=180+ri*270
            box(draw,(x1,y1,x1+440,y1+200),title,body,"#F8FAFC",color,30,24)
    image.save(ASSETS/"baseline-teknologi.png")


def make_logical_topology():
    image, draw = canvas("Topologi Teknologi Logis", "Aliran trafik dan dependensi service full Docker", (1800, 1120))
    box(draw,(70,190,360,390),"Client","Browser\nDesktop / mobile","#F8FAFC",PRIMARY)
    box(draw,(500,190,820,390),"Web","Nginx\npublic root","#EAF4FB",PRIMARY)
    box(draw,(970,190,1290,390),"App","PHP-FPM\nLaravel","#ECFDF5",GREEN)
    arrow(draw,(360,290),(500,290)); arrow(draw,(820,290),(970,290))
    services=[
        (80,690,350,900,"MySQL","Data relasional\nvolume",AMBER),
        (420,690,690,900,"Redis","Cache / session / queue\nproduksi",ROSE),
        (760,690,1030,900,"Queue","Job async\nfailed jobs",PURPLE),
        (1100,690,1370,900,"Scheduler","Reminder\nsync session",PRIMARY_DARK),
        (1440,690,1730,900,"WAHA","GOWS\nWhatsApp session",GREEN),
    ]
    for x1,y1,x2,y2,t,l,c in services: box(draw,(x1,y1,x2,y2),t,l,"#F8FAFC",c,25,19)
    for x in (215,555,895,1235,1585): arrow(draw,(1130,390),(x,690),color="#94A3B8",width=3)
    draw.rounded_rectangle((450,960,1350,1050),radius=18,fill="#EAF4FB",outline=PRIMARY,width=3)
    draw.text((900,1005),"Docker bridge network + DNS service name; hanya port yang diperlukan diekspos",anchor="mm",font=font("arialbd.ttf",22),fill=INK)
    image.save(ASSETS/"topologi-logis.png")


def make_environments():
    image, draw = canvas("Pemisahan Lingkungan", "Development, testing, staging, dan production memiliki state serta kontrol berbeda", (1800, 1080))
    envs=[
        (50,180,430,920,"Development lokal","Compose development\nSource bind mount\nVite HMR\nAPP_DEBUG=true\nDB/session/queue/cache database\nPort 8000, 3307, 5173, 3000",PRIMARY),
        (490,180,870,920,"Testing","Container app\nAPP_ENV=testing\nSQLite :memory:\nCache/session array\nQueue sync\nTidak menyentuh DB lokal",GREEN),
        (930,180,1310,920,"Staging (target)","Image produksi\nHTTPS + secret test\nData sintetis\nWorker/scheduler aktif\nBackup dan restore drill\nUAT lintas perangkat",AMBER),
        (1370,180,1750,920,"Production (target)","Immutable image\nAPP_DEBUG=false\nRedis state\nMySQL persisten\nWAHA dashboard off\nTLS, backup, monitoring\nRelease terkontrol",ROSE),
    ]
    for x1,y1,x2,y2,title,lines,color in envs:
        box(draw,(x1,y1,x2,y2),title,lines,"#F8FAFC",color,27,21)
    image.save(ASSETS/"lingkungan.png")


def make_state():
    image, draw = canvas("Session, Cookie, dan Cache", "State request dipisahkan dari data transaksi dan dapat dipindahkan ke Redis", (1800, 1050))
    box(draw,(80,250,400,520),"Browser","Cookie session ID\nSecure + SameSite\nCSRF token","#F8FAFC",PRIMARY)
    box(draw,(560,210,940,560),"Laravel","Session middleware\nCookie encryption\nAuth + CSRF\nCrmSetting cache","#EAF4FB",PRIMARY)
    arrow(draw,(400,385),(560,385)); arrow(draw,(560,445),(400,445),color=GREEN)
    box(draw,(1120,170,1640,430),"Development","Session database\nCache database\nQueue database","#FFFBEB",AMBER)
    box(draw,(1120,580,1640,840),"Production target","Session Redis + encrypt\nCache Redis\nQueue Redis","#ECFDF5",GREEN)
    arrow(draw,(940,330),(1120,300)); arrow(draw,(940,475),(1120,700))
    draw.rounded_rectangle((210,720,880,900),radius=22,fill="#FFF1F2",outline=ROSE,width=3)
    draw.multiline_text((545,810),"Data bisnis tetap di MySQL.\nCache boleh hilang; session memiliki masa berlaku;\nqueue harus dipantau dan dapat di-retry.",anchor="mm",align="center",font=font("arialbd.ttf",23),fill=INK)
    image.save(ASSETS/"state-session-cache.png")


def make_network():
    image, draw = canvas("Zona Jaringan dan Komunikasi", "Target VPS meminimalkan permukaan publik dan menjaga trafik internal", (1800, 1100))
    draw.rounded_rectangle((50,160,1750,1000),radius=30,fill="#F8FAFC",outline="#CBD5E1",width=4)
    draw.text((130,190),"VPS / HOST",font=font("arialbd.ttf",28),fill=PRIMARY_DARK)
    box(draw,(90,300,390,520),"Internet","User browser\nWhatsApp cloud","white",PRIMARY)
    box(draw,(530,270,870,550),"Edge target","DNS + TLS\nFirewall/WAF\nReverse proxy","#EAF4FB",PRIMARY)
    box(draw,(1020,270,1360,550),"Web/App","Nginx + PHP-FPM\nport internal 9000","#ECFDF5",GREEN)
    arrow(draw,(390,410),(530,410)); arrow(draw,(870,410),(1020,410))
    box(draw,(260,690,560,890),"Data zone","MySQL 3306\nRedis 6379","#FFFBEB",AMBER)
    box(draw,(750,690,1050,890),"Worker zone","Queue\nScheduler","#FAF5FF",PURPLE)
    box(draw,(1240,690,1540,890),"Integration","WAHA 3000\nOutbound WhatsApp","#FFF1F2",ROSE)
    for x in (410,900,1390): arrow(draw,(1190,550),(x,690),color="#94A3B8",width=3)
    image.save(ASSETS/"zona-jaringan.png")


def make_backup():
    image, draw = canvas("Strategi Backup dan Restore", "Target kontrol 3-2-1 dengan verifikasi dan restore rehearsal", (1800, 1080))
    sources=[("MySQL","Dump konsisten\n+ binary log target",120,220,PRIMARY),
             ("Storage","File aplikasi\nprivate/public",120,500,GREEN),
             ("WAHA","Session volume\nmetadata akun",120,780,PURPLE)]
    for t,l,x,y,c in sources: box(draw,(x,y,x+330,y+190),t,l,"#F8FAFC",c,27,20)
    box(draw,(650,350,1050,750),"Backup pipeline","Snapshot/dump\nCompress\nEncrypt\nChecksum\nRetention policy\nJob result + alert","#EAF4FB",PRIMARY_DARK,31,22)
    for _,_,x,y,_ in sources: arrow(draw,(x+330,y+95),(650,500+(y-500)//3),color="#94A3B8",width=4)
    box(draw,(1260,220,1690,440),"Copy 1","Local/attached\nrecovery cepat","#FFFBEB",AMBER)
    box(draw,(1260,520,1690,740),"Copy 2","Off-site/object storage\nencrypted","#ECFDF5",GREEN)
    box(draw,(1260,820,1690,1000),"Restore drill","Staging terisolasi\nukur RPO/RTO","#FFF1F2",ROSE)
    arrow(draw,(1050,460),(1260,330)); arrow(draw,(1050,550),(1260,630)); arrow(draw,(1475,740),(1475,820))
    image.save(ASSETS/"backup-restore.png")


def make_observability():
    image, draw = canvas("Logging, Monitoring, dan Health Check", "Dari sinyal teknis menuju alert dan tindakan operasional", (1800, 1050))
    signals=[
        (70,200,400,390,"Application","Laravel log\nexception/audit",PRIMARY),
        (70,440,400,630,"Runtime","Nginx/PHP\ncontainer state",GREEN),
        (70,680,400,870,"Data/async","MySQL/Redis\nqueue/scheduler",AMBER),
    ]
    for x1,y1,x2,y2,t,l,c in signals: box(draw,(x1,y1,x2,y2),t,l,"#F8FAFC",c,26,20)
    box(draw,(620,290,1060,780),"Collector target","Structured log\nMetric\nHealth probe\nQueue age\nError rate\nDisk/capacity","#EAF4FB",PRIMARY_DARK,31,22)
    for _,y1,_,y2,_,_,_ in signals: arrow(draw,(400,(y1+y2)//2),(620,535),color="#94A3B8",width=4)
    box(draw,(1280,220,1710,450),"Dashboard","Availability\nlatency\ncapacity","#ECFDF5",GREEN)
    box(draw,(1280,570,1710,800),"Alert","Owner + threshold\nrunbook + escalation","#FFF1F2",ROSE)
    arrow(draw,(1060,450),(1280,335)); arrow(draw,(1060,620),(1280,685))
    draw.rounded_rectangle((520,860,1220,970),radius=20,fill="#FFFBEB",outline=AMBER,width=3)
    draw.text((870,915),"Health bukan observability: probe sehat belum membuktikan queue tidak macet",anchor="mm",font=font("arialbd.ttf",21),fill=INK)
    image.save(ASSETS/"observability.png")


def make_scalability():
    image, draw = canvas("Ketersediaan dan Skalabilitas", "Scale secara bertahap berdasarkan bottleneck dan bukti metrik", (1800, 1080))
    stages=[
        (50,220,390,840,"Tahap 1\nSingle VPS","Web + app\nworker + scheduler\nMySQL + Redis\nWAHA\nbackup off-site",PRIMARY),
        (480,220,820,840,"Tahap 2\nPisah data","Web/app tetap VPS\nDB managed/dedicated\nRedis terpisah\nobject storage\nmonitoring",GREEN),
        (910,220,1250,840,"Tahap 3\nHorizontal web","Load balancer\nN app stateless\nshared Redis/storage\nworker pool\nDB primary",AMBER),
        (1340,220,1750,840,"Tahap 4\nHA terukur","Multi-zone bila perlu\nDB replica/failover\nredundant edge\nDR site\nSLO formal",PURPLE),
    ]
    for x1,y1,x2,y2,t,l,c in stages:
        box(draw,(x1,y1,x2,y2),t,l,"#F8FAFC",c,28,22)
    for x in (390,820,1250): arrow(draw,(x,530),(x+90,530))
    image.save(ASSETS/"skalabilitas.png")


def add_figure(doc,name,caption,width=6.15): b.add_figure(doc,ASSETS/name,caption,width)
def table(doc,caption,headers,rows,widths): b.add_table(doc,headers,rows,widths,caption)
def page(doc): doc.add_page_break()


def add_steps(doc,items):
    table(doc,"Urutan proses",["Langkah","Uraian"],[[str(i),x] for i,x in enumerate(items,1)],[0.75,5.75])


def build_document():
    for fn in (make_cover,make_stack,make_logical_topology,make_environments,make_state,make_network,make_backup,make_observability,make_scalability): fn()
    doc=Document(); sec=doc.sections[0]
    sec.page_width=Inches(8.5); sec.page_height=Inches(11)
    sec.top_margin=Inches(0.85); sec.bottom_margin=Inches(0.85); sec.left_margin=Inches(1); sec.right_margin=Inches(1)
    sec.header_distance=Inches(0.4); sec.footer_distance=Inches(0.4)
    b.configure_styles(doc)
    sec.header.paragraphs[0].text="SIVMI CRM  |  BAB VII: ARSITEKTUR TEKNOLOGI"
    for run in sec.header.paragraphs[0].runs: b.set_run_font(run,8.5,True,b.PRIMARY)
    b.add_page_number(sec.footer.paragraphs[0])
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(0)
    pic=p.add_run().add_picture(str(COVER),width=Inches(6.1)); pic._inline.docPr.set("descr","Sampul BAB VII Arsitektur Teknologi SIVMI CRM")
    page(doc)

    b.add_heading(doc,"KENDALI DOKUMEN",1)
    table(doc,"Tabel 0-1. Kendali dan baseline dokumen",["Atribut","Nilai"],[
        ["Judul","BAB VII: Arsitektur Teknologi SIVMI CRM"],["Versi","1.0"],["Tanggal baseline","25 Agustus 2026"],
        ["Status","Dokumen kerja terpisah - siap validasi teknologi dan kesiapan VPS"],
        ["Ruang lingkup","Baseline stack, topologi, lingkungan, session/cookie/cache, queue, storage, backup, jaringan, observability, availability, scalability, dan verifikasi"],
        ["Basis implementasi","PHP 8.4 FPM, Laravel 12, Filament 5, Livewire 4, Nginx 1.27, MySQL 8.4, Redis 7.4, Node 22, WAHA CORE/GOWS"],
        ["Deployment","Full Docker untuk lokal; compose produksi sebagai baseline target VPS"],
        ["Sumber kebenaran","Dockerfile, Compose, env example, entrypoint, Nginx/PHP config, Laravel config, scripts start/status/test, routes, dan test"],
    ],[1.55,4.95])
    b.add_callout(doc,"Kedudukan dokumen","BAB VII memisahkan fakta as-built dari target. Kontrol yang belum memiliki otomatisasi atau bukti operasional tidak dinyatakan selesai, meskipun telah tersedia sebagai konfigurasi contoh atau rekomendasi.")
    b.add_heading(doc,"DAFTAR ISI",1); b.add_toc(doc); page(doc)

    b.add_heading(doc,"RINGKASAN EKSEKUTIF",1)
    b.add_body(doc,"SIVMI CRM telah memiliki baseline teknologi containerized yang dapat berjalan tanpa Laragon. Compose development mengorkestrasi Nginx, PHP-FPM, MySQL, Redis, Node/Vite, queue worker, scheduler, dan WAHA. Compose produksi membangun image immutable untuk app dan web, memindahkan state Laravel ke Redis, mempertahankan volume data, serta memasang pemeriksaan secret dan mode produksi pada entrypoint.")
    b.add_body(doc,"Kesiapan VPS belum hanya ditentukan oleh keberhasilan container. TLS termination, firewall, secret management, backup otomatis dan restore drill, monitoring terpusat, alert queue/scheduler, pinning image, kapasitas, dan prosedur incident recovery masih perlu diimplementasikan serta dibuktikan. Arsitektur yang disarankan dimulai dari satu VPS yang di-hardening, lalu dipisah atau diskalakan hanya berdasarkan metrik.")
    b.add_callout(doc,"Keputusan utama","Full Docker menjadi baseline pengembangan dan deployment. Laragon tetap dapat dipakai sebagai fallback developer, tetapi bukan dependensi arsitektur target. VPS awal memakai single-node terkontrol; HA dan microservice tidak dipaksakan sebelum kebutuhan terukur.",b.LIGHT_GREEN)
    table(doc,"Tabel 0-2. Status ringkas",["Area","As-built","Target VPS"],[
        ["Runtime","Compose full stack berjalan lokal","Image immutable + registry/tag"],
        ["State","MySQL volume; Redis tersedia","Redis session/cache/queue; volume/managed data"],
        ["Security","Secret lokal digenerate; HMAC/API key","TLS, firewall, secret store, rotation"],
        ["Recovery","Volume persisten","Backup off-site otomatis + restore drill"],
        ["Observability","Log dan health dasar","Structured log, metric, alert, SLO"],
        ["Availability","Restart + health dependency","Runbook, backup, optional redundancy"],
    ],[1.2,2.55,2.75])

    b.add_heading(doc,"BAB VII: ARSITEKTUR TEKNOLOGI",1)
    b.add_heading(doc,"7.1 Pengertian",2)
    b.add_body(doc,"Arsitektur teknologi mendefinisikan platform runtime, jaringan, penyimpanan, mekanisme state, deployment, keamanan infrastruktur, observability, backup, dan kapasitas yang menopang arsitektur aplikasi. Tujuannya memastikan kode yang benar dapat dijalankan secara aman, konsisten, dapat dipulihkan, dan dapat dikembangkan.")
    table(doc,"Tabel 7-1. Domain teknologi",["Domain","Pertanyaan kontrol","Artefak"],[
        ["Compute/runtime","Di mana proses berjalan?","Image, container, PHP-FPM, worker"],
        ["Network","Siapa dapat terhubung ke port apa?","Bridge, DNS, firewall, TLS"],
        ["State/storage","Di mana data dan session bertahan?","MySQL, Redis, volume, filesystem"],
        ["Operations","Bagaimana diketahui sehat/gagal?","Health, log, metric, alert"],
        ["Continuity","Bagaimana pulih dari kehilangan?","Backup, restore, RPO/RTO, runbook"],
    ],[1.25,2.7,2.55])

    b.add_heading(doc,"7.2 Baseline Teknologi",2)
    add_figure(doc,"baseline-teknologi.png","Gambar 7-1. Baseline versi teknologi utama")
    table(doc,"Tabel 7-2. Inventaris baseline",["Lapisan","Teknologi/versi","Fungsi","Sumber"],[
        ["Runtime","PHP 8.4 FPM; requirement >=8.3","Eksekusi Laravel","Dockerfile/composer.json"],
        ["Framework","Laravel 12; Eloquent","Web, domain, queue, scheduler","composer.json"],
        ["Admin UI","Filament 5; Livewire 4","Panel server-driven","composer.lock/provider"],
        ["Frontend","Tailwind CSS 4; Vite 8","Theme dan build asset","package.json"],
        ["Web server","Nginx 1.27 Alpine","Static/public dan FastCGI","Compose/Dockerfile"],
        ["Database","MySQL 8.4","Data transaksi","Compose"],
        ["State","Redis 7.4 + phpredis","Cache/session/queue target","Compose/Dockerfile"],
        ["Build","Node 22; Composer 2","Dependency/build multi-stage","Dockerfile"],
        ["WhatsApp","WAHA CORE/GOWS","Session, API, webhook","Compose/services config"],
        ["Test","PHPUnit 12.5; SQLite memory","Regression terisolasi","composer/test script"],
    ],[1.05,1.75,2.1,1.6])
    b.add_heading(doc,"7.2.1 Parameter PHP dan Web",3)
    table(doc,"Tabel 7-3. Parameter runtime",["Parameter","Nilai","Implikasi"],[
        ["memory_limit","512 MB","Batas per proses PHP; monitor concurrency"],
        ["upload/post max","20 MB","Selaras client_max_body_size Nginx"],
        ["max_execution_time","120 detik","Request panjang harus dipindah ke queue"],
        ["timezone","Asia/Jakarta","Konsistensi jadwal dan laporan"],
        ["opcache","Aktif; timestamp validation aktif","Cocok dev; produksi dapat dituning"],
        ["FastCGI read timeout","120 detik","Bukan pengganti queue"],
    ],[1.7,1.65,3.15])

    page(doc); b.add_heading(doc,"7.3 Topologi Teknologi Logis",2)
    add_figure(doc,"topologi-logis.png","Gambar 7-2. Topologi logis full Docker")
    table(doc,"Tabel 7-4. Service, state, dan dependency",["Service","Dependency","State","Restart/health"],[
        ["web","app healthy","Stateless + public build","restart; /up produksi"],
        ["app","database + redis healthy","storage/cache runtime","FPM socket health"],
        ["database","-","mysql_data","mysqladmin ping"],
        ["redis","-","redis_data + AOF","redis-cli ping"],
        ["queue","database + redis","jobs/failed jobs","worker; health belum aktif"],
        ["scheduler","database + redis","schedule process","health belum aktif"],
        ["waha","Network/outbound internet","waha_sessions","/health"],
        ["node (dev)","Source/package lock","node_modules","Vite HMR"],
    ],[1.1,1.85,1.65,1.9])

    b.add_heading(doc,"7.4 Lingkungan",2)
    add_figure(doc,"lingkungan.png","Gambar 7-3. Pemisahan lingkungan dan kontrol")
    table(doc,"Tabel 7-5. Matriks environment",["Aspek","Development","Testing","Production target"],[
        ["Code","Bind mount host","Image/source test","Immutable image"],
        ["Debug","true","false","false wajib"],
        ["Database","MySQL volume","SQLite :memory:","MySQL volume/managed"],
        ["Session/cache/queue","Database","Array/array/sync","Redis"],
        ["Asset","Vite HMR 5173","Build bila perlu","public/build baked"],
        ["WAHA","Dashboard aktif, GOWS","Fake/mock HTTP","Dashboard off"],
        ["Data","Seeder demo","Factory/fixture","Migrasi tanpa demo seeder"],
        ["Secrets","Generated local","Ephemeral","Secret manager/deployment env"],
    ],[1.2,1.7,1.45,2.15])
    b.add_heading(doc,"7.4.1 Konfigurasi dan Secret",3)
    b.add_bullets(doc,[
        "docker/.env dibuat dari example dan tidak boleh dikomit; start-full menghasilkan APP_KEY, WAHA API key, HMAC secret, dan password dashboard lokal jika masih default.",
        "Produksi mewajibkan APP_DEBUG=false, APP_URL HTTPS, APP_KEY, DB password/root password, dan Redis password non-placeholder melalui entrypoint.",
        "WAHA API key hash, HMAC secret, SMTP credential, dan kunci lainnya harus disimpan pada secret manager atau environment deployment yang dibatasi.",
        "Staging dan production tidak boleh berbagi key, database, nomor WhatsApp, cookie domain, atau volume.",
    ])

    b.add_heading(doc,"7.5 Session, Cookie, dan Cache",2)
    add_figure(doc,"state-session-cache.png","Gambar 7-4. Arsitektur state session, cookie, dan cache")
    table(doc,"Tabel 7-6. Kebijakan state",["Objek","Lokal","Produksi target","Kontrol"],[
        ["Session","Database","Redis/default connection","TTL, encryption, logout/revoke"],
        ["Cookie","HTTP lokal","Secure=true; SameSite=lax","HTTPS, HttpOnly framework"],
        ["Cache","Database","Redis","Key prefix, invalidation, capacity"],
        ["Maintenance state","Default/file","Cache store Redis","Konsisten lintas instance"],
        ["CrmSetting","Cache rememberForever","Redis cache","Explicit Cache::forget"],
    ],[1.25,1.5,1.65,2.1])
    b.add_callout(doc,"Prinsip state","Session dan cache bukan sumber kebenaran data bisnis. Cache boleh dibangun ulang; kehilangan session memaksa login ulang; queue memerlukan delivery/retry; data transaksi tetap berada pada MySQL dan backupnya menjadi prioritas tertinggi.",b.LIGHT_GREEN)

    b.add_heading(doc,"7.6 Queue dan Scheduler",2)
    table(doc,"Tabel 7-7. Komponen asynchronous",["Komponen","Lokal","Produksi","Kontrol"],[
        ["Queue driver","Database","Redis","retry_after 90; failed_jobs"],
        ["Worker","Profile workers; --tries=3 --timeout=90","Service tetap; max-time=3600","Restart dan observasi queue age"],
        ["Webhook job","ProcessWahaWebhook","Sama","tries=5; backoff; idempotensi"],
        ["Scheduler","schedule:work","Service tetap","withoutOverlapping"],
        ["Reminder","Hourly","Hourly","chunkById 100; dedupe notification"],
        ["WAHA sync","Setiap 5 menit","Setiap 5 menit","per-account catch/report"],
    ],[1.4,1.55,1.55,2.0])
    add_steps(doc,[
        "Start-full mengaktifkan profile workers agar queue dan scheduler tidak tertinggal.",
        "Webhook masuk disimpan sebagai event lalu job dipush; request mengembalikan 202 tanpa menunggu pemrosesan.",
        "Worker menjalankan job, retry terkontrol, dan memindahkan kegagalan akhir ke failed_jobs.",
        "Scheduler memicu reminder dan sinkronisasi status session dengan lock tanpa overlap.",
        "Monitoring harus membedakan container hidup, worker aktif, queue age, failed jobs, dan schedule terakhir berhasil.",
    ])
    b.add_callout(doc,"Gap","Health check queue dan scheduler dinonaktifkan pada compose produksi. Tambahkan heartbeat/job sentinel atau process supervisor check sebelum go-live agar proses hidup tetapi macet dapat dideteksi.",b.LIGHT_AMBER)

    page(doc); b.add_heading(doc,"7.7 Penyimpanan File",2)
    table(doc,"Tabel 7-8. Peta storage",["Storage","As-built","Persistensi","Target kontrol"],[
        ["storage/app/private","Disk local","storage_data produksi","Tidak public; backup/enkripsi"],
        ["storage/app/public","Disk public + link","storage_data","Validasi MIME; CDN/object storage opsional"],
        ["public/build","Baked dalam image produksi","Image immutable","Hash asset dan registry"],
        ["storage/logs","Local volume/container","Tergantung deployment","Ship/rotate/retention"],
        ["mysql_data","Named volume","Persisten","Backup konsisten + off-site"],
        ["redis_data","Named volume + AOF","Persisten","Password, capacity, backup sesuai kebutuhan"],
        ["waha_sessions","Named volume","Persisten","Backup hati-hati; akses sangat terbatas"],
        ["S3","Config tersedia","Belum aktif","Aktif bila skala/DR membenarkan"],
    ],[1.45,1.75,1.35,1.95])

    page(doc); b.add_heading(doc,"7.8 Backup, Restore, RPO, dan RTO",2)
    add_figure(doc,"backup-restore.png","Gambar 7-5. Target pipeline backup dan restore")
    b.add_body(doc,"Codebase belum memiliki job backup otomatis yang membuktikan keberhasilan backup atau restore. Karena itu tabel berikut adalah target yang diusulkan dan harus disahkan pemilik bisnis, administrator, serta penyedia VPS. Angka tidak boleh dianggap SLA sampai diuji.")
    table(doc,"Tabel 7-9. Usulan sasaran pemulihan",["Kelas data","RPO usulan","RTO usulan","Metode"],[
        ["MySQL transaksi","<= 15 menit","<= 4 jam","Snapshot/dump berkala; binlog bila tersedia"],
        ["Storage aplikasi","<= 24 jam","<= 8 jam","Incremental harian + checksum"],
        ["WAHA session","<= 24 jam","<= 4 jam","Volume backup; opsi scan ulang terkontrol"],
        ["Redis cache/session","Best effort","<= 1 jam","Rebuild cache; pengguna login ulang"],
        ["Image/config release","RPO 0 per tag","<= 1 jam","Registry + config/secret version"],
        ["Audit/log operasional","<= 24 jam","<= 8 jam","Log shipping/backup sesuai retensi"],
    ],[1.45,1.1,1.1,2.85])
    b.add_heading(doc,"7.8.1 Prosedur Backup",3)
    add_steps(doc,[
        "Catat versi aplikasi, migration, image tag, ukuran volume, dan waktu mulai.",
        "Buat dump/snapshot konsisten MySQL tanpa menghentikan transaksi secara sembarangan.",
        "Salin storage aplikasi dan session WAHA dengan metode yang menjaga konsistensi file.",
        "Kompres, enkripsi, buat checksum, dan kirim copy off-site dengan akses terpisah.",
        "Verifikasi ukuran, checksum, umur backup, dan kirim alert bila gagal atau terlambat.",
        "Terapkan retensi harian/mingguan/bulanan yang telah disahkan dan purge aman.",
    ])
    b.add_heading(doc,"7.8.2 Restore Rehearsal",3)
    add_steps(doc,[
        "Pulihkan ke staging terisolasi, bukan menimpa production.",
        "Validasi checksum, import database, file storage, permission, dan secret dummy.",
        "Jalankan migration status, health check, login, RBAC, dashboard, transaksi, queue, dan WAHA sesuai scope drill.",
        "Ukur kehilangan data aktual dan total waktu pemulihan; bandingkan dengan RPO/RTO.",
        "Dokumentasikan temuan, perbaiki runbook, dan minta persetujuan penutupan drill.",
    ])

    page(doc); b.add_heading(doc,"7.9 Jaringan dan Komunikasi",2)
    add_figure(doc,"zona-jaringan.png","Gambar 7-6. Zona jaringan target VPS")
    table(doc,"Tabel 7-10. Matriks port dan exposure",["Port/service","Lokal","Produksi target","Aturan"],[
        ["80/443 edge","Host 8000 -> web:80","Publik 443; redirect 80","TLS termination + firewall"],
        ["PHP-FPM 9000","Internal app","Internal saja","Tidak diekspos internet"],
        ["MySQL 3306","Host 3307 untuk dev","Internal saja","Akses admin via tunnel/VPN"],
        ["Redis 6379","Internal","Internal + password","Tidak public"],
        ["Vite 5173","Publik host dev","Tidak ada","Hanya development"],
        ["WAHA 3000","Host untuk dashboard lokal","Internal; dashboard off","Akses admin via tunnel bila perlu"],
        ["Outbound 443","Composer/npm/WAHA/SMTP","Allowlist sesuai kebutuhan","DNS/egress logging"],
    ],[1.35,1.5,1.65,2.0])
    b.add_heading(doc,"7.9.1 Kontrol Nginx dan HTTP",3)
    b.add_bullets(doc,[
        "Document root hanya public; file dot ditolak; PHP harus ada sebelum diteruskan ke FastCGI.",
        "Security header as-built: X-Content-Type-Options, X-Frame-Options SAMEORIGIN, dan Referrer-Policy.",
        "Produksi perlu TLS modern, HSTS setelah validasi, Content-Security-Policy bertahap, rate limiting, dan trusted proxy yang benar.",
        "Resolver Docker 127.0.0.11 mengurangi risiko stale IP ketika container app dibuat ulang.",
        "Webhook WAHA memakai jaringan internal dan HMAC; endpoint publik pihak ketiga perlu rate limit serta audit tambahan.",
    ])

    b.add_heading(doc,"7.10 Logging, Monitoring, dan Health Check",2)
    add_figure(doc,"observability.png","Gambar 7-7. Target observability dan alur alert")
    table(doc,"Tabel 7-11. Health check as-built",["Komponen","Probe","Frekuensi/retry","Makna"],[
        ["Laravel","GET /up","Dari web/start script","Framework dapat merespons"],
        ["PHP-FPM","Socket 127.0.0.1:9000","5/10 detik sesuai compose","FPM menerima koneksi"],
        ["Nginx produksi","wget /up","10 detik; 12 retry","Web ke app berfungsi"],
        ["MySQL","mysqladmin ping","5/10 detik","Server menerima koneksi"],
        ["Redis","redis-cli ping","5/10 detik","Redis merespons"],
        ["WAHA","GET /health","10 detik","Service hidup; belum menjamin session Working"],
        ["Queue/scheduler","Belum ada probe aktif","-","Gap observability"],
    ],[1.25,1.7,1.55,2.0])
    table(doc,"Tabel 7-12. Sinyal minimum target",["Sinyal","Metric/log","Threshold awal","Respons"],[
        ["Availability","HTTP success dan latency","3 kegagalan/5 menit","Triage web/app/data"],
        ["Error","5xx/exception rate","Di atas baseline","Periksa release/log"],
        ["Queue","Oldest job + failed jobs","Age > 5 menit / failed > 0","Restart/triage worker"],
        ["Scheduler","Last success","Lewat 2 interval","Periksa lock/process"],
        ["Database","Connections, slow query, disk","80% capacity / query lambat","Tune/scale/storage"],
        ["WAHA","Service + session status","FAILED/disconnected","Reconnect/runbook"],
        ["Backup","Age + checksum + job status","Lewat jadwal/gagal","Eskalasi admin"],
    ],[1.25,1.85,1.6,1.8])
    b.add_callout(doc,"Logging aman","Produksi menggunakan stack daily pada level warning sebagai baseline example. Retensi default daily adalah 14 file; kebijakan final harus mempertimbangkan audit dan kapasitas. Secret, raw credential, payload personal, dan body chat tidak boleh dicatat pada log umum.",b.LIGHT_GREEN)

    page(doc); b.add_heading(doc,"7.11 Ketersediaan dan Skalabilitas",2)
    add_figure(doc,"skalabilitas.png","Gambar 7-8. Roadmap ketersediaan dan skalabilitas bertahap")
    table(doc,"Tabel 7-13. Kontrol availability",["Kontrol","As-built","Keterbatasan","Target"],[
        ["Restart policy","unless-stopped","Tidak menangani host failure","Host auto-start + monitoring"],
        ["Health dependency","DB/Redis/app health","Bukan business readiness","Readiness/smoke check"],
        ["Persistent volume","MySQL/Redis/storage/WAHA","Tetap satu host","Backup off-site/managed"],
        ["Stateless web","Sebagian; storage terpisah produksi","Single replica","Shared state + load balancer"],
        ["Database","Single MySQL","Single point of failure","Managed/replica bila SLA perlu"],
        ["Redis","Single Redis+AOF","Single point of failure","Managed/sentinel bila perlu"],
        ["WAHA","Single service/session volume","Kanal eksternal kritis","Runbook reconnect dan SLA"],
    ],[1.3,1.75,1.75,1.7])
    b.add_heading(doc,"7.11.1 Strategi Scale",3)
    b.add_bullets(doc,[
        "Scale vertical lebih dahulu untuk single VPS sambil mengukur CPU, RAM, I/O, disk, latency, connection, dan queue age.",
        "Pisahkan MySQL/Redis atau gunakan managed service ketika resource contention, recovery, atau SLA membenarkannya.",
        "Horizontal web memerlukan session/cache/queue bersama di Redis dan file bersama di object storage; scheduler hanya satu leader.",
        "Worker dapat diskalakan per queue, tetapi job harus idempoten dan rate limit WAHA harus dihormati.",
        "HA multi-zone baru diterapkan setelah SLO, dampak downtime, biaya, dan runbook failover disetujui.",
    ])

    b.add_heading(doc,"7.12 Status As-Built dan Target Teknologi",2)
    status_rows=[
        ["Runtime full Docker","Implemented","Pertahankan start/status/test script"],
        ["Image produksi multi-stage","Implemented","Tambahkan registry, tag immutable, SBOM/signing"],
        ["Secret gate produksi","Partial","Validasi seluruh WAHA/SMTP secret dan rotasi"],
        ["Redis production state","Configured","Buktikan di staging dan monitor capacity"],
        ["TLS/domain/firewall","Target","Implementasi pada edge VPS"],
        ["Backup otomatis/off-site","Target High","Implementasi, alert, restore drill"],
        ["Central log/metrics/alert","Target High","Pilih stack dan owner on-call"],
        ["Queue/scheduler heartbeat","Gap High","Tambahkan probe dan alert"],
        ["Staging environment","Target High","Pisahkan data/secret dan UAT"],
        ["HA/replication","Future","Aktifkan berdasarkan SLO dan metrik"],
        ["Object storage/CDN","Future","Aktifkan berdasarkan volume/file"],
        ["CI/CD + security scan","Target Medium","Build/test/image scan/release gate"],
    ]
    table(doc,"Tabel 7-14. Matriks as-built dan target",["Kapabilitas","Status","Tindakan"],status_rows,[2.2,1.25,3.05])
    table(doc,"Tabel 7-15. Prioritas sebelum VPS",["Prioritas","Deliverable","Kriteria selesai"],[
        ["P0","Domain/TLS, firewall, secret unik, APP_DEBUG=false","Security smoke test lulus"],
        ["P0","Backup MySQL/storage/WAHA + off-site","Restore rehearsal memenuhi target"],
        ["P0","Worker/scheduler aktif + heartbeat","Webhook/reminder terbukti end-to-end"],
        ["P1","Monitoring, alert, log rotation, disk/capacity","Alert test diterima owner"],
        ["P1","Registry/tag image dan rollback runbook","Rollback rehearsal berhasil"],
        ["P1","Staging dan release checklist","UAT dan approval tercatat"],
        ["P2","SBOM, image scan, dependency audit","Tidak ada critical unresolved"],
    ],[0.75,2.8,2.95])

    page(doc); b.add_heading(doc,"7.13 Dasar Verifikasi Project",2)
    b.add_body(doc,"Verifikasi teknologi harus dapat diulang dari repository tanpa bergantung pada pengetahuan operator. Bukti minimum mencakup versi file, hash/tag image, output konfigurasi Compose, health status, test/build, migration status, backup/restore, dan smoke test aplikasi.")
    table(doc,"Tabel 7-16. Sumber verifikasi",["Artefak","Yang diverifikasi","Bukti"],[
        ["composer.json/lock","PHP/Laravel/Filament/PHPUnit","Version constraint + lock"],
        ["package.json/lock","Tailwind/Vite/Node dependency","Lock reproducible"],
        ["Dockerfile","Runtime extension dan multi-stage build","Build image sukses"],
        ["compose*.yml","Service, network, volume, health, command","docker compose config/ps"],
        ["env example","Mode lokal/produksi dan secret requirement","Checklist tanpa nilai secret"],
        ["Nginx/PHP config","Public root, limit, timeout, header","Config test + HTTP probe"],
        ["Laravel config/routes","State driver, queue, log, health","config show/cache + route list"],
        ["start/status/test scripts","Operasi Windows full Docker","Output script dan exit code"],
        ["Test suite/build","Regression dan asset","Test/build report"],
        ["Backup/restore log","Continuity","Timestamp, checksum, RPO/RTO aktual"],
    ],[1.65,2.55,2.3])
    b.add_heading(doc,"7.13.1 Checklist Verifikasi Lokal",3)
    add_steps(doc,[
        "Jalankan docker/start-full.ps1 dan pastikan app, web, database, Redis, queue, scheduler, Node, dan WAHA sesuai profile berstatus berjalan/sehat.",
        "Akses /up, login admin, dashboard, serta WAHA /health; verifikasi session WhatsApp berstatus Working secara terpisah.",
        "Jalankan docker/test-full.ps1 agar test memakai SQLite memory dan tidak menyentuh database development.",
        "Jalankan build Vite produksi dan pastikan public/build tersedia tanpa error.",
        "Uji webhook HMAC, queue processing, inbox, balasan chat, reminder, audit, dan notifikasi.",
        "Catat docker compose config/ps, image digest, migration status, test summary, dan temuan visual UAT.",
    ])
    b.add_heading(doc,"7.13.2 Gate Deployment VPS",3)
    table(doc,"Tabel 7-17. Gate rilis teknologi",["Gate","Lulus bila","Stop bila"],[
        ["Security","HTTPS, secret unik, debug off, port internal tertutup","Default secret/HTTP/public DB"],
        ["Build","Image immutable dan dependency check lulus","Build drift/critical vulnerability"],
        ["Data","Backup terbaru dan migration reviewed","Backup gagal/tidak ada rollback"],
        ["Runtime","Semua health dan heartbeat lulus","Worker/scheduler/WAHA gagal"],
        ["Application","Test, smoke, RBAC, transaksi, webhook lulus","Regresi atau data scope bocor"],
        ["Recovery","Restore/rollback rehearsal tersedia","RTO/RPO belum terukur"],
        ["Operations","Monitoring, alert, owner, runbook siap","Tidak ada observasi/eskalasi"],
    ],[1.05,2.75,2.7])

    b.add_heading(doc,"LAMPIRAN A - REGISTER RISIKO TEKNOLOGI",1)
    table(doc,"Tabel A-1. Risiko dan mitigasi",["ID","Risiko","Dampak","Mitigasi"],[
        ["T-01","Single VPS/volume rusak","Downtime dan kehilangan data","Off-site backup, restore drill, provider snapshot"],
        ["T-02","Worker tidak aktif","Webhook/inbox berhenti","Heartbeat, queue age alert, restart policy"],
        ["T-03","Secret bocor/default","Akses tidak sah","Secret store, rotation, log redaction"],
        ["T-04","Disk penuh","DB/log/WAHA gagal","Capacity alert, retention, cleanup"],
        ["T-05","WAHA/session gagal","Kanal chat terputus","Session monitor, reconnect runbook, fallback operasional"],
        ["T-06","Migration merusak data","Gangguan transaksi","Review, backup, staging, backward compatibility"],
        ["T-07","Redis gagal","Session/queue/cache terganggu","AOF, restart, fallback/recovery runbook"],
        ["T-08","Image latest berubah","Deployment tidak reproducible","Pin version/digest dan tag immutable"],
    ],[0.6,2.0,1.8,2.1])

    b.add_heading(doc,"PENUTUP",1)
    b.add_body(doc,"Arsitektur teknologi SIVMI CRM telah memiliki fondasi full Docker yang cukup untuk pengembangan lokal dan baseline deployment VPS. Kesiapan produksi harus dibuktikan melalui hardening, secret management, backup/restore, worker/scheduler heartbeat, observability, dan release gate. Skalabilitas diterapkan bertahap setelah bottleneck dan kebutuhan availability terukur.")
    b.add_callout(doc,"Baseline final BAB VII","Setiap perubahan image, versi runtime, driver state, port, volume, health check, RPO/RTO, atau topologi VPS wajib memperbarui dokumen ini, Compose/runbook, diagram, serta bukti verifikasi secara bersamaan.",b.LIGHT_GREEN)

    OUTPUT.parent.mkdir(parents=True,exist_ok=True); doc.save(OUTPUT); print(f"Created: {OUTPUT}")


if __name__ == "__main__": build_document()
