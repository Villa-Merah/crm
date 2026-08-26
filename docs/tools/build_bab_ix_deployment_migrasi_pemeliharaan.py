from __future__ import annotations

from pathlib import Path
import sys

PROJECT_ROOT = Path(__file__).resolve().parents[2]
DOC_TOOLS = PROJECT_ROOT / "storage" / "app" / "doc-tools"
sys.path.insert(0, str(DOC_TOOLS))

from PIL import Image, ImageDraw
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

import build_bab_iv_arsitektur_bisnis as b
import build_bab_vii_arsitektur_teknologi as v7


OUTPUT = PROJECT_ROOT / "docs" / "BAB IX - Deployment, Migrasi, dan Pemeliharaan Sistem SIVMI CRM v1.0.docx"
ASSETS = PROJECT_ROOT / "docs" / "_bab_ix_assets"
COVER = ASSETS / "cover-bab-ix.png"

PRIMARY = "#1476B8"
PRIMARY_DARK = "#075985"
INK = "#0B2545"
MUTED = "#64748B"
BG = "#F6F8FC"
GREEN = "#10B981"
AMBER = "#F59E0B"
ROSE = "#F43F5E"
PURPLE = "#A855F7"


def font(name: str, size: int):
    return v7.font(name, size)


def canvas(title: str, subtitle: str, size=(1800, 1080)):
    return v7.canvas(title, subtitle, size)


def box(draw, xy, title, lines, fill="#F8FAFC", outline=PRIMARY,
        title_size=27, body_size=20):
    return v7.box(draw, xy, title, lines, fill, outline, title_size, body_size)


def arrow(draw, start, end, color="#64748B", width=5):
    return v7.arrow(draw, start, end, color, width)


def make_cover() -> None:
    ASSETS.mkdir(parents=True, exist_ok=True)
    image = Image.new("RGB", (1600, 2200), BG)
    draw = ImageDraw.Draw(image)
    for y in range(2200):
        ratio = y / 2200
        start = (20, 118, 184)
        end = (7, 89, 133)
        draw.line((0, y, 1600, y), fill=tuple(round(start[i] * (1 - ratio) + end[i] * ratio) for i in range(3)))
    draw.ellipse((1010, -180, 1770, 580), fill=(27, 139, 199))
    draw.ellipse((1140, 140, 1640, 640), fill=(8, 103, 151))
    draw.rounded_rectangle((120, 250, 1480, 1910), radius=44, fill="white")
    draw.text((190, 350), "SIVMI CRM", font=font("arialbd.ttf", 34), fill=PRIMARY)
    draw.text((190, 570), "BAB IX", font=font("arialbd.ttf", 48), fill=PRIMARY_DARK)
    draw.multiline_text((190, 660), "DEPLOYMENT, MIGRASI,\nDAN PEMELIHARAAN\nSISTEM", font=font("arialbd.ttf", 72), fill=INK, spacing=16)
    draw.text((190, 1060), "Pengendalian rilis, perubahan data, pemulihan,", font=font("arial.ttf", 31), fill="#475569")
    draw.text((190, 1112), "observasi, keamanan, dan operasi berkelanjutan", font=font("arial.ttf", 31), fill="#475569")
    draw.rounded_rectangle((190, 1260, 805, 1340), radius=40, fill="#EAF4FB")
    draw.text((230, 1282), "BASELINE PROYEK 25 AGUSTUS 2026", font=font("arial.ttf", 28), fill=PRIMARY_DARK)
    draw.text((190, 1610), "Versi dokumen 1.0", font=font("arial.ttf", 28), fill=MUTED)
    draw.text((190, 1665), "Status: pedoman operasional dan target VPS", font=font("arial.ttf", 28), fill=MUTED)
    draw.text((190, 1775), "SIVMI CRM - Catat, Kelola, Closing", font=font("arialbd.ttf", 34), fill=PRIMARY)
    image.save(COVER)


def make_release_lifecycle() -> None:
    image, draw = canvas("Siklus Rilis Terkendali", "Perubahan hanya dipromosikan bila bukti dan keputusan pada setiap gate lengkap", (1800, 1100))
    nodes = [
        (70, 250, 340, 600, "Rencana", "Scope, risiko, owner\nmigration & rollback", PRIMARY),
        (410, 250, 680, 600, "Build", "Test, asset, image\nimmutable tag", PURPLE),
        (750, 250, 1020, 600, "Staging", "Config production-like\nsmoke & UAT", GREEN),
        (1090, 250, 1360, 600, "Release", "Backup, maintenance\nmigrate, start", AMBER),
        (1430, 250, 1730, 600, "Observe", "Health, log, queue\ndata & WAHA", ROSE),
    ]
    for x1, y1, x2, y2, title, lines, color in nodes:
        box(draw, (x1, y1, x2, y2), title, lines, "#F8FAFC", color, 29, 20)
    for x in (340, 680, 1020, 1360):
        arrow(draw, (x, 425), (x + 70, 425))
    draw.rounded_rectangle((390, 740, 1410, 920), radius=24, fill="#EAF4FB", outline=PRIMARY, width=3)
    draw.multiline_text((900, 830), "Gate gagal -> hentikan promosi, pulihkan kondisi aman,\ncatat bukti, dan kembalikan perubahan ke owner.", anchor="mm", align="center", spacing=8,
                        font=font("arialbd.ttf", 24), fill=INK)
    image.save(ASSETS / "siklus-rilis.png")


def make_raci_flow() -> None:
    image, draw = canvas("Kendali Peran dan Keputusan", "Satu pihak accountable; eksekusi dan verifikasi dipisahkan sejauh praktis", (1800, 1080))
    roles = [
        (70, 180, 390, 390, "Pemilik Bisnis", "Scope, prioritas, UAT\ndan risk acceptance", PRIMARY),
        (510, 180, 830, 390, "Release Manager", "Rencana, gate,\ngo/no-go, timeline", PURPLE),
        (950, 180, 1270, 390, "Tim Teknis", "Build, migration,\nconfig, rollback", GREEN),
        (1390, 180, 1730, 390, "Operasi/QA", "Backup, smoke,\nmonitoring, evidence", AMBER),
    ]
    for x1, y1, x2, y2, title, lines, color in roles:
        box(draw, (x1, y1, x2, y2), title, lines, "#F8FAFC", color, 27, 20)
    for x in (390, 830, 1270):
        arrow(draw, (x, 285), (x + 120, 285), color="#94A3B8", width=4)
    box(draw, (520, 610, 1280, 860), "Catatan Rilis Tunggal", "versi & waktu | approval | backup\nmigration | hasil smoke | observasi\nincident/rollback | penutupan", "#EAF4FB", PRIMARY_DARK, 31, 22)
    for x in (230, 670, 1110, 1560):
        arrow(draw, (x, 390), (900, 610), color="#94A3B8", width=4)
    image.save(ASSETS / "kendali-peran.png")


def make_environment_topology() -> None:
    image, draw = canvas("Promosi Antarlingkungan", "Artefak yang sama bergerak ke lingkungan yang makin ketat; konfigurasi tetap terpisah", (1800, 1080))
    envs = [
        (80, 190, 460, 760, "LOCAL", "Full Docker\nsource mount + Vite\ndata demo\nlocalhost\n\nTujuan: develop & debug", PRIMARY),
        (710, 190, 1090, 760, "STAGING", "Image immutable\nconfig production-like\ndata sintetis\nTLS test\n\nTujuan: gate & UAT", PURPLE),
        (1340, 190, 1720, 760, "PRODUCTION", "Image disetujui\nsecret unik\ndata nyata\nHTTPS + backup\n\nTujuan: operasi", GREEN),
    ]
    for x1, y1, x2, y2, title, lines, color in envs:
        box(draw, (x1, y1, x2, y2), title, lines, "#F8FAFC", color, 32, 22)
    arrow(draw, (460, 475), (710, 475), color=AMBER, width=6)
    arrow(draw, (1090, 475), (1340, 475), color=AMBER, width=6)
    draw.text((585, 430), "test + build", anchor="mm", font=font("arialbd.ttf", 22), fill=MUTED)
    draw.text((1215, 430), "approval + backup", anchor="mm", font=font("arialbd.ttf", 22), fill=MUTED)
    draw.rounded_rectangle((335, 865, 1465, 995), radius=20, fill="#FFF7ED", outline=AMBER, width=3)
    draw.text((900, 930), "Kode/image dipromosikan; database, secret, session WAHA, dan volume tidak disalin sembarang.", anchor="mm", font=font("arialbd.ttf", 22), fill=INK)
    image.save(ASSETS / "promosi-lingkungan.png")


def make_deploy_sequence() -> None:
    image, draw = canvas("Urutan Maintenance dan Migrasi", "Migration adalah keputusan rilis eksplisit, bukan efek samping restart container", (1800, 1120))
    stages = [
        (70, 180, 380, 430, "Preflight", "config --quiet\nhealth lama\nbackup siap", PRIMARY),
        (500, 180, 810, 430, "Proteksi", "maintenance bila perlu\nfreeze worker/write\ncatat mulai", AMBER),
        (930, 180, 1240, 430, "Migrasi", "run --rm app\nartisan migrate --force\ncatat output", PURPLE),
        (1360, 180, 1730, 430, "Aktivasi", "up -d\nremove-orphans\noptimize", GREEN),
    ]
    for x1, y1, x2, y2, title, lines, color in stages:
        box(draw, (x1, y1, x2, y2), title, lines, "#F8FAFC", color, 28, 20)
    for x in (380, 810, 1240):
        arrow(draw, (x, 305), (x + 120, 305))
    box(draw, (210, 680, 720, 930), "Jalur Berhasil", "health -> smoke -> up\nqueue/scheduler/WAHA\nobservasi & tutup", "#ECFDF5", GREEN, 30, 21)
    box(draw, (1080, 680, 1590, 930), "Jalur Gagal", "stop promosi -> rollback image\nrestore bila migrasi tidak kompatibel\nverifikasi & incident record", "#FFF1F2", ROSE, 30, 20)
    arrow(draw, (1545, 430), (465, 680), color=GREEN, width=5)
    arrow(draw, (1085, 430), (1335, 680), color=ROSE, width=5)
    image.save(ASSETS / "urutan-maintenance-migrasi.png")


def make_runtime_state() -> None:
    image, draw = canvas("State Runtime dan Ketergantungan", "Health web saja tidak cukup; state persisten dan proses latar harus diverifikasi", (1800, 1120))
    box(draw, (690, 400, 1110, 700), "Laravel App", "PHP-FPM\nFilament + Livewire\nconfig/route/view cache", "#EAF4FB", PRIMARY_DARK, 32, 22)
    nodes = [
        (70, 130, 450, 360, "Nginx", "HTTP /up\nstatic + /storage", PRIMARY),
        (70, 760, 450, 990, "Storage", "storage_data\npublic/private file", AMBER),
        (1350, 130, 1730, 360, "MySQL", "data transaksi\nmigration state", PURPLE),
        (1350, 760, 1730, 990, "Redis", "cache/session/queue\npersistence AOF", ROSE),
        (710, 80, 1090, 300, "Queue Worker", "job, retry, failed_jobs\nrestart policy", GREEN),
        (710, 800, 1090, 1020, "Scheduler", "reminder hourly\nWAHA sync / 5 min", PRIMARY),
    ]
    for x1, y1, x2, y2, title, lines, color in nodes:
        box(draw, (x1, y1, x2, y2), title, lines, "#F8FAFC", color, 27, 20)
        arrow(draw, ((x1 + x2) // 2, (y1 + y2) // 2), (900, 550), color="#94A3B8", width=4)
    draw.text((900, 1080), "WAHA/GOWS memakai volume session terpisah dan webhook ke aplikasi", anchor="mm", font=font("arialbd.ttf", 22), fill=MUTED)
    image.save(ASSETS / "state-runtime.png")


def make_smoke_observe() -> None:
    image, draw = canvas("Smoke Test dan Observasi Pascarilis", "Pemeriksaan teknis, bisnis, dan data membentuk satu keputusan stabilitas", (1800, 1080))
    checks = [
        (70, 170, 450, 380, "Akses", "HTTPS, /up, login\nasset production", PRIMARY),
        (520, 170, 900, 380, "Otorisasi", "role & cabang\nSuperadmin/Admin/Staff/Direksi", PURPLE),
        (970, 170, 1350, 380, "Transaksi", "lead, sekolah, daftar\npembayaran, CSV", GREEN),
        (1420, 170, 1730, 380, "Integrasi", "queue, scheduler\nWAHA/GOWS", AMBER),
    ]
    for x1, y1, x2, y2, title, lines, color in checks:
        box(draw, (x1, y1, x2, y2), title, lines, "#F8FAFC", color, 27, 19)
        arrow(draw, ((x1 + x2) // 2, y2), (900, 560), color="#94A3B8", width=4)
    box(draw, (630, 500, 1170, 720), "Observation Gate", "error rate | latency | queue age\nfailed job | DB/storage | WAHA\ndata reconciliation | keluhan user", "#EAF4FB", PRIMARY_DARK, 30, 20)
    box(draw, (180, 825, 720, 1000), "Stabil", "release ditutup\nbaseline diperbarui", "#ECFDF5", GREEN, 27, 19)
    box(draw, (1080, 825, 1620, 1000), "Tidak stabil", "contain / rollback\nincident & recovery", "#FFF1F2", ROSE, 27, 19)
    arrow(draw, (760, 720), (450, 825), color=GREEN, width=5)
    arrow(draw, (1040, 720), (1350, 825), color=ROSE, width=5)
    image.save(ASSETS / "smoke-observasi.png")


def make_backup_recovery() -> None:
    image, draw = canvas("Backup, Restore, dan Rollback", "Pemulihan aplikasi dan pemulihan data adalah jalur berbeda yang harus dikoordinasikan", (1800, 1100))
    box(draw, (70, 180, 410, 430), "Sebelum Rilis", "dump MySQL konsisten\nstorage snapshot\nmanifest + checksum", PRIMARY, PRIMARY, 28, 20)
    box(draw, (540, 180, 880, 430), "Release", "image tag baru\nmigration version\ncatatan perubahan", PURPLE, PURPLE, 28, 20)
    box(draw, (1010, 180, 1350, 430), "Gangguan", "kode / config / data\nseverity & scope", ROSE, ROSE, 28, 20)
    box(draw, (1480, 180, 1730, 430), "Keputusan", "rollback app\natau restore", AMBER, AMBER, 27, 19)
    for x in (410, 880, 1350):
        arrow(draw, (x, 305), (x + 130, 305))
    box(draw, (240, 690, 780, 940), "Rollback Aplikasi", "set IMAGE_TAG sebelumnya\nup -d --remove-orphans\nsmoke & observe", "#EAF4FB", PRIMARY_DARK, 30, 21)
    box(draw, (1020, 690, 1560, 940), "Restore Data", "maintenance + approval\nrestore DB/storage terverifikasi\nreconcile & audit", "#FFF7ED", AMBER, 30, 20)
    arrow(draw, (1605, 430), (510, 690), color=PRIMARY, width=5)
    arrow(draw, (1605, 430), (1290, 690), color=ROSE, width=5)
    image.save(ASSETS / "backup-rollback.png")


def make_ops_model() -> None:
    image, draw = canvas("Model Operasi, Keamanan, dan Insiden", "Monitoring mengubah sinyal runtime menjadi tindakan terukur dan pembelajaran", (1800, 1100))
    box(draw, (690, 400, 1110, 700), "SIVMI CRM", "web + app + data\nqueue/scheduler\nWAHA/GOWS", "#EAF4FB", PRIMARY_DARK, 32, 21)
    nodes = [
        (70, 160, 470, 380, "Availability", "health, uptime\nHTTP 5xx", PRIMARY),
        (70, 750, 470, 970, "Capacity", "CPU, RAM, disk\nDB, queue age", AMBER),
        (1330, 160, 1730, 380, "Security", "auth failure, secret\nwebhook HMAC", ROSE),
        (1330, 750, 1730, 970, "Business", "lead/chat/closing\njob & data lag", GREEN),
        (700, 90, 1100, 290, "Alert & Triage", "threshold, owner\nseverity, escalation", PURPLE),
        (700, 810, 1100, 1010, "Improve", "RCA, patch, runbook\ncapacity/security plan", PRIMARY),
    ]
    for x1, y1, x2, y2, title, lines, color in nodes:
        box(draw, (x1, y1, x2, y2), title, lines, "#F8FAFC", color, 27, 20)
        arrow(draw, ((x1 + x2) // 2, (y1 + y2) // 2), (900, 550), color="#94A3B8", width=4)
    image.save(ASSETS / "model-operasi.png")


def add_figure(doc: Document, name: str, caption: str, width=6.15) -> None:
    b.add_figure(doc, ASSETS / name, caption, width)


def table(doc: Document, caption: str, headers: list[str], rows: list[list[str]], widths: list[float]) -> None:
    b.add_table(doc, headers, rows, widths, caption)


def callout(doc: Document, title: str, body: str, fill: str = b.LIGHT_BLUE) -> None:
    b.add_callout(doc, title, body, fill)
    row = doc.tables[-1].rows[0]
    row._tr.get_or_add_trPr().append(b.OxmlElement("w:cantSplit"))


def page(doc: Document) -> None:
    doc.add_page_break()


def steps(doc: Document, caption: str, items: list[str]) -> None:
    table(doc, caption, ["Langkah", "Uraian"], [[str(i), item] for i, item in enumerate(items, 1)], [0.75, 5.75])


def build_document() -> None:
    for fn in (make_cover, make_release_lifecycle, make_raci_flow, make_environment_topology,
               make_deploy_sequence, make_runtime_state, make_smoke_observe,
               make_backup_recovery, make_ops_model):
        fn()

    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = Inches(0.85)
    sec.bottom_margin = Inches(0.85)
    sec.left_margin = Inches(1)
    sec.right_margin = Inches(1)
    sec.header_distance = Inches(0.4)
    sec.footer_distance = Inches(0.4)
    b.configure_styles(doc)
    sec.header.paragraphs[0].text = "SIVMI CRM  |  BAB IX: DEPLOYMENT, MIGRASI, DAN PEMELIHARAAN SISTEM"
    for run in sec.header.paragraphs[0].runs:
        b.set_run_font(run, 8.5, True, b.PRIMARY)
    b.add_page_number(sec.footer.paragraphs[0])

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    pic = p.add_run().add_picture(str(COVER), width=Inches(6.1))
    pic._inline.docPr.set("descr", "Sampul BAB IX Deployment, Migrasi, dan Pemeliharaan Sistem SIVMI CRM")
    page(doc)

    b.add_heading(doc, "KENDALI DOKUMEN", 1)
    table(doc, "Tabel 0-1. Kendali dan baseline dokumen", ["Atribut", "Nilai"], [
        ["Judul", "BAB IX: Deployment, Migrasi, dan Pemeliharaan Sistem SIVMI CRM"],
        ["Versi", "1.0"],
        ["Tanggal baseline", "25 Agustus 2026"],
        ["Status", "Pedoman operasional terpisah - baseline lokal dan target deployment VPS"],
        ["Ruang lingkup", "Pengendalian rilis, RACI, lingkungan, maintenance, migration, runtime state, smoke test, observasi, backup, restore, rollback, pemeliharaan, keamanan, dan insiden"],
        ["Baseline aplikasi", "Laravel 12, Filament 5, Livewire 4, PHP 8.4 FPM, Nginx 1.27, MySQL 8.4, Redis 7.4, Node 22, queue, scheduler, serta WAHA CORE/GOWS"],
        ["Bukti kualitas", "72 test dan 448 assertion lulus; build Vite 8.2.1 lulus pada 25 Agustus 2026"],
        ["Sumber kebenaran", "Codebase, Docker Compose, entrypoint, konfigurasi Laravel, migration, test suite, scripts operasi, dan runbook deployment"],
    ], [1.55, 4.95])
    callout(doc, "Kedudukan dokumen", "BAB IX menerjemahkan arsitektur teknologi dan manajemen debugging menjadi prosedur rilis serta operasi. Perintah di dalam dokumen harus dijalankan dari root proyek dengan environment yang benar; secret, domain, image tag, jalur backup, dan approval wajib disesuaikan sebelum produksi.")
    b.add_heading(doc, "DAFTAR ISI", 1)
    b.add_toc(doc)
    page(doc)

    b.add_heading(doc, "RINGKASAN EKSEKUTIF", 1)
    b.add_body(doc, "SIVMI CRM telah memiliki baseline full Docker untuk pengembangan lokal dan Compose produksi satu server. Layanan utama terdiri dari Nginx, PHP-FPM/Laravel, MySQL, Redis, queue worker, scheduler, dan WAHA dengan engine GOWS. Laragon bukan dependensi deployment. Migration produksi sengaja tidak dijalankan otomatis saat container restart agar perubahan skema tetap menjadi keputusan rilis yang dapat diaudit.")
    b.add_body(doc, "Dokumen ini memisahkan kondisi as-built dari target. Local full Docker, scripts start/status/log/test, health endpoint, volume persisten, explicit migration, production secret guard, dan smoke checklist telah tersedia. Staging production-like, TLS dan firewall VPS, backup/restore drill, monitoring terpusat, alert, CI/CD, serta RPO/RTO terukur masih menjadi gate sebelum sistem dinyatakan siap produksi.")
    callout(doc, "Keputusan utama", "Produksi tidak boleh diperlakukan sebagai tempat pengujian pertama. Rilis hanya dilanjutkan setelah image immutable, backup terverifikasi, migration direview, smoke/UAT lulus di staging, rollback tersedia, dan owner bisnis serta teknis memberi keputusan go.", b.LIGHT_GREEN)
    table(doc, "Tabel 0-2. Status kesiapan ringkas", ["Area", "As-built", "Target/gate"], [
        ["Local runtime", "Full Docker; web/app/MySQL/Redis/queue/scheduler/WAHA", "Pertahankan sebagai dev baseline"],
        ["Production Compose", "Tersedia; no source mount; healthcheck; persistent volumes", "Validasi di VPS/staging"],
        ["Quality gate", "72 test/448 assertion; build lulus", "CI, image scan, browser UAT"],
        ["Migration", "Eksplisit dengan --force; seeding produksi dilarang", "Review kompatibilitas dan rehearsal"],
        ["Recovery", "Tag image dan prinsip backup/restore terdokumentasi", "Drill serta RPO/RTO aktual"],
        ["Observability", "Health dan log per service", "Central log, metrics, alert, on-call"],
        ["Security", "Secret guard, HTTPS APP_URL, HMAC webhook, header Nginx", "TLS termination, firewall, rotation, scan"],
    ], [1.35, 2.55, 2.6])

    b.add_heading(doc, "BAB IX: DEPLOYMENT, MIGRASI, DAN PEMELIHARAAN SISTEM", 1)
    b.add_heading(doc, "9.1 Prinsip Pengendalian", 2)
    b.add_body(doc, "Deployment adalah perubahan terencana terhadap aplikasi, konfigurasi, skema data, atau layanan pendukung. Pengendalian bertujuan menjaga kerahasiaan, integritas, ketersediaan, keterlacakan, dan kemampuan pulih tanpa menghambat perbaikan yang sah.")
    add_figure(doc, "siklus-rilis.png", "Gambar 9-1. Siklus rilis terkendali dan jalur saat gate gagal")
    table(doc, "Tabel 9-1. Prinsip wajib", ["Prinsip", "Penerapan SIVMI CRM", "Bukti"], [
        ["Artefak tunggal", "Image ditag immutable; artefak yang diuji dipromosikan", "IMAGE_TAG/commit SHA"],
        ["Configuration externalized", "Environment dan secret tidak ditanam dalam image", "env/secret store"],
        ["Least privilege", "Akses deployment, DB, backup, WAHA dibatasi", "akses dan audit"],
        ["Explicit change", "Migration tidak otomatis saat restart", "catatan command/output"],
        ["Reversible", "Rollback image dan jalur restore tersedia", "runbook + drill"],
        ["Evidence based", "Test, build, smoke, observasi menjadi gate", "release record"],
        ["No production seeding", "Data demo tidak dibuat di produksi", "review command"],
        ["Data preservation", "Volume tidak dihapus ketika stop/release", "larangan down -v"],
    ], [1.35, 3.15, 2.0])
    callout(doc, "Larangan utama", "Jangan menjalankan docker compose down -v, migrate:fresh, db:wipe, atau seeding demo di produksi. Jangan menerapkan migration destruktif tanpa backup, review dampak, jalur pemulihan, dan approval eksplisit.", b.LIGHT_AMBER)

    b.add_heading(doc, "9.2 Peran dan Tanggung Jawab", 2)
    add_figure(doc, "kendali-peran.png", "Gambar 9-2. Pemisahan tanggung jawab dan catatan rilis tunggal")
    table(doc, "Tabel 9-2. Peran operasional", ["Peran", "Tanggung jawab utama", "Tidak boleh diasumsikan"], [
        ["Pemilik Bisnis/Direksi", "Prioritas, window, UAT, dampak, risk acceptance", "Menjalankan command teknis"],
        ["Release Manager", "Koordinasi, checklist, go/no-go, komunikasi, penutupan", "Menyetujui sendiri seluruh risiko"],
        ["Developer", "Kode, migration, test, release note, rollback teknis", "Mengubah data produksi tanpa change record"],
        ["DBA/Data Owner", "Review migration, backup, restore, rekonsiliasi", "Rollback schema otomatis selalu aman"],
        ["DevOps/Administrator", "VPS, Compose, TLS, secret, network, capacity", "Menilai kebenaran proses bisnis sendiri"],
        ["QA/UAT", "Test, smoke, regresi, visual lintas perangkat, bukti", "Mengabaikan scope RBAC/cabang"],
        ["Security/Incident Owner", "Triage, containment, secret rotation, evidence", "Mempublikasikan payload/PII mentah"],
        ["Admin Cabang/Staff", "Validasi operasional cabang dan pelaporan anomali", "Akses data cabang lain"],
    ], [1.4, 3.05, 2.05])

    b.add_heading(doc, "9.3 Matriks RACI Ringkas", 2)
    table(doc, "Tabel 9-3. RACI rilis dan pemeliharaan", ["Aktivitas", "A", "R", "C", "I"], [
        ["Scope dan window rilis", "Pemilik Bisnis", "Release Manager", "QA, DevOps", "Pengguna"],
        ["Build/test/image", "Lead Teknis", "Developer", "QA, Security", "Release Manager"],
        ["Review migration", "Data Owner", "Developer/DBA", "QA, Bisnis", "DevOps"],
        ["Backup sebelum rilis", "Data Owner", "DBA/DevOps", "Release Manager", "Bisnis"],
        ["Deploy dan konfigurasi", "Release Manager", "DevOps", "Developer", "QA/Bisnis"],
        ["Smoke test", "QA Lead", "QA + owner modul", "DevOps, Developer", "Release Manager"],
        ["Go/no-go", "Pemilik Bisnis", "Release Manager", "QA, Teknis, Data", "Pengguna"],
        ["Rollback/restore", "Incident Commander", "DevOps/DBA", "Bisnis, Security", "Stakeholder"],
        ["Pemeliharaan rutin", "Operation Owner", "DevOps/Administrator", "Developer, Security", "Bisnis"],
    ], [2.0, 1.1, 1.4, 1.15, 0.85])
    b.add_body(doc, "A = Accountable, R = Responsible, C = Consulted, I = Informed. Nama personel dan pengganti on-call harus dicatat pada runbook operasional; matriks ini mendefinisikan fungsi, bukan menunjuk individu permanen.")

    page(doc)
    b.add_heading(doc, "9.4 Lingkungan dan Baseline Teknologi", 2)
    add_figure(doc, "promosi-lingkungan.png", "Gambar 9-3. Pola promosi local, staging, dan production")
    table(doc, "Tabel 9-4. Baseline teknologi", ["Komponen", "Baseline", "Peran deployment"], [
        ["PHP/Laravel", "PHP 8.4 FPM; Laravel 12", "Runtime aplikasi dan Artisan"],
        ["UI", "Filament 5; Livewire 4; Tailwind; Vite 8", "Panel CRM dan asset production"],
        ["Web", "Nginx 1.27", "Reverse proxy internal, static asset, /up"],
        ["Database", "MySQL 8.4", "Data transaksi dan migration"],
        ["State/async", "Redis 7.4; queue worker; scheduler", "Cache/session/queue dan proses latar"],
        ["WhatsApp", "WAHA CORE, engine GOWS", "Session, pesan, QR, webhook"],
        ["Build", "Node.js 22; Composer production", "Asset dan dependency image"],
        ["Packaging", "Docker Engine + Compose", "Isolasi service dan volume"],
    ], [1.25, 2.15, 3.1])
    table(doc, "Tabel 9-5. Perbedaan lingkungan", ["Kontrol", "Local", "Staging", "Production"], [
        ["Tujuan", "Develop/debug", "Rehearsal/UAT", "Operasi nyata"],
        ["Source", "Mount host", "Image immutable", "Image immutable disetujui"],
        ["Data", "Seeder/demo", "Sintetis/anonymized", "Data nyata"],
        ["Debug", "Dapat aktif", "False", "False"],
        ["TLS", "Opsional localhost", "Wajib menyerupai prod", "Wajib HTTPS"],
        ["Secret", "Lokal unik", "Terpisah", "Terpisah dan dikelola"],
        ["WAHA", "Nomor uji", "Nomor uji/staging", "Nomor operasional"],
        ["Backup", "Opsional", "Uji restore", "Wajib dan termonitor"],
    ], [1.25, 1.65, 1.75, 1.85])
    callout(doc, "Status target", "Compose produksi telah tersedia, tetapi production readiness belum terbukti sampai deployment staging/VPS, TLS, firewall, backup/restore drill, monitoring, dan UAT berhasil. Jangan menulis status Production Ready hanya berdasarkan container lokal yang sehat.", b.LIGHT_AMBER)

    b.add_heading(doc, "9.5 Maintenance Mode", 2)
    b.add_body(doc, "Maintenance mode dipakai bila perubahan dapat menghasilkan skema campuran, write yang tidak kompatibel, rekonsiliasi data, atau outage terencana. Perubahan yang benar-benar backward-compatible dapat memakai rolling/recreate terkontrol, tetapi keputusan tersebut wajib tercatat.")
    table(doc, "Tabel 9-6. Keputusan maintenance", ["Kondisi", "Mode", "Kontrol"], [
        ["Perubahan UI/kode kompatibel tanpa migration", "Dapat tanpa maintenance", "Health + smoke cepat"],
        ["Additive migration kompatibel", "Bersyarat", "Deploy urutan expand/contract"],
        ["Rename/drop/transform besar", "Maintenance wajib", "Freeze write, backup, rehearsal"],
        ["Restore/repair data", "Maintenance wajib", "Approval dan rekonsiliasi"],
        ["Security incident", "Sesuai containment", "Batasi akses dan simpan bukti"],
        ["WAHA maintenance saja", "Aplikasi dapat tetap hidup", "Tandai kanal tidak tersedia; queue/retry"],
    ], [2.35, 1.35, 2.8])
    steps(doc, "Tabel 9-7. Prosedur maintenance mode", [
        "Umumkan window, dampak, owner, kanal komunikasi, dan estimasi pemulihan.",
        "Pastikan backup, rollback, akses break-glass, dan checklist tersedia.",
        "Aktifkan maintenance melalui container aplikasi: php artisan down --retry=60; gunakan secret bypass hanya bila benar-benar diperlukan dan jangan dipublikasikan.",
        "Pause atau awasi worker yang dapat menulis data; dokumentasikan job in-flight.",
        "Lakukan perubahan, migration, dan verifikasi teknis dari jaringan internal.",
        "Aktifkan aplikasi dengan php artisan up, lalu jalankan smoke dan observasi.",
        "Komunikasikan selesai/gagal, waktu aktual, dan tindakan lanjutan.",
    ])

    b.add_heading(doc, "9.6 Menjalankan Migrasi", 2)
    add_figure(doc, "urutan-maintenance-migrasi.png", "Gambar 9-4. Urutan maintenance, migration, aktivasi, dan pemulihan")
    b.add_heading(doc, "9.6.1 Pra-migrasi", 3)
    b.add_bullets(doc, [
        "Review setiap migration: operasi DDL/DML, index, constraint, nullable/default, volume data, lock, dan kompatibilitas versi aplikasi lama/baru.",
        "Pastikan backup berhasil, ruang disk cukup, image tag lama tersedia, serta restore path diuji.",
        "Jalankan migration pada salinan data staging dengan engine MySQL yang setara; ukur durasi dan lock.",
        "Pastikan migration pending memang bagian release candidate; jangan mengandalkan migration lokal yang tidak masuk artefak.",
    ])
    table(doc, "Tabel 9-8. Perintah migration produksi", ["Tahap", "Perintah ringkas", "Gate"], [
        ["Validasi Compose", "docker compose --env-file docker/.env.production -f docker/compose.production.yml config --quiet", "Exit 0"],
        ["Start dependency", "docker compose ... up -d database redis", "Healthy"],
        ["Status migration", "docker compose ... run --rm app php artisan migrate:status", "Expected pending"],
        ["Eksekusi", "docker compose ... run --rm app php artisan migrate --force", "Exit 0 + output disimpan"],
        ["Start release", "docker compose ... up -d --remove-orphans", "Service sehat"],
    ], [1.25, 3.95, 1.3])
    callout(doc, "Prinsip rollback migration", "Jangan menganggap php artisan migrate:rollback sebagai pemulihan data. Migration yang menghapus/mengubah data dapat memerlukan restore dari backup sebelum rilis. Gunakan pola expand-contract untuk meminimalkan ketidakcocokan antarversi.", b.LIGHT_AMBER)

    b.add_heading(doc, "9.7 Cache, Storage, Queue, dan Scheduler", 2)
    add_figure(doc, "state-runtime.png", "Gambar 9-5. State persisten dan proses latar yang harus diverifikasi")
    table(doc, "Tabel 9-9. Kendali state runtime", ["Komponen", "As-built", "Tindakan rilis/pemeliharaan"], [
        ["Config/route/view", "php artisan optimize pada start production", "Clear/rebuild saat config berubah; jangan cache env salah"],
        ["Application cache", "Store dapat database/Redis via env", "Purge hanya key/scope relevan; hindari flush global"],
        ["Session", "Database/Redis via env; secure cookie wajib prod", "Perubahan APP_KEY mengakhiri/dekripsi session"],
        ["Storage", "storage_data; public link dibuat entrypoint", "Backup file; verifikasi /storage dan permission"],
        ["Queue", "queue:work, tries 3, timeout 90, max-time 3600", "Pantau failed_jobs, age, retry; restart worker setelah code"],
        ["Scheduler", "schedule:work", "Verifikasi process dan jadwal"],
        ["Reminder", "crm:send-reminders hourly, no overlap", "Cek timestamp terakhir dan notifikasi"],
        ["WAHA sync", "crm:sync-whatsapp-sessions / 5 menit", "Cek status session, no overlap, background"],
        ["WAHA session", "Volume waha_sessions", "Jangan hapus; backup sesuai kebutuhan risiko"],
    ], [1.35, 2.55, 2.6])
    table(doc, "Tabel 9-10. Pemeriksaan proses latar", ["Pemeriksaan", "Contoh", "Respons gagal"], [
        ["Container", "compose ps; logs queue scheduler", "Restart terkontrol dan root cause"],
        ["Failed jobs", "php artisan queue:failed", "Review payload aman; retry selektif"],
        ["Schedule", "php artisan schedule:list", "Cek timezone, lock, process"],
        ["Backlog", "queue depth/age", "Scale worker atau perbaiki job lambat"],
        ["Idempotensi", "Webhook/reminder tidak duplikat", "Dedupe key dan transaction boundary"],
    ], [1.4, 2.4, 2.7])

    page(doc)
    b.add_heading(doc, "9.8 Smoke Test Pasca-Deployment", 2)
    add_figure(doc, "smoke-observasi.png", "Gambar 9-6. Smoke test berlapis dan observation gate")
    table(doc, "Tabel 9-11. Checklist smoke wajib", ["Area", "Skenario", "Hasil yang diharapkan"], [
        ["Platform", "Compose ps, /up, HTTPS, asset", "Healthy; tanpa Vite dev server"],
        ["Login", "Akun produksi non-demo", "Masuk sekali; session/cookie aman"],
        ["RBAC", "Superadmin, Admin Cabang, Staff, Direksi", "Menu dan aksi sesuai role"],
        ["Cabang", "Bandung, JakSel, JakPus", "Scope baca/tulis tidak bocor"],
        ["Dashboard", "Filter cabang/tahun/bulan/user", "KPI dan grafik mengikuti filter"],
        ["Lead", "Buat/ubah lead uji", "Status/source/program/PIC tersimpan"],
        ["CSV", "Import kecil dan export", "Validasi, scope, formula safety"],
        ["Transaksi", "Pendaftaran dan pembayaran uji", "Status dan audit konsisten"],
        ["Async", "Notifikasi, reminder, email", "Queue/scheduler memproses"],
        ["WAHA", "Session WORKING; pesan masuk/keluar", "Inbox, matching lead, ACK"],
        ["Storage", "Upload dan buka /storage", "File tersedia dan permission benar"],
        ["UI", "Mobile, tablet, laptop, desktop", "Responsif; tanpa console error blocker"],
    ], [1.25, 2.65, 2.6])
    callout(doc, "Data uji produksi", "Gunakan record bertanda jelas, seminimal mungkin, dan hapus/arsipkan sesuai prosedur setelah verifikasi. Jangan menggunakan identitas siswa nyata yang tidak diperlukan untuk smoke test.")

    b.add_heading(doc, "9.9 Observasi Setelah Release", 2)
    b.add_body(doc, "Release tidak selesai saat container berstatus running. Periode observasi dimulai setelah smoke lulus dan berakhir ketika indikator teknis, integrasi, data, serta pengalaman pengguna stabil sesuai risiko perubahan.")
    table(doc, "Tabel 9-12. Observation window", ["Periode", "Fokus", "Keputusan"], [
        ["0-15 menit", "Health, 5xx, login, asset, DB, restart loop", "Rollback cepat bila layanan inti gagal"],
        ["15-60 menit", "Queue age, failed job, scheduler, WAHA, latency", "Contain/scale/fix atau lanjut"],
        ["Hari kerja pertama", "Data reconciliation, KPI, import/export, laporan, keluhan", "Tutup release atau incident"],
        ["Siklus terjadwal", "Reminder hourly dan WAHA sync 5 menit", "Bukti jadwal benar-benar berjalan"],
    ], [1.4, 3.45, 1.65])
    table(doc, "Tabel 9-13. Catatan observasi", ["Field", "Isi minimum"], [
        ["Release", "image tag, commit, migration batch, waktu WIB/UTC"],
        ["Baseline", "request/error/latency, queue, CPU/RAM/disk, DB connections"],
        ["Business signal", "lead/chat/pendaftaran/pembayaran yang masuk"],
        ["Anomali", "waktu, scope, severity, owner, evidence tersanitasi"],
        ["Keputusan", "continue, contain, rollback, atau extend observation"],
        ["Penutupan", "approver, waktu, follow-up, baseline baru"],
    ], [1.35, 5.15])

    b.add_heading(doc, "9.10 Migrasi Basis Data dan Data", 2)
    b.add_body(doc, "Migrasi basis data mengubah struktur dan/atau isi data. Migration schema Laravel, backfill, normalisasi nomor WhatsApp, import CSV, perubahan enum, dan rekonsiliasi relasi harus diperlakukan sebagai change terukur dengan validasi sebelum dan sesudah.")
    table(doc, "Tabel 9-14. Klasifikasi perubahan data", ["Jenis", "Contoh SIVMI", "Kontrol minimum"], [
        ["Additive schema", "Kolom nullable/index baru", "Backward compatible + index review"],
        ["Constraint", "Unique/foreign key/enum", "Duplicate/orphan scan sebelum apply"],
        ["Transform", "Normalisasi status/source/nomor WA", "Mapping, dry-run, count, exception file"],
        ["Backfill", "Mengisi field baru dari data lama", "Batch, idempotent, resume, metric"],
        ["Import", "Leads/sekolah CSV", "Template, validation, branch scope, formula safety"],
        ["Delete/archive", "Data uji atau retention", "Approval, retention, audit, backup"],
    ], [1.35, 2.4, 2.75])
    steps(doc, "Tabel 9-15. Prosedur migrasi data", [
        "Definisikan sumber, pemilik, mapping, key pencocokan, scope cabang, dan acceptance criteria.",
        "Profilkan jumlah, null, duplikasi, orphan, encoding, tanggal, telepon, enum, dan referensi program/sekolah.",
        "Buat backup dan dry-run pada salinan data; simpan count sebelum/sesudah serta exception.",
        "Gunakan transaksi/batch/idempotency yang sesuai; jangan memuat seluruh data besar ke memory.",
        "Validasi constraint, aggregate dashboard, sampling record, RBAC, dan audit trail.",
        "Rekonsiliasi dengan data owner; arsipkan mapping, checksum, error, dan approval.",
    ])
    callout(doc, "Duplikasi lead WhatsApp", "Normalisasi nomor dan pencocokan lintas pesan harus menggunakan aturan aplikasi yang sama. Data yang sudah ada diberi indikator sumber/duplikat sesuai rule; jangan membuat lead baru hanya karena format nomor, LID, atau SenderAlt berbeda.", b.LIGHT_GREEN)

    b.add_heading(doc, "9.11 Rollback dan Pemulihan Release", 2)
    add_figure(doc, "backup-rollback.png", "Gambar 9-7. Pemisahan rollback aplikasi dan restore data")
    table(doc, "Tabel 9-16. Pemilihan tindakan pemulihan", ["Kondisi", "Tindakan utama", "Validasi"], [
        ["Kode/config gagal; schema kompatibel", "Kembalikan IMAGE_TAG lama", "Health + smoke + queue"],
        ["Migration additive; aplikasi lama kompatibel", "Rollback image; schema dapat tetap", "Fungsi lama dan data"],
        ["Migration destruktif/tidak kompatibel", "Maintenance + restore backup", "Integrity + reconciliation"],
        ["Secret bocor", "Revoke/rotate dan restart terkontrol", "Akses lama ditolak"],
        ["Queue menghasilkan duplikasi", "Pause worker; contain; reconcile", "Dedupe dan count"],
        ["WAHA session bermasalah", "Isolasi session/reconnect sesuai runbook", "WORKING + pesan uji"],
    ], [2.15, 2.35, 2.0])
    steps(doc, "Tabel 9-17. Urutan rollback aplikasi", [
        "Nyatakan incident, hentikan perubahan tambahan, dan lindungi bukti.",
        "Tentukan apakah masalah hanya kode/config atau juga data/schema.",
        "Aktifkan maintenance bila write berisiko dan pause worker yang relevan.",
        "Set IMAGE_TAG ke artefak sebelumnya yang telah diketahui baik lalu jalankan Compose secara terkontrol.",
        "Jika data tidak kompatibel, lakukan restore sesuai approval dan runbook teruji.",
        "Jalankan health, migration status, smoke, data reconciliation, dan observasi.",
        "Komunikasikan pemulihan dan buka corrective action/RCA.",
    ])

    b.add_heading(doc, "9.12 Backup dan Recovery Strategy", 2)
    b.add_body(doc, "Backup dinilai berhasil hanya bila dapat direstore dan hasilnya tervalidasi. Cakupan minimal meliputi MySQL, storage aplikasi, konfigurasi non-secret yang diperlukan, inventory secret/referensi secret, image tag, migration version, dan session WAHA sesuai penilaian risiko.")
    table(doc, "Tabel 9-18. Cakupan backup", ["Objek", "Metode target", "Validasi"], [
        ["MySQL", "Logical dump konsisten dan/atau snapshot volume", "Restore, row count, integrity"],
        ["Storage", "Snapshot/copy storage_data", "Checksum dan file sampling"],
        ["Redis", "AOF/snapshot sesuai kebutuhan", "Start dan state kritis; bukan sumber utama transaksi"],
        ["WAHA session", "Volume waha_sessions sesuai dukungan/risiko", "Session dapat dipulihkan atau pairing ulang"],
        ["Image", "Registry dengan tag immutable", "Digest dan pull test"],
        ["Config", "Versioned template + secret reference", "Recreate environment tanpa membocorkan secret"],
        ["Dokumen operasi", "Runbook, release record, inventory", "Akses saat insiden"],
    ], [1.4, 3.0, 2.1])
    table(doc, "Tabel 9-19. Usulan retensi dan pemulihan", ["Kelas", "Usulan", "Catatan"], [
        ["Pre-release", "Sebelum setiap migration/perubahan data", "Wajib"],
        ["Harian", "7-14 salinan", "Sesuaikan kapasitas dan regulasi"],
        ["Mingguan", "4-8 salinan", "Media/lokasi terpisah"],
        ["Bulanan", "6-12 salinan", "Untuk kebutuhan historis"],
        ["RPO", "Usulan <= 24 jam; lebih kecil untuk transaksi penting", "Belum disahkan/diuji"],
        ["RTO", "Usulan <= 4 jam untuk layanan inti", "Belum disahkan/diuji"],
        ["Restore drill", "Minimal triwulanan dan sebelum perubahan besar", "Catat RPO/RTO aktual"],
    ], [1.4, 2.75, 2.35])
    callout(doc, "RPO dan RTO", "Nilai pada tabel adalah proposal awal, bukan SLA. Direksi, pemilik bisnis, data owner, dan operasi harus menyetujui nilai berdasarkan dampak bisnis; kemudian nilainya dibuktikan melalui restore drill.", b.LIGHT_AMBER)

    b.add_heading(doc, "9.13 Pemeliharaan Sistem", 2)
    table(doc, "Tabel 9-20. Kalender pemeliharaan", ["Frekuensi", "Aktivitas", "Bukti"], [
        ["Harian", "Health, failed jobs, disk, backup status, WAHA session, error kritis", "Dashboard/checklist"],
        ["Mingguan", "Review log, queue trend, scheduler, user/cabang, data exception", "Operation report"],
        ["Bulanan", "Patch dependency/image, kapasitas, akun, secret age, restore sample", "Maintenance record"],
        ["Triwulanan", "Restore drill, access review, incident exercise, performance baseline", "Drill report"],
        ["Tahunan", "Architecture/security review, retention, SLA/RPO/RTO, DR plan", "Approval dan roadmap"],
        ["Ad hoc", "Critical CVE, outage, perubahan API WAHA, sertifikat", "Emergency change record"],
    ], [1.25, 3.8, 1.45])
    table(doc, "Tabel 9-21. Pemeliharaan aplikasi dan data", ["Area", "Aktivitas"], [
        ["Dependency", "Review Composer/npm/image; test kompatibilitas; scan; update bertahap"],
        ["Database", "Index/slow query, growth, orphan/duplicate, migration history, backup"],
        ["Queue/scheduler", "Age, failure, overlap lock, retry policy, heartbeat"],
        ["WAHA/GOWS", "Version compatibility, session state, webhook delivery, HMAC, pairing plan"],
        ["Storage", "Growth, permission, orphan file, retention, restore sampling"],
        ["RBAC", "User aktif, role, cabang, least privilege, akun tidak terpakai"],
        ["Dokumentasi", "Runbook, diagram, env template, test baseline, release record"],
    ], [1.5, 5.0])

    b.add_heading(doc, "9.14 Monitoring, Performa, dan Kapasitas", 2)
    add_figure(doc, "model-operasi.png", "Gambar 9-8. Model monitoring, triage, keamanan, dan improvement")
    table(doc, "Tabel 9-22. Indikator minimum", ["Domain", "Indikator", "Aksi saat anomali"], [
        ["Availability", "Uptime, /up, restart count, HTTP 5xx", "Triage service dan dependency"],
        ["Performance", "p50/p95 latency, slow query, build/runtime", "Profiling dan query/index review"],
        ["Capacity", "CPU, memory, disk, inode, DB size/connections", "Scale/cleanup terkontrol"],
        ["Queue", "Depth, oldest age, failed rate, processing time", "Tambah worker/perbaiki job"],
        ["Scheduler", "Last success, duration, overlap", "Periksa lock/timezone/process"],
        ["WAHA", "Health, session state, webhook lag/error", "Reconnect/contain/API review"],
        ["Database", "Lock, connection, replication/backup status", "DBA response"],
        ["Business", "Lead/chat/pendaftaran/pembayaran lag", "Reconcile pipeline"],
    ], [1.25, 3.05, 2.2])
    callout(doc, "Threshold berbasis baseline", "Angka threshold tidak boleh dibuat sembarang. Kumpulkan baseline staging dan periode lokal representatif, tetapkan SLO bersama pemilik bisnis, lalu uji alert agar tidak terlalu sunyi atau bising.")

    b.add_heading(doc, "9.15 Pemeliharaan Keamanan", 2)
    table(doc, "Tabel 9-23. Kontrol keamanan operasional", ["Kontrol", "Penerapan/target", "Frekuensi"], [
        ["HTTPS", "TLS pada reverse proxy/load balancer; secure cookie", "Kontinu/renewal"],
        ["Secret", "APP_KEY, DB, Redis, SMTP, WAHA API/HMAC unik; secret store", "Review/rotate"],
        ["Webhook", "HMAC wajib; request tanpa signature ditolak; replay/dedupe", "Setiap request"],
        ["Network", "Expose hanya port perlu; DB/Redis/WAHA internal", "Review firewall"],
        ["Image/dependency", "Pin/tag, scan CVE, patch teruji", "Setiap build/bulanan"],
        ["RBAC/cabang", "Least privilege, access review, audit aktivitas", "Bulanan/triwulanan"],
        ["Logging", "Redaction secret/PII; retensi; akses terbatas", "Kontinu"],
        ["Backup", "Enkripsi, lokasi terpisah, akses dan restore drill", "Setiap backup/drill"],
        ["Hardening", "APP_DEBUG=false, dashboard Swagger WAHA off, header Nginx", "Setiap deploy"],
    ], [1.3, 3.75, 1.45])
    b.add_heading(doc, "9.15.1 Rotasi Secret", 3)
    steps(doc, "Tabel 9-24. Prosedur rotasi", [
        "Identifikasi consumer, owner, expiry, dan blast radius secret.",
        "Buat secret baru melalui kanal aman; jangan kirim di chat/tiket/log.",
        "Perbarui provider dan environment dengan strategi overlap bila didukung.",
        "Restart/reload service terkontrol dan jalankan health serta alur integrasi.",
        "Cabut secret lama, verifikasi penolakan, dan catat waktu/owner.",
        "Periksa log untuk penggunaan lama dan tutup change record.",
    ])

    b.add_heading(doc, "9.16 Penanganan Insiden dan Pemeliharaan Darurat", 2)
    table(doc, "Tabel 9-25. Tahapan insiden", ["Tahap", "Aktivitas", "Output"], [
        ["Detect", "Alert/laporan, waktu, environment, release", "Incident ID"],
        ["Triage", "Severity, scope, data/security, owner", "Prioritas"],
        ["Contain", "Maintenance, pause worker, block/revoke, isolate", "Dampak dibatasi"],
        ["Recover", "Rollback, restart, restore, reconnect", "Layanan minimum"],
        ["Reconcile", "Cek data, job, chat, transaksi", "Integritas terkonfirmasi"],
        ["Observe", "Metric/log/user confirmation", "Stabilitas"],
        ["Review", "RCA, corrective/preventive action", "Pembelajaran"],
    ], [1.1, 3.4, 2.0])
    table(doc, "Tabel 9-26. Emergency change", ["Syarat", "Ketentuan"], [
        ["Alasan", "Risiko menunggu lebih besar daripada risiko perubahan"],
        ["Approval", "Incident commander dan owner berwenang; dicatat walau dipercepat"],
        ["Scope", "Perubahan minimum untuk containment/recovery"],
        ["Bukti", "Timeline, command, image/config, backup, hasil"],
        ["Verifikasi", "Targeted test/smoke dan observasi sesuai keadaan"],
        ["Pasca-insiden", "Review lengkap, test/regression, dokumentasi, permanent fix"],
    ], [1.45, 5.05])
    callout(doc, "Keamanan dan privasi", "Saat insiden, jangan menyalin docker/.env, token, cookie, kredensial, file CSV, atau payload chat mentah ke kanal umum. Gunakan ID internal, redaction, akses least privilege, dan penyimpanan bukti yang disetujui.", b.LIGHT_AMBER)

    b.add_heading(doc, "9.17 Rekomendasi Pengembangan Operasional", 2)
    table(doc, "Tabel 9-27. Roadmap menuju VPS", ["Prioritas", "Rekomendasi", "Kriteria selesai"], [
        ["P0", "Staging full Docker production-like", "Image yang sama, TLS, data sintetis, smoke/UAT"],
        ["P0", "CI test, build, image, scan, artifact", "Merge/release diblokir bila gate gagal"],
        ["P0", "Registry dan IMAGE_TAG immutable", "Digest tercatat dan rollback pull-tested"],
        ["P0", "Backup otomatis dan restore drill", "RPO/RTO aktual serta evidence"],
        ["P0", "TLS, firewall, secret management", "DB/Redis/WAHA tidak terekspos publik"],
        ["P0", "Central logging dan alert", "5xx, queue, disk, DB, WAHA memiliki owner"],
        ["P1", "Zero/minimal-downtime migration pattern", "Expand-contract dan rehearsal"],
        ["P1", "Queue/scheduler heartbeat", "Missed run dan backlog menghasilkan alert"],
        ["P1", "Capacity/SLO baseline", "Threshold disahkan berbasis data"],
        ["P1", "Automated smoke dan synthetic check", "Alur kritis diuji pascadeploy"],
        ["P2", "Object storage/offsite backup", "Lifecycle, encryption, restore"],
        ["P2", "High availability bila dibutuhkan", "DB/Redis managed dan multi-instance plan"],
    ], [0.75, 3.15, 2.6])
    b.add_heading(doc, "9.17.1 Urutan Implementasi", 3)
    steps(doc, "Tabel 9-28. Tahap operasional", [
        "Bekukan baseline versi, test, build, diagram, dan dependency.",
        "Siapkan VPS staging, DNS/TLS test, secret terpisah, registry, dan Compose production.",
        "Otomatiskan CI serta hasil test/build/scan; deploy staging dengan image immutable.",
        "Uji migration, smoke, WAHA, backup, restore, rollback, dan visual UAT.",
        "Tetapkan RACI, window, on-call, severity, komunikasi, RPO/RTO, serta risk acceptance.",
        "Siapkan VPS production, firewall, TLS, logging/monitoring, backup offsite, dan runbook.",
        "Lakukan rilis terkontrol, observation window, reconciliation, dan post-release review.",
    ])

    page(doc)
    b.add_heading(doc, "9.18 Kriteria Dokumen Dinyatakan Diterapkan", 2)
    table(doc, "Tabel 9-29. Definition of Implemented", ["Kriteria", "Bukti minimum", "Status baseline"], [
        ["Owner dan RACI disahkan", "Nama/peran/pengganti", "Target"],
        ["Environment terpisah", "Local, staging, production inventory", "Local as-built; lainnya target"],
        ["Image immutable", "Tag/digest dan registry", "Target"],
        ["Test/build/scan gate", "CI artifact dan exit code", "Test/build lokal as-built; CI/scan target"],
        ["Migration rehearsal", "Output staging, durasi, rollback", "Target"],
        ["Backup/restore", "Backup sukses dan restore drill", "Prosedur ada; drill target"],
        ["Smoke/UAT", "Checklist dan approval lintas role/device", "Automated sebagian; VPS target"],
        ["Observability", "Health, metric, log, alert, on-call", "Health/log dasar; central alert target"],
        ["Security", "TLS, firewall, secret rotation, webhook HMAC", "App guard as-built; infra target"],
        ["Incident/rollback", "Runbook dan simulasi", "Runbook as-built; drill target"],
        ["Dokumentasi", "Versi, diagram, release record, review date", "Dokumen ini tersedia"],
        ["Persetujuan", "Bisnis, teknis, data, operasi, security", "Target"],
    ], [2.4, 2.65, 1.45])
    callout(doc, "Pernyataan penerapan", "Dokumen hanya boleh diberi status Diterapkan setelah seluruh kontrol wajib memiliki owner, bukti, hasil uji, dan persetujuan. Keberadaan file Compose atau dokumen saja tidak membuktikan kesiapan produksi.", b.LIGHT_GREEN)

    b.add_heading(doc, "LAMPIRAN A - CHECKLIST RELEASE", 1)
    table(doc, "Tabel A-1. Checklist ringkas sebelum, saat, dan sesudah rilis", ["Fase", "Checklist"], [
        ["Sebelum", "Scope/owner/window; test/build/scan; image tag; env; migration review; backup; rollback; komunikasi"],
        ["Saat", "Maintenance bila perlu; status service; migrate --force; start; health; command/output tercatat"],
        ["Smoke", "Login; RBAC/cabang; dashboard; lead/CSV; transaksi; queue/scheduler; WAHA; storage; responsive UI"],
        ["Observasi", "5xx/latency; queue/failed job; DB/disk; scheduler; WAHA; data/business signal; user report"],
        ["Penutupan", "Go/rollback result; reconciliation; approval; incident/follow-up; baseline dan dokumentasi diperbarui"],
    ], [1.15, 5.35])

    b.add_heading(doc, "LAMPIRAN B - TEMPLATE CATATAN RILIS", 1)
    table(doc, "Tabel B-1. Field catatan rilis", ["Field", "Isi"], [
        ["Release ID / versi / image digest", ""],
        ["Environment / domain / waktu", ""],
        ["Scope / issue / migration", ""],
        ["Owner / RACI / approval", ""],
        ["Test / build / scan", ""],
        ["Backup / restore point", ""],
        ["Command deployment", ""],
        ["Smoke / UAT", ""],
        ["Observasi / metric / anomali", ""],
        ["Rollback / incident", ""],
        ["Data reconciliation", ""],
        ["Keputusan penutupan / follow-up", ""],
    ], [2.35, 4.15])

    b.add_heading(doc, "LAMPIRAN C - REGISTER RISIKO OPERASIONAL", 1)
    table(doc, "Tabel C-1. Risiko dan kontrol utama", ["ID", "Risiko", "Dampak", "Kontrol"], [
        ["O-01", "Migration tanpa backup/rehearsal", "Kehilangan data/outage", "Review, backup, staging, restore path"],
        ["O-02", "Image latest berubah", "Rollback tidak deterministik", "Tag/digest immutable"],
        ["O-03", "Secret development dipakai prod", "Kompromi akses", "Secret unik, guard, rotation"],
        ["O-04", "Volume terhapus", "Data/session hilang", "Larangan down -v, backup"],
        ["O-05", "Queue/scheduler tidak terlihat", "Pesan/reminder tertunda", "Heartbeat, age, alert"],
        ["O-06", "WAHA session gagal", "Kanal WhatsApp berhenti", "Persist session, health, reconnect runbook"],
        ["O-07", "Scope cabang bocor", "Kebocoran data", "RBAC test dan smoke lintas role"],
        ["O-08", "Disk/log tumbuh", "Service berhenti", "Capacity alert dan retention"],
        ["O-09", "Backup tidak dapat direstore", "Pemulihan gagal", "Restore drill dan checksum"],
        ["O-10", "Rilis ditutup terlalu cepat", "Regresi terlambat", "Observation gate dan owner approval"],
    ], [0.65, 2.0, 1.75, 2.1])

    b.add_heading(doc, "PENUTUP", 1)
    b.add_body(doc, "BAB IX menetapkan bahwa deployment SIVMI CRM adalah proses bisnis dan teknis yang terkendali, bukan sekadar menjalankan container. Baseline full Docker lokal dan Compose produksi memberikan fondasi yang baik; kesiapan VPS harus dibuktikan melalui staging, image immutable, CI, migration rehearsal, backup/restore, keamanan jaringan, observability, smoke/UAT, serta keputusan pemilik yang dapat diaudit.")
    callout(doc, "Baseline final BAB IX", "Jangan menyatakan production ready sebelum semua gate pada 9.18 memiliki bukti. Setelah setiap rilis, perbarui baseline versi, test, migration, kapasitas, risiko, runbook, diagram, dan catatan pemeliharaan.", b.LIGHT_GREEN)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(f"Created: {OUTPUT}")


if __name__ == "__main__":
    build_document()
