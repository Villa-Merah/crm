from __future__ import annotations

from pathlib import Path
import sys

PROJECT_ROOT = Path(__file__).resolve().parents[2]
DOC_TOOLS = PROJECT_ROOT / "storage" / "app" / "doc-tools"
sys.path.insert(0, str(DOC_TOOLS))

from PIL import Image, ImageDraw
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

import build_bab_iv_arsitektur_bisnis as b
import build_bab_vii_arsitektur_teknologi as v7


OUTPUT = PROJECT_ROOT / "docs" / "BAB XI - Panduan Administrator SIVMI CRM v1.0.docx"
ASSETS = PROJECT_ROOT / "docs" / "_bab_xi_assets"
COVER = ASSETS / "cover-bab-xi.png"

PRIMARY = "#1476B8"
PRIMARY_DARK = "#075985"
INK = "#0B2545"
MUTED = "#64748B"
BG = "#F6F8FC"
GREEN = "#10B981"
AMBER = "#F59E0B"
ROSE = "#F43F5E"
PURPLE = "#A855F7"


def font(name: str, size: int):
    return v7.font(name, size)


def canvas(title: str, subtitle: str, size=(1800, 1080)):
    return v7.canvas(title, subtitle, size)


def box(draw, xy, title, lines, fill="#F8FAFC", outline=PRIMARY, title_size=27, body_size=20):
    return v7.box(draw, xy, title, lines, fill, outline, title_size, body_size)


def arrow(draw, start, end, color="#64748B", width=5):
    return v7.arrow(draw, start, end, color, width)


def make_cover() -> None:
    ASSETS.mkdir(parents=True, exist_ok=True)
    image = Image.new("RGB", (1600, 2200), BG)
    draw = ImageDraw.Draw(image)
    for y in range(2200):
        ratio = y / 2200
        start, end = (20, 118, 184), (7, 89, 133)
        draw.line((0, y, 1600, y), fill=tuple(round(start[i] * (1 - ratio) + end[i] * ratio) for i in range(3)))
    draw.ellipse((1010, -180, 1770, 580), fill=(27, 139, 199))
    draw.ellipse((1140, 140, 1640, 640), fill=(8, 103, 151))
    draw.rounded_rectangle((120, 250, 1480, 1910), radius=44, fill="white")
    draw.text((190, 350), "SIVMI CRM", font=font("arialbd.ttf", 34), fill=PRIMARY)
    draw.text((190, 610), "BAB XI", font=font("arialbd.ttf", 48), fill=PRIMARY_DARK)
    draw.multiline_text((190, 700), "PANDUAN\nADMINISTRATOR", font=font("arialbd.ttf", 82), fill=INK, spacing=18)
    draw.text((190, 1045), "Instalasi, konfigurasi, operasi, keamanan,", font=font("arial.ttf", 31), fill="#475569")
    draw.text((190, 1097), "backup, monitoring, WAHA, dan pemulihan", font=font("arial.ttf", 31), fill="#475569")
    draw.rounded_rectangle((190, 1245, 805, 1325), radius=40, fill="#EAF4FB")
    draw.text((230, 1267), "BASELINE APLIKASI 25 AGUSTUS 2026", font=font("arial.ttf", 28), fill=PRIMARY_DARK)
    draw.text((190, 1610), "Versi dokumen 1.0", font=font("arial.ttf", 28), fill=MUTED)
    draw.text((190, 1665), "Sasaran: administrator aplikasi dan infrastruktur", font=font("arial.ttf", 28), fill=MUTED)
    draw.text((190, 1775), "SIVMI CRM - Catat, Kelola, Closing", font=font("arialbd.ttf", 34), fill=PRIMARY)
    image.save(COVER)


def make_admin_map() -> None:
    image, draw = canvas("Peta Tanggung Jawab Administrator", "Administrasi aplikasi, data, integrasi, dan infrastruktur dipisahkan tetapi terkoordinasi")
    box(draw, (650, 100, 1150, 300), "ADMINISTRATOR", "Kontrol akses | Operasi | Bukti", "#EAF4FB", PRIMARY_DARK, 31, 22)
    items = [
        (70, 500, 430, 880, "APLIKASI", "akun & role\nmaster data\nkonfigurasi\naudit", PRIMARY),
        (500, 500, 860, 880, "DATA", "migration\nbackup/restore\nretensi\nrekonsiliasi", PURPLE),
        (930, 500, 1290, 880, "INTEGRASI", "WAHA/GOWS\nwebhook HMAC\nqueue\nscheduler", GREEN),
        (1360, 500, 1730, 880, "PLATFORM", "Docker\nTLS/firewall\nlog/health\nkapasitas", AMBER),
    ]
    for x1, y1, x2, y2, title, lines, color in items:
        box(draw, (x1, y1, x2, y2), title, lines, "#F8FAFC", color, 28, 20)
        arrow(draw, (900, 300), ((x1 + x2) // 2, y1), "#94A3B8", 4)
    image.save(ASSETS / "peta-tanggung-jawab.png")


def make_topology() -> None:
    image, draw = canvas("Topologi Full Docker", "Laragon tidak diperlukan pada baseline utama; seluruh layanan berjalan dalam satu Compose")
    stages = [
        (55, 230, 330, 560, "BROWSER", "localhost:8000\nlocalhost:3000", PRIMARY),
        (405, 230, 680, 560, "WEB", "Nginx 1.27\n/up | /storage", PRIMARY_DARK),
        (755, 230, 1030, 560, "APP", "PHP 8.4\nLaravel 12\nFilament 5", PURPLE),
        (1105, 150, 1390, 430, "DATA", "MySQL 8.4\nRedis 7.4", AMBER),
        (1105, 500, 1390, 780, "WORKER", "Queue\nScheduler", GREEN),
        (1465, 230, 1745, 560, "WAHA", "CORE / GOWS\nSession WA", ROSE),
    ]
    for x1, y1, x2, y2, title, lines, color in stages:
        box(draw, (x1, y1, x2, y2), title, lines, "#F8FAFC", color, 27, 19)
    arrow(draw, (330, 395), (405, 395)); arrow(draw, (680, 395), (755, 395))
    arrow(draw, (1030, 340), (1105, 290)); arrow(draw, (1030, 470), (1105, 610))
    arrow(draw, (1390, 395), (1465, 395))
    draw.text((900, 920), "Volume persisten: mysql_data | redis_data | storage_data (production) | waha_sessions", anchor="mm", font=font("arialbd.ttf", 23), fill=MUTED)
    image.save(ASSETS / "topologi-full-docker.png")


def make_install_flow() -> None:
    image, draw = canvas("Alur Instalasi Pertama", "Script start-full menyiapkan environment lokal, membangun image, dan menunggu health check")
    labels = [
        (60, 270, 340, 620, "1. PRASYARAT", "Docker Desktop\nLinux containers\nrepo lokal", PRIMARY),
        (410, 270, 690, 620, "2. START", "ExecutionPolicy\nstart-full.ps1", PURPLE),
        (760, 270, 1040, 620, "3. OTOMATIS", ".env lokal\nAPP_KEY\nWAHA secrets", AMBER),
        (1110, 270, 1390, 620, "4. BUILD", "image + volume\nmigration lokal", GREEN),
        (1460, 270, 1740, 620, "5. VALIDASI", "/up sehat\nWAHA /health\nlogin", PRIMARY_DARK),
    ]
    for x1, y1, x2, y2, title, lines, color in labels:
        box(draw, (x1, y1, x2, y2), title, lines, "#F8FAFC", color, 25, 19)
    for x in (340, 690, 1040, 1390): arrow(draw, (x, 445), (x + 70, 445))
    image.save(ASSETS / "alur-instalasi.png")


def make_secret_flow() -> None:
    image, draw = canvas("Kepemilikan Environment dan Secret", "Nilai secret tidak masuk Git, log, screenshot, atau dokumen operasional")
    box(draw, (80, 200, 500, 720), "DEVELOPMENT", "docker/.env\nAPP_DEBUG=true\nlocalhost\ncredential lokal unik", "#F8FAFC", PRIMARY, 29, 20)
    box(draw, (690, 200, 1110, 720), "SUMBER CONTOH", "docker/.env.example\ndocker/.env.production.example\ntanpa secret nyata", "#FFF7ED", AMBER, 29, 20)
    box(draw, (1300, 200, 1720, 720), "PRODUCTION", "secret manager / file terlindungi\nAPP_DEBUG=false\nHTTPS\ncredential unik", "#F8FAFC", GREEN, 29, 20)
    arrow(draw, (690, 460), (500, 460)); arrow(draw, (1110, 460), (1300, 460))
    draw.rounded_rectangle((400, 820, 1400, 970), radius=22, fill="#FFF1F2", outline=ROSE, width=3)
    draw.text((900, 895), "API key WAHA, HMAC webhook, APP_KEY, DB, Redis, dan SMTP harus dipisahkan.", anchor="mm", font=font("arialbd.ttf", 23), fill=INK)
    image.save(ASSETS / "environment-secret.png")


def make_data_flow() -> None:
    image, draw = canvas("Kontrol Database dan Storage", "Migration eksplisit di production; seeder hanya untuk database development kosong")
    nodes = [
        (70, 260, 370, 650, "BACKUP", "database\nstorage\nsession WAHA\nmanifest", PRIMARY),
        (470, 260, 770, 650, "MIGRATION", "review status\nmigrate --force\ncatat versi", PURPLE),
        (870, 260, 1170, 650, "VALIDASI", "schema\nlogin\nCRUD sampel\nrekonsiliasi", GREEN),
        (1270, 260, 1730, 650, "OPERASI", "queue/scheduler\nhealth check\nmonitoring\nretensi backup", AMBER),
    ]
    for x1, y1, x2, y2, title, lines, color in nodes:
        box(draw, (x1, y1, x2, y2), title, lines, "#F8FAFC", color, 28, 20)
    for x in (370, 770, 1170): arrow(draw, (x, 455), (x + 100, 455))
    draw.text((900, 850), "Larangan: db:seed di production dan docker compose down -v pada volume aktif.", anchor="mm", font=font("arialbd.ttf", 23), fill=ROSE)
    image.save(ASSETS / "kontrol-data.png")


def make_async_flow() -> None:
    image, draw = canvas("Queue, Scheduler, dan WAHA", "Pekerjaan asinkron harus hidup terus, termonitor, dan dapat dipulihkan")
    box(draw, (70, 300, 380, 650), "EVENT", "Webhook WAHA\nimport/export\nnotifikasi", "#F8FAFC", PRIMARY, 29, 20)
    box(draw, (480, 180, 790, 520), "QUEUE", "queue:work\nsleep=3 | tries=3\ntimeout=90", "#F8FAFC", PURPLE, 29, 20)
    box(draw, (480, 590, 790, 900), "SCHEDULER", "schedule:work\nreminder hourly\nsync WA 5 menit", "#F8FAFC", AMBER, 29, 20)
    box(draw, (920, 300, 1230, 650), "LARAVEL", "proses event\ndedupe\naudit\nnotification", "#F8FAFC", GREEN, 29, 20)
    box(draw, (1330, 300, 1730, 650), "HASIL", "Inbox WhatsApp\nlead/update\nstatus session\nfailed job", "#F8FAFC", PRIMARY_DARK, 29, 20)
    arrow(draw, (380, 450), (480, 350)); arrow(draw, (790, 350), (920, 420)); arrow(draw, (790, 745), (920, 530)); arrow(draw, (1230, 475), (1330, 475))
    image.save(ASSETS / "alur-asinkron.png")


def make_backup_flow() -> None:
    image, draw = canvas("Alur Backup, Restore, dan Rekonsiliasi", "Restore dinyatakan selesai hanya setelah data dan layanan diverifikasi")
    steps = [
        (60, 270, 330, 620, "INVENTARIS", "DB\nstorage\nWA session\nversi image", PRIMARY),
        (400, 270, 670, 620, "BACKUP", "dump konsisten\narsip volume\nenkripsi\nchecksum", PURPLE),
        (740, 270, 1010, 620, "SIMPAN", "lokasi terpisah\nretensi\nakses terbatas", AMBER),
        (1080, 270, 1350, 620, "RESTORE", "isolasi\nimport\npermission\nrestart", GREEN),
        (1420, 270, 1740, 620, "REKONSILIASI", "jumlah record\nfile/session\nsmoke test\nbukti", PRIMARY_DARK),
    ]
    for x1, y1, x2, y2, title, lines, color in steps:
        box(draw, (x1, y1, x2, y2), title, lines, "#F8FAFC", color, 24, 18)
    for x in (330, 670, 1010, 1350): arrow(draw, (x, 445), (x + 70, 445))
    image.save(ASSETS / "backup-restore.png")


def make_incident_flow() -> None:
    image, draw = canvas("Monitoring, Insiden, dan Maintenance", "Utamakan containment dan bukti; perubahan darurat tetap dicatat")
    nodes = [
        (60, 260, 350, 650, "DETEKSI", "health/log\nalert pengguna\ncapacity", PRIMARY),
        (430, 260, 720, 650, "TRIAGE", "scope\ndampak\nseverity\nownership", AMBER),
        (800, 260, 1090, 650, "KENDALIKAN", "maintenance\nisolasi\nrollback\nrestart aman", ROSE),
        (1170, 260, 1460, 650, "PULIHKAN", "health\nqueue\ndata\nsmoke test", GREEN),
        (1540, 260, 1750, 650, "TUTUP", "RCA\nbukti\nfollow-up", PRIMARY_DARK),
    ]
    for x1, y1, x2, y2, title, lines, color in nodes:
        box(draw, (x1, y1, x2, y2), title, lines, "#F8FAFC", color, 24, 18)
    for x in (350, 720, 1090, 1460): arrow(draw, (x, 455), (x + 80, 455))
    image.save(ASSETS / "monitoring-insiden.png")


def add_figure(doc: Document, name: str, caption: str, width=6.15) -> None:
    b.add_figure(doc, ASSETS / name, caption, width)


def table(doc: Document, caption: str, headers: list[str], rows: list[list[str]], widths: list[float]) -> None:
    b.add_table(doc, headers, rows, widths, caption)


def callout(doc: Document, title: str, body: str, fill: str = b.LIGHT_BLUE) -> None:
    b.add_callout(doc, title, body, fill)
    doc.tables[-1].rows[0]._tr.get_or_add_trPr().append(b.OxmlElement("w:cantSplit"))


def steps(doc: Document, caption: str, items: list[str]) -> None:
    table(doc, caption, ["Langkah", "Tindakan"], [[str(i), item] for i, item in enumerate(items, 1)], [0.75, 5.75])


def command(doc: Document, title: str, lines: list[str]) -> None:
    callout(doc, title, "\n".join(lines), b.LIGHT_BLUE)
    for paragraph in doc.tables[-1].cell(0, 0).paragraphs:
        for run in paragraph.runs:
            if not run.bold:
                run.font.name = "Consolas"
                run.font.size = Pt(8.5)


def page(doc: Document) -> None:
    doc.add_page_break()


def build_document() -> None:
    for fn in (make_cover, make_admin_map, make_topology, make_install_flow, make_secret_flow, make_data_flow, make_async_flow, make_backup_flow, make_incident_flow):
        fn()

    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.5); sec.page_height = Inches(11)
    sec.top_margin = Inches(0.85); sec.bottom_margin = Inches(0.85)
    sec.left_margin = Inches(1); sec.right_margin = Inches(1)
    sec.header_distance = Inches(0.4); sec.footer_distance = Inches(0.4)
    b.configure_styles(doc)
    sec.header.paragraphs[0].text = "SIVMI CRM  |  BAB XI: PANDUAN ADMINISTRATOR"
    for run in sec.header.paragraphs[0].runs:
        b.set_run_font(run, 8.5, True, b.PRIMARY)
    b.add_page_number(sec.footer.paragraphs[0])

    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(0)
    pic = p.add_run().add_picture(str(COVER), width=Inches(6.1))
    pic._inline.docPr.set("descr", "Sampul BAB XI Panduan Administrator SIVMI CRM")
    page(doc)

    b.add_heading(doc, "KENDALI DOKUMEN", 1)
    table(doc, "Tabel 0-1. Kendali dokumen", ["Atribut", "Nilai"], [
        ["Judul", "BAB XI: Panduan Administrator SIVMI CRM"],
        ["Versi", "1.0"],
        ["Tanggal baseline", "25 Agustus 2026"],
        ["Sasaran", "Superadmin, administrator aplikasi, administrator database, DevOps/infrastruktur, dan petugas integrasi WhatsApp"],
        ["Baseline utama", "Full Docker lokal dan target Docker Compose pada VPS; Laragon/hybrid hanya kompatibilitas legacy"],
        ["Teknologi", "PHP 8.4, Laravel 12, Filament 5, Livewire 4, Vite 8, Nginx, MySQL 8.4, Redis 7.4, queue, scheduler, WAHA CORE/GOWS"],
        ["Quality gate tercatat", "72 test, 448 assertion, dan build Vite production lulus pada baseline proyek"],
        ["Sumber kebenaran", "docker/compose*.yml, Dockerfile, entrypoint, Nginx config, environment example, routes, config Laravel, runbook deployment/WAHA, migration, test, dan kode aplikasi"],
    ], [1.55, 4.95])
    callout(doc, "Kedudukan dokumen", "Panduan ini menjelaskan pekerjaan administrator. Nilai credential, token, API key, HMAC secret, password, dan dump data tidak boleh dimasukkan ke dokumen, Git, tiket terbuka, atau screenshot.")
    b.add_heading(doc, "DAFTAR ISI", 1); b.add_toc(doc); page(doc)

    b.add_heading(doc, "RINGKASAN ADMINISTRATOR", 1)
    b.add_body(doc, "Administrator memastikan aplikasi tersedia, data terlindungi, proses asinkron berjalan, integrasi WAHA sehat, dan akses mengikuti role serta cabang. Kewenangan teknis tidak menghapus kebutuhan persetujuan bisnis: perubahan data master, role global, restore, migration produksi, dan pemutusan session WhatsApp harus memiliki pemilik keputusan dan bukti pelaksanaan.")
    add_figure(doc, "peta-tanggung-jawab.png", "Gambar 11-1. Peta tanggung jawab administrator SIVMI CRM")
    table(doc, "Tabel 0-2. Prinsip administrasi", ["Prinsip", "Penerapan"], [
        ["Least privilege", "Berikan role dan scope cabang minimum yang diperlukan"],
        ["Reversible", "Backup dan rencana rollback tersedia sebelum perubahan berisiko"],
        ["Repeatable", "Gunakan script, Compose, migration, dan checklist yang berversi"],
        ["Observable", "Health, log, queue, scheduler, kapasitas, dan audit dapat diperiksa"],
        ["Secret hygiene", "Secret unik per lingkungan dan tidak tersimpan pada repositori"],
        ["Evidence based", "Catat waktu, pelaksana, versi, hasil, dan bukti verifikasi"],
    ], [1.55, 4.95])

    b.add_heading(doc, "BAB XI: PANDUAN ADMINISTRATOR", 1)
    b.add_heading(doc, "11.1 Peran Administrator dan Ruang Lingkup", 2)
    table(doc, "Tabel 11-1. Pembagian peran", ["Peran", "Tanggung jawab", "Batas"], [
        ["Superadmin", "Akun, role, cabang, pengaturan, audit, koneksi WhatsApp", "Tidak membagikan akun/secret"],
        ["Admin aplikasi", "Konfigurasi Laravel/Filament, release, test, log aplikasi", "Tidak mengubah data bisnis tanpa otorisasi"],
        ["Admin database", "Backup, restore, migration, integritas, kapasitas", "Tidak menjalankan seeder produksi"],
        ["Admin platform", "Docker, Nginx, TLS, firewall, volume, monitoring", "Tidak menghapus volume aktif"],
        ["Admin WAHA", "API, session, webhook, HMAC, sinkronisasi", "Satu session aktif per pengguna"],
        ["Pemilik bisnis", "Persetujuan role global, data master, downtime, restore", "Bukan pengganti kontrol teknis"],
    ], [1.3, 3.2, 2.0])
    callout(doc, "Pemisahan tugas", "Untuk perubahan produksi berisiko tinggi, pelaksana dan pemeriksa sebaiknya berbeda. Minimal, bukti backup, migration, smoke test, dan persetujuan harus dapat ditelusuri melalui tiket atau catatan rilis.", b.LIGHT_AMBER)

    b.add_heading(doc, "11.2 Prasyarat dan Baseline Sistem", 2)
    add_figure(doc, "topologi-full-docker.png", "Gambar 11-2. Topologi layanan full Docker")
    table(doc, "Tabel 11-2. Baseline layanan", ["Layanan", "Fungsi", "Persistensi/health"], [
        ["web", "Nginx reverse proxy aplikasi", "GET /up; /storage alias"],
        ["app", "PHP-FPM dan Laravel", "Port internal 9000; storage/cache"],
        ["database", "MySQL 8.4", "mysql_data; mysqladmin ping"],
        ["redis", "Cache/queue/session production", "redis_data; PING"],
        ["node", "Vite development lokal", "node_modules; port 5173"],
        ["queue", "Job asinkron", "Proses harus terus hidup"],
        ["scheduler", "Jadwal Laravel", "schedule:work terus hidup"],
        ["waha", "Gateway WhatsApp GOWS", "waha_sessions; GET /health"],
    ], [1.2, 3.1, 2.2])
    table(doc, "Tabel 11-3. Perbedaan lingkungan", ["Kontrol", "Development full Docker", "Production/VPS"], [
        ["URL", "localhost:8000", "Domain HTTPS final"],
        ["Debug", "APP_DEBUG=true lokal", "APP_DEBUG=false wajib"],
        ["Aset", "Node/Vite dev server", "Aset hasil build pada image"],
        ["Data", "Volume lokal; boleh seeder awal", "Data nyata; tanpa seeder"],
        ["Session/cache/queue", "Database sesuai contoh lokal", "Redis dengan password"],
        ["WAHA dashboard", "Boleh aktif terbatas", "Dinonaktifkan secara default"],
        ["TLS/firewall", "Loopback/local", "Wajib di edge/reverse proxy"],
    ], [1.55, 2.45, 2.5])

    b.add_heading(doc, "11.3 Instalasi Development Full Docker", 2)
    add_figure(doc, "alur-instalasi.png", "Gambar 11-3. Alur instalasi pertama full Docker")
    steps(doc, "Tabel 11-4. Instalasi pertama pada Windows", [
        "Pastikan Docker Desktop aktif, Engine running, dan menggunakan Linux containers.",
        "Buka PowerShell pada root proyek C:\\laragon\\www\\sivmi.crm; lokasi folder boleh dipindah karena Laragon tidak digunakan sebagai runtime.",
        "Jalankan Set-ExecutionPolicy -Scope Process Bypass untuk sesi PowerShell tersebut bila script lokal diblokir.",
        "Jalankan .\\docker\\start-full.ps1. Script membuat docker/.env dari contoh bila belum ada serta membangkitkan APP_KEY dan secret WAHA lokal.",
        "Tunggu health check. Buka http://localhost:8000/admin/login dan dashboard WAHA development di http://localhost:3000/dashboard.",
        "Gunakan db:seed hanya pada database lokal yang benar-benar baru dan kosong, kemudian login dengan akun bootstrap yang ditetapkan tim.",
        "Jalankan .\\docker\\test-full.ps1 dan verifikasi halaman utama sebelum mulai pengembangan.",
    ])
    command(doc, "Perintah operasional lokal", [
        ".\\docker\\start-full.ps1", ".\\docker\\status-full.ps1", ".\\docker\\logs-full.ps1 app",
        ".\\docker\\logs-full.ps1 waha", ".\\docker\\test-full.ps1", ".\\docker\\stop-full.ps1",
    ])
    callout(doc, "Volume tetap disimpan", "stop-full.ps1 menghentikan container tanpa menghapus volume database dan session WAHA. Jangan menambahkan opsi -v kecuali penghapusan seluruh data memang disetujui dan backup telah diverifikasi.", b.LIGHT_AMBER)

    b.add_heading(doc, "11.4 Konfigurasi Environment dan Secret", 2)
    add_figure(doc, "environment-secret.png", "Gambar 11-4. Pemisahan environment dan secret")
    table(doc, "Tabel 11-5. Inventaris konfigurasi kritis", ["Kelompok", "Variabel", "Kontrol"], [
        ["Aplikasi", "APP_ENV, APP_KEY, APP_DEBUG, APP_URL", "APP_KEY unik; production HTTPS dan debug false"],
        ["Database", "DB_HOST, DB_DATABASE, DB_USERNAME, DB_PASSWORD", "User terbatas; password unik"],
        ["Redis", "REDIS_HOST, REDIS_PASSWORD", "Password wajib production"],
        ["Session", "SESSION_DRIVER, SESSION_ENCRYPT, SESSION_SECURE_COOKIE", "Redis/enkripsi/cookie aman production"],
        ["WAHA", "WAHA_BASE_URL, WAHA_API_KEY, WAHA_API_KEY_HASH", "Key asli untuk CRM; hash sesuai image WAHA"],
        ["Webhook", "WAHA_WEBHOOK_URL, WAHA_WEBHOOK_HMAC_SECRET", "HTTPS publik; secret berbeda dari API key"],
        ["Email", "MAIL_HOST, MAIL_USERNAME, MAIL_PASSWORD", "Akun SMTP khusus; rotasi"],
    ], [1.1, 2.75, 2.65])
    steps(doc, "Tabel 11-6. Perubahan environment yang aman", [
        "Salin file example menjadi file environment yang diabaikan Git; jangan mengedit contoh untuk menyimpan secret nyata.",
        "Isi nilai per lingkungan melalui secret manager atau file dengan permission terbatas.",
        "Validasi Compose dengan docker compose ... config --quiet tanpa mencetak secret ke tiket atau chat.",
        "Setelah perubahan Laravel, muat ulang/recreate service dan bersihkan cache konfigurasi bila diperlukan.",
        "Uji health, login, queue, scheduler, WAHA, dan fungsi terkait; catat versi serta waktu rotasi.",
    ])

    b.add_heading(doc, "11.5 Database, Migration, Seeder, dan Storage", 2)
    add_figure(doc, "kontrol-data.png", "Gambar 11-5. Kontrol perubahan database dan storage")
    table(doc, "Tabel 11-7. Aturan perubahan data", ["Aktivitas", "Development", "Production"], [
        ["Migration", "Dapat otomatis saat app start", "Eksplisit dengan migrate --force"],
        ["Seeder", "Hanya database kosong", "Dilarang"],
        ["Rollback schema", "Boleh setelah tinjauan", "Utamakan restore/forward fix yang disetujui"],
        ["Storage link", "Verifikasi /storage", "Volume persisten dan permission terjaga"],
        ["Penghapusan volume", "Hanya reset lokal sadar risiko", "Dilarang pada volume aktif"],
    ], [1.55, 2.35, 2.6])
    steps(doc, "Tabel 11-8. Migration production", [
        "Catat migration pending dan tinjau kompatibilitas aplikasi serta kebutuhan downtime.",
        "Buat backup database/storage dan verifikasi checksum serta keterbacaan arsip.",
        "Bangun image/tag release yang immutable dan validasi Compose.",
        "Jalankan database dan Redis, lalu run --rm app php artisan migrate --force.",
        "Naikkan seluruh service, periksa status migration, health, log, dan smoke test.",
        "Jika gagal, hentikan dampak, pilih forward fix atau restore sesuai keputusan rilis; jangan menebak rollback data.",
    ])

    b.add_heading(doc, "11.6 Queue dan Scheduler", 2)
    add_figure(doc, "alur-asinkron.png", "Gambar 11-6. Alur kerja queue, scheduler, dan integrasi WAHA")
    table(doc, "Tabel 11-9. Proses asinkron", ["Proses", "Baseline", "Pemeriksaan"], [
        ["Queue", "queue:work --sleep=3 --tries=3 --timeout=90; production max-time=3600", "Log, failed jobs, throughput"],
        ["Scheduler", "schedule:work", "schedule:list dan log"],
        ["Reminder", "crm:send-reminders setiap jam, tanpa overlap", "Notifikasi jatuh tempo"],
        ["Sync WA", "crm:sync-whatsapp-sessions setiap 5 menit, tanpa overlap, background", "Status koneksi terbaru"],
    ], [1.4, 3.35, 1.75])
    steps(doc, "Tabel 11-10. Penanganan job gagal", [
        "Periksa status container queue dan scheduler, kemudian lihat log layanan terkait.",
        "Gunakan queue:failed untuk mengidentifikasi job, waktu, exception, dan payload yang aman.",
        "Perbaiki akar masalah seperti jaringan, credential, data referensi, atau service yang tidak sehat.",
        "Retry hanya setelah dampak idempotensi/duplikasi dipahami; pantau hasil pada tabel bisnis dan audit.",
        "Restart worker secara terkendali setelah release/config change agar kode dan konfigurasi baru digunakan.",
    ])

    b.add_heading(doc, "11.7 WAHA/GOWS dan Webhook", 2)
    table(doc, "Tabel 11-11. Konfigurasi WAHA", ["Aspek", "Lokal full Docker", "Production"], [
        ["Base URL", "http://waha:3000", "http://waha:3000 pada network internal"],
        ["Webhook URL", "http://web/api/webhooks/waha", "URL HTTPS aplikasi"],
        ["Engine", "GOWS", "GOWS sesuai image yang diuji"],
        ["Dashboard", "http://localhost:3000/dashboard", "Nonaktif/default; akses sangat dibatasi"],
        ["API security", "X-Api-Key unik lokal", "Key/hash kuat dan secret manager"],
        ["Webhook security", "HMAC secret unik", "HMAC unik, rotasi, HTTPS"],
        ["Persistensi", "waha_sessions", "Volume persisten terenkripsi/terbackup"],
    ], [1.35, 2.55, 2.6])
    steps(doc, "Tabel 11-12. Menyiapkan akun WhatsApp", [
        "Buat satu Koneksi WhatsApp untuk setiap admin/staff yang berwenang dan pastikan nomor serta cabangnya benar.",
        "Siapkan session, Mulai, tampilkan QR dari CRM, lalu scan melalui menu Perangkat tertaut pada WhatsApp personal atau Business.",
        "Sinkronkan sampai status Working, kemudian Aktifkan chat agar event dikirim ke CRM.",
        "Kirim pesan dari nomor lain; pastikan webhook event, conversation, message, unread, lead matching, dan ACK berubah.",
        "Jika nomor sama ditemukan lintas cabang/lebih dari satu lead, pastikan status Conflict dan jangan membuat duplikat manual.",
    ])
    callout(doc, "Normalisasi GOWS", "Chat personal dapat memakai ID @lid. Sistem membaca nomor alternatif dari SenderAlt, membuang suffix perangkat, menormalisasi menjadi 62..., lalu mencocokkannya dalam scope cabang.", b.LIGHT_GREEN)

    b.add_heading(doc, "11.8 Client API dan Endpoint Integrasi", 2)
    table(doc, "Tabel 11-13. Endpoint administratif", ["Endpoint/kanal", "Kegunaan", "Kontrol"], [
        ["GET /up", "Health aplikasi melalui Nginx/Laravel", "Tanpa data sensitif"],
        ["GET WAHA /health", "Health gateway", "Boleh dikecualikan dari API key sesuai config"],
        ["WAHA /api/sessions", "Status dan lifecycle session", "X-Api-Key; jaringan terbatas"],
        ["POST /api/webhooks/waha", "Event pesan/session masuk", "Signature HMAC; queue; audit"],
        ["Panel Filament", "Operasi admin terautentikasi", "Session, role, cabang, policy/scope"],
    ], [1.65, 2.75, 2.1])
    callout(doc, "Batas API", "Baseline tidak mendokumentasikan API publik umum untuk pihak ketiga. Setiap client baru harus memiliki kontrak, autentikasi, rate limit, idempotency, logging aman, versi endpoint, dan test sebelum diaktifkan.", b.LIGHT_AMBER)

    b.add_heading(doc, "11.9 Administrasi User, Role, dan Cabang", 2)
    table(doc, "Tabel 11-14. Matriks kewenangan ringkas", ["Role", "Scope", "Administrasi"], [
        ["Superadmin", "Semua cabang", "Kelola sistem, master data, koneksi, audit"],
        ["Admin Cabang", "Cabang sendiri", "Kelola tim dan operasi cabang sesuai policy"],
        ["Staff", "Cabang/assignment sendiri", "Operasi lead, follow-up, chat, transaksi sesuai izin"],
        ["Direksi", "Semua cabang", "Monitoring global; tidak mengubah operasional"],
    ], [1.25, 2.25, 3.0])
    steps(doc, "Tabel 11-15. Siklus akun", [
        "Terima permintaan dengan identitas, role, cabang, atasan/pemilik persetujuan, dan tanggal berlaku.",
        "Pastikan email unik; pilih role berdasarkan tugas, bukan jabatan informal.",
        "Tetapkan Bandung, JakSel, atau JakPus untuk Admin Cabang/Staff; akses global hanya Superadmin/Direksi.",
        "Sampaikan credential awal melalui kanal aman dan minta perubahan password bila kebijakan mengharuskan.",
        "Review berkala akun aktif, role global, cabang, dan koneksi WhatsApp.",
        "Nonaktifkan akses segera ketika mutasi/keluar; pertahankan jejak audit dan alihkan pekerjaan secara resmi.",
    ])

    b.add_heading(doc, "11.10 Konfigurasi Aplikasi dan Program", 2)
    table(doc, "Tabel 11-16. Konfigurasi yang dikelola", ["Kelompok", "Contoh", "Kontrol perubahan"], [
        ["Identitas", "Nama aplikasi, organisasi, tagline", "Review visual dan komunikasi"],
        ["Dukungan", "Email/nomor dukungan", "Uji kanal sebelum publikasi"],
        ["Laporan", "Izin export Staff", "Persetujuan dan uji scope"],
        ["Program", "Kategori, kelas, tipe, metode, jenjang, tingkat", "Jangan hapus opsi yang dipakai"],
        ["Sekolah", "Prefix, status, Grade A-E, kota", "NPSN opsional; hindari duplikasi"],
    ], [1.3, 2.55, 2.65])
    callout(doc, "Nama program", "Nama program otomatis hanya memakai Kategori + Kelas program + Metode belajar, misalnya SR GOLD ONLINE. Tipe, jenjang, dan tingkat kelas tetap menjadi dimensi data tetapi bukan bagian nama.")

    b.add_heading(doc, "11.11 Backup", 2)
    add_figure(doc, "backup-restore.png", "Gambar 11-7. Alur backup, restore, dan rekonsiliasi")
    table(doc, "Tabel 11-17. Cakupan backup", ["Objek", "Isi", "Catatan"], [
        ["Database", "Schema dan data MySQL", "Dump konsisten, terenkripsi, checksum"],
        ["Storage", "Upload dan file aplikasi", "Selaras dengan titik waktu database"],
        ["WAHA session", "State perangkat tertaut", "Sangat sensitif; akses minimum"],
        ["Konfigurasi", "Compose, image tag, config non-secret", "Secret dicatat melalui referensi vault"],
        ["Bukti", "Waktu, ukuran, checksum, lokasi, operator", "Tanpa memasukkan isi secret"],
    ], [1.3, 2.9, 2.3])
    steps(doc, "Tabel 11-18. Prosedur backup", [
        "Tentukan cakupan, titik waktu, retensi, lokasi tujuan, RPO/RTO, dan pemilik persetujuan.",
        "Pastikan kapasitas tujuan cukup serta aksesnya dibatasi dan terenkripsi.",
        "Buat dump database konsisten; arsipkan storage dan session WAHA sesuai kebijakan.",
        "Hitung checksum, catat manifest, pindahkan salinan ke lokasi terpisah dari host aplikasi.",
        "Uji pemulihan secara berkala pada lingkungan terisolasi; backup tanpa uji restore belum dianggap andal.",
    ])

    b.add_heading(doc, "11.12 Restore dan Rekonsiliasi", 2)
    steps(doc, "Tabel 11-19. Prosedur restore", [
        "Nyatakan insiden/permintaan, scope data, recovery point, dampak, dan persetujuan.",
        "Hentikan penulisan atau aktifkan maintenance mode bila diperlukan; simpan snapshot kondisi gagal.",
        "Verifikasi checksum backup dan lakukan restore pada environment terisolasi bila waktu memungkinkan.",
        "Pulihkan database, storage, dan session yang relevan dengan permission benar; jalankan migration hanya sesuai versi image.",
        "Naikkan service berurutan dan lakukan rekonsiliasi jumlah record, file, relasi, transaksi, inbox WA, serta audit.",
        "Lakukan smoke test, buka layanan, pantau, dan catat gap data terhadap RPO sebelum menutup tindakan.",
    ])
    table(doc, "Tabel 11-20. Rekonsiliasi minimum", ["Area", "Bukti"], [
        ["Schema", "Migration sesuai release; tidak ada pending tak terencana"],
        ["Data bisnis", "Sampel lead, follow-up, pendaftaran, pembayaran, target"],
        ["File", "Upload dapat dibuka melalui /storage"],
        ["WAHA", "Session/health; pesan uji; webhook dan queue"],
        ["Akses", "Login dan scope role/cabang"],
        ["Pelaporan", "KPI/omzet konsisten dengan transaksi sampel"],
    ], [1.45, 5.05])

    b.add_heading(doc, "11.13 Permission Server, File, dan Volume", 2)
    table(doc, "Tabel 11-21. Kontrol permission", ["Objek", "Kebutuhan", "Larangan"], [
        ["storage", "Dapat ditulis user proses aplikasi", "World-writable tanpa alasan"],
        ["bootstrap/cache", "Dapat ditulis user proses aplikasi", "Owner root permanen yang memblokir app"],
        ["source/image", "Read-only pada runtime bila memungkinkan", "Edit manual di container production"],
        ["environment", "Hanya akun deployment/operator", "Commit atau tampilkan isi ke log"],
        ["volume DB/WA", "Service owner, backup terbatas", "Mount publik atau penghapusan tanpa backup"],
        ["TLS key", "Reverse proxy/secret store", "Disimpan bersama source"],
    ], [1.45, 2.75, 2.3])
    callout(doc, "Entry point", "Container membuat direktori storage/framework dan bootstrap/cache lalu menyesuaikan owner ke www-data. Jika permission kembali rusak setelah deploy, periksa mount/volume dan UID/GID alih-alih menjalankan chmod luas.", b.LIGHT_AMBER)

    b.add_heading(doc, "11.14 Maintenance Mode", 2)
    steps(doc, "Tabel 11-22. Mengaktifkan maintenance", [
        "Umumkan waktu, dampak, pemilik, dan kriteria selesai; pastikan backup terbaru tersedia.",
        "Aktifkan php artisan down dari service app. Pada multi-instance gunakan maintenance store bersama sesuai konfigurasi.",
        "Lakukan perubahan/migration dan simpan log perintah serta hasil tanpa secret.",
        "Jalankan health dan smoke test internal; pastikan queue/scheduler tidak menghasilkan efek yang bertentangan.",
        "Aktifkan kembali php artisan up, verifikasi akses pengguna, lalu pantau error dan kapasitas.",
    ])
    command(doc, "Perintah maintenance di Compose production", [
        "docker compose --env-file docker/.env.production -f docker/compose.production.yml exec app php artisan down",
        "docker compose --env-file docker/.env.production -f docker/compose.production.yml exec app php artisan up",
    ])

    b.add_heading(doc, "11.15 Deployment dan Update", 2)
    table(doc, "Tabel 11-23. Gerbang rilis", ["Tahap", "Kriteria"], [
        ["Persiapan", "Scope, risiko, persetujuan, changelog, rollback, backup"],
        ["Build", "Compose valid; image/tag immutable; dependency dan aset terbangun"],
        ["Test", "Automated test dan build lulus; migration ditinjau"],
        ["Deploy", "Secret production; migration eksplisit; service sehat"],
        ["Smoke", "Login, dashboard, scope, CRUD sampel, file, queue, WAHA"],
        ["Observasi", "Log/error/latency/kapasitas stabil; incident owner siaga"],
        ["Penutupan", "Bukti, versi, hasil, isu tersisa, dokumentasi diperbarui"],
    ], [1.3, 5.2])
    command(doc, "Urutan produksi ringkas", [
        "docker compose --env-file docker/.env.production -f docker/compose.production.yml config --quiet",
        "docker compose --env-file docker/.env.production -f docker/compose.production.yml build --pull",
        "docker compose --env-file docker/.env.production -f docker/compose.production.yml up -d database redis",
        "docker compose --env-file docker/.env.production -f docker/compose.production.yml run --rm app php artisan migrate --force",
        "docker compose --env-file docker/.env.production -f docker/compose.production.yml up -d --remove-orphans",
    ])
    callout(doc, "Target VPS", "Compose production adalah baseline satu server. Sebelum go-live, tambahkan DNS, TLS, firewall, SMTP, backup off-host, monitoring, rotasi secret, batas akses WAHA, dan drill restore. Untuk skala lebih besar, database/Redis terkelola dapat dipertimbangkan.", b.LIGHT_GREEN)

    b.add_heading(doc, "11.16 Monitoring, Logging, dan Health Check", 2)
    add_figure(doc, "monitoring-insiden.png", "Gambar 11-8. Siklus monitoring, insiden, dan maintenance")
    table(doc, "Tabel 11-24. Sinyal operasional", ["Sinyal", "Pemeriksaan", "Respons awal"], [
        ["HTTP /up", "Status 200 dan waktu respons", "Periksa web/app/database"],
        ["Container", "Running/healthy/restart count", "Lihat log dan dependency"],
        ["Laravel log", "Error, exception, rate", "Triage request/job terkait"],
        ["Queue", "Worker, failed jobs, backlog", "Perbaiki dependency lalu retry selektif"],
        ["Scheduler", "Proses hidup dan jadwal berjalan", "Restart dan cek overlap/lock"],
        ["WAHA", "Health, session Working, webhook", "Key/network/session/log"],
        ["Kapasitas", "Disk, RAM, CPU, DB/Redis", "Scale/cleanup aman sebelum penuh"],
        ["Keamanan", "Login gagal, perubahan role, anomaly", "Containment dan audit"],
    ], [1.25, 2.8, 2.45])
    callout(doc, "Logging aman", "Jangan menulis password, APP_KEY, API key, HMAC signature/secret, cookie, isi dump, QR WhatsApp, atau isi chat sensitif ke log. Masking harus diterapkan sebelum log dipusatkan.", b.LIGHT_AMBER)

    b.add_heading(doc, "11.17 Troubleshooting dan Pemulihan", 2)
    table(doc, "Tabel 11-25. Diagnostik umum", ["Gejala", "Penyebab umum", "Tindakan"], [
        ["502 Bad Gateway", "app/PHP-FPM belum sehat atau IP upstream berubah", "Log app/web; recreate web; cek port 9000"],
        ["DB connection refused", "MySQL belum hidup/host salah", "Gunakan host database di Docker; cek health"],
        ["CSS tidak tampil", "Node/Vite/build atau URL aset", "Cek node, Vite, build, dan mixed content"],
        ["Inbox WA kosong", "Webhook/queue/HMAC/session", "Cek event, queue, log WAHA, session Working"],
        ["WAHA 401", "API key tidak sama", "Samakan key CRM dan service; recreate"],
        ["Session Failed", "GOWS/session/jaringan", "Log WAHA; restart tanpa hapus volume; scan ulang bila perlu"],
        ["Job gagal", "Dependency/data/API", "queue:failed; perbaiki akar; retry selektif"],
        ["Disk penuh", "Log/image/backup/volume", "Hentikan pertumbuhan; identifikasi; cleanup terkontrol"],
    ], [1.45, 2.5, 2.55])
    steps(doc, "Tabel 11-26. Format triage insiden", [
        "Catat waktu mulai, pelapor, lingkungan, URL/service, gejala, dan dampak bisnis.",
        "Reproduksi dengan langkah minimum dan kumpulkan status, log relevan, versi image/commit, serta perubahan terakhir.",
        "Tentukan severity dan pemilik; lakukan containment tanpa menghapus bukti atau volume.",
        "Pulihkan dengan restart, rollback, forward fix, atau restore yang disetujui; verifikasi end-to-end.",
        "Dokumentasikan akar masalah, data terdampak, tindakan pencegahan, pemilik, dan tenggat.",
    ])

    b.add_heading(doc, "11.18 Checklist Penerimaan dan Serah Terima", 2)
    table(doc, "Tabel 11-27. Checklist administrator", ["Area", "Kriteria lulus"], [
        ["Runtime", "Seluruh service running/healthy dan restart policy aktif"],
        ["Aplikasi", "Login, dashboard, menu, search, notification, profile berfungsi"],
        ["RBAC", "Empat role dan scope Bandung/JakSel/JakPus diuji"],
        ["Data", "Migration tepat; seeder tidak dijalankan di production; storage terbaca"],
        ["Async", "Queue, scheduler, reminder, sync WA berjalan"],
        ["WAHA", "GOWS, session Working, webhook HMAC, inbox/chat/ACK/matching lead"],
        ["Keamanan", "HTTPS, firewall, debug false, secret unik, cookie/session aman"],
        ["Backup", "Backup off-host dan uji restore terdokumentasi"],
        ["Observabilitas", "Health, log, alert, kapasitas, dan incident path tersedia"],
        ["Quality", "Test/build/smoke/UAT lulus dan bukti tersimpan"],
        ["Dokumen", "Runbook, daftar owner, kontak eskalasi, dan changelog terbaru"],
    ], [1.4, 5.1])
    callout(doc, "Kriteria siap operasional", "Sistem tidak dinyatakan siap hanya karena halaman login terbuka. Data, RBAC, queue, scheduler, WAHA, backup/restore, keamanan, monitoring, dan jalur insiden harus dibuktikan bersama.", b.LIGHT_GREEN)

    b.add_heading(doc, "LAMPIRAN A - REFERENSI PERINTAH", 1)
    table(doc, "Tabel A-1. Perintah cepat full Docker", ["Tujuan", "Perintah"], [
        ["Start", ".\\docker\\start-full.ps1"], ["Status", ".\\docker\\status-full.ps1"],
        ["Log app", ".\\docker\\logs-full.ps1 app"], ["Log WAHA", ".\\docker\\logs-full.ps1 waha"],
        ["Test", ".\\docker\\test-full.ps1"], ["Stop", ".\\docker\\stop-full.ps1"],
        ["Seeder lokal", "docker compose --env-file docker\\.env -f docker\\compose.yml exec app php artisan db:seed"],
        ["Health lokal", "Invoke-WebRequest -UseBasicParsing http://127.0.0.1:8000/up"],
    ], [1.5, 5.0])

    b.add_heading(doc, "LAMPIRAN B - CATATAN ADMINISTRASI", 1)
    table(doc, "Tabel B-1. Template catatan perubahan", ["Field", "Isi yang dicatat"], [
        ["Identitas", "ID tiket/perubahan, judul, lingkungan, tanggal/waktu"],
        ["Otorisasi", "Pemohon, pemberi persetujuan, pelaksana, pemeriksa"],
        ["Scope", "Service, data, cabang, pengguna, downtime"],
        ["Baseline", "Versi image/commit, migration, konfigurasi non-secret"],
        ["Risiko", "Dampak, mitigasi, rollback, recovery point"],
        ["Pelaksanaan", "Perintah/tindakan dan waktu; secret dimasking"],
        ["Verifikasi", "Health, test, smoke, rekonsiliasi, UAT"],
        ["Hasil", "Berhasil/gagal, isu tersisa, owner, tenggat"],
    ], [1.45, 5.05])

    b.add_heading(doc, "LAMPIRAN C - LARANGAN OPERASIONAL", 1)
    table(doc, "Tabel C-1. Tindakan yang dilarang", ["Tindakan", "Risiko", "Alternatif aman"], [
        ["docker compose down -v di production", "Menghapus volume data/session", "down tanpa -v; backup dan approval"],
        ["db:seed di production", "Data demo/duplikat", "Migration eksplisit; import terkontrol"],
        ["APP_DEBUG=true di production", "Kebocoran detail", "Debug false; log terlindungi"],
        ["Commit file environment", "Kebocoran secret", "Secret manager/file terproteksi"],
        ["Satu akun/session dibagi", "Audit dan ownership salah", "Akun dan session per orang"],
        ["Retry job massal tanpa analisis", "Duplikasi/efek ganda", "Retry selektif setelah RCA"],
        ["chmod luas tanpa diagnosis", "Eskalasi akses", "Perbaiki owner/mount/UID secara tepat"],
    ], [2.25, 2.0, 2.25])

    b.add_heading(doc, "PENUTUP", 1)
    b.add_body(doc, "Panduan ini merekam baseline administrator SIVMI CRM per 25 Agustus 2026. Setiap perubahan versi framework, image, migration, konfigurasi WAHA, jadwal, role, cabang, dan strategi deployment harus diikuti pembaruan runbook, test, bukti restore, serta dokumen ini.")
    callout(doc, "Prinsip akhir", "Administrator yang baik tidak hanya membuat layanan hidup, tetapi juga menjaga perubahan dapat diulang, akses dapat dijelaskan, data dapat dipulihkan, dan setiap keputusan dapat diaudit.", b.LIGHT_GREEN)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(f"Created: {OUTPUT}")


if __name__ == "__main__":
    build_document()
