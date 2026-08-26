from __future__ import annotations

from pathlib import Path
import sys

PROJECT_ROOT = Path(__file__).resolve().parents[2]
DOC_TOOLS = PROJECT_ROOT / "storage" / "app" / "doc-tools"
sys.path.insert(0, str(DOC_TOOLS))

from PIL import Image, ImageDraw
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

import build_bab_iv_arsitektur_bisnis as b
import build_bab_vii_arsitektur_teknologi as v7


OUTPUT = PROJECT_ROOT / "docs" / "BAB I - Pendahuluan SIVMI CRM v1.0.docx"
ASSETS = PROJECT_ROOT / "docs" / "_bab_i_assets"
COVER = ASSETS / "cover-bab-i.png"

PRIMARY = "#1476B8"
PRIMARY_DARK = "#075985"
INK = "#0B2545"
MUTED = "#64748B"
BG = "#F6F8FC"
GREEN = "#10B981"
AMBER = "#F59E0B"
ROSE = "#F43F5E"
PURPLE = "#A855F7"


def font(name: str, size: int):
    return v7.font(name, size)


def canvas(title: str, subtitle: str, size=(1800, 1080)):
    return v7.canvas(title, subtitle, size)


def box(draw, xy, title, lines, fill="#F8FAFC", outline=PRIMARY, title_size=27, body_size=20):
    return v7.box(draw, xy, title, lines, fill, outline, title_size, body_size)


def arrow(draw, start, end, color="#64748B", width=5):
    return v7.arrow(draw, start, end, color, width)


def make_cover() -> None:
    ASSETS.mkdir(parents=True, exist_ok=True)
    image = Image.new("RGB", (1600, 2200), BG)
    draw = ImageDraw.Draw(image)
    for y in range(2200):
        ratio = y / 2200
        start, end = (20, 118, 184), (7, 89, 133)
        draw.line((0, y, 1600, y), fill=tuple(round(start[i] * (1 - ratio) + end[i] * ratio) for i in range(3)))
    draw.ellipse((1010, -180, 1770, 580), fill=(27, 139, 199))
    draw.ellipse((1140, 140, 1640, 640), fill=(8, 103, 151))
    draw.rounded_rectangle((120, 250, 1480, 1910), radius=44, fill="white")
    draw.text((190, 350), "SIVMI CRM", font=font("arialbd.ttf", 34), fill=PRIMARY)
    draw.text((190, 610), "BAB I", font=font("arialbd.ttf", 48), fill=PRIMARY_DARK)
    draw.text((190, 700), "PENDAHULUAN", font=font("arialbd.ttf", 82), fill=INK)
    draw.text((190, 900), "Konteks, masalah, tujuan, manfaat,", font=font("arial.ttf", 31), fill="#475569")
    draw.text((190, 952), "dan ruang lingkup pengembangan SIVMI CRM", font=font("arial.ttf", 31), fill="#475569")
    draw.rounded_rectangle((190, 1160, 805, 1240), radius=40, fill="#EAF4FB")
    draw.text((230, 1182), "BASELINE APLIKASI 25 AGUSTUS 2026", font=font("arial.ttf", 28), fill=PRIMARY_DARK)
    draw.text((190, 1610), "Versi dokumen 1.0", font=font("arial.ttf", 28), fill=MUTED)
    draw.text((190, 1665), "Sasaran: pemilik bisnis, manajemen, analis, pengembang, QA, dan administrator", font=font("arial.ttf", 24), fill=MUTED)
    draw.text((190, 1775), "SIVMI CRM - Catat, Kelola, Closing", font=font("arialbd.ttf", 34), fill=PRIMARY)
    image.save(COVER)


def make_background() -> None:
    image, draw = canvas("Konteks Transformasi CRM Bimbel", "Data calon siswa perlu bergerak dari kanal masuk menjadi keputusan dan layanan yang terukur")
    nodes = [
        (50, 245, 330, 650, "KANAL MASUK", "WhatsApp\nIklan & media sosial\nKunjungan\nRujukan", PRIMARY),
        (390, 245, 670, 650, "PENGELOLAAN", "Lead & sumber\nPIC & cabang\nFollow-up\nStatus", PURPLE),
        (730, 245, 1010, 650, "KONVERSI", "Pendaftaran\nProgram\nPembayaran\nSiswa aktif", GREEN),
        (1070, 245, 1350, 650, "KENDALI", "Target\nKPI\nLaporan\nAudit", AMBER),
        (1410, 245, 1750, 650, "KEPUTUSAN", "Prioritas follow-up\nKinerja tim\nProgram diminati\nWilayah potensial", PRIMARY_DARK),
    ]
    for x1, y1, x2, y2, title, lines, color in nodes:
        box(draw, (x1, y1, x2, y2), title, lines, "#F8FAFC", color, 24, 18)
    for x in (330, 670, 1010, 1350):
        arrow(draw, (x, 450), (x + 60, 450))
    draw.text((900, 850), "SIVMI CRM memusatkan data, aktivitas, komunikasi, transaksi, dan analitik sesuai role serta cabang.", anchor="mm", font=font("arialbd.ttf", 23), fill=MUTED)
    image.save(ASSETS / "konteks-transformasi.png")


def make_problem_map() -> None:
    image, draw = canvas("Peta Masalah Operasional", "Masalah saling terkait dan tidak dapat diselesaikan hanya dengan menambah form")
    box(draw, (650, 365, 1150, 700), "AKAR MASALAH", "Data dan komunikasi tidak terpusat\nKontrol akses belum konsisten\nBukti proses sulit ditelusuri", "#FFF7ED", AMBER, 29, 20)
    items = [
        (55, 145, 430, 425, "DATA", "Duplikat nomor\nformat tidak konsisten\nriwayat tersebar", PRIMARY),
        (55, 675, 430, 955, "PROSES", "Follow-up terlambat\nstatus tidak seragam\nclosing sulit dilacak", PURPLE),
        (1370, 145, 1745, 425, "AKSES", "Risiko lintas cabang\nhak menu/aksi berbeda\nexport tidak terscope", ROSE),
        (1370, 675, 1745, 955, "KEPUTUSAN", "KPI tidak real-time\ntarget tidak terukur\ntren tidak terlihat", GREEN),
    ]
    for x1, y1, x2, y2, title, lines, color in items:
        box(draw, (x1, y1, x2, y2), title, lines, "#F8FAFC", color, 25, 19)
    arrow(draw, (430, 285), (650, 455)); arrow(draw, (430, 815), (650, 610))
    arrow(draw, (1370, 285), (1150, 455)); arrow(draw, (1370, 815), (1150, 610))
    image.save(ASSETS / "peta-masalah.png")


def make_objective_chain() -> None:
    image, draw = canvas("Rantai Sasaran SIVMI CRM", "Tujuan sistem diturunkan dari masalah menjadi kapabilitas dan hasil yang dapat diamati")
    nodes = [
        (55, 270, 340, 670, "MASALAH", "Data tersebar\nFollow-up tidak terukur\nChat terpisah\nAkses rawan", ROSE),
        (410, 270, 695, 670, "KAPABILITAS", "CRM terpusat\nworkflow\nRBAC + branch scope\nWAHA/GOWS", PRIMARY),
        (765, 270, 1050, 670, "OUTPUT", "Lead bersih\nriwayat lengkap\ninbox terkait lead\nKPI terfilter", PURPLE),
        (1120, 270, 1405, 670, "OUTCOME", "Respons cepat\nduplikasi berkurang\nkontrol membaik\nkeputusan berbasis data", GREEN),
        (1475, 270, 1745, 670, "DAMPAK", "Konversi terpantau\nakuntabilitas tim\nlayanan konsisten\nsiap berkembang", AMBER),
    ]
    for x1, y1, x2, y2, title, lines, color in nodes:
        box(draw, (x1, y1, x2, y2), title, lines, "#F8FAFC", color, 24, 18)
    for x in (340, 695, 1050, 1405):
        arrow(draw, (x, 470), (x + 70, 470))
    image.save(ASSETS / "rantai-sasaran.png")


def make_benefits() -> None:
    image, draw = canvas("Peta Manfaat Sistem", "Manfaat dibedakan menurut penerima agar dapat dievaluasi pada UAT dan operasi")
    groups = [
        (55, 210, 430, 825, "STAFF", "Satu daftar kerja\nriwayat lead dan chat\nquick action\nreminder follow-up", PRIMARY),
        (490, 210, 865, 825, "ADMIN CABANG", "Kontrol distribusi\nkinerja tim\ntarget cabang\nkonflik/duplikat", PURPLE),
        (925, 210, 1300, 825, "DIREKSI", "KPI lintas cabang\ntren dan ranking\nprogram/sekolah/wilayah\ndasar keputusan", GREEN),
        (1360, 210, 1745, 825, "ADMINISTRATOR", "RBAC dan audit\nhealth/log\nbackup/restore\nkonfigurasi terpusat", AMBER),
    ]
    for x1, y1, x2, y2, title, lines, color in groups:
        box(draw, (x1, y1, x2, y2), title, lines, "#F8FAFC", color, 27, 20)
    image.save(ASSETS / "peta-manfaat.png")


def make_scope() -> None:
    image, draw = canvas("Batas Ruang Lingkup SIVMI CRM", "Ruang lingkup membedakan fungsi inti, integrasi, operasi, dan hal di luar baseline")
    box(draw, (120, 160, 1680, 900), "DALAM RUANG LINGKUP", "", "#EFF6FF", PRIMARY_DARK, 30, 20)
    areas = [
        (175, 300, 520, 775, "CRM INTI", "Dashboard\nLead & CSV\nFollow-up\nPendaftaran\nPembayaran", PRIMARY),
        (565, 300, 910, 775, "MASTER & KONTROL", "Cabang\nSekolah\nPengguna\nProgram\nTarget/laporan/audit", PURPLE),
        (955, 300, 1300, 775, "WHATSAPP", "Nomor per user\nWAHA/GOWS\nWebhook\nInbox/chat\nmatching & dedupe", GREEN),
        (1345, 300, 1625, 775, "OPERASI", "Full Docker lokal\nqueue/scheduler\nhealth/log\nbackup\ntarget VPS", AMBER),
    ]
    for x1, y1, x2, y2, title, lines, color in areas:
        box(draw, (x1, y1, x2, y2), title, lines, "white", color, 23, 18)
    image.save(ASSETS / "batas-ruang-lingkup.png")


def make_scope_boundary() -> None:
    image, draw = canvas("Batas Sistem dan Ketergantungan", "SIVMI mengendalikan data internal, tetapi bergantung pada layanan dan tata kelola eksternal")
    box(draw, (600, 260, 1200, 740), "SIVMI CRM", "Laravel 12 | Filament 5 | Livewire 4\nMySQL | Redis | Queue\nRBAC | Branch scope | Audit\nDashboard | CRM | WhatsApp inbox", "#EAF4FB", PRIMARY_DARK, 31, 21)
    externals = [
        (55, 220, 435, 510, "PENGGUNA", "Superadmin\nAdmin Cabang\nStaff\nDireksi", PRIMARY),
        (55, 590, 435, 880, "KANAL DATA", "CSV\nform manual\nWhatsApp\nsumber lead", PURPLE),
        (1365, 220, 1745, 510, "WAHA/GOWS", "Session dan QR\nAPI key\nwebhook\nstatus pesan", GREEN),
        (1365, 590, 1745, 880, "INFRASTRUKTUR", "Docker host\nDNS/TLS\nbackup\nmonitoring/VPS", AMBER),
    ]
    for x1, y1, x2, y2, title, lines, color in externals:
        box(draw, (x1, y1, x2, y2), title, lines, "#F8FAFC", color, 25, 19)
    arrow(draw, (435, 365), (600, 430)); arrow(draw, (435, 735), (600, 570))
    arrow(draw, (1200, 430), (1365, 365)); arrow(draw, (1200, 570), (1365, 735))
    image.save(ASSETS / "batas-sistem.png")


def add_figure(doc: Document, name: str, caption: str, width=6.15) -> None:
    b.add_figure(doc, ASSETS / name, caption, width)


def table(doc: Document, caption: str, headers: list[str], rows: list[list[str]], widths: list[float]) -> None:
    b.add_table(doc, headers, rows, widths, caption)


def callout(doc: Document, title: str, body: str, fill: str = b.LIGHT_BLUE) -> None:
    b.add_callout(doc, title, body, fill)
    doc.tables[-1].rows[0]._tr.get_or_add_trPr().append(b.OxmlElement("w:cantSplit"))


def bullets(doc: Document, items: list[str]) -> None:
    b.add_bullets(doc, items)


def page(doc: Document) -> None:
    doc.add_page_break()


def build_document() -> None:
    for fn in (make_cover, make_background, make_problem_map, make_objective_chain, make_benefits, make_scope, make_scope_boundary):
        fn()

    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.5); sec.page_height = Inches(11)
    sec.top_margin = Inches(0.85); sec.bottom_margin = Inches(0.85)
    sec.left_margin = Inches(1); sec.right_margin = Inches(1)
    sec.header_distance = Inches(0.4); sec.footer_distance = Inches(0.4)
    b.configure_styles(doc)
    sec.header.paragraphs[0].text = "SIVMI CRM  |  BAB I: PENDAHULUAN"
    for run in sec.header.paragraphs[0].runs:
        b.set_run_font(run, 8.5, True, b.PRIMARY)
    b.add_page_number(sec.footer.paragraphs[0])

    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(0)
    pic = p.add_run().add_picture(str(COVER), width=Inches(6.1))
    pic._inline.docPr.set("descr", "Sampul BAB I Pendahuluan SIVMI CRM")
    page(doc)

    b.add_heading(doc, "KENDALI DOKUMEN", 1)
    table(doc, "Tabel 0-1. Kendali dokumen", ["Atribut", "Nilai"], [
        ["Judul", "BAB I: Pendahuluan SIVMI CRM"],
        ["Versi", "1.0"],
        ["Tanggal baseline", "25 Agustus 2026"],
        ["Tujuan", "Menjelaskan alasan pengembangan, masalah, pertanyaan, manfaat, tujuan, dan batas ruang lingkup SIVMI CRM"],
        ["Pemangku kepentingan", "Pemilik bisnis, Direksi, Superadmin, Admin Cabang, Staff, administrator, analis, pengembang, QA, dan pengguna operasional"],
        ["Cabang baseline", "Bandung, Jakarta Selatan (JakSel), dan Jakarta Pusat (JakPus)"],
        ["Lingkungan", "Full Docker sebagai baseline lokal; Docker Compose production sebagai target VPS; Laragon/hybrid dipertahankan sebagai opsi legacy"],
        ["Sumber kebenaran", "Implementasi aplikasi, enum/model/migration, resource/page Filament, policy/scope, integrasi WAHA/GOWS, Docker Compose, seeder, test, dan dokumen arsitektur terbaru"],
    ], [1.55, 4.95])
    callout(doc, "Tanpa daftar isi", "Sesuai pola dokumen BAB II dan BAB III, file ini tidak memuat daftar isi. Pembahasan langsung menuju ringkasan eksekutif dan BAB I.")

    b.add_heading(doc, "RINGKASAN EKSEKUTIF", 1)
    b.add_body(doc, "SIVMI CRM dirancang sebagai ruang kerja terpusat bagi bimbingan belajar untuk mengelola calon siswa sejak lead masuk, tindak lanjut, pendaftaran, pembayaran, hingga evaluasi kinerja. Sistem menyatukan data operasional, master data, dashboard KPI, notifikasi, audit aktivitas, dan komunikasi WhatsApp melalui WAHA dengan engine GOWS.")
    b.add_body(doc, "Kebutuhan utama muncul karena data dan percakapan dapat berasal dari banyak sumber serta dikelola oleh pengguna pada cabang berbeda. Tanpa normalisasi nomor, pencocokan identitas, pembatasan cabang, histori proses, dan indikator yang konsisten, organisasi menghadapi risiko duplikasi lead, follow-up terlewat, akses lintas cabang, serta keputusan yang tidak berbasis bukti.")
    add_figure(doc, "konteks-transformasi.png", "Gambar 1-1. Konteks transformasi CRM bimbel")
    table(doc, "Tabel 0-2. Ringkasan arah pengembangan", ["Aspek", "Arah", "Bukti keberhasilan"], [
        ["Operasi", "Satu siklus lead sampai pembayaran", "Status, PIC, cabang, riwayat, dan transaksi dapat ditelusuri"],
        ["Komunikasi", "Nomor WhatsApp per admin/staff terhubung WAHA/GOWS", "Pesan tersimpan, terkait percakapan/lead, dan mempunyai status"],
        ["Kontrol", "RBAC dan branch scope", "Menu, aksi, query, widget, import/export, dan notifikasi mengikuti hak akses"],
        ["Analitik", "KPI, target, tren, program, sekolah, dan wilayah", "Filter global menghasilkan angka konsisten tanpa data dummy produksi"],
        ["Operasi TI", "Lingkungan dapat diulang melalui Docker", "Build, migration, queue, scheduler, health, backup, test, dan log dapat dijalankan"],
    ], [1.25, 3.2, 2.05])
    callout(doc, "Batas klaim", "Dokumen membedakan As-built (tersedia pada baseline yang diaudit) dan Target (harus dituntaskan atau disahkan melalui UAT/operasi). Angka dampak bisnis tidak dinyatakan sebagai realisasi sebelum tersedia data pembanding yang disetujui.", b.LIGHT_AMBER)

    b.add_heading(doc, "BAB I: PENDAHULUAN", 1)
    b.add_heading(doc, "1.1 Latar Belakang", 2)
    b.add_body(doc, "Layanan bimbingan belajar bergantung pada kemampuan organisasi mengenali calon siswa, merespons minat secara tepat waktu, menawarkan program yang sesuai, serta mempertahankan keterlacakan sampai pendaftaran dan pembayaran. Proses tersebut melibatkan interaksi manusia, data sekolah dan wilayah, variasi program, target cabang, serta banyak kanal komunikasi.")
    b.add_body(doc, "Pada kondisi operasional multi-cabang, informasi yang tersimpan terpisah pada percakapan personal, spreadsheet, catatan manual, atau aplikasi berbeda akan menyulitkan koordinasi. Nomor WhatsApp yang sama dapat masuk melalui sumber berbeda, lead dapat ditangani tanpa histori yang lengkap, dan manajemen tidak memperoleh gambaran konversi secara konsisten. Karena itu CRM tidak cukup hanya menjadi tempat menyimpan kontak; CRM harus menjadi sistem kendali proses dan sumber data keputusan.")
    b.add_body(doc, "SIVMI CRM dibangun sebagai aplikasi web berbahasa Indonesia menggunakan Laravel 12, Filament 5, Livewire 4, Tailwind CSS 4, MySQL, Redis, dan Docker. Integrasi WAHA/GOWS menghubungkan nomor WhatsApp masing-masing admin atau staff sehingga pesan dapat diterima, dinormalisasi, dicocokkan, disimpan, dan ditampilkan pada inbox CRM sesuai cabang dan pemilik session.")
    table(doc, "Tabel 1-1. Faktor pendorong pengembangan", ["Faktor", "Kondisi", "Implikasi sistem"], [
        ["Multi-kanal", "Lead berasal dari WhatsApp, iklan, Instagram, kunjungan, saudara, Facebook, dan TikTok", "Sumber lead dan identitas harus dinormalisasi"],
        ["Multi-cabang", "Bandung, JakSel, dan JakPus mempunyai tim dan target", "Semua query dan aksi harus mengikuti scope cabang"],
        ["Multi-role", "Superadmin, Admin Cabang, Staff, dan Direksi membutuhkan fungsi berbeda", "RBAC harus diterapkan hingga tingkat data dan widget"],
        ["Produk bervariasi", "Program dibentuk dari kategori, kelas program, dan metode belajar", "Nama program konsisten, contoh SR GOLD ONLINE"],
        ["Komunikasi intensif", "Chat menjadi sumber lead dan bukti follow-up", "Webhook, queue, dedupe, inbox, status pesan, dan audit diperlukan"],
        ["Keputusan cepat", "Target, tren, ranking, sekolah, dan wilayah perlu dipantau", "Dashboard harus memakai data nyata dan filter global"],
    ], [1.25, 2.75, 2.5])
    callout(doc, "Prinsip latar belakang", "Nilai SIVMI CRM terletak pada konsistensi siklus data dan kontrol proses, bukan sekadar jumlah menu atau form yang tersedia.")

    b.add_heading(doc, "1.2 Identifikasi Masalah", 2)
    b.add_body(doc, "Identifikasi masalah disusun berdasarkan alur kerja bisnis, data, akses, integrasi, pengalaman pengguna, dan operasional teknologi. Setiap masalah perlu ditangani pada lapisan yang tepat agar solusi tidak hanya memperbaiki tampilan tetapi juga mencegah inkonsistensi data dan risiko akses.")
    add_figure(doc, "peta-masalah.png", "Gambar 1-2. Peta masalah operasional SIVMI CRM")
    table(doc, "Tabel 1-2. Identifikasi masalah utama", ["ID", "Masalah", "Risiko", "Kebutuhan penanganan"], [
        ["M-01", "Data lead tersebar dan format nomor berbeda", "Duplikasi atau identitas gagal dicocokkan", "Normalisasi telepon, canonical identity, unique constraint, dan konflik eksplisit"],
        ["M-02", "Status lead tidak seragam", "Pipeline dan laporan tidak dapat dibandingkan", "Enum Cold, Warm, Hot, Daftar, dan Tidak Berminat"],
        ["M-03", "Follow-up tidak mempunyai histori/jadwal konsisten", "Calon siswa terlambat dihubungi", "Catatan, metode, hasil, status kontak, dan jadwal berikutnya"],
        ["M-04", "Percakapan WhatsApp terpisah dari CRM", "Bukti interaksi dan pemilik penanganan hilang", "Session per user, webhook, inbox, chat, delivery status, dan audit"],
        ["M-05", "Data lintas cabang berpotensi tercampur", "Kebocoran data dan salah tindak lanjut", "Policy, query scope, filter, export, notification, dan widget sesuai cabang"],
        ["M-06", "Sekolah baru belum ada di master", "Input lead tertunda atau data bebas tidak konsisten", "Create-on-demand dengan prefix SMAN/SMAS/SMKN/SMKS/MAN dan Alumni/Gapyear"],
        ["M-07", "Indikator kinerja tidak terpusat", "Target dan prioritas sulit dievaluasi", "KPI, tren, progress target, ranking, program, kategori, tipe, metode, kelas, jenjang, sekolah, dan wilayah"],
        ["M-08", "Lingkungan pengembangan sulit direplikasi", "Perbedaan runtime dan kegagalan deployment", "Full Docker lokal, konfigurasi environment, health, migration, queue, scheduler, dan test"],
    ], [0.6, 2.2, 1.65, 2.05])
    bullets(doc, [
        "Masalah data harus dicegah melalui constraint, normalisasi, validasi, dan transaksi; bukan hanya pesan peringatan pada UI.",
        "Masalah akses harus diuji pada menu, aksi, query Eloquent, relasi, widget dashboard, file import/export, notifikasi, dan endpoint integrasi.",
        "Masalah integrasi harus mempertimbangkan retry, event duplikat, urutan pesan, session gagal, QR kedaluwarsa, dan provider tidak tersedia.",
        "Masalah visual harus mempunyai skeleton loading, empty state, error state, layout responsif, dan transisi yang tidak menghambat operasi.",
    ])
    callout(doc, "Prioritas risiko", "Kebocoran lintas cabang, duplikasi lead, kehilangan event WhatsApp, dan transaksi yang tidak dapat diaudit merupakan risiko prioritas tinggi karena memengaruhi data serta kepercayaan pengguna.", b.LIGHT_AMBER)

    b.add_heading(doc, "1.3 Rumusan Masalah", 2)
    b.add_body(doc, "Rumusan masalah mengubah temuan operasional menjadi pertanyaan desain dan evaluasi. Pertanyaan ini menjadi penghubung antara pendahuluan, kebutuhan pada BAB III, arsitektur pada BAB IV-VII, pengujian pada BAB VIII, dan operasi pada BAB IX-XI.")
    table(doc, "Tabel 1-3. Rumusan masalah", ["No.", "Pertanyaan"], [
        ["1", "Bagaimana membangun satu siklus data calon siswa dari lead, follow-up, pendaftaran, pembayaran, hingga pelaporan yang konsisten dan dapat ditelusuri?"],
        ["2", "Bagaimana memastikan Superadmin, Admin Cabang, Staff, dan Direksi memperoleh menu, aksi, data, widget, export, serta notifikasi sesuai role dan cabang?"],
        ["3", "Bagaimana menghubungkan nomor WhatsApp masing-masing admin/staff melalui WAHA/GOWS sehingga pesan terekam otomatis dan terkait dengan lead yang benar?"],
        ["4", "Bagaimana mendeteksi nomor yang sudah terdaftar, mencegah duplikasi lead, dan menangani konflik identitas atau cabang tanpa kehilangan pesan?"],
        ["5", "Bagaimana menyediakan dashboard KPI dan target yang mengikuti filter cabang, tahun, bulan termasuk Semua Bulan, serta pengguna dengan query yang efisien?"],
        ["6", "Bagaimana mendukung input sekolah baru secara terkendali, NPSN opsional, grade A-E, prefix sekolah yang valid, serta pilihan Alumni/Gapyear?"],
        ["7", "Bagaimana memastikan aplikasi responsif, memiliki loading/empty/error state, mudah dioperasikan, dan tidak bergantung pada angka dummy di produksi?"],
        ["8", "Bagaimana menyiapkan lingkungan full Docker yang dapat diuji di lokal dan dipindahkan ke VPS dengan prosedur migration, backup, observasi, serta rollback?"],
    ], [0.55, 5.95])
    callout(doc, "Kriteria jawaban", "Setiap pertanyaan dinyatakan terjawab apabila terdapat implementasi, aturan data, bukti pengujian, hasil UAT, dan prosedur operasional yang saling dapat ditelusuri.")

    b.add_heading(doc, "1.4 Manfaat Sistem", 2)
    b.add_body(doc, "Manfaat SIVMI CRM dikelompokkan menjadi manfaat operasional, manajerial, data, layanan, teknologi, dan tata kelola. Manfaat yang bersifat kuantitatif harus diukur menggunakan baseline dan periode pembanding yang disepakati; dokumen ini tidak menetapkan persentase keberhasilan tanpa data aktual.")
    add_figure(doc, "peta-manfaat.png", "Gambar 1-3. Peta manfaat menurut pemangku kepentingan")
    table(doc, "Tabel 1-4. Manfaat sistem dan indikator", ["Dimensi", "Manfaat", "Indikator evaluasi"], [
        ["Operasional", "Lead dan follow-up berada dalam satu daftar kerja", "Jumlah tindakan tertunda, waktu respons, dan kelengkapan histori"],
        ["Layanan", "Chat dan identitas calon siswa terhubung", "Pesan tersimpan, status delivery, percakapan terhubung, dan konflik tertangani"],
        ["Manajerial", "Kinerja cabang/user dan target terlihat", "KPI terfilter, tren, ranking, progress target, serta top program/sekolah/wilayah"],
        ["Data", "Data lebih konsisten dan tidak ganda", "Normalisasi nomor, unique constraint, validasi enum, dan audit perubahan"],
        ["Keamanan", "Pengguna hanya mengakses ruang lingkup yang diizinkan", "Uji policy/scope negatif, audit login/aksi, dan export terscope"],
        ["Teknologi", "Lingkungan dapat diulang dan dipulihkan", "Build/test lulus, health service, backup dapat direstore, dan deployment terdokumentasi"],
    ], [1.2, 3.0, 2.3])
    bullets(doc, [
        "Staff memperoleh konteks lead, sekolah, minat program, percakapan, dan tindak lanjut pada ruang kerja yang sama.",
        "Admin Cabang dapat memantau distribusi tugas, performa tim, session WhatsApp, target, dan konflik data pada cabangnya.",
        "Direksi memperoleh ringkasan lintas cabang tanpa perlu mengubah data operasional.",
        "Superadmin dan administrator mempunyai kontrol atas master data, RBAC, pengaturan, audit, kesehatan layanan, dan konfigurasi integrasi.",
        "Calon siswa menerima tindak lanjut yang lebih konsisten karena histori tidak bergantung pada perangkat pribadi atau ingatan petugas.",
    ])
    callout(doc, "Pengukuran manfaat", "Manfaat dinilai melalui KPI operasional, UAT per role, audit kualitas data, observasi performa, serta keberhasilan backup/restore. Target numerik ditetapkan oleh pemilik bisnis pada dokumen KPI atau rencana rilis.", b.LIGHT_GREEN)

    b.add_heading(doc, "1.5 Tujuan Pembuatan", 2)
    b.add_body(doc, "Tujuan umum pembuatan SIVMI CRM adalah menyediakan sistem informasi CRM bimbel yang terpusat, aman, dapat ditelusuri, responsif, dan siap dioperasikan untuk mendukung pertumbuhan calon siswa serta pengendalian kinerja organisasi.")
    add_figure(doc, "rantai-sasaran.png", "Gambar 1-4. Rantai sasaran SIVMI CRM")
    table(doc, "Tabel 1-5. Tujuan khusus", ["ID", "Tujuan", "Keluaran utama", "Kriteria penerimaan ringkas"], [
        ["T-01", "Memusatkan siklus calon siswa", "Lead, follow-up, pendaftaran, pembayaran", "Relasi dan status dapat ditelusuri tanpa pencatatan ganda"],
        ["T-02", "Menerapkan kontrol akses", "RBAC dan branch scope", "Uji positif/negatif sesuai empat role dan tiga cabang"],
        ["T-03", "Mengintegrasikan WhatsApp", "Session per user, webhook, inbox, chat", "Pesan masuk/keluar dan statusnya tersimpan secara idempotent"],
        ["T-04", "Meningkatkan kualitas data", "Normalisasi, enum, validation, dedupe", "Nomor identik tidak otomatis menciptakan lead baru"],
        ["T-05", "Menyediakan analitik", "KPI, target, tren, ranking, top list", "Semua widget mengikuti filter dan scope pengguna"],
        ["T-06", "Mendukung interoperabilitas data", "Import/export CSV sekolah dan lead", "Template, validasi, error report, dan export terscope tersedia"],
        ["T-07", "Menjamin kualitas antarmuka", "Responsive UI dan state visual", "Mobile-tablet-desktop, skeleton, empty, error, dan transisi diuji"],
        ["T-08", "Menyiapkan operasi berulang", "Docker, test, build, health, backup", "Instalasi lokal dapat direplikasi dan target VPS memiliki runbook"],
    ], [0.55, 2.1, 1.85, 2.0])
    callout(doc, "Tujuan bukan klaim hasil", "Tujuan menyatakan keadaan yang ingin dicapai. Pencapaian akhir harus dibuktikan melalui hasil UAT, metrik operasi, dan persetujuan pemilik sistem.")

    b.add_heading(doc, "1.6 Ruang Lingkup Sistem", 2)
    b.add_body(doc, "Ruang lingkup menetapkan fungsi yang dibahas dan batas tanggung jawab sistem. Penetapan ini menghindari perluasan kebutuhan tanpa kendali serta membantu tim menentukan acceptance criteria, risiko integrasi, dan rencana rilis.")
    add_figure(doc, "batas-ruang-lingkup.png", "Gambar 1-5. Batas ruang lingkup fungsional")
    table(doc, "Tabel 1-6. Modul dalam ruang lingkup", ["Kelompok", "Cakupan"], [
        ["Akses dan identitas", "Login, session, profil, role, branch scope, notifikasi, dan audit aktivitas"],
        ["Dashboard", "Filter cabang/tahun/bulan/user, ringkasan hari ini, KPI, target, tren, aktivitas, quick action, ranking, pending action, KPI program, kategori, tipe, metode, kelas, jenjang, sekolah, dan wilayah"],
        ["Lead", "CRUD terbatas sesuai role, status Cold/Warm/Hot/Daftar/Tidak Berminat, sumber lead, PIC, sekolah, minat program, import/export CSV, dan dedupe"],
        ["Follow-up", "Waktu, metode, status kontak, catatan hasil, jadwal berikutnya, tindakan tertunda, dan keterkaitan lead"],
        ["Pendaftaran dan pembayaran", "Pendaftaran siswa, pemilihan program, status aktif, pencatatan/validasi pembayaran, dan omzet"],
        ["Program bimbel", "Kategori, kelas program, metode belajar, tipe, jenjang, tingkat kelas, biaya, diskon, pertemuan, fasilitas, status aktif, serta nama otomatis tiga unsur"],
        ["Master sekolah", "NPSN opsional, jenis/prefix SMAN/SMAS/SMKN/SMKS/MAN, Alumni/Gapyear, grade A-E, wilayah, status, create-on-demand, serta import/export CSV"],
        ["WhatsApp", "Nomor per admin/staff, session WAHA/GOWS, QR, webhook, penyimpanan event, matching lead, konflik duplikat, inbox, chat, dan status pesan"],
        ["Kontrol dan laporan", "Target kinerja, laporan, pengaturan termasuk program, pengguna, cabang, sekolah, audit, dan pembatasan role"],
        ["Operasi teknologi", "Full Docker lokal, Nginx/PHP-FPM, MySQL, Redis, queue, scheduler, storage, health check, log, backup/restore, test/build, dan target deployment VPS"],
    ], [1.55, 4.95])
    add_figure(doc, "batas-sistem.png", "Gambar 1-6. Batas sistem dan ketergantungan eksternal")
    table(doc, "Tabel 1-7. Di luar ruang lingkup baseline", ["Area", "Batasan"], [
        ["Seluruh CRUD organisasi", "Pengembangan difokuskan pada shell dashboard dan modul prioritas; tidak semua proses pendidikan/akademik dibangun pada baseline"],
        ["Learning management system", "Materi, kelas virtual, ujian daring, penilaian akademik, dan presensi pembelajaran penuh bukan fungsi inti CRM saat ini"],
        ["Akuntansi lengkap", "Jurnal umum, buku besar, pajak, payroll, dan rekonsiliasi bank bukan cakupan pembayaran CRM"],
        ["WhatsApp Cloud API resmi", "Baseline menggunakan WAHA/GOWS; migrasi provider memerlukan analisis API, biaya, template, consent, dan kebijakan terpisah"],
        ["AI otomatis mengambil keputusan", "AI hanya direkomendasikan sebagai alat bantu; assignment, status, respons, dan keputusan bisnis tetap dikendalikan manusia serta aturan sistem"],
        ["Produksi shared hosting", "Target operasional diarahkan ke VPS/container; shared hosting hanya mungkin bila memenuhi kebutuhan worker, scheduler, webhook, storage, dan konfigurasi runtime"],
        ["Angka manfaat yang dijamin", "Peningkatan konversi atau penurunan waktu respons harus diukur dari data aktual, bukan diasumsikan dalam kode atau dokumen"],
    ], [1.85, 4.65])
    table(doc, "Tabel 1-8. Asumsi dan ketergantungan", ["Jenis", "Asumsi/ketergantungan", "Konsekuensi"], [
        ["Organisasi", "Role, cabang, PIC, target, dan aturan bisnis disahkan pemilik sistem", "Perubahan kebijakan memerlukan update konfigurasi, policy, data, test, dan dokumen"],
        ["Data", "Nomor WhatsApp dapat dinormalisasi dan data master mempunyai nilai yang valid", "Data lama perlu cleansing sebelum import produksi"],
        ["WAHA", "Session GOWS Working, API key benar, webhook dapat mencapai aplikasi", "Gangguan provider ditangani retry, queue, health, dan rekonsiliasi"],
        ["Infrastruktur", "Docker host menyediakan CPU, RAM, disk, port, DNS, TLS, dan backup memadai", "Sizing dan monitoring ditentukan sebelum go-live VPS"],
        ["Keamanan", "Secret tidak disimpan di repository dan akses produksi dibatasi", "Rotasi key, HTTPS, firewall, log aman, dan audit menjadi syarat operasi"],
        ["Operasi", "Queue worker dan scheduler berjalan terus", "Reminder, webhook processing, dan pekerjaan asinkron gagal bila worker berhenti"],
    ], [1.05, 3.5, 1.95])
    callout(doc, "Batas perubahan", "Permintaan baru yang mengubah aktor, proses inti, model data, provider eksternal, keamanan, biaya, atau target deployment harus melalui analisis dampak dan persetujuan sebelum dimasukkan ke baseline.", b.LIGHT_AMBER)

    b.add_heading(doc, "PENUTUP", 1)
    b.add_body(doc, "BAB I menetapkan alasan, masalah, pertanyaan, manfaat, tujuan, dan batas pengembangan SIVMI CRM. Landasan ini menjadi konteks untuk teori pada BAB II, analisa kebutuhan pada BAB III, arsitektur bisnis/data/aplikasi/teknologi pada BAB IV-VII, pengujian dan debugging pada BAB VIII, serta deployment dan operasi pada bab berikutnya.")
    callout(doc, "Prinsip kesinambungan", "Setiap perubahan kebutuhan harus memperbarui dokumen terkait, diagram, migration/model, policy, test, UAT, dan runbook agar implementasi serta dokumentasi tetap sinkron.", b.LIGHT_GREEN)

    doc.core_properties.title = "BAB I - Pendahuluan SIVMI CRM"
    doc.core_properties.subject = "Pendahuluan sistem CRM bimbingan belajar berdasarkan baseline aplikasi terbaru"
    doc.core_properties.author = "SIVMI CRM"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(f"Created: {OUTPUT}")


if __name__ == "__main__":
    build_document()
